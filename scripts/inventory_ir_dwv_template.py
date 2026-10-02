"""Emit a deterministic JSON inventory for the IR/DWV workbook template."""

from __future__ import annotations

import argparse
from copy import copy
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def purify_template(source: Path, output: Path) -> Path:
    """Create a blank first-block asset; never replace the source or another file."""
    from backend.infrastructure.office.ir_dwv_record_workbook_gateway import (
        IrDwvRecordWorkbookGateway, copy_block, copy_block_logos, template_block_logos,
    )
    from backend.infrastructure.office.ir_dwv_record_workbook_layout import block_layout
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve():
        raise ValueError("Template source and purified output must be different files.")
    if output.exists():
        raise FileExistsError("Purified output already exists; choose a new path.")
    IrDwvRecordWorkbookGateway(source).template_fingerprint()
    original = load_workbook(source, data_only=False)
    # Keep the source style table in memory: raw StyleArray IDs are workbook-local.
    purified = original
    try:
        prototype, sheet = original.worksheets[0], purified.create_sheet("__purified")
        sheet.title = "IR&DWV Template"
        for attribute in ("sheet_format", "sheet_properties", "page_margins", "page_setup", "print_options", "oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"):
            setattr(sheet, attribute, copy(getattr(prototype, attribute)))
        purified.calculation = copy(original.calculation)
        sheet.sheet_view.showGridLines = prototype.sheet_view.showGridLines
        sheet.freeze_panes = prototype.freeze_panes
        copy_block(prototype, sheet, source_min_row=1, source_max_row=40,
                   source_min_column=1, source_max_column=14, target_min_row=1, target_min_column=1)
        copy_block_logos(template_block_logos(prototype), sheet,
                         block_layout(origin_column=2, sample_count=5, pair_count=0))
        for row in range(3, 8):
            for column in (5, 8, 9, 10):
                if row != 3 or column not in (5, 8):
                    sheet.cell(row, column).value = None
        for coordinate, label in {"B8": "Start Date:", "B9": "Finish Date:", "K8": "Amb Temp:", "K9": "Rel. Hum.:",
                                  "B10": "Item/Process:", "I10": "Request  No.:", "B11": "Remarks:", "I11": "Product Name:",
                                  "B12": "Tested By:", "E12": "Checked By:", "H12": "Approved By:", "K12": "Requestor:"}.items():
            sheet[coordinate].value = label
        for row in range(14, 19):
            sheet.cell(row, 2).value = None
            sheet.cell(row, 8).value = None
        for row in range(22, 34):
            for column in range(2, 13):
                sheet.cell(row, column).value = None
        sheet.print_area = "B1:L40"
        for other in tuple(purified.worksheets):
            if other is not sheet:
                purified.remove(other)
        output.parent.mkdir(parents=True, exist_ok=True)
        # Exclusive creation prevents an unnoticed race from overwriting any file.
        with output.open("xb") as stream:
            try:
                purified.save(stream)
            except Exception:
                stream.close()
                output.unlink(missing_ok=True)
                raise
    finally:
        purified.close()
    return output


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _cell_inventory(cell: Any) -> dict[str, Any]:
    return {
        "coordinate": cell.coordinate,
        "value": _json_value(cell.value),
        "data_type": cell.data_type,
        "style_id": cell.style_id,
        "number_format": cell.number_format,
    }


def _conditional_formatting(sheet: Any) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for item in sheet.conditional_formatting:
        entries.append(
            {
                "sqref": str(item.sqref),
                "rules": [
                    {
                        "type": rule.type,
                        "priority": rule.priority,
                        "operator": rule.operator,
                        "formula": list(rule.formula or ()),
                        "stop_if_true": rule.stopIfTrue,
                    }
                    for rule in item.rules
                ],
            }
        )
    return entries


def _data_validations(sheet: Any) -> list[dict[str, Any]]:
    return [
        {
            "sqref": str(validation.sqref),
            "type": validation.type,
            "operator": validation.operator,
            "formula1": validation.formula1,
            "formula2": validation.formula2,
            "allow_blank": validation.allowBlank,
            "show_error_message": validation.showErrorMessage,
            "show_input_message": validation.showInputMessage,
        }
        for validation in sheet.data_validations.dataValidation
    ]


def inventory_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            cells = [
                _cell_inventory(sheet.cell(row=row, column=column))
                for row in range(1, sheet.max_row + 1)
                for column in range(1, sheet.max_column + 1)
            ]
            sheets.append(
                {
                    "title": sheet.title,
                    "state": sheet.sheet_state,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "cells": cells,
                    "formulas": [
                        {"coordinate": cell["coordinate"], "formula": cell["value"]}
                        for cell in cells
                        if isinstance(cell["value"], str) and cell["value"].startswith("=")
                    ],
                    "merged_cells": sorted(str(item) for item in sheet.merged_cells.ranges),
                    "row_heights": {
                        str(index): dimension.height
                        for index, dimension in sheet.row_dimensions.items()
                        if dimension.height is not None
                    },
                    "column_widths": {
                        label: dimension.width
                        for label, dimension in sheet.column_dimensions.items()
                        if dimension.width is not None
                    },
                    "conditional_formatting": _conditional_formatting(sheet),
                    "data_validations": _data_validations(sheet),
                    "print": {
                        "print_area": str(sheet.print_area or ""),
                        "print_title_rows": sheet.print_title_rows,
                        "print_title_cols": sheet.print_title_cols,
                        "freeze_panes": str(sheet.freeze_panes or ""),
                        "sheet_view_show_grid_lines": sheet.sheet_view.showGridLines,
                        "page_setup": {
                            "orientation": sheet.page_setup.orientation,
                            "paper_size": sheet.page_setup.paperSize,
                            "fit_to_width": sheet.page_setup.fitToWidth,
                            "fit_to_height": sheet.page_setup.fitToHeight,
                            "scale": sheet.page_setup.scale,
                        },
                        "page_margins": {
                            "left": sheet.page_margins.left,
                            "right": sheet.page_margins.right,
                            "top": sheet.page_margins.top,
                            "bottom": sheet.page_margins.bottom,
                            "header": sheet.page_margins.header,
                            "footer": sheet.page_margins.footer,
                        },
                        "print_options": {
                            "horizontal_centered": sheet.print_options.horizontalCentered,
                            "vertical_centered": sheet.print_options.verticalCentered,
                            "headings": sheet.print_options.headings,
                            "grid_lines": sheet.print_options.gridLines,
                        },
                    },
                }
            )
        return {
            "schema": "connlab.ir-dwv-template-inventory",
            "version": 1,
            "workbook": {
                "sheet_names": workbook.sheetnames,
                "calculation_mode": workbook.calculation.calcMode,
                "full_calc_on_load": workbook.calculation.fullCalcOnLoad,
                "force_full_calc": workbook.calculation.forceFullCalc,
            },
            "sheets": sheets,
        }
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--purify-output", type=Path, help="Create a separate blank template without overwriting any file.")
    args = parser.parse_args()
    if args.purify_output is not None:
        print(purify_template(args.workbook, args.purify_output))
        return
    payload = json.dumps(
        inventory_workbook(args.workbook), ensure_ascii=False, indent=2, sort_keys=True
    )
    if args.output is None:
        print(payload)
        return
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(payload + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
