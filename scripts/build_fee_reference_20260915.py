"""Capture the approved September Fee reference and compile its immutable runtime seed.

Run once at rule-version review time, never during Fee Evaluation generation.
"""

from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from backend.modules.fee_evaluation.fee_rule_seed_compiler import compile_fee_rule_seed_files


ROOT = Path(__file__).resolve().parents[1]
SEEDS = ROOT / "backend" / "modules" / "fee_evaluation" / "seeds"
SOURCE_HASH = "c230445ad8620ee4d6ad178717d45f0a7d75b249342db3e98b752aa5db56d27f"
SOURCE_NAME = "FDQF-E-176 Testing Fee Evaluation_Rev_F_20260915.xlsx"
VERSION_ID = "fee_rules_v2026_09_15"
INSERTED_ROWS = {
    18: ("fee_rule_shock_half_sine_high_g", ["Mechanical Shock >50G"], None, "60", "time", "manual_required", "Confirm >50G condition, occurrences and base fee."),
    33: ("fee_rule_impulse_voltage", ["Impulse Voltage Test"], None, "10", "reading", "manual_required", "Confirm the number of readings and specimen-preparation base fee."),
    41: ("fee_rule_fluid_resistance", ["Fluid resistance testing"], None, "700", "reagent", "manual_required", "Confirm reagents, units and any additional chemical cost."),
    44: ("fee_rule_ct_analysis", ["Computed tomography analysis"], "0", "3000", "specimen", "manual_required", "Confirm specimen count, scan duration and discount."),
    45: ("fee_rule_dsc_tga", ["DSC", "TGA", "DSC analysis", "TGA analysis"], "0", "650", "specimen", "manual_required", "Confirm DSC or TGA test and specimen count."),
}


def _source_row_for_old(row: int) -> int:
    return row + (row >= 18) + (row >= 32) + (row >= 39) + 2 * (row >= 41)


def _cell_text(value: object) -> str:
    return "" if value is None else str(value)


def main(workbook_path: Path) -> None:
    if hashlib.sha256(workbook_path.read_bytes()).hexdigest() != SOURCE_HASH:
        raise ValueError("Fee reference workbook differs from the reviewed 2026-09-15 source.")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Unit Price Reference"]
        rows = [
            {
                "source_row": row,
                "english_description": _cell_text(sheet.cell(row, 2).value),
                "chinese_description": _cell_text(sheet.cell(row, 3).value),
                "base_fee_text": _cell_text(sheet.cell(row, 4).value),
                "unit_price_text": _cell_text(sheet.cell(row, 5).value),
                "applicable_standard": _cell_text(sheet.cell(row, 6).value),
                "range_condition": _cell_text(sheet.cell(row, 7).value),
                "chamber_or_note": _cell_text(sheet.cell(row, 8).value),
            }
            for row in range(4, 53)
        ]
        policy_text = str(sheet.cell(54, 2).value or "")
    finally:
        workbook.close()

    snapshot = {
        "source": {
            "source_file_name": SOURCE_NAME,
            "source_sheet": "Unit Price Reference",
            "source_hash": f"sha256:{SOURCE_HASH}",
            "captured_at": "2026-09-15T07:04:23+08:00",
        },
        "rows": rows,
        "policies": [{"source_row": 54, "policy_type": "discount_principles", "text": policy_text}],
    }
    reviewed = json.loads((SEEDS / "fee_rule_extensions_v2026_08_23_r11.json").read_text(encoding="utf-8"))
    reviewed["version"] = {
        "version_id": VERSION_ID,
        "source_file_name": SOURCE_NAME,
        "source_sheet": "Unit Price Reference",
        "source_hash": f"sha256:{SOURCE_HASH}",
        "effective_from_basis": "project.sample_received_date",
        "created_at": "2026-09-29T00:00:00+08:00",
    }
    rules = reviewed["source_rules"]
    for rule in rules:
        rule["source_row"] = _source_row_for_old(rule["source_row"])
    by_id = {rule["rule_id"]: rule for rule in rules}
    for rule_id in (
        "fee_rule_high_temperature_life", "fee_rule_low_temperature_life",
        "fee_rule_temperature_humidity", "fee_rule_steam_aging",
        "fee_rule_thermal_shock", "fee_rule_thermal_cycling_3_5c",
        "fee_rule_thermal_cycling_5c",
    ):
        by_id[rule_id]["base_fee_amount"] = "200"
    for rule_id, amount in {
        "fee_rule_thermal_cycling_3_5c": "30",
        "fee_rule_mfg_class_iiia": "1500",
        "fee_rule_dust_benign": "3000",
        "fee_rule_plating_thickness": "15",
        "fee_rule_visual_exam": "15",
        "fee_rule_pcb_fixture_design": "200",
    }.items():
        by_id[rule_id]["unit_price_amount"] = amount
    by_id["fee_rule_dust_benign"].update(
        unit_label="time", calculation_strategy="manual_required", review_required=True,
        review_reason="Confirm one-hour run count and dust preparation fee.",
    )
    by_id["fee_rule_report_preparation"].update(
        unit_price_amount=None, calculation_strategy="manual_required", review_required=True,
        review_reason="Confirm multi-group waiver or single-group page tier (500/800 RMB).",
    )
    by_id["fee_rule_insulation_resistance"]["review_reason"] = (
        "Confirm test-point readings per specimen and the 1-minute/2-minute price."
    )
    by_id["fee_rule_dielectric_withstanding_voltage"]["review_reason"] = (
        "Confirm test-point readings per specimen and the 1-minute/2-minute price."
    )
    by_id["fee_rule_shock_half_sine"]["review_reason"] = (
        "Confirm <=50G condition, occurrences and base fee."
    )
    for row, (rule_id, aliases, base_fee, unit_price, unit, strategy, reason) in INSERTED_ROWS.items():
        rules.append({
            "source_row": row, "rule_id": rule_id, "aliases": aliases,
            "base_fee_amount": base_fee, "unit_price_amount": unit_price,
            "unit_label": unit, "calculation_strategy": strategy,
            "review_required": True, "review_reason": reason,
        })
    rules.sort(key=lambda rule: rule["source_row"])
    snapshot_path = SEEDS / "fee_reference_rows_v2026_09_15.json"
    extensions_path = SEEDS / "fee_rule_extensions_v2026_09_15.json"
    output_path = SEEDS / f"{VERSION_ID}.json"
    snapshot_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    extensions_path.write_text(json.dumps(reviewed, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    compile_fee_rule_seed_files(snapshot_path, extensions_path, output_path)
    print(f"Compiled {VERSION_ID} from {SOURCE_NAME}.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python scripts/build_fee_reference_20260915.py <approved-workbook.xlsx>")
    main(Path(sys.argv[1]))
