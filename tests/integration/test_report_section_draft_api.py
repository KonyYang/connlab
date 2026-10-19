"""Public report update requests with real files/DB and an isolated Word boundary."""
from dataclasses import replace
from hashlib import sha256
from types import SimpleNamespace
from io import BytesIO
import shutil

from docx import Document
import pytest

from backend.api import dependencies as deps
from backend.infrastructure.storage.repositories.result_dataset import ResultDatasetRepository
from backend.infrastructure.storage.repositories import ConfirmedMatrixAuthorityRepository
from tests.integration.test_internal_report_generation_api import lab, official_report
from tests.unit.test_downloaded_internal_report_source import draft
from tests.unit.test_result_dataset_repository import _dataset


BASE = "/api/projects/P1/report-workspace/update-source"
DL = "DL-2026-05-001"


@pytest.fixture
def section_lab(lab, monkeypatch):
    downloads = lab["root"] / "downloads"
    downloads.mkdir()
    original = draft(downloads, dl=DL)
    def synchronize(*, source_path, output_path, dataset):
        document = Document(source_path)
        document.add_paragraph("Confirmed LLCR results")
        document.save(output_path)
        return output_path
    monkeypatch.setattr(lab["writer"], "synchronize_llcr_results", synchronize, raising=False)
    with lab["factory"]() as session:
        matrix = ConfirmedMatrixAuthorityRepository(session).get_active_by_project("P1")
        dataset = replace(_dataset("dataset-1", 1), confirmed_matrix_id=matrix.version.confirmed_matrix_id,
                          confirmed_matrix_revision=matrix.version.confirmed_revision)
        ResultDatasetRepository(session).create_dataset(dataset)
        session.commit()
    return {**lab, "original": original, "downloads": downloads}


def test_llcr_updates_a_new_draft_and_customer_projection_uses_it(section_lab):
    lab = section_lab
    before = lab["original"].read_bytes()
    source = lab["client"].get(BASE)
    assert source.status_code == 200, source.text
    assert source.json()["mode"] == "download_draft"
    result = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1",
        "expected_report_sha256": source.json()["file_sha256"]})
    assert result.status_code == 200, result.text
    assert result.json()["mode"] == "download_draft"
    artifact = lab["client"].get(result.json()["download_url"])
    assert artifact.status_code == 200
    assert "draft" in artifact.headers["content-disposition"].lower()
    assert lab["original"].read_bytes() == before
    current = lab["client"].get(BASE).json()
    assert current["file_sha256"] == sha256(artifact.content).hexdigest()
    customer = lab["client"].get("/api/projects/P1/report-workspace/current-customer-report").json()
    assert customer["internal_report_sha256"] == current["file_sha256"]
    assert customer["can_generate"]
    assert lab["client"].get("/api/projects/P1/report-workspace").json()["report_revisions"] == []
    assert lab["client"].get("/api/projects/P1/report-workspace/current-report").json()["status"] == "missing"


def test_equipment_chains_after_llcr_and_downloads_do_not_modify_original(section_lab, monkeypatch):
    lab = section_lab
    catalog = lab["root"] / "catalog.xls"
    catalog.write_bytes(b"calibration boundary")
    reader = SimpleNamespace(read_equipment_calibrations=lambda: SimpleNamespace(resource_path=catalog, rows=(
        SimpleNamespace(equipment_id="DG-Q-0033", equipment_name="Meter", manufacturer="Maker",
                        last_calibration_date="2026-01-01", calibration_due_date="2027-01-01", source_sheet="All"),)))
    monkeypatch.setattr(deps, "ExternalExcelReadService", lambda resources: reader)
    def equipment(*, source_path, output_path, rows):
        document = Document(source_path)
        document.add_paragraph("Equipment: " + rows[0].id_number)
        document.save(output_path)
        return output_path
    monkeypatch.setattr(lab["writer"], "synchronize_equipment_list", equipment, raising=False)
    initial = lab["original"].read_bytes()
    first = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1",
        "expected_report_sha256": lab["client"].get(BASE).json()["file_sha256"]}).json()
    missing = lab["client"].post(BASE + "/equipment", data={"expected_report_sha256": first["current_sha256"]})
    assert missing.status_code == 200, missing.text
    assert missing.json()["status"] == "source_required"
    result = lab["client"].post(BASE + "/equipment", data={"expected_report_sha256": first["current_sha256"],
        "references_text": "DG-Q-0033"})
    assert result.status_code == 200, result.text
    assert result.json()["mode"] == "download_draft"
    download = lab["client"].get(result.json()["download_url"])
    document = Document(BytesIO(download.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Confirmed LLCR results" in text and "Equipment: DG-Q-0033" in text
    assert lab["client"].get(first["download_url"]).status_code == 200
    assert lab["original"].read_bytes() == initial
    assert not list(lab["downloads"].glob("EquipmentID.docx"))
    customer = lab["client"].get("/api/projects/P1/report-workspace/current-customer-report").json()
    assert customer["internal_report_sha256"] == sha256(download.content).hexdigest()


def test_missing_source_accepts_a_matching_uploaded_report_but_not_another_project(section_lab):
    lab = section_lab
    original = lab["original"].read_bytes()
    lab["original"].rename(lab["original"].with_suffix(".retained"))
    assert lab["client"].get(BASE).json()["status"] == "missing"
    other = draft(lab["downloads"], dl="DL-2026-05-999").read_bytes()
    rejected = lab["client"].post(BASE + "/select", files={"file": ("Other.docx", other)})
    assert rejected.status_code == 409
    selected = lab["client"].post(BASE + "/select", files={"file": ("Manually named.docx", original)})
    assert selected.status_code == 200, selected.text
    assert selected.json()["mode"] == "download_draft"
    response = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1",
        "expected_report_sha256": selected.json()["current_sha256"]})
    assert response.status_code == 200, response.text
    assert lab["original"].with_suffix(".retained").read_bytes() == original


def test_draft_update_rejects_changed_source_and_stale_data_without_advancing(section_lab):
    from backend.infrastructure.storage.models_confirmed_matrix_authority import ConfirmedMatrixVersionModel
    lab = section_lab
    old = lab["client"].get(BASE).json()
    document = Document(lab["original"])
    document.add_paragraph("Manual edit")
    document.save(lab["original"])
    response = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1", "expected_report_sha256": old["file_sha256"]})
    assert response.status_code == 409
    current = lab["client"].get(BASE).json()
    with lab["factory"]() as session:
        session.get(ConfirmedMatrixVersionModel, "cmv-1").confirmed_revision += 1
        session.commit()
    response = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1", "expected_report_sha256": current["file_sha256"]})
    assert response.status_code == 409
    assert "stale" in response.json()["detail"].lower()
    assert lab["client"].get(BASE).json()["file_sha256"] == current["file_sha256"]


def test_unchanged_copy_is_reported_as_no_change(section_lab, monkeypatch):
    lab = section_lab
    monkeypatch.setattr(lab["writer"], "synchronize_llcr_results",
        lambda *, source_path, output_path, dataset: shutil.copy2(source_path, output_path))
    source = lab["client"].get(BASE).json()
    result = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1",
        "expected_report_sha256": source["file_sha256"]})
    assert result.status_code == 200, result.text
    assert result.json()["changed"] is False
    assert lab["client"].get(result.json()["download_url"]).content == lab["original"].read_bytes()


def test_folder_appearing_during_generation_does_not_publish_or_replace_draft(section_lab, monkeypatch):
    lab = section_lab
    source = lab["client"].get(BASE).json()
    previous = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1",
        "expected_report_sha256": source["file_sha256"]}).json()
    # Register the location first; only the report appears while Word is running.
    official, local = official_report(lab)
    preserved = official.read_bytes()
    official.unlink()
    def update(*, source_path, output_path, dataset):
        shutil.copy2(source_path, output_path)
        official.write_bytes(preserved)
    monkeypatch.setattr(lab["writer"], "synchronize_llcr_results", update)
    result = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1",
        "expected_report_sha256": previous["current_sha256"]})
    assert result.status_code == 409, result.text
    assert official.read_bytes() == preserved
    assert not (local / "History").exists()
    assert lab["client"].get(previous["download_url"]).status_code == 200


def test_formal_mode_keeps_archive_behavior_and_draft_mode_cannot_write_it(section_lab, monkeypatch):
    lab = section_lab
    official, local = official_report(lab)
    original = lab["original"].read_bytes()
    before = official.read_bytes()
    def update(*, source_path, output_path, dataset):
        output_path.write_bytes(source_path.read_bytes() + b"|LLCR")
        return output_path
    monkeypatch.setattr(lab["writer"], "synchronize_llcr_results", update)
    source = lab["client"].get(BASE).json()
    assert source["mode"] == "official"
    payload = {"dataset_id": "dataset-1", "expected_report_sha256": source["file_sha256"]}
    rejected = lab["client"].post(BASE + "/llcr", json=payload)
    assert rejected.status_code == 409
    assert official.read_bytes() == before
    result = lab["client"].post(BASE + "/llcr", json={**payload, "expected_mode": "official"})
    assert result.status_code == 200, result.text
    assert result.json()["mode"] == "official"
    assert official.read_bytes() == before + b"|LLCR"
    assert list((local / "History" / "Report").glob("*.docx"))[0].read_bytes() == before
    assert lab["original"].read_bytes() == original


def test_pending_folder_operation_blocks_draft_writes_but_not_retained_download(section_lab):
    from backend.infrastructure.files.generation_journal import GenerationJournal
    lab = section_lab
    source = lab["client"].get(BASE).json()
    previous = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1",
        "expected_report_sha256": source["file_sha256"]}).json()
    GenerationJournal(lab["settings"].data_dir / "project_folder_generation").create("P1", None, "test")
    blocked = lab["client"].post(BASE + "/llcr", json={"dataset_id": "dataset-1",
        "expected_report_sha256": previous["current_sha256"]})
    assert blocked.status_code == 409
    assert "unfinished" in blocked.json()["detail"]
    assert lab["client"].get(previous["download_url"]).status_code == 200
    assert lab["client"].get(previous["download_url"].replace("/P1/", "/P2/")).status_code != 200
