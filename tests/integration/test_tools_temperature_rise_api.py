from io import BytesIO
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.api import dependencies as deps
from backend.api.main import app
from backend.shared.config import Settings


@pytest.fixture
def client(tmp_path):
    data = tmp_path / "data"
    data.mkdir()
    settings = Settings(data_dir=data, templates_dir=tmp_path / "templates", projects_dir=tmp_path / "projects", database_path=tmp_path / "test.sqlite")
    app.dependency_overrides[deps.get_settings] = lambda: settings
    try:
        yield TestClient(app), data
    finally:
        app.dependency_overrides.pop(deps.get_settings)


CSV = b"Scan,Time,1_HS,Ambient,Current\n1,12:00,22,20,10\n2,12:01,26,20,20\n3,12:02,32,20,30\n"


def test_preview_analyze_export_and_cleanup(client):
    http, data = client
    files = {"file": ("../readings.csv", CSV, "text/csv")}
    preview = http.post("/api/tools/temperature-rise/preview", files=files)
    assert preview.status_code == 200, preview.text
    block = preview.json()["blocks"][0]
    options = {"block_id": block["id"], "mapping": block["suggested_mapping"], "selected_rows": [2, 3, 4]}
    analysis = http.post("/api/tools/temperature-rise/analyze", files=files, data={"options": json.dumps(options)})
    assert analysis.status_code == 200, analysis.text
    assert analysis.json()["max_curve"]["a"] == pytest.approx(.01)
    export = http.post("/api/tools/temperature-rise/export", files=files, data={"options": json.dumps(options)})
    assert export.status_code == 200, export.text
    assert "readings_TemperatureRise.xlsx" in export.headers["content-disposition"]
    book = load_workbook(BytesIO(export.content))
    assert book.sheetnames == ["Setup", "T-riseChart", "Derating", "Initial Data"]
    book.close()
    assert list(data.iterdir()) == []


@pytest.mark.parametrize("content,options", [(b"invalid", "{}"), (CSV, "bad json"), (CSV, "[]"), (CSV, '{"block_id":"0"}')])
def test_invalid_input_returns_guidance_without_download_and_cleans(client, content, options):
    http, data = client
    response = http.post("/api/tools/temperature-rise/export", files={"file": ("readings.csv", content)}, data={"options": options})
    assert response.status_code == 422
    assert "content-disposition" not in response.headers
    assert list(data.iterdir()) == []
