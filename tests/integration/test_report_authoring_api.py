"""Report authoring is persistent, project-scoped and copy-only."""
from hashlib import sha256
from io import BytesIO

from docx import Document

from tests.integration.test_internal_report_generation_api import lab


BASE = "/api/projects/P1/report-workspace/authoring"


def docx_bytes(dl="DL-2026-05-001"):
    document = Document()
    document.sections[0].header.paragraphs[0].text = dl
    document.sections[0].footer.paragraphs[0].text = "E-3707"
    document.add_paragraph("Approved manual wording")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def upload(lab, *, role="original", revision=0, content=None, name="Report.docx"):
    return lab["client"].post(BASE + "/sources", data={"role": role, "expected_revision": revision},
        files={"file": (name, content or docx_bytes())})


def test_source_and_confirmed_mapping_survive_new_requests_without_real_file_writes(lab):
    assert lab["client"].get(BASE).json()["revision"] == 0
    original = docx_bytes()
    result = upload(lab, content=original)
    assert result.status_code == 200, result.text
    state = result.json()
    source = state["sources"][0]
    assert source["sha256"] == sha256(original).hexdigest()
    assert source["role"] == "original"
    assert "path" not in source
    response = lab["client"].post(BASE + "/mappings", json={
        "expected_revision": 1, "source_id": source["source_id"], "kind": "specification",
        "source_locator": "revision record page 27", "label": "GS-12-2181 Rev.2",
        "value": "2026-06-18", "status": "confirmed", "reason": "Revision record, not authorization header",
    })
    assert response.status_code == 200, response.text
    reopened = lab["client"].get(BASE).json()
    assert reopened["revision"] == 2
    assert reopened["mappings"][0]["value"] == "2026-06-18"
    assert reopened["mappings"][0]["status"] == "confirmed"
    download = lab["client"].get(source["download_url"])
    assert download.content == original
    assert not list(lab["root"].glob("projects/**/*.docx"))


def test_stale_changes_reject_without_losing_mapping_or_advancing_revision(lab):
    state = upload(lab).json()
    result = upload(lab, revision=0)
    assert result.status_code == 409
    assert lab["client"].get(BASE).json() == state


def test_cross_project_original_is_rejected_but_explicit_layout_reference_is_retained(lab):
    other = docx_bytes("DL-2026-05-999")
    assert upload(lab, content=other).status_code == 422
    reference = upload(lab, content=other, role="layout_reference")
    assert reference.status_code == 200, reference.text
    assert reference.json()["sources"][0]["role"] == "layout_reference"
    assert reference.json()["mappings"] == []


def test_blank_dl_approved_template_is_only_allowed_as_explicit_layout_reference(lab):
    template = docx_bytes('')
    assert upload(lab, content=template, role='original').status_code == 422
    reference = upload(lab, content=template, role='layout_reference', name='Approved template.docx')
    assert reference.status_code == 200, reference.text
    source = reference.json()['sources'][0]
    assert source['report_dl'] == ''
    prepared = lab['client'].post(BASE+'/base', json={'expected_revision': 1, 'source_id': source['source_id']})
    assert prepared.status_code == 200, prepared.text
    output = prepared.json()['outputs'][-1]
    assert any('identity' in message.casefold() for message in output['issues'])
    copy = Document(BytesIO(lab['client'].get(output['download_url']).content))
    assert copy.paragraphs[0].text == 'Approved manual wording'


def test_missing_business_decision_stays_pending_and_cross_project_source_is_not_available(lab):
    state = upload(lab).json()
    response = lab["client"].post(BASE + "/mappings", json={
        "expected_revision": 1, "source_id": state["sources"][0]["source_id"],
        "kind": "acceptance", "source_locator": "MFG report", "label": "Corrosion acceptance",
    })
    assert response.status_code == 200, response.text
    assert response.json()["mappings"][0]["status"] == "pending"
    assert response.json()["mappings"][0]["value"] == ""
    assert lab["client"].get(state["sources"][0]["download_url"].replace("/P1/", "/P2/")).status_code == 404


def test_using_an_approved_reference_creates_a_separate_draft_without_inheriting_pass(lab):
    document = Document(BytesIO(docx_bytes("DL-2026-05-999")))
    document.add_paragraph("Group A Test Results")
    table = document.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Step", "Result", "Comment"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["1", "Pass", "Old corrosion exception"]):
        cell.text = text
    stream = BytesIO()
    document.save(stream)
    original = stream.getvalue()
    state = upload(lab, content=original, role="layout_reference").json()
    source = state["sources"][0]
    response = lab["client"].post(BASE + "/base", json={"source_id": source["source_id"], "expected_revision": 1})
    assert response.status_code == 200, response.text
    state = response.json()
    assert state["base_source_id"] == source["source_id"]
    output = state["outputs"][-1]
    assert output["status"] == "needs_review"
    generated = Document(BytesIO(lab["client"].get(output["download_url"]).content))
    assert generated.paragraphs[0].text == "Approved manual wording"
    assert generated.sections[0].header.paragraphs[0].text == "DL-2026-05-001"
    assert [cell.text for cell in generated.tables[0].rows[1].cells] == ["1", "", ""]
    assert lab["client"].get(source["download_url"]).content == original
    assert lab["client"].get("/api/projects/P1/report-workspace/current-report").json()["status"] == "missing"


def test_cr_preview_uses_measured_current_and_persists_only_reviewed_results(lab):
    response = upload(lab, role="cr", content=b"SN,Voltage_mV,Current_A\nSN1,38.48,450\n", name="CR.csv")
    assert response.status_code == 200, response.text
    state = response.json()
    step = state["context"]["steps"][0]
    request = {"expected_revision": 1, "source_id": state["sources"][0]["source_id"],
        "kind": "cr", "sheet_name": "", "start_row": 2, "end_row": 2,
        "columns": {"sample": 1, "voltage": 2, "current": 3}, "unit": "mV",
        "group_id": step["group_id"], "row_id": step["row_id"], "step_token": step["step_token"]}
    preview = lab["client"].post(BASE + "/results/preview", json=request)
    assert preview.status_code == 200, preview.text
    assert abs(float(preview.json()["data"]["rows"][0]["resistance_mohm"]) - 0.0855111111111111) < 1e-12
    assert preview.json()["outcome"] == "pending"
    confirmed = lab["client"].post(BASE + "/results/confirm", json={**request,
        "preview_sha256": preview.json()["preview_sha256"]})
    assert confirmed.status_code == 200, confirmed.text
    reopened = lab["client"].get(BASE).json()
    assert reopened["results"][0]["source_sha256"] == state["sources"][0]["sha256"]
    assert reopened["results"][0]["outcome"] == "pending"
    assert lab["client"].post(BASE + "/results/confirm", json={**request,
        "preview_sha256": preview.json()["preview_sha256"]}).status_code == 409


def test_temperature_import_retains_measured_vs_fitted_values_and_channel_statistics(lab):
    source = upload(lab, role="temperature", name="temperature.csv", content=(
        b"Ambient,Current,T1,T2\n20,100,25,27\n20,100,26,28\n20,200,40,44\n20,200,41,45\n")).json()
    step = source["context"]["steps"][0]
    request = {"expected_revision": 1, "source_id": source["sources"][0]["source_id"], "kind": "temperature",
        "group_id": step["group_id"], "row_id": step["row_id"], "step_token": step["step_token"],
        "temperature_selection": {"header_row": 1, "start_row": 2, "end_row": 5, "ambient_column": 1,
            "current_column": 2, "temperature_columns": [3, 4], "thermocouples_per_sample": 1, "sample_count": 2}}
    preview = lab["client"].post(BASE + "/results/preview", json=request)
    assert preview.status_code == 200, preview.text
    data = preview.json()["data"]
    assert data["analysis"]["points"][-1]["maximum"] == 25
    assert data["analysis"]["points"][-1]["average"] == 23
    assert data["analysis"]["points"][-1]["current"] == 200
    assert len(data["raw_rows"]) == 4
    assert data["analysis"]["points"][0]["source_row"] is None  # synthetic fit origin is not a measurement


def test_mechanical_summary_retains_both_colors_and_six_second_hold_semantics(lab):
    source = upload(lab, role="mechanical", name="Mechanical.csv", content=(
        b"SN,Instrument,Color,Force,Seconds\nSN1,rec9,Orange,410,6\nSN2,rec10,Black,420,6\n")).json()
    step = source["context"]["steps"][0]
    request = {"expected_revision": 1, "source_id": source["sources"][0]["source_id"], "kind": "mechanical",
        "group_id": step["group_id"], "row_id": step["row_id"], "step_token": step["step_token"],
        "unit": "N", "start_row": 2, "end_row": 3, "measurement_kind": "hold", "test_type": "retention",
        "columns": {"sample": 1, "record_id": 2, "color": 3, "value": 4, "hold_seconds": 5}}
    preview = lab["client"].post(BASE + "/results/preview", json=request)
    assert preview.status_code == 200, preview.text
    data = preview.json()["data"]
    assert data["test_type"] == "retention" and data["quantity"] == "hold"
    assert [row["color"] for row in data["rows"]] == ["Orange", "Black"]
    assert data["rows"][0]["sample_id"] == "SN1" and data["rows"][0]["record_id"] == "rec9"
    assert data["summary"]["average"] == "415"
    assert data["rows"][0]["hold_seconds"] == "6"


def test_confirmed_fact_replacement_preserves_fixed_sentence_and_other_zip_parts(lab):
    from zipfile import ZipFile
    document = Document(BytesIO(docx_bytes()))
    paragraph = document.add_paragraph("The purpose is qualification to Rev.2 ")
    paragraph.add_run("2025/").bold = True
    paragraph.add_run("12/18").italic = True
    original_stream = BytesIO()
    document.save(original_stream)
    original = original_stream.getvalue()
    state = upload(lab, content=original).json()
    state = lab["client"].post(BASE + "/base", json={"source_id": state["sources"][0]["source_id"], "expected_revision": 1}).json()
    mapped = lab["client"].post(BASE + "/mappings", json={"expected_revision": 2,
        "source_id": state["sources"][0]["source_id"], "kind": "specification", "label": "Revision date",
        "source_locator": "Revision table page 27", "value": "2026/06/18", "report_text": "2025/12/18",
        "status": "confirmed", "reason": "Actual revision date, not authorization date"})
    assert mapped.status_code == 200, mapped.text
    state = mapped.json()
    request = {"expected_revision": 3, "output_id": state["outputs"][-1]["output_id"],
               "mapping_ids": [state["mappings"][0]["mapping_id"]]}
    preview = lab["client"].post(BASE + "/facts/preview", json=request)
    assert preview.status_code == 200, preview.text
    assert preview.json()["changes"][0]["occurrences"] == 1
    applied = lab["client"].post(BASE + "/facts/apply", json={**request, "preview_sha256": preview.json()["preview_sha256"]})
    assert applied.status_code == 200, applied.text
    output = lab["client"].get(applied.json()["outputs"][-1]["download_url"]).content
    after = Document(BytesIO(output))
    assert after.paragraphs[-1].text == "The purpose is qualification to Rev.2 2026/06/18"
    assert after.paragraphs[-1].runs[1].bold
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(output)) as result:
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == result.read(name)
    assert lab["client"].get(state["sources"][0]["download_url"]).content == original
    updated = applied.json()
    repeat_request = {**request, 'expected_revision': updated['revision'], 'output_id': updated['outputs'][-1]['output_id']}
    repeat = lab['client'].post(BASE+'/facts/preview', json=repeat_request)
    assert repeat.status_code == 200, repeat.text
    assert repeat.json()['changes'] == []
    unchanged = lab['client'].post(BASE+'/facts/apply', json={**repeat_request, 'preview_sha256': repeat.json()['preview_sha256']})
    assert unchanged.status_code == 200 and unchanged.json()['revision'] == updated['revision']


def test_confirmed_cr_result_updates_only_new_report_copy_and_rejects_replaced_source(lab):
    source = upload(lab, role='cr', name='CR.csv', content=b'SN,V,I\nSN1,38.48,450\n').json()
    step = source['context']['steps'][0]
    request = {'expected_revision': 1, 'source_id': source['sources'][0]['source_id'], 'kind': 'cr',
        'columns': {'sample': 1, 'voltage': 2, 'current': 3}, **{key: step[key] for key in ('group_id', 'row_id', 'step_token')}}
    preview = lab['client'].post(BASE+'/results/preview', json=request).json()
    state = lab['client'].post(BASE+'/results/confirm', json={**request, 'preview_sha256': preview['preview_sha256']}).json()
    doc = Document(BytesIO(docx_bytes()))
    doc.add_paragraph(f"Group {step['group_label']} Test Results")
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['Step', 'Result', 'Comment']): cell.text = value
    table.cell(1, 0).text = step['step_token']
    table.cell(1, 2).text = 'Pass'
    buffer = BytesIO(); doc.save(buffer)
    state = upload(lab, revision=2, content=buffer.getvalue()).json()
    state = lab['client'].post(BASE+'/base', json={'expected_revision': 3, 'source_id': state['sources'][-1]['source_id']}).json()
    result_id, output_id = state['results'][0]['result_id'], state['outputs'][-1]['output_id']
    targets = lab['client'].get(BASE+f'/outputs/{output_id}/result-targets/{result_id}?expected_revision=4')
    assert targets.status_code == 200, targets.text
    applied = lab['client'].post(BASE+'/results/apply', json={'expected_revision': 4, 'output_id': output_id,
        'output_sha256': targets.json()['output_sha256'], 'result_id': result_id, 'result_table': 0,
        'result_row': 1, 'appendix_title': 'Appendix C: Contact resistance'})
    assert applied.status_code == 200, applied.text
    state = applied.json()
    result = Document(BytesIO(lab['client'].get(state['outputs'][-1]['download_url']).content))
    assert '0.085511' in result.tables[0].cell(1, 1).text
    assert result.tables[0].cell(1, 2).text == ''
    state = upload(lab, role='cr', name='CR.csv', revision=5, content=b'SN,V,I\nSN1,40,450\n').json()
    check = lab['client'].get(BASE+f"/outputs/{state['outputs'][-1]['output_id']}/result-targets/{result_id}?expected_revision=6")
    assert check.status_code == 409


def test_photo_replacement_binds_exact_copy_and_retains_original(lab):
    from PIL import Image
    from zipfile import ZipFile
    def png(color):
        stream = BytesIO(); Image.new('RGB', (40, 20), color).save(stream, 'PNG'); return stream.getvalue()
    old_image, new_image = png('red'), png('blue')
    doc = Document(BytesIO(docx_bytes()))
    doc.add_picture(BytesIO(old_image))
    stream = BytesIO(); doc.save(stream)
    original = stream.getvalue()
    state = upload(lab, content=original).json()
    state = lab['client'].post(BASE+'/base', json={'source_id': state['sources'][0]['source_id'], 'expected_revision': 1}).json()
    state = upload(lab, role='photo', name='sample.png', revision=2, content=new_image).json()
    output = state['outputs'][-1]
    inspected = lab['client'].get(BASE+f"/outputs/{output['output_id']}/photos?expected_revision=3")
    assert inspected.status_code == 200, inspected.text
    slots = inspected.json()
    request = {'expected_revision': 3, 'output_id': output['output_id'], 'output_sha256': slots['output_sha256'],
        'assignments': [{'slot_id': slots['slots'][0]['slot_id'], 'source_id': state['sources'][-1]['source_id']}]}
    bad = lab['client'].post(BASE+'/photos/apply', json={**request, 'output_sha256': '0'*64})
    assert bad.status_code == 409
    applied = lab['client'].post(BASE+'/photos/apply', json=request)
    assert applied.status_code == 200, applied.text
    with ZipFile(BytesIO(lab['client'].get(applied.json()['outputs'][-1]['download_url']).content)) as package:
        assert any(package.read(name) == new_image for name in package.namelist() if name.startswith('word/media/'))
    assert lab['client'].get(state['sources'][0]['download_url']).content == original
    updated = applied.json()
    inspected = lab['client'].get(BASE+f"/outputs/{updated['outputs'][-1]['output_id']}/photos?expected_revision=4").json()
    again = lab['client'].post(BASE+'/photos/apply', json={**request, 'expected_revision': 4,
        'output_id': updated['outputs'][-1]['output_id'], 'output_sha256': inspected['output_sha256']})
    assert again.status_code == 200, again.text
    assert again.json()['revision'] == 4 and len(again.json()['outputs']) == 2
    # Modify only the isolated retained file: metadata remains unchanged, so a metadata-only check is insufficient.
    artifact = updated['outputs'][-1]['artifact']
    asset = lab['settings'].data_dir/'report_authoring_assets'/sha256(b'P1').hexdigest()/(artifact['sha256']+'.docx')
    modified = Document(asset); modified.add_paragraph('UNREVIEWED FOREIGN EDIT'); modified.save(asset)
    blocked = lab['client'].post(BASE+'/photos/apply', json={**request, 'expected_revision': 4,
        'output_id': updated['outputs'][-1]['output_id'], 'output_sha256': inspected['output_sha256']})
    assert blocked.status_code == 422, blocked.text
    assert lab['client'].get(BASE).json()['revision'] == 4


def test_customer_copy_is_separate_and_package_checks_require_explicit_scope(lab, monkeypatch):
    from backend.api import dependencies
    from zipfile import ZipFile
    template = lab['settings'].templates_dir/'E-4515 Customer Report.docx'
    template.write_bytes(docx_bytes())
    class Converter:
        def generate_customer_report(self, *, source_path, template_path, output_path):
            output_path.write_bytes(source_path.read_bytes())
    monkeypatch.setattr(dependencies, 'CustomerReportDocumentGateway', Converter)
    state = upload(lab).json()
    state = lab['client'].post(BASE+'/base', json={'source_id': state['sources'][0]['source_id'], 'expected_revision': 1}).json()
    internal = state['outputs'][-1]
    generated = lab['client'].post(BASE+'/customer/generate', json={'expected_revision': 2,
        'output_id': internal['output_id'], 'output_sha256': internal['artifact']['sha256']})
    assert generated.status_code == 200, generated.text
    state = generated.json()
    assert state['outputs'][-1] == internal
    customer = state['customer_outputs'][-1]
    assert customer['status'] == 'needs_review'
    check = {'expected_revision': 3, 'output_id': customer['output_id'], 'output_sha256': customer['artifact']['sha256']}
    result = lab['client'].post(BASE+'/customer/check', json=check)
    assert result.status_code == 200, result.text
    audit = result.json()['customer_outputs'][-1]['audit']
    assert not audit['passed']
    assert any(item['code'] == 'disclosure_scope_unconfirmed' for item in audit['issues'])
    reviewed = {}
    for issue in audit['issues']:
        if issue['code'] == 'unreviewed_package_part':
            part = lab['client'].get(BASE+f"/customer/{customer['output_id']}/part", params={'part': issue['location']})
            assert part.status_code == 200
            assert sha256(part.content).hexdigest() == issue['sha256']
            reviewed[issue['location']] = issue['sha256']
    checked = lab['client'].post(BASE+'/customer/check', json={**check, 'expected_revision': 4,
        'disclosure_scope_confirmed': True, 'reviewed_parts': reviewed, 'content_review_sha256': internal['artifact']['sha256']})
    assert checked.status_code == 200, checked.text
    assert checked.json()['customer_outputs'][-1]['audit']['passed']
    assert checked.json()['customer_outputs'][-1]['status'] == 'package_checked'
    assert not list(lab['root'].glob('projects/**/*.docx'))
    pending = lab['client'].post(BASE+'/mappings', json={'expected_revision': 5, 'source_id': state['sources'][0]['source_id'],
        'kind': 'specification', 'label': 'New revision decision', 'source_locator': 'Revision table'})
    assert pending.status_code == 200, pending.text
    reopened = lab['client'].get(BASE).json()['customer_outputs'][-1]
    assert reopened['status'] == 'needs_review' and not reopened['audit']['passed']
    assert reopened['recorded_audit']['passed']
    assert any(item['code'] == 'review_context_changed' for item in reopened['audit']['issues'])


def test_temperature_region_requires_explicit_confirmation_and_customer_does_not_duplicate_it(lab, monkeypatch):
    from backend.api import dependencies
    from zipfile import ZipFile
    template = lab['settings'].templates_dir/'E-4515 Customer Report.docx'
    template.write_bytes(docx_bytes())
    class Converter:
        def generate_customer_report(self, *, source_path, template_path, output_path):
            output_path.write_bytes(source_path.read_bytes())
    monkeypatch.setattr(dependencies, 'CustomerReportDocumentGateway', Converter)
    state = upload(lab, role='temperature', name='temperature.csv', content=
        b'Ambient,I,T1\n20,100,26\n20,100,26\n20,200,45\n20,200,45\n').json()
    step = state['context']['steps'][0]
    request = {'expected_revision': 1, 'source_id': state['sources'][0]['source_id'], 'kind': 'temperature',
        **{key: step[key] for key in ('group_id', 'row_id', 'step_token')}, 'sample_ids': ['SN1'],
        'temperature_selection': {'header_row': 1, 'start_row': 2, 'end_row': 5, 'ambient_column': 1,
            'current_column': 2, 'temperature_columns': [3], 'thermocouples_per_sample': 1, 'sample_count': 1}}
    preview = lab['client'].post(BASE+'/results/preview', json=request).json()
    state = lab['client'].post(BASE+'/results/confirm', json={**request, 'preview_sha256': preview['preview_sha256']}).json()
    doc = Document(BytesIO(docx_bytes()))
    doc.add_paragraph(f"Group {step['group_label']} Test Results")
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['Step', 'Result', 'Comment']): cell.text = value
    table.cell(1, 0).text = step['step_token']; table.cell(1, 2).text = 'Pass'
    doc.add_paragraph('Figure 14: Setup photo caption remains')
    doc.add_paragraph('Table-1: Temperature Rise Measurements')
    table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = 'Applied current'; table.cell(0, 1).text = 'OLD VALUES'
    doc.add_paragraph('The maximum temperature rise reached 45°C at 496.9A.')
    doc.add_paragraph('Group NEXT Test Results'); doc.add_paragraph('Protected next Group')
    buffer = BytesIO(); doc.save(buffer)
    state = upload(lab, revision=2, content=buffer.getvalue()).json()
    state = lab['client'].post(BASE+'/base', json={'expected_revision': 3, 'source_id': state['sources'][-1]['source_id']}).json()
    result_id = state['results'][0]['result_id']; output = state['outputs'][-1]
    targets = lab['client'].get(BASE+f"/outputs/{output['output_id']}/result-targets/{result_id}?expected_revision=4").json()
    blocks = targets['evidence_blocks']
    start = next(block['block_index'] for block in blocks if block['text'].startswith('Table-1'))
    end = next(block['block_index'] for block in blocks if '496.9A' in block['text'])
    update = {'expected_revision': 4, 'output_id': output['output_id'], 'output_sha256': targets['output_sha256'],
        'result_id': result_id, 'result_table': 0, 'result_row': 1, 'evidence_start': start, 'evidence_end': end}
    assert lab['client'].post(BASE+'/results/apply', json=update).status_code == 422
    applied = lab['client'].post(BASE+'/results/apply', json={**update, 'evidence_selection_confirmed': True})
    assert applied.status_code == 200, applied.text
    state = applied.json(); output = state['outputs'][-1]
    customer = lab['client'].post(BASE+'/customer/generate', json={'expected_revision': 5, 'output_id': output['output_id'],
        'output_sha256': output['artifact']['sha256'], 'temperature_result_ids': [result_id]})
    assert customer.status_code == 200, customer.text
    with ZipFile(BytesIO(lab['client'].get(customer.json()['customer_outputs'][-1]['download_url']).content)) as package:
        body = package.read('word/document.xml').decode()
        assert '496.9A' not in body and 'OLD VALUES' not in body
        assert 'Figure 14' in body and 'Protected next Group' in body
        assert body.count('connlab-report-result:' + result_id) == 1
        assert 'connlab-customer-temperature:' not in body
