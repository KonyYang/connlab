from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from PIL import Image as PILImage

from backend.infrastructure.office.mechanical_summary_excel_copy import copy_instrument_workbooks, add_reviewed_results


@pytest.mark.office_integration
def test_excel_copy_preserves_source_values_and_drawing():
    source = Workbook()
    sheet = source.active
    sheet['D13'] = ' 20.25'
    png = BytesIO()
    PILImage.new('RGB', (40, 20), 'blue').save(png, 'PNG')
    png.seek(0)
    sheet.add_image(Image(png), 'A20')
    stream = BytesIO()
    source.save(stream)
    original = stream.getvalue()
    result = copy_instrument_workbooks([('instrument.xlsx', original)])
    result_book = load_workbook(BytesIO(result))
    try:
        sheet = result_book.worksheets[0]
        assert sheet['D13'].value == ' 20.25'
        assert len(sheet._images) == 1
        assert stream.getvalue() == original
    finally:
        result_book.close()


@pytest.mark.office_integration
def test_reviewed_results_are_editable_and_statistics_use_selected_values():
    source = Workbook()
    source.active['A1'] = 'Original instrument value'
    png = BytesIO()
    PILImage.new('RGB', (40, 20), 'blue').save(png, 'PNG')
    png.seek(0)
    source.active.add_image(Image(png), 'A20')
    stream = BytesIO()
    source.save(stream)
    source.close()
    copied = copy_instrument_workbooks([('instrument.xlsx', stream.getvalue())])
    result = add_reviewed_results(copied, [{
        'name': 'Hold force', 'kind': 'hold', 'group': 'Ab',
        'rows': [{'curve': '1.20', 'sample': '3-Rec', 'value': 241.07},
                 {'curve': '1.21', 'sample': '4-Rec', 'value': 243.0}],
        'information': {'LTR': 'DL-1'},
    }])
    workbook = load_workbook(BytesIO(result), data_only=False)
    try:
        assert workbook.worksheets[0]['A1'].value == 'Original instrument value'
        assert len(workbook.worksheets[0]._images) == 1
        sheet = workbook['Hold force']
        assert sheet['A4'].value == '1.20'
        assert sheet['C4'].value == 241.07
        assert sheet['C7'].value == '=MIN(C4:C5)'
        assert sheet['B12'].value == 'DL-1'
        assert sheet['B3'].fill.fgColor.rgb[-6:] == 'BFBFBF'
        assert sheet['C7'].fill.fgColor.rgb[-6:] == 'DAEEF3'
        assert sheet.column_dimensions['A'].width <= 30
        assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
        assert sheet.page_setup.fitToWidth in (None, 1)  # XLSX default width is one page.
    finally:
        workbook.close()


@pytest.mark.office_integration
def test_qualitative_results_live_on_original_sheet_without_changing_data_or_picture():
    source = Workbook()
    source.active['D13'] = ' 20.25'
    png = BytesIO()
    PILImage.new('RGB', (40, 20), 'blue').save(png, 'PNG')
    png.seek(0)
    source.active.add_image(Image(png), 'F20')
    stream = BytesIO()
    source.save(stream)
    source.close()
    copied = copy_instrument_workbooks([('instrument.xlsx', stream.getvalue())])
    result = add_reviewed_results(copied, [{
        'target_sheet_index': 1, 'kind': 'qualitative', 'group': 'F',
        'rows': [{'curve': '1.20', 'sample': 'sample 1', 'result': 'Pass'}],
        'information': {'LTR': 'DL-1'},
    }])
    book = load_workbook(BytesIO(result), data_only=False)
    try:
        assert len(book.worksheets) == 1
        sheet = book.worksheets[0]
        assert sheet['D13'].value == ' 20.25'
        assert len(sheet._images) == 1
        matches = [(cell.coordinate, cell.value) for row in sheet for cell in row if cell.value == 'Pass']
        assert len(matches) == 1
        assert sheet[matches[0][0]].column > sheet._images[0].anchor._from.col + 1
    finally:
        book.close()


@pytest.mark.office_integration
def test_reviewed_text_is_not_executed_as_an_excel_formula():
    source = Workbook()
    stream = BytesIO()
    source.save(stream)
    source.close()
    copied = copy_instrument_workbooks([('instrument.xlsx', stream.getvalue())])
    result = add_reviewed_results(copied, [{
        'name': 'Result', 'kind': 'retention',
        'rows': [{'curve': '1', 'sample': '=1+1', 'value': 12}],
    }])
    book = load_workbook(BytesIO(result), data_only=False)
    try:
        assert book['Result']['B4'].value == '=1+1'
        assert book['Result']['B4'].data_type == 's'
    finally:
        book.close()


@pytest.mark.office_integration
def test_unknown_qualitative_result_is_not_assumed_to_pass():
    source = Workbook()
    stream = BytesIO()
    source.save(stream)
    source.close()
    copied = copy_instrument_workbooks([('instrument.xlsx', stream.getvalue())])
    with pytest.raises(ValueError, match='explicitly selected'):
        add_reviewed_results(copied, [{
            'target_sheet_index': 1, 'kind': 'qualitative',
            'rows': [{'curve': '1', 'sample': 'sample 1', 'result': ''}],
        }])
