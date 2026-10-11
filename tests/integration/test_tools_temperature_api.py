from hashlib import sha256
from io import BytesIO, StringIO
import csv

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from backend.api.main import app
from tests.unit.test_temperature_rise_calculations import baseline_rows
from tests.unit.test_temperature_data_preparation import grouped_scanner


def scanner_bytes():
    book = Workbook()
    book.active.title = 'Initial Data'
    book.active.append(['Scan', 'Time', *(f'TC{i}' for i in range(1, 21)), 'Ambient', 'Current'])
    for row in baseline_rows():
        book.active.append([row.source_row, '17:17', *row.temperatures, row.ambient, row.current])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_report_metadata_exports_literal_compact_header_without_cr_or_calculation_changes():
    client = TestClient(app)
    data = client.post('/api/tools/temperature-rise/import', files={'file': ('scanner.xlsx', scanner_bytes())}).json()
    request = {'table': data['table'], 'selection': data['selection'], 'zero_intercept': True,
               'target_rise': 30}
    baseline = client.post('/api/tools/temperature-rise/download', json=request)
    metadata = {'ltr_number': 'DL-2026-09-002A', 'tested_by': '=1+1',
                'test_date': '2026-09-10 ~ 2026-09-11', 'ambient_temperature': '22.0–23.5',
                'relative_humidity': 60.9, 'equipment_id': 'DG-Q-0801/0269',
                'test_condition': '10 / 20 / 30 A', 'remark': 'Manual test record'}
    response = client.post('/api/tools/temperature-rise/download', json={**request, 'report_metadata': metadata})
    assert response.status_code == 200, response.text[:300]
    plain = load_workbook(BytesIO(baseline.content), data_only=True)
    book = load_workbook(BytesIO(response.content), data_only=False)
    cached = load_workbook(BytesIO(response.content), data_only=True)
    try:
        assert book.sheetnames == ['Setup', 'Initial Data', 'T-riseChart', 'Derating']
        assert book.active.title == 'T-riseChart'
        sheet = book['T-riseChart']
        header_range = next(book.defined_names['TemperatureTestInfo'].destinations)[1].replace('$', '')
        assert sheet[header_range][0][0].coordinate == 'A1'
        assert [[cell.value for cell in row if cell.value is not None] for row in sheet[header_range]] == [
            ['LTR', 'DL-2026-09-002A'],
            ['Test Date', '2026-09-10 ~ 2026-09-11'],
            ['Tested By', '=1+1'],
            ['Amb Temp. and Rel. Hum.', '22.0–23.5 °C / 60.9 %RH'],
            ['Test Equipment ID', 'DG-Q-0801/0269'],
            ['Test Condition', '10 / 20 / 30 A'],
            ['Remark', 'Manual test record'],
        ]
        assert all(cell.data_type != 'f' for row in sheet[header_range] for cell in row)
        summary_range = next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')
        plain_range = next(plain.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')
        assert [[cell.value for cell in row] for row in cached['T-riseChart'][summary_range]] == [
            [cell.value for cell in row] for row in plain['T-riseChart'][plain_range]]
        assert list(cached['Initial Data'].values) == list(plain['Initial Data'].values)
        assert sheet._charts[0].anchor._from.row >= sheet[header_range][-1][0].row
        assert '$A$1:' in str(sheet.print_area)
        raw_range = next(book.defined_names['KeyStageReadings'].destinations)[1].replace('$', '')
        assert sheet[raw_range][0][0].row > sheet[summary_range][-1][0].row
        baseline_raw = next(plain.defined_names['KeyStageReadings'].destinations)[1].replace('$', '')
        assert [[cell.value for cell in row] for row in cached['T-riseChart'][raw_range]] == [
            [cell.value for cell in row] for row in plain['T-riseChart'][baseline_raw]]
    finally:
        plain.close()
        book.close()
        cached.close()


@pytest.mark.parametrize(('metadata', 'environment'), [({}, None), ({'relative_humidity': 0}, '0 %RH')])
def test_empty_report_fields_remain_blank_at_top_left_and_zero_humidity_is_preserved(metadata, environment):
    client = TestClient(app)
    data = client.post('/api/tools/temperature-rise/import', files={'file': ('scanner.xlsx', scanner_bytes())}).json()
    response = client.post('/api/tools/temperature-rise/download', json={
        'table': data['table'], 'selection': data['selection'], 'report_metadata': metadata,
    })
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content), data_only=True)
    try:
        sheet = book['T-riseChart']
        info = sheet[next(book.defined_names['TemperatureTestInfo'].destinations)[1].replace('$', '')]
        assert info[0][0].coordinate == 'A1'
        assert len(info) == 5
        assert [row[0].value for row in info] == [
            'LTR', 'Test Date', 'Tested By', 'Amb Temp. and Rel. Hum.', 'Test Equipment ID',
        ]
        assert [row[2].value for row in info] == [None, None, None, environment, None]
    finally:
        book.close()


@pytest.mark.parametrize('humidity', [-1, 101, 'nan'])
def test_download_rejects_invalid_report_humidity(humidity):
    client = TestClient(app)
    data = client.post('/api/tools/temperature-rise/import', files={'file': ('scanner.xlsx', scanner_bytes())}).json()
    response = client.post('/api/tools/temperature-rise/download', json={
        'table': data['table'], 'selection': data['selection'],
        'report_metadata': {'relative_humidity': humidity},
    })
    assert response.status_code == 422


@pytest.mark.parametrize(('names', 'labels'), [
    ({'3': 'Tail-R', '4': '=1+1'}, ['=1+1', 'Tail-R']),
    ({'3': 'Tail-R'}, ['Crimp', 'Tail-R']),
    ({'3': '', '4': ' Crimp '}, ['Crimp', 'T2']),
    ({'3': '1021_R1A', '4': ' 1_Header Solder '}, ['Header Solder', 'R1A']),
    ({'3': '1021_2_R1A', '4': '1_2_Header Solder'}, ['Header Solder', 'R1A']),
    ({'3': '1_HS', '4': 'Contact'}, ['Contact', 'HS']),
    ({'3': '1_hS', '4': 'R2_Contact'}, ['R2_Contact', 'hS']),
    ({'3': '未命名测点', '4': 'LPH'}, ['LPH', '未命名测点']),
    ({'3': '123', '4': '-Tail R'}, ['-Tail R', 'T2']),
])
def test_edited_channel_names_follow_source_ids_into_reordered_summary_without_changing_readings(names, labels):
    rows = [['Scan', 'Time', '101 <1_HS> (C)', '102 <1_Crimp> (C)', 'Ambient (C)', 'Current (VDC)']]
    for current, temperature in ((10, 25), (20, 35), (30, 45)):
        rows.extend([[len(rows), '18:00', temperature, temperature + 1, 20, current],
                     [len(rows) + 1, '18:01', temperature, temperature + 1, 20, current]])
    request = {'table': {'file_name': 'names.csv', 'sheet_names': ['Data'], 'sheet_name': 'Data', 'rows': rows},
               'selection': {'header_row': 1, 'start_row': 2, 'end_row': 7, 'ambient_column': 5,
                             'current_column': 6, 'temperature_columns': [4, 3], 'thermocouples_per_sample': 2,
                             'sample_groups': [[4, 3]], 'channel_names': names}}
    client = TestClient(app)
    analysis = client.post('/api/tools/temperature-rise/analyze', json=request)
    assert analysis.status_code == 200, analysis.text
    assert analysis.json()['points'][1]['rises'] == [6, 5]
    download = client.post('/api/tools/temperature-rise/download', json=request)
    assert download.status_code == 200, download.text[:300]
    book = load_workbook(BytesIO(download.content), data_only=False)
    try:
        sheet = book['T-riseChart']
        summary = sheet[next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
        assert [row[0].value for row in summary[1:3]] == labels
        assert all(row[0].data_type == 's' for row in summary[1:3])
        raw = sheet[next(book.defined_names['KeyStageReadings'].destinations)[1].replace('$', '')]
        assert raw[0][2].value == '101 <1_HS> (C)'
        assert request['table']['rows'][0][2] == '101 <1_HS> (C)'
    finally:
        book.close()


@pytest.mark.parametrize('names', [
    {'1': 'Index'}, {'25': 'Outside table'}, {'999': 'Missing'}, {'3': 'a' * 201}, {'3': 'Line\nBreak'}, {'3': 'Null\x00Byte'},
])
def test_measurement_name_overrides_reject_nonmeasurement_sources_and_invalid_text(names):
    client = TestClient(app)
    data = client.post('/api/tools/temperature-rise/import', files={'file': ('scanner.xlsx', scanner_bytes())}).json()
    request = {'table': data['table'], 'selection': {**data['selection'], 'channel_names': names}}
    assert client.post('/api/tools/temperature-rise/prepare', json=request).status_code == 422


def test_balanced_thermocouple_exclusion_preserves_sample_maxima_in_analysis_and_download():
    headers = ['Scan', 'Time', *(f'{101 + i} <{i // 4 + 1}_T{i % 4 + 1}> (C)' for i in range(12)),
               'Ambient (C)', 'Current (VDC)']
    rows = [headers]
    for current in (10, 20, 30):
        for _ in range(2):
            rows.append([len(rows), '18:00', 21, 99, 23, 24, 31, 99, 33, 34, 41, 99, 43, 44, 20, current])
    groups = [[3, 5, 6], [7, 9, 10], [11, 13, 14]]
    request = {'table': {'file_name': 'balanced.csv', 'sheet_names': ['Data'], 'sheet_name': 'Data', 'rows': rows},
               'selection': {'header_row': 1, 'start_row': 2, 'end_row': 7, 'ambient_column': 15,
                             'current_column': 16, 'temperature_columns': [3, 5, 6, 7, 9, 10, 11, 13, 14],
                             'thermocouples_per_sample': 3, 'sample_groups': groups}}
    client = TestClient(app)
    response = client.post('/api/tools/temperature-rise/analyze', json=request)
    assert response.status_code == 200, response.text
    point = response.json()['points'][1]
    assert point['sample_maxima'] == [4, 14, 24]
    assert point['maximum'] == 24
    assert point['average'] == 14  # Average of sample maxima, not all thermocouples.
    downloaded = client.post('/api/tools/temperature-rise/download', json=request)
    assert downloaded.status_code == 200, downloaded.text[:300]
    book = load_workbook(BytesIO(downloaded.content), data_only=True)
    assert book['Initial Data'].cell(2, 3).value == 21
    assert book['Initial Data'].cell(2, 4).value == 99  # Excluded channel still preserved in raw source.
    assert book['Initial Data'].cell(2, 5).value == 23
    book.close()
    # A divisible total is insufficient: the first original sample cannot disappear.
    uneven = {**request, 'selection': {**request['selection'], 'temperature_columns': list(range(7, 15)),
                                      'thermocouples_per_sample': 4,
                                      'sample_groups': [[], [7, 8, 9, 10], [11, 12, 13, 14]]}}
    rejected = client.post('/api/tools/temperature-rise/prepare', json=uneven)
    assert rejected.status_code == 422
    assert 'sample' in rejected.text.lower()


def test_grouped_csv_import_reports_auxiliary_currents_and_nonzero_intercept_default():
    source_table = grouped_scanner()
    stream = StringIO(newline='')
    csv.writer(stream, delimiter='\t').writerows(source_table.rows)
    source = stream.getvalue().encode('utf-16')
    client = TestClient(app)
    response = client.post('/api/tools/temperature-rise/import', files={'file': ('scanner.csv', source)})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['selection']['current_column'] == 35
    layout = result['channel_layout']
    assert layout['current_columns'] == [34, 35, 36, 37]
    assert layout['stable_current_columns'] == [34, 36, 37]
    assert [group['sample_id'] for group in layout['sample_groups']] == ['1', '2', '3']
    assert layout['zero_intercept_default'] is False
    request = {'table': result['table'], 'selection': result['selection']}
    review = client.post('/api/tools/temperature-rise/prepare', json=request).json()
    assert not review['ready']
    assert [row['current'] for row in review['measurements'][:3]] == [0, 0, 10]
    assert result['table']['rows'][1][34] == '0.009999'
    download = client.post('/api/tools/temperature-rise/download', json={
        **request, 'zero_intercept': False, 'acknowledge_warnings': True,
    })
    assert download.status_code == 200, download.text[:300]
    book = load_workbook(BytesIO(download.content), data_only=True)
    initial = book['Initial Data']
    assert initial.max_column == 37
    assert initial.cell(1, 34).value == source_table.rows[0][33]
    assert [float(initial.cell(2, column).value) for column in (34, 36, 37)] == [125.001, 3.001, 1.02]
    assert float(initial.cell(2, 35).value) == .009999  # Normalization belongs only in analysis.
    book.close()


def test_import_prepare_analyze_and_download_round_trip_without_project_authority():
    client = TestClient(app)
    source = scanner_bytes()
    digest = sha256(source).hexdigest()
    imported = client.post('/api/tools/temperature-rise/import', files={'file': ('scanner.xlsx', source)})
    assert imported.status_code == 200, imported.text
    data = imported.json()
    assert data['region_issue'] is None
    assert data['selection']['header_row'] == 1
    assert data['selection']['channel_names'] == {}
    assert data['selection']['start_row'] == 2
    assert data['selection']['end_row'] == 301
    request = {'table': data['table'], 'selection': data['selection'], 'acknowledge_warnings': False, 'zero_intercept': True}
    confirmed = client.post('/api/tools/temperature-rise/prepare', json=request)
    assert confirmed.status_code == 200
    assert confirmed.json()['ready']
    analyzed = client.post('/api/tools/temperature-rise/analyze', json=request)
    assert analyzed.status_code == 200, analyzed.text
    analysis = analyzed.json()
    assert analysis['maximum_fit']['coefficients']['a'] == .004677
    maximum, average = analysis['maximum_fit']['coefficients'], analysis['average_fit']['coefficients']
    calculated = client.post('/api/tools/temperature-rise/current', json={'coefficients': maximum, 'target_rise': 30})
    assert calculated.json()['current'] == pytest.approx(66.54445742154181)
    parameters = {'max_temperature': 105, 'step': 2.5, 'ambient_point': 75}
    derating = client.post('/api/tools/temperature-rise/derating', json={'coefficients': average, **parameters})
    assert derating.json()['annotation']['basic'] == pytest.approx(68.38881593205294)
    download = client.post('/api/tools/temperature-rise/download', json={**request,
        'maximum_coefficients': maximum, 'average_coefficients': average, 'target_rise': 30, 'derating': parameters})
    assert download.status_code == 200, download.text[:300]
    assert 'scanner_T-rise_Derating.xlsx' in download.headers['content-disposition']
    book = load_workbook(BytesIO(download.content), data_only=True)
    assert book['Derating']['C36'].value == pytest.approx(54.71105274564235)
    book.close()
    assert sha256(source).hexdigest() == digest


def test_unchecked_intercept_round_trip_retains_confirmed_background_baseline():
    stream = StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow(['Scan', 'Time', 'TC1', 'TC2', 'Ambient', 'Current'])
    writer.writerow([1, '10:00', 29, 28, 20, .009999])
    writer.writerow([2, '10:01', 24, 23, 22, .011427])
    for current in (10, 20, 30):
        for _ in range(2):
            writer.writerow([current, '10:02', 25 + .01 * current ** 2 + .2 * current + 2,
                             25 + .005 * current ** 2 + .1 * current + 1, 25, current])
    source = stream.getvalue().encode('utf-8')
    client = TestClient(app)
    imported = client.post('/api/tools/temperature-rise/import', files={'file': ('background.csv', source)})
    assert imported.status_code == 200, imported.text
    result = imported.json()
    request = {'table': result['table'], 'selection': {
        **result['selection'], 'temperature_columns': [3, 4], 'thermocouples_per_sample': 1,
        'ambient_column': 5, 'current_column': 6,
    }, 'zero_intercept': False, 'acknowledge_warnings': True}
    analyzed = client.post('/api/tools/temperature-rise/analyze', json=request)
    assert analyzed.status_code == 200, analyzed.text
    analysis = analyzed.json()
    assert [point['source_row'] for point in analysis['points']] == [3, 5, 7, 9]
    assert analysis['points'][0]['maximum'] == 2
    assert analysis['points'][0]['average'] == 1.5
    assert analysis['maximum_fit']['coefficients'] == {'a': .01, 'b': .2, 'c': 2}
    assert analysis['average_fit']['coefficients'] == {'a': .0075, 'b': .15, 'c': 1.5}
    downloaded = client.post('/api/tools/temperature-rise/download', json=request)
    assert downloaded.status_code == 200, downloaded.text[:300]
    book = load_workbook(BytesIO(downloaded.content), data_only=True)
    area = next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')
    assert [cell.value for cell in book['T-riseChart'][area][-2][1:]] == [2, 5, 10, 17]
    book.close()


def test_utf16_scanner_csv_uses_the_same_confirmation_calculation_and_excel_download_flow():
    stream = StringIO(newline='')
    writer = csv.writer(stream, delimiter='\t')
    writer.writerow(['名称:', '扫描仪数据'])
    writer.writerow(['扫描', '时间', *(f'TC{i}' for i in range(1, 21)), 'Ambient', 'Current'])
    for row in baseline_rows():
        writer.writerow([row.source_row, '2024/10/28 17:10:16:391', *row.temperatures, row.ambient, row.current])
    source = stream.getvalue().encode('utf-16')
    digest = sha256(source).hexdigest()
    client = TestClient(app)
    imported = client.post('/api/tools/temperature-rise/import', files={'file': ('../scanner.csv', source)})
    assert imported.status_code == 200, imported.text
    data = imported.json()
    assert data['table']['file_name'] == 'scanner.csv'
    assert data['table']['rows'][0] == ['名称:', '扫描仪数据']
    assert data['region_issue'] is None
    assert data['selection']['header_row'] == 2
    request = {'table': data['table'], 'selection': data['selection'], 'zero_intercept': True}
    assert client.post('/api/tools/temperature-rise/prepare', json=request).json()['ready']
    analyzed = client.post('/api/tools/temperature-rise/analyze', json=request)
    assert analyzed.status_code == 200, analyzed.text
    assert analyzed.json()['maximum_fit']['coefficients']['a'] == .004677
    download = client.post('/api/tools/temperature-rise/download', json=request)
    assert download.status_code == 200, download.text[:300]
    assert 'scanner_T-rise_Derating.xlsx' in download.headers['content-disposition']
    book = load_workbook(BytesIO(download.content), data_only=False)
    assert len(book['T-riseChart']._charts) == 1
    assert book['Initial Data']['B3'].value == '2024/10/28 17:10:16:391'
    assert float(book['Initial Data']['C3'].value) == baseline_rows()[0].temperatures[0]
    assert book['Initial Data']['A1'].value == '名称:'
    book.close()
    assert sha256(source).hexdigest() == digest


def test_analyze_requires_warning_confirmation_and_invalid_cells_remain_blockers():
    client = TestClient(app)
    data = client.post('/api/tools/temperature-rise/import', files={'file': ('scanner.xlsx', scanner_bytes())}).json()
    request = {'table': data['table'], 'selection': data['selection'], 'zero_intercept': True}
    request['table']['rows'][-1][-1] = 0
    review = client.post('/api/tools/temperature-rise/prepare', json=request).json()
    assert not review['ready']
    assert any(issue['code'] == 'zero_current' for issue in review['issues'])
    assert client.post('/api/tools/temperature-rise/analyze', json=request).status_code == 422
    request['acknowledge_warnings'] = True
    assert client.post('/api/tools/temperature-rise/analyze', json=request).status_code == 200
    request['table']['rows'][-1][2] = None
    assert client.post('/api/tools/temperature-rise/analyze', json=request).status_code == 422


def test_import_rejects_unreadable_files_and_download_name_uses_no_client_path():
    client = TestClient(app)
    response = client.post('/api/tools/temperature-rise/import', files={'file': ('bad.xlsx', b'not a workbook')})
    assert response.status_code == 422
    data = client.post('/api/tools/temperature-rise/import', files={'file': ('../scanner.xlsx', scanner_bytes())}).json()
    assert data['table']['file_name'] == 'scanner.xlsx'


def test_unidentified_region_retains_source_for_manual_recovery_and_still_requires_data_validation():
    book = Workbook()
    book.active.title = 'Data'
    book.active.append(['Probe', 'Room', 'Supply'])
    book.active.append([25, 20, 10])
    book.active.append([26, 20, 0])
    stream = BytesIO()
    book.save(stream)
    book.close()
    content = stream.getvalue()
    client = TestClient(app)
    imported = client.post('/api/tools/temperature-rise/import', files={'file': ('custom.xlsx', content)})
    assert imported.status_code == 200
    data = imported.json()
    assert data['selection'] is None
    assert 'data region' in data['region_issue']
    assert data['table']['rows'] == [['Probe', 'Room', 'Supply'], [25, 20, 10], [26, 20, 0]]
    assert client.post('/api/tools/temperature-rise/analyze', json={
        'table': data['table'], 'selection': None,
    }).status_code == 422

    recovered = client.post('/api/tools/temperature-rise/import', files={'file': ('custom.xlsx', content)},
                            data={'header_row': 1, 'start_row': 2, 'end_row': 3})
    assert recovered.status_code == 200
    data = recovered.json()
    assert data['region_issue'] is None
    assert data['selection']['start_row'] == 2
    assert data['selection']['excluded_rows'] == []
    review = client.post('/api/tools/temperature-rise/prepare', json={
        'table': data['table'], 'selection': data['selection'],
    }).json()
    assert not review['ready']
    assert any(issue['code'] == 'zero_current' for issue in review['issues'])
    for override in ({'header_row': 1}, {'header_row': 2, 'start_row': 2, 'end_row': 3},
                     {'header_row': 1, 'start_row': 2, 'end_row': 99}):
        assert client.post('/api/tools/temperature-rise/import', files={'file': ('custom.xlsx', content)},
                           data=override).status_code == 422


def test_multiple_regions_require_operator_choice_and_a_different_sheet_can_recover_automatically():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Ambiguous'
    header = ['Scan', 'Time', 'TC1', 'Ambient', 'Current']
    sheet.append(header)
    sheet.append([1, '17:17', 25, 20, 10])
    sheet.append(header)
    sheet.append([2, '17:18', 30, 20, 20])
    other = book.create_sheet('Complete')
    other.append(header)
    other.append([1, '17:17', 25, 20, 10])
    other.append([2, '17:18', 25, 20, 0])
    stream = BytesIO()
    book.save(stream)
    book.close()
    content = stream.getvalue()
    client = TestClient(app)
    def import_sheet(**data):
        response = client.post('/api/tools/temperature-rise/import', files={'file': ('scanner.xlsx', content)}, data=data)
        assert response.status_code == 200, response.text
        return response.json()

    ambiguous = import_sheet(sheet_name='Ambiguous')
    assert ambiguous['selection'] is None
    assert 'Multiple' in ambiguous['region_issue']
    assert len(ambiguous['table']['rows']) == 4
    manual = import_sheet(sheet_name='Ambiguous', header_row=3, start_row=4, end_row=4)
    assert manual['selection']['header_row'] == 3
    assert manual['region_issue'] is None
    automatic = import_sheet(sheet_name='Complete')
    assert automatic['region_issue'] is None
    assert automatic['selection']['end_row'] == 3
    assert automatic['selection']['excluded_rows'] == []
    assert automatic['table']['rows'][-1][-1] == 0
