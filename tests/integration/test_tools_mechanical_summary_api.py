from io import BytesIO
import hashlib
import json

from fastapi.testclient import TestClient
from fastapi import FastAPI
from reportlab.pdfgen import canvas
from openpyxl import Workbook, load_workbook
import pytest

from backend.api.routes_tools_mechanical_summary import router

app = FastAPI()
app.include_router(router)


def test_standalone_pdf_preview_keeps_curve_id_and_requires_review():
    stream = BytesIO()
    page = canvas.Canvas(stream)
    page.drawString(40, 760, '1.20 3#-Rec 241.07 241.07 241 5.2')
    page.save()
    result = TestClient(app).post(
        '/api/tools/mechanical-summary/preview',
        files={'file': ('instrument.pdf', stream.getvalue(), 'application/pdf')},
        data={'kind': 'hold'},
    )
    assert result.status_code == 200, result.text
    assert result.json()['rows'][0]['curve'] == '1.20'
    assert result.json()['requires_review'] is True


def test_invalid_workbook_is_reported_as_an_input_error():
    result = TestClient(app).post(
        '/api/tools/mechanical-summary/preview',
        files={'file': ('broken.xls', b'not a workbook', 'application/vnd.ms-excel')},
        data={'kind': 'qualitative'},
    )
    assert result.status_code == 422
    assert 'workbook' in result.json()['detail'].lower()


def test_invalid_pdf_is_reported_as_an_input_error():
    result = TestClient(app).post(
        '/api/tools/mechanical-summary/preview',
        files={'file': ('broken.pdf', b'not a PDF', 'application/pdf')},
        data={'kind': 'hold'},
    )
    assert result.status_code == 422


def test_generation_rejects_source_changed_since_review():
    selection = [{'kind': 'hold', 'sha256': 'wrong', 'confirmed': True, 'rows': []}]
    response = TestClient(app).post('/api/tools/mechanical-summary/generate',
        files=[('files', ('source.pdf', b'new contents', 'application/pdf'))],
        data={'manifest': json.dumps(selection)})
    assert response.status_code == 422
    assert 'changed' in response.json()['detail']


def test_generation_rejects_unconfirmed_source():
    selection = [{'kind': 'hold', 'sha256': 'unused', 'confirmed': False, 'rows': []}]
    response = TestClient(app).post('/api/tools/mechanical-summary/generate',
        files=[('files', ('source.pdf', b'contents', 'application/pdf'))],
        data={'manifest': json.dumps(selection)})
    assert response.status_code == 422
    assert 'Confirm' in response.json()['detail']


@pytest.mark.office_integration
def test_pdf_only_summary_has_no_empty_placeholder_sheet():
    stream = BytesIO()
    pdf = canvas.Canvas(stream)
    pdf.drawString(40, 760, '1.20 3#-Rec 241.07 241.07 241 5.2')
    pdf.save()
    content = stream.getvalue()
    manifest = [{'kind': 'hold', 'sha256': hashlib.sha256(content).hexdigest(),
                 'confirmed': True, 'rows': [{'curve': '1.20', 'sample': '3#-Rec', 'value': 241.07}]}]
    response = TestClient(app).post('/api/tools/mechanical-summary/generate',
        files=[('files', ('hold.pdf', content, 'application/pdf'))], data={'manifest': json.dumps(manifest)})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    try:
        assert len(book.worksheets) == 1
        assert book.worksheets[0]['C4'].value == 241.07
    finally:
        book.close()


@pytest.mark.office_integration
def test_generate_from_reviewed_inputs_without_project():
    book = Workbook()
    book.active['D13'] = ' 20.25'
    stream = BytesIO()
    book.save(stream)
    book.close()
    xls = stream.getvalue()
    pdf_stream = BytesIO()
    pdf = canvas.Canvas(pdf_stream)
    pdf.drawString(40, 760, '1.20 3#-Rec 241.07 241.07 241 5.2')
    pdf.save()
    pdf_bytes = pdf_stream.getvalue()
    manifest = [
        {'kind': 'qualitative', 'sha256': hashlib.sha256(xls).hexdigest(), 'sheet': 1,
         'group': 'F', 'confirmed': True,
         'rows': [{'curve': '1', 'sample': '3', 'result': 'Pass'}], 'information': {}},
        {'kind': 'hold', 'sha256': hashlib.sha256(pdf_bytes).hexdigest(),
         'group': 'Ab', 'confirmed': True,
         'rows': [{'curve': '1.20', 'sample': '3#-Rec', 'value': 241.07}], 'information': {}},
    ]
    response = TestClient(app).post('/api/tools/mechanical-summary/generate',
        files=[('files', ('instrument.xlsx', xls, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')),
               ('files', ('hold.pdf', pdf_bytes, 'application/pdf'))],
        data={'manifest': json.dumps(manifest)})
    assert response.status_code == 200, response.text
    result = load_workbook(BytesIO(response.content), data_only=False)
    try:
        assert result.worksheets[0]['D13'].value == ' 20.25'
        assert any(cell.value == 'Pass' for row in result.worksheets[0] for cell in row)
        assert result.worksheets[-1]['C4'].value == 241.07
        assert len(result.worksheets[-1]._images) == 1
        assert result.worksheets[-1].page_setup.orientation == 'portrait'
        assert result.worksheets[-1].row_breaks.count >= 1
    finally:
        result.close()
