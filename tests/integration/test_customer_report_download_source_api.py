"""Real request/worker/filesystem/DB composition; replace only the Office writer."""
from hashlib import sha256

from docx import Document
import pytest

from backend.api import dependencies as deps
from backend.api.main import app
from backend.api.project_customer_report_runner import ProjectCustomerReportRunner
from backend.application.project_customer_report_job_service import ProjectCustomerReportJobService
from backend.infrastructure.storage.models import ProjectModel
from tests.integration.test_internal_report_generation_api import lab, official_report
from tests.unit.test_downloaded_internal_report_source import draft


BASE = "/api/projects/P1/report-workspace/current-customer-report"
DL = "DL-2026-05-001"


@pytest.fixture
def customer_lab(lab):
    downloads = lab["root"] / "downloads"
    downloads.mkdir()
    template = lab["settings"].templates_dir / "E-4515_F Customer Report.docx"
    template.write_bytes(b"customer template boundary")
    pending = []
    sources = []
    after_write = []
    class Writer:
        def generate_customer_report(self, *, source_path, template_path, output_path, progress):
            sources.append(source_path)
            progress("saving_document")
            document = Document(source_path)
            document.add_paragraph("Customer projection boundary")
            document.save(output_path)
            if after_write:
                after_write[0]()
            return output_path
    runner = ProjectCustomerReportRunner(session_factory=lab["factory"], settings=lab["settings"],
                                         writer_factory=lambda root: Writer())
    jobs = ProjectCustomerReportJobService(root=lab["root"] / "jobs", generate=runner, dispatch=pending.append)
    app.dependency_overrides[deps.get_project_customer_report_job_service] = lambda: jobs
    yield {**lab, "downloads": downloads, "pending": pending, "sources": sources,
           "after_write": after_write, "jobs": jobs}
    jobs.close()


def state(lab):
    response = lab["client"].get(BASE)
    assert response.status_code == 200, response.text
    return response.json()


def start_and_run(lab, expected):
    response = lab["client"].post(BASE + "/jobs", json={"expected_internal_report_sha256": expected})
    assert response.status_code == 202, response.text
    operation = response.json()["operation_id"]
    lab["pending"].pop(0)()
    return operation, lab["client"].get(BASE + "/jobs/" + operation).json()


@pytest.mark.parametrize("registered", [False, True])
def test_latest_download_generates_customer_draft_without_registering_or_modifying_sources(customer_lab, registered):
    lab = customer_lab
    retained = None
    if registered:
        retained, local = official_report(lab)
        moved = retained.parent.with_name("operator moved folder")
        retained.parent.rename(moved)
        retained = moved / retained.name
        retained_bytes = retained.read_bytes()
    old = draft(lab["downloads"], dl=DL)
    latest = draft(lab["downloads"], " (1)", dl=DL, modified=200)
    draft(lab["downloads"], dl=DL + "-CR", modified=300, form="E-4515")
    before = {path: path.read_bytes() for path in lab["downloads"].iterdir()}
    preview = state(lab)
    assert preview["can_generate"] and preview["mode"] == "managed_download"
    assert preview["internal_report_sha256"] == sha256(before[latest]).hexdigest()
    assert lab["client"].get(BASE.replace("current-customer-report", "current-report")).json()["status"] == "missing"
    for _ in range(2):
        operation, result = start_and_run(lab, preview["internal_report_sha256"])
        assert result["status"] == "completed", result
        assert result["result"]["mode"] == "managed_download"
        assert "-CR " in result["result"]["file_name"]
        assert "draft" in result["result"]["file_name"].casefold()
        assert result["result"]["archive_path"] is None
        for _ in range(2):
            response = lab["client"].get(BASE + "/jobs/" + operation + "/download")
            assert response.status_code == 200, response.text
            assert response.content.startswith(b"PK")
            assert "draft" in response.headers["content-disposition"].casefold()
    assert lab["sources"] == [latest, latest]
    assert {path: path.read_bytes() for path in lab["downloads"].iterdir()} == before
    assert old.exists()
    assert lab["client"].get("/api/projects/P1/report-workspace").json()["report_revisions"] == []
    with lab["factory"]() as session:
        assert session.get(ProjectModel, "P1").registry_revision == 0
    if retained:
        assert retained.read_bytes() == retained_bytes
        assert not (local / "History").exists()


@pytest.mark.parametrize("when", ["after_preview", "during_generation"])
def test_newer_draft_cannot_silently_replace_reviewed_source(customer_lab, when):
    lab = customer_lab
    original = draft(lab["downloads"], dl=DL)
    expected = state(lab)["internal_report_sha256"]
    def new_source():
        path = draft(lab["downloads"], " (1)", dl=DL, modified=200)
        document = Document(path)
        document.add_paragraph("New manual revision")
        document.save(path)
    if when == "after_preview":
        new_source()
    else:
        lab["after_write"].append(new_source)
    operation, result = start_and_run(lab, expected)
    assert result["status"] == "failed", result
    assert "changed" in result["message"]
    assert lab["client"].get(BASE + "/jobs/" + operation + "/download").status_code != 200
    assert original.exists()
    assert not list((lab["root"] / "jobs").rglob("*.docx"))


def test_invalid_latest_returns_manual_selection_state_instead_of_server_error(customer_lab):
    lab = customer_lab
    draft(lab["downloads"], dl=DL)
    corrupt = draft(lab["downloads"], " (1)", dl=DL, modified=200)
    corrupt.write_bytes(b"incomplete")
    preview = state(lab)
    assert not preview["can_generate"]
    assert preview["status"] == "blocked"
    assert not lab["sources"]
