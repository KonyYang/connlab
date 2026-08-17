from __future__ import annotations

from collections.abc import Generator
from datetime import date
import os
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from backend.api.dependencies import (
    get_project_test_plan_matrix_preview_service,
    get_session,
    get_settings,
)
from backend.api.main import app
from backend.domain import FileAsset, FileAssetType, Project, ProjectStatus
from backend.infrastructure.storage.database import (
    create_database_engine,
    create_session_factory,
    init_db,
)
from backend.application.official_project_workspace_service import OfficialWorkspaceRecord
from backend.infrastructure.storage.repositories import (
    FileAssetRepository,
    ProjectOfficialWorkspaceRepository,
    ProjectRepository,
)
from backend.application.project_test_plan_matrix_preview_service import (
    MatrixPreviewFromPathCommand,
    ProjectTestPlanMatrixPreview,
)
from backend.shared.config import Settings


def test_source_candidates_api_lists_project_docx_and_preview_by_candidate(
    tmp_path: Path,
) -> None:
    client, engine = _client(tmp_path)
    try:
        _create_project("P1", tmp_path)
        spec = tmp_path / "product_matrix_spec.docx"
        _write_product_spec_docx(spec)
        note = tmp_path / "notes.docx"
        note.write_bytes(b"x")
        _create_file_asset("P1", "A1", spec, "Product Matrix Spec.docx", tmp_path)
        _create_file_asset("P1", "A2", note, "notes.docx", tmp_path)

        listed = client.get("/api/projects/P1/test-plan/source-candidates")
        assert listed.status_code == 200
        payload = listed.json()
        assert payload["project_id"] == "P1"
        assert payload["candidates"][0]["source_asset_id"] == "A1"
        assert payload["candidates"][0]["candidate_kind"] == "likely_spec_or_matrix"

        preview = client.post("/api/projects/P1/test-plan/source-candidates/A1/matrix-preview")
        assert preview.status_code == 200
        preview_body = preview.json()
        assert preview_body["project_id"] == "P1"
        assert preview_body["source_document_name"] == "product_matrix_spec.docx"
        assert preview_body["source_format"] == ".docx"
        assert preview_body["groups"][0]["group_label"] == "Group 1"
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_source_candidate_preview_rejects_cross_project_asset(tmp_path: Path) -> None:
    client, engine = _client(tmp_path)
    try:
        _create_project("P1", tmp_path)
        _create_project("P2", tmp_path)
        spec = tmp_path / "spec.docx"
        _write_product_spec_docx(spec)
        _create_file_asset("P1", "A1", spec, "spec.docx", tmp_path)

        response = client.post("/api/projects/P2/test-plan/source-candidates/A1/matrix-preview")
        assert response.status_code == 404
        assert "not found" in response.json()["detail"]
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_source_candidates_api_projects_submitted_material_directory(tmp_path: Path) -> None:
    client, engine = _client(tmp_path)
    try:
        _create_project("P1", tmp_path)
        official = tmp_path / "official"
        submitted = official / "Submitted Material"
        submitted.mkdir(parents=True)
        _create_workspace("P1", official, tmp_path)

        response = client.get("/api/projects/P1/test-plan/source-candidates")

        assert response.status_code == 200
        assert response.json()["preferred_import_directory"] == str(submitted)
        assert response.json()["preferred_import_directory_source"] == "submitted_material"
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_resolved_directory_view_lists_path_free_direct_files_and_previews_current_id(
    tmp_path: Path,
) -> None:
    client, engine = _client(tmp_path)
    try:
        _create_project("P1", tmp_path)
        official = tmp_path / "official"
        submitted = official / "Submitted Material"
        submitted.mkdir(parents=True)
        spec = submitted / "Alpha Matrix.docx"
        _write_product_spec_docx(spec)
        original_bytes = spec.read_bytes()
        (submitted / "beta.PDF").write_bytes(b"%PDF-1.4")
        (submitted / "legacy.doc").write_bytes(b"legacy")
        (submitted / "ignore.txt").write_bytes(b"ignore")
        nested = submitted / "nested"
        nested.mkdir()
        (nested / "hidden.docx").write_bytes(b"hidden")
        _create_workspace("P1", official, tmp_path)

        listed = client.get(
            "/api/projects/P1/test-plan/source-candidates",
            params={"view": "resolved_directory"},
        )

        assert listed.status_code == 200
        payload = listed.json()
        assert payload["view"] == "resolved_directory"
        assert payload["source_title"] == "Submitted Material files"
        assert payload["preferred_import_directory"] is None
        assert [item["file_name"] for item in payload["candidates"]] == [
            "Alpha Matrix.docx",
            "beta.PDF",
            "legacy.doc",
        ]
        assert set(payload["candidates"][0]) == {"candidate_id", "file_name"}
        assert str(submitted) not in listed.text
        assert spec.read_bytes() == original_bytes

        candidate_id = payload["candidates"][0]["candidate_id"]
        preview = client.post(
            f"/api/projects/P1/test-plan/source-candidates/{candidate_id}/matrix-preview",
            params={"view": "resolved_directory"},
        )
        assert preview.status_code == 200
        assert preview.json()["source_document_name"] == "Alpha Matrix.docx"
        assert preview.json()["source_document_path"] == "Alpha Matrix.docx"
        assert spec.read_bytes() == original_bytes
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def test_resolved_directory_preview_generates_pdf_token_for_docx(
    tmp_path: Path,
) -> None:
    client, engine = _client(tmp_path)
    fake_service = _FakeResolvedDirectoryPreviewService()
    app.dependency_overrides[get_project_test_plan_matrix_preview_service] = (
        lambda: fake_service
    )
    try:
        _create_project("P1", tmp_path)
        official = tmp_path / "official"
        submitted = official / "Submitted Material"
        submitted.mkdir(parents=True)
        spec = submitted / "Alpha Matrix.docx"
        _write_product_spec_docx(spec)
        _create_workspace("P1", official, tmp_path)

        listed = client.get(
            "/api/projects/P1/test-plan/source-candidates",
            params={"view": "resolved_directory"},
        )
        candidate_id = listed.json()["candidates"][0]["candidate_id"]

        preview = client.post(
            f"/api/projects/P1/test-plan/source-candidates/{candidate_id}/matrix-preview",
            params={"view": "resolved_directory"},
        )
        assert preview.status_code == 200
        payload = preview.json()

        token = payload["preview_pdf_token"]
        assert token is not None
        assert payload["source_document_name"] == "Alpha Matrix.docx"
        assert payload["source_document_path"] == "Alpha Matrix.docx"
        assert fake_service.previewed_source == spec
        assert fake_service.office.word_locations_requested == [spec]
        assert len(fake_service.office.word_pdf_exports) == 1

        preview_pdf = client.get(f"/api/test-plan/matrix-preview-pdf/{token}")
        assert preview_pdf.status_code == 200
        assert preview_pdf.headers["content-type"].startswith("application/pdf")
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.pop(get_session, None)
        app.dependency_overrides.pop(get_settings, None)
        app.dependency_overrides.pop(get_project_test_plan_matrix_preview_service, None)
        engine.dispose()


def test_resolved_directory_preview_rejects_in_place_same_name_replacement(
    tmp_path: Path,
) -> None:
    client, engine = _client(tmp_path)
    try:
        _create_project("P1", tmp_path)
        official = tmp_path / "official"
        submitted = official / "Submitted Material"
        submitted.mkdir(parents=True)
        source = submitted / "matrix.docx"
        source.write_bytes(b"first-content")
        _create_workspace("P1", official, tmp_path)
        listed = client.get(
            "/api/projects/P1/test-plan/source-candidates",
            params={"view": "resolved_directory"},
        )
        candidate_id = listed.json()["candidates"][0]["candidate_id"]
        original_times = (source.stat().st_atime_ns, source.stat().st_mtime_ns)

        source.write_bytes(b"other-content")
        os.utime(source, ns=original_times)

        response = client.post(
            f"/api/projects/P1/test-plan/source-candidates/{candidate_id}/matrix-preview",
            params={"view": "resolved_directory"},
        )
        assert response.status_code == 404
        assert "no longer available" in response.json()["detail"]
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def _client(tmp_path: Path) -> tuple[TestClient, object]:
    settings = Settings(
        data_dir=tmp_path / "data",
        projects_dir=tmp_path / "projects",
        templates_dir=tmp_path / "templates",
        database_path=tmp_path / "connlab.sqlite3",
    )
    engine = create_database_engine(settings)
    init_db(engine)
    session_factory = create_session_factory(engine)

    def override_session() -> Generator[Session, None, None]:
        with session_factory() as session:
            try:
                yield session
                session.commit()
            except Exception:
                session.rollback()
                raise

    app.dependency_overrides[get_session] = override_session
    app.dependency_overrides[get_settings] = lambda: settings
    return TestClient(app), engine


def _create_project(project_id: str, tmp_path: Path) -> None:
    settings = Settings(
        data_dir=tmp_path / "data",
        projects_dir=tmp_path / "projects",
        templates_dir=tmp_path / "templates",
        database_path=tmp_path / "connlab.sqlite3",
    )
    engine = create_database_engine(settings)
    session_factory = create_session_factory(engine)
    with session_factory() as session:
        ProjectRepository(session).create(
            Project(
                project_id=project_id,
                project_no=f"DL-2026-05-{project_id}",
                product_name="Connector",
                requestor="Alice",
                status=ProjectStatus.LTR_REGISTERED,
                created_on=date(2026, 5, 14),
            )
        )
        session.commit()
    engine.dispose()


def _create_file_asset(
    project_id: str,
    asset_id: str,
    path: Path,
    original_name: str,
    tmp_path: Path,
) -> None:
    settings = Settings(
        data_dir=tmp_path / "data",
        projects_dir=tmp_path / "projects",
        templates_dir=tmp_path / "templates",
        database_path=tmp_path / "connlab.sqlite3",
    )
    engine = create_database_engine(settings)
    session_factory = create_session_factory(engine)
    with session_factory() as session:
        FileAssetRepository(session).create(
            FileAsset(
                asset_id=asset_id,
                project_id=project_id,
                asset_type=FileAssetType.ATTACHMENT,
                path=path,
                original_name=original_name,
                registered_on=date(2026, 5, 14),
            )
        )
        session.commit()
    engine.dispose()


def _create_workspace(project_id: str, official: Path, tmp_path: Path) -> None:
    settings = Settings(
        data_dir=tmp_path / "data",
        projects_dir=tmp_path / "projects",
        templates_dir=tmp_path / "templates",
        database_path=tmp_path / "connlab.sqlite3",
    )
    engine = create_database_engine(settings)
    session_factory = create_session_factory(engine)
    with session_factory() as session:
        ProjectOfficialWorkspaceRepository(session).save(
            OfficialWorkspaceRecord(
                workspace_id="W1",
                project_id=project_id,
                dl_number="DL-2026-05-P1",
                local_workspace_path=tmp_path / "local",
                source_book_path=tmp_path / "source.xlsx",
                official_folder_path=official,
                manifest_path=tmp_path / "manifest.json",
                template_source_path=tmp_path / "template",
                created_at="2026-05-14T00:00:00Z",
            )
        )
        session.commit()
    engine.dispose()


def _write_product_spec_docx(path: Path) -> None:
    document = Document()
    table = document.add_table(rows=4, cols=4)
    rows = [
        ["test Items", "Section", "Group 1", "Group 2"],
        ["Examination of Product", "5.4", "1,10", "1,13"],
        ["Contact Resistance (Low Level)", "6.1", "2,5,8", "2,5,10"],
        ["Durability", "7.1", "", "3"],
    ]
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    document.save(path)


class _FakeResolvedDirectoryPreviewOffice:
    def __init__(self) -> None:
        self.word_locations_requested: list[Path] = []
        self.word_pdf_exports: list[tuple[Path, Path]] = []

    def read_word_table_locations(self, source_path: Path) -> tuple:
        self.word_locations_requested.append(Path(source_path))
        return ()

    def export_word_preview_pdf(self, source_path: Path, output_pdf_path: Path) -> Path:
        self.word_pdf_exports.append((Path(source_path), Path(output_pdf_path)))
        output_pdf_path.write_bytes(b"%PDF-1.4")
        return output_pdf_path


class _FakeResolvedDirectoryPreviewService:
    def __init__(self) -> None:
        self.office = _FakeResolvedDirectoryPreviewOffice()
        self.previewed_source: Path | None = None
        self.previewed_locator: tuple[int | None, int | None, str | None] | None = None

    def read_word_table_locations(self, source_path: Path) -> tuple:
        return self.office.read_word_table_locations(source_path)

    def export_word_preview_pdf(self, source_path: Path, output_pdf_path: Path) -> Path:
        return self.office.export_word_preview_pdf(source_path, output_pdf_path)

    def preview_from_path(
        self,
        command: MatrixPreviewFromPathCommand,
        *,
        preview_pdf_token: str | None = None,
        table_locations: tuple | None = None,
    ) -> ProjectTestPlanMatrixPreview:
        self.previewed_source = Path(command.source_path)
        self.previewed_locator = (
            command.page_number,
            command.page_table_index,
            command.table_text_query,
        )
        return ProjectTestPlanMatrixPreview(
            project_id=command.project_id,
            source_document_path=self.previewed_source,
            source_document_name=self.previewed_source.name,
            source_format=self.previewed_source.suffix.lower(),
            capability_status="supported",
            generated_at="2026-07-04T00:00:00+00:00",
            preview_pdf_token=preview_pdf_token,
            rows=(),
            groups=(),
            warnings=(),
            blockers=(),
            selected_table_index=None,
            selected_page_number=None,
            selected_page_table_index=None,
            candidate_tables=(),
        )
