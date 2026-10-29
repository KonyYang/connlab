from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
import pytest

from backend.api import dependencies as deps
from backend.api.main import app
from backend.application.external_excel_read_service import EquipmentCalibrationReadResult, EquipmentCalibrationRow
from backend.application.tools_equipment_report_service import ToolsEquipmentReportService
from backend.infrastructure.office.equipment_id_document_reader import EquipmentIdDocumentReader
from backend.infrastructure.office.test_report_document_gateway import TestReportDocumentGateway
from backend.shared.config import Settings


class Catalog:
    def __init__(self, path):
        self.path = path

    def read_equipment_calibrations(self):
        return EquipmentCalibrationReadResult(str(self.path), ("All Equip.",), (
            EquipmentCalibrationRow("DG-Q-0033", "Multimeter", "Keysight", "2024-01-01", "2025-01-01", "All Equip."),
            EquipmentCalibrationRow("DG-L-0002", "Microscope", "Nikon", "Not calibrated", "Not applicable", "All Equip."),
        ))


def report_bytes():
    document = Document()
    document.add_paragraph("LABORATORY TEST REPORT")
    document.add_paragraph("Reviewer manual content must remain.")
    document.add_paragraph("7. EQUIPMENT")
    table = document.add_table(rows=2, cols=5)
    for cell, value in zip(table.rows[0].cells, ("Item", "Manufacturer", "ID Number", "Last Cal.", "Cal. Due")):
        cell.text = value
    table.cell(1, 0).text = "old equipment"
    data = BytesIO()
    document.save(data)
    return data.getvalue()


@pytest.fixture
def client(tmp_path):
    data = tmp_path / "data"
    data.mkdir()
    workbook = tmp_path / "catalog.xlsx"
    workbook.write_bytes(b"fixture")
    settings = Settings(data_dir=data, templates_dir=tmp_path / "templates",
                        projects_dir=tmp_path / "projects", database_path=tmp_path / "test.sqlite")
    instance = ToolsEquipmentReportService(source_reader=EquipmentIdDocumentReader(),
        catalog_reader=Catalog(workbook), report_writer=TestReportDocumentGateway())
    app.dependency_overrides[deps.get_settings] = lambda: settings
    app.dependency_overrides[deps.get_tools_equipment_report_service] = lambda: instance
    try:
        yield TestClient(app), data
    finally:
        app.dependency_overrides.pop(deps.get_settings, None)
        app.dependency_overrides.pop(deps.get_tools_equipment_report_service, None)


def test_download_updates_equipment_preserves_manual_content_and_cleans_temporary_files(client):
    test_client, data = client
    source = report_bytes()
    digest = sha256(source).hexdigest()
    response = test_client.post("/api/tools/equipment-list",
        files={"file": ("../Internal.docx", source, "application/octet-stream")},
        data={"references_text": "Q-0033, DG-Q-0033; L-0002; DG-Q-9999"})
    assert response.status_code == 200, response.text
    assert "Internal_EquipmentUpdated.docx" in response.headers["content-disposition"]
    review = json.loads(response.headers["x-equipment-review"])
    assert review["unmatched"] == ["DG-Q-9999"]
    assert review["expired"] == ["Q-0033"]
    assert review["incomplete"] == []
    updated = Document(BytesIO(response.content))
    assert updated.paragraphs[1].text == "Reviewer manual content must remain."
    assert [cell.text for cell in updated.tables[0].rows[1].cells] == [
        "Multimeter", "Keysight", "DG-Q-0033", "01-Jan-2024", "01-Jan-2025"]
    assert updated.tables[0].cell(2, 3).text == "Not calibrated"
    assert str(updated.tables[0].cell(1, 4).paragraphs[0].runs[0].font.color.rgb) == "FF0000"
    assert updated.tables[0].cell(3, 2).text == "DG-Q-9999"
    assert sha256(source).hexdigest() == digest
    assert list(data.iterdir()) == []


def test_equipment_upload_can_share_report_filename_without_overwriting_it(client):
    test_client, data = client
    selection = Document()
    selection.add_paragraph("L-0002")
    equipment = BytesIO()
    selection.save(equipment)
    response = test_client.post("/api/tools/equipment-list", files={
        "file": ("same.docx", report_bytes()), "equipment_file": ("same.docx", equipment.getvalue()),
    })
    assert response.status_code == 200, response.text
    assert Document(BytesIO(response.content)).tables[0].cell(1, 2).text == "DG-L-0002"
    assert list(data.iterdir()) == []


@pytest.mark.parametrize("report,selection,expected", [
    (b"not a report", {"references_text": "Q-0033"}, 422),
    (report_bytes(), {}, 422),
], ids=["invalid-docx", "missing-selection"])
def test_invalid_inputs_do_not_download_and_cleanup_uploads(client, report, selection, expected):
    test_client, data = client
    response = test_client.post("/api/tools/equipment-list",
        files={"file": ("report.docx", report)}, data=selection)
    assert response.status_code == expected, response.text
    assert "content-disposition" not in response.headers
    assert list(data.iterdir()) == []


def test_invalid_equipment_upload_returns_actionable_error_and_no_download(client):
    test_client, data = client
    response = test_client.post("/api/tools/equipment-list", files={
        "file": ("report.docx", report_bytes()), "equipment_file": ("EquipmentID.docx", b"bad file"),
    })
    assert response.status_code == 422
    assert "valid .docx" in response.json()["detail"]
    assert list(data.iterdir()) == []


def test_both_equipment_sources_are_rejected(client):
    test_client, data = client
    response = test_client.post("/api/tools/equipment-list", files={
        "file": ("report.docx", report_bytes()), "equipment_file": ("EquipmentID.docx", b"unused"),
    }, data={"references_text": "Q-0033"})
    assert response.status_code == 422
    assert "not both" in response.json()["detail"]
    assert list(data.iterdir()) == []


def test_large_review_is_bounded_and_reports_omitted_items(client):
    test_client, _data = client
    references = [f"未知设备{i}-{'器' * 60}" for i in range(80)]
    response = test_client.post("/api/tools/equipment-list",
        files={"file": ("report.docx", report_bytes())}, data={"references_text": "\n".join(references)})
    assert response.status_code == 200, response.text
    encoded = response.headers["x-equipment-review"]
    assert len(encoded) <= 7000
    review = json.loads(encoded)
    assert len(review["unmatched"]) + review["omitted"]["unmatched"] == 80
    assert review["omitted"]["unmatched"] > 0


@pytest.mark.parametrize("merged_column", [0, 1, 4], ids=["item", "manufacturer", "due-date"])
@pytest.mark.parametrize("mixed_spans", [False, True], ids=["same-spans", "mixed-spans"])
def test_equipment_upload_updates_horizontally_merged_fields_without_shifting_values(client, merged_column, mixed_spans):
    test_client, data = client
    document = Document()
    document.add_paragraph("Customer or Internal Report with an existing Equipment List")
    table = document.add_table(rows=2, cols=6)
    body_column = (1 if merged_column == 0 else 0) if mixed_spans else merged_column
    for row, column in zip(table.rows, (merged_column, body_column), strict=True):
        row.cells[column].merge(row.cells[column + 1])
    headers = ("Item", "Manufacturer", "ID Number", "Last Cal.", "Cal. Due")
    # Use the document API to set each visible merged field once.
    for row, values, column in zip(table.rows, (headers, ("old", "old maker", "old ID", "old last", "old due")), (merged_column, body_column), strict=True):
        index = 0
        for field, value in enumerate(values):
            row.cells[index].text = value
            index += 2 if field == column else 1
    source = BytesIO()
    document.save(source)
    original = source.getvalue()
    original_grid = table._tbl.tblGrid.xml
    selection = Document()
    selection.add_paragraph("Q-0033")
    selection.add_paragraph("DG-Q-0033")
    selection.add_paragraph("L-0002")
    equipment = BytesIO()
    selection.save(equipment)
    response = test_client.post("/api/tools/equipment-list", files={
        "file": ("Report-CR.docx", original),
        "equipment_file": ("EquipmentID - copy.docx", equipment.getvalue()),
    })
    assert response.status_code == 200, response.text
    result = Document(BytesIO(response.content))
    assert result.paragraphs[0].text == document.paragraphs[0].text
    updated = result.tables[0]
    assert updated._tbl.tblGrid.xml == original_grid
    assert [tc.grid_span for tc in updated.rows[0]._tr.tc_lst] == [tc.grid_span for tc in table.rows[0]._tr.tc_lst]
    assert [tc.grid_span for tc in updated.rows[1]._tr.tc_lst] == [tc.grid_span for tc in table.rows[1]._tr.tc_lst]
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(response.content)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)
    before_body = deepcopy(document.element.body)
    after_body = deepcopy(result.element.body)
    before_body.remove(before_body.find(qn("w:tbl")))
    after_body.remove(after_body.find(qn("w:tbl")))
    assert before_body.xml == after_body.xml
    expected = ["Multimeter", "Keysight", "DG-Q-0033", "01-Jan-2024", "01-Jan-2025"]
    expected.insert(body_column, expected[body_column])
    assert [cell.text for cell in updated.rows[1].cells] == expected
    assert len(updated.rows) == 3  # Q-0033 and DG-Q-0033 are one equipment ID.
    second = ["Microscope", "Nikon", "DG-L-0002", "Not calibrated", "Not applicable"]
    second.insert(body_column, second[body_column])
    assert [cell.text for cell in updated.rows[2].cells] == second
    assert all(str(run.font.color.rgb) == "FF0000"
               for p in updated.rows[1].cells[-1].paragraphs for run in p.runs if run.text)
    assert all(str(run.font.color.rgb) == "000000"
               for p in updated.rows[2].cells[-1].paragraphs for run in p.runs if run.text)
    assert json.loads(response.headers["x-equipment-review"])["filled"] == 2
    repeated = test_client.post("/api/tools/equipment-list", files={
        "file": ("updated.docx", response.content),
        "equipment_file": ("EquipmentID.docx", equipment.getvalue()),
    })
    assert repeated.status_code == 200, repeated.text
    assert Document(BytesIO(repeated.content)).element.body.xml == result.element.body.xml
    assert source.getvalue() == original
    assert list(data.iterdir()) == []


def test_customer_report_without_equipment_table_returns_actionable_error(client):
    test_client, data = client
    document = Document()
    document.add_paragraph("Customer Test Report")
    document.add_paragraph("REVISION RECORD")
    source = BytesIO()
    document.save(source)
    response = test_client.post("/api/tools/equipment-list", files={
        "file": ("Customer-CR.docx", source.getvalue()),
    }, data={"references_text": "Q-0033"})
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert "Equipment List" in detail
    assert "Customer reports without" in detail
    assert "ID Number" in detail
    assert "approved table contract" not in detail
    assert "content-disposition" not in response.headers
    assert list(data.iterdir()) == []


@pytest.mark.parametrize("layout", ["vertical-merge", "merged-fields"])
def test_ambiguous_equipment_body_layout_is_rejected_without_a_download(client, layout):
    test_client, data = client
    document = Document(BytesIO(report_bytes()))
    table = document.tables[0]
    table.add_row()
    if layout == "vertical-merge":
        table.cell(1, 0).merge(table.cell(2, 0))
    elif layout == "merged-fields":
        table.cell(1, 0).merge(table.cell(1, 1))
    source = BytesIO()
    document.save(source)
    response = test_client.post("/api/tools/equipment-list", files={
        "file": ("report.docx", source.getvalue()),
    }, data={"references_text": "Q-0033"})
    assert response.status_code == 422
    assert "merged cells" in response.json()["detail"]
    assert "content-disposition" not in response.headers
    assert list(data.iterdir()) == []


def test_empty_merged_equipment_table_keeps_its_layout_when_first_row_is_written(client):
    test_client, data = client
    document = Document()
    table = document.add_table(rows=1, cols=6)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Item"
    for index, value in enumerate(("Manufacturer", "ID Number", "Last Cal.", "Cal. Due"), 2):
        table.cell(0, index).text = value
    source = BytesIO()
    document.save(source)
    response = test_client.post("/api/tools/equipment-list", files={
        "file": ("empty.docx", source.getvalue()),
    }, data={"references_text": "Q-0033"})
    assert response.status_code == 200, response.text
    updated = Document(BytesIO(response.content)).tables[0]
    assert [c.text for c in updated.rows[1].cells] == [
        "Multimeter", "Multimeter", "Keysight", "DG-Q-0033", "01-Jan-2024", "01-Jan-2025"]
    assert updated._tbl.tblGrid.xml == table._tbl.tblGrid.xml
    assert list(data.iterdir()) == []
