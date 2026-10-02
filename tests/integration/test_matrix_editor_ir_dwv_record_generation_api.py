from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from dataclasses import replace
from datetime import date
import shutil

import pytest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.api.main import app


def seed_ir_dwv_project(session, root: Path, *, project_id="IR-DWV-SMOKE", official=True, samples="5+5(d)"):
    """Reusable isolated repository fixture; never bind a business folder or template."""
    from backend.domain import Project, LtrRecord, LtrStatus, ConfirmedMatrixCell, ConfirmedMatrixStatus, ExternalResourceType
    from backend.application.external_resource_service import ExternalResourceService
    from backend.application.matrix_editor_llcr_cr_record_projection import _draft_snapshot, MatrixEditorLlcrCrRecordGroupInput, MatrixEditorLlcrCrRecordRowInput
    from backend.application.official_project_workspace_service import OfficialWorkspaceRecord
    from backend.application.project_basic_information_service import ProjectBasicInformationRecord
    from backend.application.source_matrix_import_persistence_service import SourceMatrixImportPersistenceService, PersistSourceMatrixFromPreviewCommand
    from backend.domain.matrix_contact_measurement_models import MatrixPointProfile
    from backend.infrastructure.storage.repositories.project import ProjectRepository
    from backend.infrastructure.storage.repositories.records import LtrRecordRepository
    from backend.infrastructure.storage.repositories.external_resources import ExternalResourceRepository
    from backend.infrastructure.storage.repositories.confirmed_matrix_authority import ConfirmedMatrixAuthorityRepository
    from backend.infrastructure.storage.repositories.project_basic_information import ProjectBasicInformationRepository
    from backend.infrastructure.storage.repositories.official_workspace import ProjectOfficialWorkspaceRepository
    from backend.infrastructure.storage.repositories import SourceMatrixImportRepository
    from backend.infrastructure.official_workspace_manifest import OfficialWorkspaceManifest, OfficialWorkspaceManifestGateway, stable_folder_identity
    root = Path(root)
    template = root / "templates" / "IR&DWV Template.xlsx"
    template.parent.mkdir(parents=True)
    shutil.copyfile(Path(__file__).parents[1] / "fixtures" / "ir_dwv" / "IR_DWV_Template.xlsx", template)
    ExternalResourceService(ExternalResourceRepository(session)).upsert_resource(ExternalResourceType.IR_DWV_RECORD_TEMPLATE, template, True)
    dl = "DL-2026-10-IRDWV"
    ProjectRepository(session).create(Project(project_id=project_id, project_no=dl, product_name="Isolated IR DWV connector", requestor="Requestor"))
    LtrRecordRepository(session).create(LtrRecord("smoke-ltr", project_id, dl, LtrStatus.REGISTERED, date(2026, 10, 1)))
    request = _request()
    request["groups"][0]["sample_quantity_expression"] = samples
    source_payload = {
        "groups": [{**request["groups"][0], "steps": []}],
        "rows": [{"source_row_index": index, "test_item": row["test_item"], "condition": row["condition"],
                  "group_tokens": {"6": "2,6"}, "is_sample_row": False} for index, row in enumerate(request["rows"], start=1)],
        "selected_group_keys_at_import": ["g1"], "warnings": [], "blockers": [],
    }
    source_store = SourceMatrixImportRepository(session)
    lineage = SourceMatrixImportPersistenceService(store=source_store).persist_from_preview(PersistSourceMatrixFromPreviewCommand(
        project_id=project_id, source_document_path=str(root / "isolated-spec.xlsx"), source_document_name="isolated-spec.xlsx",
        source_format=".xlsx", payload=source_payload, selected_group_keys=("g1",), created_at="2026-10-01"))
    source = source_store.get_snapshot(lineage.snapshot_id)
    groups = tuple(MatrixEditorLlcrCrRecordGroupInput(**item) for item in request["groups"])
    rows = tuple(MatrixEditorLlcrCrRecordRowInput(**item) for item in request["rows"])
    snapshot = _draft_snapshot(project_id=project_id, groups=groups, rows=rows, point_profile=MatrixPointProfile((), electrical_point_pairs="Odd&Even，P1&P2"))
    matrix_id = "smoke-ir-dwv-matrix"
    snapshot = replace(snapshot,
        version=replace(snapshot.version, confirmed_matrix_id=matrix_id, confirmed_revision=1, is_active_authority=True, status=ConfirmedMatrixStatus.CONFIRMED,
                        source_import_id=lineage.import_id, source_snapshot_id=lineage.snapshot_id,
                        planned_test_start_date="2026-10-01", planned_test_complete_date="2026-10-02"),
        groups=tuple(replace(item, confirmed_matrix_id=matrix_id, source_group_snapshot_id=source.groups[index].group_snapshot_id) for index, item in enumerate(snapshot.groups)),
        rows=tuple(replace(item, confirmed_matrix_id=matrix_id, source_row_snapshot_id=source.rows[index].row_snapshot_id) for index, item in enumerate(snapshot.rows)),
        step_quantities=tuple(replace(item, confirmed_matrix_id=matrix_id) for item in snapshot.step_quantities),
        cells=tuple(ConfirmedMatrixCell(f"smoke-cell-{index}", matrix_id, row.confirmed_row_id, snapshot.groups[0].confirmed_group_id,
                                       row.draft_row_id, snapshot.groups[0].draft_group_id, "2,6") for index, row in enumerate(snapshot.rows)),
    )
    ConfirmedMatrixAuthorityRepository(session).create_snapshot(snapshot)
    ProjectBasicInformationRepository(session).create_confirmed(ProjectBasicInformationRecord(
        "smoke-basic", project_id, "confirmed", 1, {"product_description": "Isolated authoritative connector", "requested_by": "Confirmed requestor"},
        "smoke", "2026-10-01", "2026-10-01", "2026-10-01", "smoke"))
    target = None
    if official:
        local = root / "projects" / dl
        folder = local / f"{dl} Isolated connector"
        (folder / "Test results").mkdir(parents=True)
        manifest = local / ".connlab.json"
        OfficialWorkspaceManifestGateway().write(manifest, OfficialWorkspaceManifest(1, project_id, dl, str(local), "", str(folder), str(template.parent), "2026-10-01", stable_folder_identity(folder)))
        ProjectOfficialWorkspaceRepository(session).save(OfficialWorkspaceRecord("smoke-workspace", project_id, dl, local, root / "LTR.xlsx", folder, manifest, template.parent, "2026-10-01"))
        target = folder / "Test results" / f"{dl} IR&DWV Record.xlsx"
    session.commit()
    return SimpleNamespace(project_id=project_id, request=request, target=target, template=template)


@pytest.fixture
def real_project(tmp_path, request):
    from sqlalchemy import create_engine
    from backend.api.dependencies import get_session, get_settings
    from backend.infrastructure.storage.database import init_db, create_session_factory
    from backend.shared.config import Settings
    engine = create_engine(f"sqlite:///{(tmp_path / 'smoke.sqlite').as_posix()}")
    init_db(engine)
    sessions = create_session_factory(engine)
    with sessions() as session:
        options = getattr(request, "param", "5+5(d)")
        options = options if isinstance(options, dict) else {"samples": options}
        seeded = seed_ir_dwv_project(session, tmp_path, **options)
    settings = Settings(data_dir=tmp_path / "data", database_path=tmp_path / "smoke.sqlite", templates_dir=tmp_path / "templates", projects_dir=tmp_path / "projects")
    def isolated_session():
        with sessions() as session:
            yield session
            session.commit()
    app.dependency_overrides[get_session] = isolated_session
    app.dependency_overrides[get_settings] = lambda: settings
    try:
        yield TestClient(app), seeded, settings
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


@pytest.mark.parametrize("real_project,pending,edited_samples,authority,file_name", [
    ({"official": False}, False, None, "confirmed", "DL-2026-10-IRDWV IR&DWV Record.xlsx"),
    ({"official": False}, True, None, "unconfirmed", "DL-2026-10-IRDWV IR&DWV Record draft.xlsx"),
    ({"official": True}, True, None, "unconfirmed", "DL-2026-10-IRDWV IR&DWV Record draft.xlsx"),
    ({"official": True}, False, "7", "unconfirmed", "DL-2026-10-IRDWV IR&DWV Record draft.xlsx"),
], indirect=["real_project"])
def test_download_name_uses_registered_ltr_and_verified_matrix_authority(
    real_project, pending, edited_samples, authority, file_name,
):
    from urllib.parse import unquote
    client, seeded, settings = real_project
    request = {**seeded.request, "matrix_has_pending_changes": pending}
    if edited_samples is not None:
        request["groups"][0]["sample_quantity_expression"] = edited_samples
    base = f"/api/projects/{seeded.project_id}/matrix-editor"
    preview = client.post(base + "/ir-dwv-record-publication/preview", json=request)
    assert preview.status_code == 200, preview.text
    reviewed = preview.json()
    assert (reviewed["mode"], reviewed["authority_status"]) == ("download", authority)
    for _ in range(2):
        downloaded = client.post(base + "/ir-dwv-record-draft/generate",
                                 json={**request, "preview_token": reviewed["preview_token"]})
        assert downloaded.status_code == 200, downloaded.text
        assert downloaded.content.startswith(b"PK")
        assert unquote(downloaded.headers["content-disposition"]).endswith(file_name)
    artifacts = tuple((settings.data_dir / "generated_ir_dwv_record_drafts" / seeded.project_id).glob("*.xlsx"))
    assert len(artifacts) == 2 and artifacts[0] != artifacts[1]
    if seeded.target is not None:
        assert not seeded.target.exists()


def test_real_di_creates_and_archives_official_ir_dwv_forms(real_project):
    client, seeded, _settings = real_project
    base = f"/api/projects/{seeded.project_id}/matrix-editor/ir-dwv-record-publication"
    preview = client.post(base + "/preview", json=seeded.request)
    assert preview.status_code == 200, preview.text
    assert preview.json()["mode"] == "official" and preview.json()["status"] == "ready"
    created = client.post(base + "/publish", json={**seeded.request, "preview_token": preview.json()["preview_token"], "conflict_action": "none"})
    assert created.status_code == 200, created.text
    assert created.json()["file_name"] == "DL-2026-10-IRDWV IR&DWV Record.xlsx"
    assert Path(created.json()["target_path"]) == seeded.target
    workbook = load_workbook(seeded.target)
    try:
        assert "Isolated authoritative connector" in str(workbook.worksheets[0]["I11"].value)
        assert workbook.worksheets[0]["C22"].value is None
    finally:
        workbook.close()
    seeded.target.write_bytes(b"operator measurements kept exactly")
    conflict = client.post(base + "/preview", json=seeded.request)
    assert conflict.json()["status"] == "conflict"
    denied = client.post(base + "/publish", json={**seeded.request, "preview_token": conflict.json()["preview_token"], "conflict_action": "none"})
    assert denied.status_code == 409
    replaced = client.post(base + "/publish", json={**seeded.request, "preview_token": conflict.json()["preview_token"], "conflict_action": "archive"})
    assert replaced.status_code == 200, replaced.text
    assert Path(replaced.json()["archive_path"]).read_bytes() == b"operator measurements kept exactly"


def test_real_di_editor_session_rebuilds_matching_current_ui_request(real_project):
    client, seeded, _settings = real_project
    response = client.get(f"/api/projects/{seeded.project_id}/matrix-editor/session")
    assert response.status_code == 200, response.text
    session = response.json()
    assert session["source_status"] == "available"
    assert session["loaded_source"] == "authority"
    draft = session["editor_draft"]
    groups = draft["groups"]
    cells = {(cell["draft_row_id"], cell["draft_group_id"]): cell["cell_value"] for cell in draft["cells"]}
    request = {**seeded.request,
        "groups": [{key: group[key] for key in ("group_key", "group_label", "sample_quantity_expression", "draft_group_id")} for group in groups],
        "rows": [{**{key: row[key] for key in ("test_item", "condition", "draft_row_id")},
                  "group_values": {group["group_key"]: cells[(row["draft_row_id"], group["draft_group_id"])] for group in groups}} for row in draft["rows"]],
        "point_profile": draft["point_profile"],
        "point_overrides": draft["point_overrides"],
    }
    preview = client.post(f"/api/projects/{seeded.project_id}/matrix-editor/ir-dwv-record-publication/preview", json=request)
    assert preview.status_code == 200, preview.text
    assert preview.json()["mode"] == "official"
    assert preview.json()["authority_status"] == "confirmed"


@pytest.mark.parametrize("real_project", ["18"], indirect=True)
def test_real_di_blocks_over_capacity_before_publication_operation(real_project):
    from backend.infrastructure.files.generation_journal import GenerationJournal
    client, seeded, settings = real_project
    seeded.target.write_bytes(b"measured original must remain")
    base = f"/api/projects/{seeded.project_id}/matrix-editor/ir-dwv-record-publication"
    for pending in (False, True):
        request = {**seeded.request, "matrix_has_pending_changes": pending}
        preview = client.post(base + "/preview", json=request)
        assert preview.status_code == 200, preview.text
        assert preview.json()["status"] == "blocked"
        assert any("cannot fit" in reason for reason in preview.json()["blockers"])
        attempted = client.post(base + "/publish", json={**request, "preview_token": preview.json()["preview_token"], "conflict_action": "archive"})
        assert attempted.status_code == 409
        assert GenerationJournal(settings.data_dir / "ir_dwv_record_publication").read(seeded.project_id) is None
        assert seeded.target.read_bytes() == b"measured original must remain"
        assert not list((seeded.target.parent.parent / "History").glob("**/*.xlsx"))


def test_real_di_rejects_changed_template_target_and_concurrent_writer(real_project):
    from backend.infrastructure.files.generation_journal import GenerationJournal
    client, seeded, settings = real_project
    base = f"/api/projects/{seeded.project_id}/matrix-editor/ir-dwv-record-publication"
    seeded.target.write_bytes(b"existing measured report")
    preview = client.post(base + "/preview", json=seeded.request).json()
    seeded.template.write_bytes(seeded.template.read_bytes() + b"changed-template")
    changed = client.post(base + "/publish", json={**seeded.request, "preview_token": preview["preview_token"], "conflict_action": "archive"})
    assert changed.status_code == 409 and "changed" in changed.json()["detail"]
    assert seeded.target.read_bytes() == b"existing measured report"
    preview = client.post(base + "/preview", json=seeded.request).json()
    seeded.target.write_bytes(b"changed target")
    changed = client.post(base + "/publish", json={**seeded.request, "preview_token": preview["preview_token"], "conflict_action": "archive"})
    assert changed.status_code == 409
    with GenerationJournal(settings.data_dir / "project_folder_generation").lock(seeded.project_id):
        concurrent = client.post(base + "/publish", json={**seeded.request, "preview_token": preview["preview_token"], "conflict_action": "archive"})
    assert concurrent.status_code == 409 and "already running" in concurrent.json()["detail"]
    assert seeded.target.read_bytes() == b"changed target"


def test_real_di_pending_matrix_download_and_writer_failure_preserve_official_file(real_project, monkeypatch):
    from backend.infrastructure.office.ir_dwv_record_workbook_gateway import IrDwvRecordWorkbookGateway
    client, seeded, _settings = real_project
    base = f"/api/projects/{seeded.project_id}/matrix-editor/ir-dwv-record-publication"
    request = {**seeded.request, "matrix_has_pending_changes": True}
    preview = client.post(base + "/preview", json=request)
    assert preview.json()["mode"] == "download" and preview.json()["authority_status"] == "unconfirmed"
    downloaded = client.post(f"/api/projects/{seeded.project_id}/matrix-editor/ir-dwv-record-draft/generate",
                             json={**request, "preview_token": preview.json()["preview_token"]})
    assert downloaded.status_code == 200 and downloaded.content.startswith(b"PK")
    assert not seeded.target.exists()
    seeded.target.write_bytes(b"existing measured report")
    preview = client.post(base + "/preview", json=seeded.request).json()
    def locked_writer(*args, **kwargs):
        raise PermissionError("simulated XLSX write denied")
    monkeypatch.setattr(IrDwvRecordWorkbookGateway, "write", locked_writer)
    failed = client.post(base + "/publish", json={**seeded.request, "preview_token": preview["preview_token"], "conflict_action": "archive"})
    assert failed.status_code == 409 and "write denied" in failed.json()["detail"]
    assert seeded.target.read_bytes() == b"existing measured report"


def test_real_di_template_change_during_generation_stops_archive(real_project, monkeypatch):
    from backend.infrastructure.office.ir_dwv_record_workbook_gateway import IrDwvRecordWorkbookGateway
    client, seeded, _settings = real_project
    base = f"/api/projects/{seeded.project_id}/matrix-editor/ir-dwv-record-publication"
    seeded.target.write_bytes(b"measured original")
    preview = client.post(base + "/preview", json=seeded.request).json()
    real_write = IrDwvRecordWorkbookGateway.write
    def changing_writer(self, **kwargs):
        result = real_write(self, **kwargs)
        seeded.template.write_bytes(seeded.template.read_bytes() + b"source changed while generating")
        return result
    monkeypatch.setattr(IrDwvRecordWorkbookGateway, "write", changing_writer)
    rejected = client.post(base + "/publish", json={**seeded.request, "preview_token": preview["preview_token"], "conflict_action": "archive"})
    assert rejected.status_code == 409 and "changed" in rejected.json()["detail"]
    assert seeded.target.read_bytes() == b"measured original"


def _request():
    return {"source": "matrix_editor_current_ui_state", "record_type": "ir_dwv",
            "groups": [{"group_key": "g1", "group_label": "6", "sample_quantity_expression": "5+5(d)"}],
            "rows": [{"test_item": "IR", "condition": "500 VDC", "group_values": {"g1": "2,6"}},
                     {"test_item": "DWV", "condition": "1500 VDC", "group_values": {"g1": "2,6"}}],
            "point_profile": {"categories": [], "electrical_point_pairs": "Odd&Even，P1&P2"}}


@pytest.mark.parametrize("samples", [3, 7])
def test_ir_dwv_download_uses_current_ui_rounds_pairs_and_only_authoritative_header(tmp_path, samples):
    from backend.api.dependencies import get_matrix_editor_ir_dwv_record_generation_service, get_matrix_editor_ir_dwv_record_publication_service
    from backend.application.matrix_editor_ir_dwv_record_generation_service import MatrixEditorIrDwvRecordGenerationService
    from backend.infrastructure.files.ir_dwv_record_artifact_store import IrDwvRecordArtifactStore
    from backend.infrastructure.office.ir_dwv_record_workbook_gateway import IrDwvRecordWorkbookGateway
    service = MatrixEditorIrDwvRecordGenerationService(
        confirmed_store=SimpleNamespace(get_active_by_project=lambda _: None),
        basic_information_store=SimpleNamespace(get_latest_confirmed=lambda _: SimpleNamespace(version=1, values={"product_description": "Actual connector", "requested_by": "Requestor"})),
        ltr_store=SimpleNamespace(list_by_project=lambda _: [SimpleNamespace(status="registered", registered_on="2026-10-01", ltr_number="DL-001")]),
        workbook_gateway=IrDwvRecordWorkbookGateway(Path(__file__).parents[1] / "fixtures" / "ir_dwv" / "IR_DWV_Template.xlsx"),
        artifact_store=IrDwvRecordArtifactStore(tmp_path / "drafts"),
    )
    class Publication:
        def validate_download(self, command, token):
            assert token == "reviewed" and command.draft.record_type == "ir_dwv"
            return SimpleNamespace(authority_status="unconfirmed")
    app.dependency_overrides[get_matrix_editor_ir_dwv_record_generation_service] = lambda: service
    app.dependency_overrides[get_matrix_editor_ir_dwv_record_publication_service] = lambda: Publication()
    request = _request()
    request["groups"][0]["sample_quantity_expression"] = str(samples)
    request["rows"] = [
        {"test_item": "Visual", "group_values": {"g1": "1,10"}},
        {"test_item": "IR", "condition": "500 VDC 2 min", "requirement": "1000 MΩ minimum", "group_values": {"g1": "2,5,8"}},
        {"test_item": "DWV", "condition": "1500 VDC 1 min", "requirement": "1 mA maximum", "group_values": {"g1": "3,6,9"}},
        {"test_item": "Thermal Shock", "group_values": {"g1": "4"}},
        {"test_item": "Humidity", "group_values": {"g1": "7"}},
    ]
    try:
        response = TestClient(app).post("/api/projects/P1/matrix-editor/ir-dwv-record-draft/generate", json={**request, "preview_token": "reviewed"})
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 200, response.text
    with BytesIO(response.content) as stream:
        workbook = load_workbook(stream)
        try:
            sheet = workbook.worksheets[0]
            from backend.infrastructure.office.ir_dwv_record_workbook_layout import block_layout, plan_block_placements
            placements = plan_block_placements(sample_count=samples, pair_count=2, step_count=3)
            assert sum(len(sheet._images) for sheet in workbook.worksheets) == 3
            for placement, label in zip(placements, ("Initial", "After Thermal Shock", "Final"), strict=True):
                current_sheet = workbook.worksheets[placement.sheet_index]
                layout = block_layout(origin_column=placement.origin_column, sample_count=samples, pair_count=2)
                assert str(current_sheet.cell(10, layout.origin_column).value).endswith(f"IR&DWV-{label}")
                assert "500 VDC 2 min" in current_sheet.cell(14, layout.origin_column).value
                assert "1500 VDC 1 min" in current_sheet.cell(14, layout.dwv_start_column).value
                assert [current_sheet.cell(layout.sample_id_row, layout.ir_start_column + index).value for index in range(layout.sample_slot_count)] == [f"{index + 1}#" if index < samples else None for index in range(layout.sample_slot_count)]
                assert current_sheet.cell(layout.sample_id_row, layout.dwv_start_column + samples - 1).value == f"{samples}#"
                assert current_sheet.cell(3, layout.origin_column + 3).value == "Eleactrical Safety Compliance Analyzer"
                assert current_sheet.cell(3, layout.origin_column + 6).value == "DG-Q-0624"
                assert current_sheet.cell(3, layout.origin_column + 7).value is None
                assert current_sheet.cell(3, layout.origin_column + 8).value is None
            assert any("Actual connector" in str(cell.value) for cell in sheet[11])
            first_layout = block_layout(origin_column=2, sample_count=samples, pair_count=2)
            assert sheet.cell(first_layout.data_start_row, 2).value == "Odd&Even"
            assert sheet.cell(first_layout.data_start_row + 1, 2).value == "P1&P2"
            assert sheet.cell(first_layout.data_start_row, 3).value is None
            assert "5+5" not in str(sheet.cell(first_layout.sample_id_row, 3).value)
        finally:
            workbook.close()


def test_ir_dwv_preview_and_stale_download_contract():
    from backend.api.dependencies import get_matrix_editor_ir_dwv_record_publication_service
    from backend.application.confirmed_matrix_llcr_cr_record_generation_service import MatrixEditorLlcrCrPublicationPreview
    class Publication:
        def preview(self, command):
            assert command.draft.record_type == "ir_dwv"
            return MatrixEditorLlcrCrPublicationPreview("P1", "download", "ready", "unconfirmed", None, False, (), "reviewed", information=("No equipment selected",))
        def validate_download(self, command, token):
            raise ValueError("Template or authority changed; preview again")
    app.dependency_overrides[get_matrix_editor_ir_dwv_record_publication_service] = lambda: Publication()
    try:
        client = TestClient(app)
        preview = client.post("/api/projects/P1/matrix-editor/ir-dwv-record-publication/preview", json=_request())
        stale = client.post("/api/projects/P1/matrix-editor/ir-dwv-record-draft/generate", json={**_request(), "preview_token": "old"})
        no_token = client.post("/api/projects/P1/matrix-editor/ir-dwv-record-draft/generate", json=_request())
    finally:
        app.dependency_overrides.clear()
    assert preview.status_code == 200
    assert preview.json()["information"] == ["No equipment selected"]
    assert stale.status_code == 422 and "changed" in stale.json()["detail"]
    assert no_token.status_code == 422
