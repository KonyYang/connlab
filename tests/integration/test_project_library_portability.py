import pytest
from hashlib import sha256
import json
from pathlib import Path
import shutil
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from backend.shared.config import Settings
from backend.infrastructure.storage.database import init_db
from backend.infrastructure.storage.models import ProjectModel
from backend.infrastructure.storage.models_report_authoring import ReportAuthoringRevisionModel
from backend.api.routes_project_library import router, get_project_library_service
from backend.application.project_library.service import ProjectLibraryService
from backend.infrastructure.project_library.files import LibraryFiles
from backend.infrastructure.project_library.portability import ProjectPortability
from backend.infrastructure.storage.repositories.project_library import ProjectLibraryRepository


def app_at(root, *, managed=True):
    settings = Settings(root / "Data", root / "Projects", root / "Templates", root / "Data" / "connlab.sqlite3", library_root=root if managed else None)
    settings.ensure_directories()
    engine = create_engine(f"sqlite:///{settings.database_path}", connect_args={"check_same_thread": False})
    init_db(engine)
    app = FastAPI()
    app.include_router(router)
    def service():
        with Session(engine) as session:
            store = ProjectLibraryRepository(session)
            files = LibraryFiles(settings, destination_configured=True)
            yield ProjectLibraryService(store, settings, files=files,
                portability=ProjectPortability(store, files, settings), opener=lambda path: None)
    app.dependency_overrides[get_project_library_service] = service
    return TestClient(app), settings, engine


def confirm(client, url, plan):
    response = client.post(url, json={"plan_id": plan["plan_id"], "fingerprint": plan["fingerprint"]})
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize("route", ["inspect-portable", "restore-plan"])
@pytest.mark.parametrize("source_state", ["no_metadata", "empty_portable", "unavailable", "damaged"])
def test_restore_review_distinguishes_missing_snapshot_and_keeps_restore_blocked(tmp_path, route, source_state):
    client, settings, engine = app_at(tmp_path / "library")
    source = tmp_path / "old project"
    if source_state != "unavailable":
        source.mkdir()
        (source / "note.txt").write_text("Original project material", encoding="utf-8")
    if source_state in {"empty_portable", "damaged"}:
        portable = source / ".connlab/portable"
        portable.mkdir(parents=True)
        if source_state == "damaged":
            (portable / "snapshot.json").write_text("{", encoding="utf-8")
    original = {path.relative_to(source): path.read_bytes() for path in source.rglob("*") if path.is_file()}
    try:
        response = client.post(f"/api/project-library/{route}", json={"folder": str(source)})
        if source_state == "unavailable":
            assert response.status_code == 409, response.text
            assert response.json()["detail"] == "The selected location is unavailable."
        else:
            assert response.status_code == 200, response.text
            plan = response.json()
            assert plan["status"] == "unavailable", plan
            assert plan["resolution_options"] == [], plan
            if source_state in {"no_metadata", "empty_portable"}:
                assert plan["blockers"] == ["未找到 ConnLab 数据快照"]
            else:
                assert plan["blockers"] == ["Project metadata is damaged. Restore a verified copy."]
            assert client.post("/api/project-library/restore", json={key: plan[key] for key in ("plan_id", "fingerprint")}).status_code == 409
        assert client.get("/api/project-library/members").json()["items"] == []
        assert list(settings.projects_dir.iterdir()) == []
        assert {path.relative_to(source): path.read_bytes() for path in source.rglob("*") if path.is_file()} == original
        assert source.exists() is (source_state != "unavailable")
    finally:
        client.close()
        engine.dispose()


def test_project_travels_without_source_database_and_preserves_report_evidence(tmp_path):
    a, settings_a, engine_a = app_at(tmp_path / "A")
    b, settings_b, engine_b = app_at(tmp_path / "B")
    source = tmp_path / "old" / "DL-2026-09-032 Manual project"
    source.mkdir(parents=True)
    (source / "result.csv").write_text("P1,12.34", encoding="utf-8")
    plan = a.post("/api/project-library/intake-plan", json={"sources": [str(source)]}).json()
    batch = a.post("/api/project-library/batches", json={**{k: plan[k] for k in ("plan_id", "fingerprint")}, "selected": [str(source)]}).json()
    member_id = batch["items"][0]["member_id"]
    take_plan = a.get(f"/api/project-library/members/{member_id}/takeover-plan").json()
    managed = confirm(a, f"/api/project-library/members/{member_id}/takeover", take_plan)
    project_id = managed["project_id"]
    content = b"evidence from computer A"
    content_hash = sha256(content).hexdigest()
    evidence = settings_a.data_dir / "report_authoring_assets" / sha256(project_id.encode()).hexdigest() / (content_hash + ".txt")
    evidence.parent.mkdir(parents=True)
    evidence.write_bytes(content)
    payload = {"project_id": project_id, "revision": 1, "sources": [{"source_id": "source-A", "role": "evidence", "file_name": "evidence.txt", "sha256": content_hash, "size_bytes": len(content), "created_at": "2026-10-27T00:00:00Z", "report_dl": ""}], "mappings": [], "base_source_id": None, "outputs": [], "results": [], "customer_outputs": []}
    with Session(engine_a) as session:
        session.add(ReportAuthoringRevisionModel(project_id=project_id, revision=1, payload_json=json.dumps(payload)))
        second = {**payload, "revision": 2, "base_source_id": "source-A"}
        session.add(ReportAuthoringRevisionModel(project_id=project_id, revision=2, payload_json=json.dumps(second)))
        session.commit()
    with Session(engine_b) as session:
        session.add(ProjectModel(project_id="unrelated", project_no=None, product_name="Other", requestor="Other", status="draft"))
        session.flush()
        session.add(ReportAuthoringRevisionModel(project_id="unrelated", revision=1, payload_json=json.dumps({"sources": []})))
        session.commit()
    info = a.get(f"/api/project-library/members/{member_id}/portable").json()
    prepared = a.post(f"/api/project-library/members/{member_id}/prepare", json={"fingerprint": info["fingerprint"]})
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()["record_counts"]["report_authoring_revisions"] == 2
    member = a.get(f"/api/project-library/members/{member_id}").json()
    travelling = tmp_path / "copied to another disk" / Path(member["folder_path"]).name
    shutil.copytree(member["folder_path"], travelling)
    engine_a.dispose()
    settings_a.database_path.rename(settings_a.database_path.with_suffix(".inaccessible"))
    shutil.rmtree(settings_a.projects_dir)
    restore_plan = b.post("/api/project-library/restore-plan", json={"folder": str(travelling)}).json()
    assert restore_plan["status"] == "ready", restore_plan
    restored = confirm(b, "/api/project-library/restore", restore_plan)
    assert restored["member"]["project_id"] == project_id
    assert restored["member"]["can_continue"] is True
    assert restored["member"]["snapshot_status"] == "current"
    assert (Path(restored["member"]["folder_path"]) / "result.csv").read_text() == "P1,12.34"
    restored_evidence = settings_b.data_dir / "report_authoring_assets" / sha256(project_id.encode()).hexdigest() / (content_hash + ".txt")
    assert restored_evidence.read_bytes() == content
    with Session(engine_b) as session:
        rows = list(session.scalars(select(ReportAuthoringRevisionModel).where(ReportAuthoringRevisionModel.project_id == project_id)))
        assert len(rows) == 2
        assert json.loads(rows[0].payload_json)["sources"][0]["source_id"] == "source-A"
        assert rows[0].id != 1
    repeated = b.post("/api/project-library/restore-plan", json={"folder": str(travelling)}).json()
    assert repeated["status"] == "already_added", repeated
    assert confirm(b, "/api/project-library/restore", repeated)["status"] == "already_added"
    engine_b.dispose()


def prepared_manual_project(client, root, ltr="DL-2026-09-032"):
    source = root / f"{ltr} History"
    source.mkdir(parents=True)
    (source / "result.txt").write_text("source result", encoding="utf-8")
    plan = client.post("/api/project-library/intake-plan", json={"sources": [str(source)]}).json()
    batch = client.post("/api/project-library/batches", json={**{k: plan[k] for k in ("plan_id", "fingerprint")}, "selected": [str(source)]}).json()
    member_id = batch["items"][0]["member_id"]
    plan = client.get(f"/api/project-library/members/{member_id}/takeover-plan").json()
    member = confirm(client, f"/api/project-library/members/{member_id}/takeover", plan)["member"]
    info = client.get(f"/api/project-library/members/{member_id}/portable").json()
    response = client.post(f"/api/project-library/members/{member_id}/prepare", json={"fingerprint": info["fingerprint"]})
    assert response.status_code == 200, response.text
    return member


@pytest.mark.parametrize("new_owner_ltr", ["DL-2026-09-033", "DL-2026-09-032"])
def test_restore_review_tracks_relevant_ltr_owners_only(tmp_path, new_owner_ltr):
    from backend.infrastructure.storage.models import LtrRecordModel
    a, _, engine_a = app_at(tmp_path / "A")
    b, settings_b, engine_b = app_at(tmp_path / "B", managed=False)
    try:
        first = prepared_manual_project(a, tmp_path / "sources")
        second = prepared_manual_project(a, tmp_path / "sources", "DL-2026-09-033")
        with Session(engine_a) as session:
            for index, member in enumerate((first, second)):
                session.add(LtrRecordModel(ltr_id=f"source-ltr-{index}", project_id=member["project_id"], ltr_number=member["ltr_number"], status="registered", is_current_owner=True))
            session.commit()
        for member in (first, second):
            url = f"/api/project-library/members/{member['member_id']}"
            info = a.get(url + "/portable").json()
            assert a.post(url + "/prepare", json={"fingerprint": info["fingerprint"]}).status_code == 200
        reviews = [b.post("/api/project-library/restore-plan", json={"folder": member["folder_path"]}).json() for member in (first, second)]
        assert all(plan["status"] == "ready" for plan in reviews)
        confirm(b, "/api/project-library/restore", reviews[0])
        if new_owner_ltr == "DL-2026-09-033":
            restored = confirm(b, "/api/project-library/restore", reviews[1])
            assert restored["member"]["ltr_number"] == new_owner_ltr
            assert len(b.get("/api/project-library/members").json()["items"]) == 2
        else:
            # A same-LTR decision must still be invalidated by a new owner.
            with Session(engine_b) as session:
                session.add(ProjectModel(project_id="new-owner", product_name="Other", requestor="Lab", status="draft"))
                session.flush()
                session.add(LtrRecordModel(ltr_id="new-ltr", project_id="new-owner", ltr_number="DL-2026-09-033", status="registered", is_current_owner=True))
                session.commit()
            response = b.post("/api/project-library/restore", json={key: reviews[1][key] for key in ("plan_id", "fingerprint")})
            assert response.status_code == 409, response.text
            assert "LTR ownership changed" in response.json()["detail"]
            assert len(b.get("/api/project-library/members").json()["items"]) == 1
            assert not Path(reviews[1]["target_path"]).exists()
    finally:
        a.close(); b.close(); engine_a.dispose(); engine_b.dispose()


def test_group_completion_events_travel_with_project(tmp_path):
    from backend.infrastructure.storage.models_group_completion import GroupCompletionEventModel
    a, _, engine_a = app_at(tmp_path / "A")
    b, _, engine_b = app_at(tmp_path / "B")
    try:
        member = prepared_manual_project(a, tmp_path / "old")
        with Session(engine_a) as session:
            session.add(GroupCompletionEventModel(event_id="event-1", project_id=member["project_id"],
                group_lineage="group-1", group_key="g1", matrix_revision=1, signature="a" * 64, completed=True,
                completed_at="2026-10-28T00:00:00Z", recorded_at="2026-10-28T00:00:00Z"))
            session.commit()
        info = a.get(f"/api/project-library/members/{member['member_id']}/portable").json()
        prepared = a.post(f"/api/project-library/members/{member['member_id']}/prepare", json={"fingerprint": info["fingerprint"]})
        assert prepared.status_code == 200, prepared.text
        plan = b.post("/api/project-library/restore-plan", json={"folder": member["folder_path"]}).json()
        assert plan["status"] == "ready", plan
        confirm(b, "/api/project-library/restore", plan)
        with Session(engine_b) as session:
            event = session.get(GroupCompletionEventModel, "event-1")
            assert event is not None and event.completed and event.signature == "a" * 64
    finally:
        engine_a.dispose()
        engine_b.dispose()


def test_legacy_portable_snapshot_without_group_completion_restores(tmp_path):
    from backend.infrastructure.project_library.portability import fingerprint
    a, _, engine_a = app_at(tmp_path / "A")
    b, _, engine_b = app_at(tmp_path / "B")
    try:
        member = prepared_manual_project(a, tmp_path / "old")
        folder = Path(member["folder_path"])
        packet_path = folder / ".connlab/portable/snapshot.json"
        packet = json.loads(packet_path.read_text(encoding="utf-8"))
        body = packet["body"]
        body["rows"].pop("group_completion_events")
        content = {key: body[key] for key in ("project_id", "ltr_number", "title", "rows", "files", "assets", "operation_history")}
        body["source_signature"] = fingerprint(content)
        packet["checksum"] = fingerprint(body)
        packet_path.write_text(json.dumps(packet), encoding="utf-8")
        plan = b.post("/api/project-library/restore-plan", json={"folder": str(folder)}).json()
        assert plan["status"] == "ready", plan
        result = confirm(b, "/api/project-library/restore", plan)
        assert result["member"]["project_id"] == member["project_id"]
        repeated = b.post("/api/project-library/restore-plan", json={"folder": str(folder)}).json()
        assert repeated["status"] == "already_added", repeated
    finally:
        engine_a.dispose()
        engine_b.dispose()


def reject_unselected_restore(client, review):
    assert review["status"] == "conflict", review
    assert review["blockers"] == [], review
    assert review["warnings"], review
    assert review["resolution_options"], review
    response = client.post("/api/project-library/restore", json={"plan_id": review["plan_id"], "fingerprint": review["fingerprint"]})
    assert response.status_code == 409, response.text


@pytest.mark.parametrize("managed", [True, False])
def test_explicit_replacement_preserves_target_recovery_copy_and_binds_revision(tmp_path, managed):
    a, _, ea = app_at(tmp_path / "A")
    b, settings_b, eb = app_at(tmp_path / "B", managed=managed)
    member_a = prepared_manual_project(a, tmp_path / "old")
    plan = b.post("/api/project-library/restore-plan", json={"folder": member_a["folder_path"]}).json()
    member_b = confirm(b, "/api/project-library/restore", plan)["member"]
    (Path(member_b["folder_path"]) / "target-only.txt").write_text("retain target evidence", encoding="utf-8")
    with Session(eb) as session:
        session.get(ProjectModel, member_b["project_id"]).product_name = "Target revision"
        session.commit()
    review = b.post("/api/project-library/restore-plan", json={"folder": member_a["folder_path"]}).json()
    assert review["status"] == "conflict"
    assert "replace_working_copy" in review["resolution_options"]
    reject_unselected_restore(b, review)
    assert (Path(member_b["folder_path"]) / "target-only.txt").read_text() == "retain target evidence"
    with Session(eb) as session:
        assert session.get(ProjectModel, member_b["project_id"]).product_name == "Target revision"
    plan = b.post("/api/project-library/restore-plan", json={"folder": member_a["folder_path"], "resolution": "replace_working_copy"}).json()
    assert plan["status"] == "ready", plan
    assert plan["backup_path"]
    response = b.post("/api/project-library/restore", json={"plan_id": plan["plan_id"], "fingerprint": plan["fingerprint"], "resolution": "replace_working_copy"})
    assert response.status_code == 200, response.text
    restored = response.json()
    assert restored["member"]["member_id"] == member_b["member_id"]
    assert restored["member"]["can_continue"] is True
    backup = b.get("/api/project-library/members/" + restored["backup_member_id"]).json()
    expected_recovery = settings_b.library_root / "Backups/Project Recovery" if managed else settings_b.data_dir / "project_library/recovery"
    assert Path(backup["folder_path"]).is_relative_to(expected_recovery)
    assert backup["is_history"] is True
    assert backup["can_continue"] is False
    assert (Path(backup["folder_path"]) / "target-only.txt").read_text() == "retain target evidence"
    packet = json.loads((Path(backup["folder_path"]) / ".connlab/portable/snapshot.json").read_text())
    assert packet["body"]["rows"]["projects"][0]["product_name"] == "Target revision"
    assert not (Path(restored["member"]["folder_path"]) / "target-only.txt").exists()
    ea.dispose()
    eb.dispose()


def test_same_ltr_different_identity_can_be_retained_as_read_only_history(tmp_path):
    from backend.infrastructure.storage.models import LtrRecordModel
    a, _, ea = app_at(tmp_path / "A")
    b, _, eb = app_at(tmp_path / "B")
    member_a = prepared_manual_project(a, tmp_path / "old")
    with Session(eb) as session:
        session.add(ProjectModel(project_id="target-owner", product_name="Existing owner", requestor="Lab", status="draft"))
        session.flush()
        session.add(LtrRecordModel(ltr_id="target-ltr", project_id="target-owner", ltr_number="DL-2026-09-032", status="registered", is_current_owner=True))
        session.commit()
    review = b.post("/api/project-library/restore-plan", json={"folder": member_a["folder_path"]}).json()
    assert "keep_as_history" in review["resolution_options"]
    reject_unselected_restore(b, review)
    with Session(eb) as session:
        assert session.get(ProjectModel, member_a["project_id"]) is None
        assert session.get(LtrRecordModel, "target-ltr").is_current_owner is True
    plan = b.post("/api/project-library/restore-plan", json={"folder": member_a["folder_path"], "resolution": "keep_as_history"}).json()
    assert plan["status"] == "ready", plan
    response = b.post("/api/project-library/restore", json={"plan_id": plan["plan_id"], "fingerprint": plan["fingerprint"], "resolution": "keep_as_history"})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "history_kept"
    assert result["member"]["is_history"] is True
    assert result["member"]["can_continue"] is False
    assert b.get(f"/api/project-library/members/{result['member']['member_id']}/takeover-plan").json()["blockers"]
    assert b.get("/api/project-library/members", params={"q":"History"}).json()["items"]
    with Session(eb) as session:
        assert session.get(LtrRecordModel, "target-ltr").is_current_owner is True
        assert session.get(ProjectModel, member_a["project_id"]).registry_state == "history"
    ea.dispose()
    eb.dispose()


def test_nonempty_complete_graph_restores_and_continues_without_computer_a(tmp_path, monkeypatch):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from backend.infrastructure.project_library.ownership import TABLES
    from sqlalchemy import text
    with seed_full_graph(tmp_path / "A") as graph:
        lab = graph["lab"]
        client_a = lab["client"]
        source = str(graph["project_root"])
        plan = client_a.post("/api/project-library/intake-plan", json={"sources": [source]}).json()
        batch = client_a.post("/api/project-library/batches", json={**{k: plan[k] for k in ("plan_id", "fingerprint")}, "selected": [source]}).json()
        assert batch["counts"]["added"] == 1, batch
        member_id = batch["items"][0]["member_id"]
        member = client_a.get(f"/api/project-library/members/{member_id}").json()
        assert member["project_id"] == "P1"
        info = client_a.get(f"/api/project-library/members/{member_id}/portable").json()
        assert not info["blockers"], info
        prepared = client_a.post(f"/api/project-library/members/{member_id}/prepare", json={"fingerprint": info["fingerprint"]})
        assert prepared.status_code == 200, prepared.text
        assert prepared.json()["status"] == "current", prepared.text
        assert prepared.json()["record_counts"]["precheck_issues"] == 10
        travelling = tmp_path / "USB" / Path(source).name
        shutil.copytree(source, travelling)
        expected = graph["expected"]
        packet = json.loads((travelling / ".connlab/portable/snapshot.json").read_text(encoding="utf-8"))
        assert any(row["kind"] == "ltr_decision_evidence" and row["replayable"] is False for row in packet["body"]["operation_history"])
        assert any(row["kind"] == "frozen_field_decision_evidence" for row in packet["body"]["operation_history"])
        assert b"QA-NOT-A-REAL-CREDENTIAL-SENTINEL" not in (travelling / ".connlab/portable/snapshot.json").read_bytes()
    # The original DB and every original project/data path are now unavailable.
    for name in ("connlab.sqlite3", "data", "projects", "templates", "public-history"):
        old = tmp_path / "A" / name
        if old.exists():
            old.rename(old.with_name(name + "-unavailable"))
    b, settings_b, engine_b = app_at(tmp_path / "B")
    try:
        plan = b.post("/api/project-library/restore-plan", json={"folder": str(travelling)}).json()
        assert plan["status"] == "ready", plan
        restored = confirm(b, "/api/project-library/restore", plan)
        member_b = restored["member"]
        assert member_b["snapshot_status"] == "current", member_b
        assert member_b["can_continue"] is True
        manifest = json.loads((Path(member_b["folder_path"]) / ".connlab/manifest.json").read_text(encoding="utf-8"))
        assert Path(manifest["official_project_folder_path"]).is_relative_to(settings_b.projects_dir)
        assert Path(manifest["official_project_folder_path"]).is_dir()
        with Session(engine_b) as session:
            for name in TABLES:
                count = session.execute(text('SELECT COUNT(*) FROM "' + name + '"')).scalar_one()
                assert count == expected["source_counts"].get(name, 0), name
            assert session.execute(text("PRAGMA foreign_key_check")).all() == []
            assert session.execute(text("SELECT auto_sync_enabled FROM project_public_folder_workflow_states")).scalar_one() == 0
            assert session.execute(text("SELECT COUNT(*) FROM ltr_duplicate_resolution_tokens")).scalar_one() == 0
            assert session.execute(text("SELECT COUNT(*) FROM frozen_field_revision_requests")).scalar_one() == 0
            assert str(session.execute(text("SELECT duration_value FROM confirmed_matrix_duration_authorities WHERE confirmed_duration_authority_id='duration-1'")).scalar_one()) == "2.5"
        from backend.api import dependencies as deps
        from backend.api.main import app
        def session_b():
            with Session(engine_b) as session:
                yield session
                session.commit()
        app.dependency_overrides[deps.get_settings] = lambda: settings_b
        app.dependency_overrides[deps.get_session] = session_b
        from backend.api.project_folder_generation_composition import ProjectFolderGenerationRunner
        runner = ProjectFolderGenerationRunner(lambda: Session(engine_b), settings_b)
        app.dependency_overrides[deps.get_project_folder_generation_service] = runner.service
        client = TestClient(app, raise_server_exceptions=False)
        try:
            outputs = client.get("/api/projects/P1/output-records/status")
            assert outputs.status_code == 200, outputs.text
            assert any(item["output_kind"] == "test_record_form" for item in outputs.json()["items"])
            folder_preview = client.get("/api/projects/P1/project-folder/generation/preview?intent=create")
            assert folder_preview.status_code == 200, folder_preview.text
            check = client.get("/api/projects/P1/prechecks/latest")
            assert check.status_code == 200, check.text
            assert len(check.json()["issues"]) == 10
            workspace = client.get("/api/projects/P1/report-workspace")
            assert workspace.status_code == 200, workspace.text
            assert len(workspace.json()["datasets"]) == 2
            assert "0.198" in workspace.text and "qa-report-1" in workspace.text
            authoring = client.get("/api/projects/P1/report-workspace/authoring")
            assert authoring.status_code == 200, authoring.text
            state = authoring.json()
            assert state["mappings"][0]["value"] == "2026-06-18"
            assert "0.0855111111111111" in authoring.text
            source = state["sources"][0]
            download = client.get(source["download_url"])
            assert download.status_code == 200
            assert sha256(download.content).hexdigest() == source["sha256"]
            updated = client.post("/api/projects/P1/report-workspace/authoring/mappings", json={"expected_revision": state["revision"], "source_id": source["source_id"], "kind":"specification", "source_locator":"Target computer review", "label":"Restored specification", "value":"Review continued on B", "status":"confirmed", "reason":"Operator continued from carried history"})
            assert updated.status_code == 200, updated.text
            assert updated.json()["revision"] == state["revision"] + 1
            draft = client.get("/api/projects/P1/matrix-drafts/qa-working-matrix")
            assert draft.status_code == 200, draft.text
            data = draft.json()
            data["rows"][1]["requirement"] = "Continue Matrix review on B"
            data["duration_authorities"] = [{**row, "draft_duration_authority_id": row["duration_authority_id"], "draft_group_id": row["group_id"], "draft_row_id": row["row_id"]} for row in data["duration_authorities"]]
            saved = client.put("/api/projects/P1/matrix-drafts/qa-working-matrix", json=data)
            assert saved.status_code == 200, saved.text
            assert saved.json()["rows"][1]["requirement"] == "Continue Matrix review on B"
        finally:
            runner.pool.shutdown(wait=True)
            app.dependency_overrides.clear()
    finally:
        engine_b.dispose()


@pytest.mark.parametrize("gap", ["non_fk_report_base", "json_source", "superseded_fk"])
def test_portable_preflight_blocks_missing_history_without_repairing_source(tmp_path, gap):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from sqlalchemy import text
    with seed_full_graph(tmp_path / "source") as graph:
        lab = graph["lab"]
        source = str(graph["project_root"])
        plan = lab["client"].post("/api/project-library/intake-plan", json={"sources": [source]}).json()
        batch = lab["client"].post("/api/project-library/batches", json={**{k:plan[k] for k in ("plan_id","fingerprint")}, "selected":[source]}).json()
        member_id = batch["items"][0]["member_id"]
        with lab["factory"]() as session:
            if gap == "non_fk_report_base":
                session.execute(text("UPDATE report_draft_revisions SET base_report_revision_id='missing-history' WHERE report_revision_id='qa-report-2'"))
            elif gap == "json_source":
                record = session.execute(text("SELECT id,payload_json FROM report_authoring_revisions ORDER BY revision DESC LIMIT 1")).first()
                payload = json.loads(record[1])
                payload["base_source_id"] = "missing-source"
                session.execute(text("UPDATE report_authoring_revisions SET payload_json=:payload WHERE id=:id"), {"payload":json.dumps(payload),"id":record[0]})
            else:
                session.execute(text("UPDATE project_matrix_draft_rows SET source_row_snapshot_id='missing-superseded-source' WHERE draft_row_id='pmdr-1'"))
            session.commit()
        status = lab["client"].get(f"/api/project-library/members/{member_id}/portable").json()
        assert status["status"] == "unavailable", status
        assert status["blockers"], status
        assert not (Path(source) / ".connlab/portable/snapshot.json").exists()
        with lab["factory"]() as session:
            if gap == "non_fk_report_base":
                assert session.execute(text("SELECT base_report_revision_id FROM report_draft_revisions WHERE report_revision_id='qa-report-2'")).scalar_one() == "missing-history"
            elif gap == "json_source":
                assert "missing-source" in session.execute(text("SELECT payload_json FROM report_authoring_revisions ORDER BY revision DESC LIMIT 1")).scalar_one()
            else:
                assert session.execute(text("SELECT COUNT(*) FROM project_matrix_draft_rows WHERE source_row_snapshot_id='missing-superseded-source'")).scalar_one() > 0


@pytest.mark.parametrize("managed", [True, False])
def test_replacement_failure_rolls_back_and_same_review_can_resume(tmp_path, monkeypatch, managed):
    a, _, ea = app_at(tmp_path / "A")
    b, _, eb = app_at(tmp_path / "B", managed=managed)
    member_a = prepared_manual_project(a, tmp_path / "old")
    plan = b.post("/api/project-library/restore-plan", json={"folder":member_a["folder_path"]}).json()
    member_b = confirm(b, "/api/project-library/restore", plan)["member"]
    folder = Path(member_b["folder_path"])
    (folder / "B-only.txt").write_text("target remains recoverable", encoding="utf-8")
    with Session(eb) as session:
        session.get(ProjectModel, member_b["project_id"]).product_name = "B historical test"
        session.commit()
    plan = b.post("/api/project-library/restore-plan", json={"folder":member_a["folder_path"], "resolution":"replace_working_copy"}).json()
    original = ProjectPortability._insert
    def disk_failure(self, *args, **kwargs):
        if kwargs.get("replace_rows") is not None:
            raise OSError("simulated publication failure")
        return original(self, *args, **kwargs)
    monkeypatch.setattr(ProjectPortability, "_insert", disk_failure)
    request = {"plan_id":plan["plan_id"], "fingerprint":plan["fingerprint"], "resolution":"replace_working_copy"}
    failed = b.post("/api/project-library/restore", json=request)
    assert failed.status_code == 422, failed.text
    assert (folder / "B-only.txt").read_text() == "target remains recoverable"
    with Session(eb) as session:
        assert session.get(ProjectModel, member_b["project_id"]).product_name == "B historical test"
    monkeypatch.setattr(ProjectPortability, "_insert", original)
    resumed = b.post("/api/project-library/restore", json=request)
    assert resumed.status_code == 200, resumed.text
    backup = b.get("/api/project-library/members/" + resumed.json()["backup_member_id"]).json()
    assert backup["is_history"] is True
    assert backup["snapshot_status"] == "current"
    assert (Path(backup["folder_path"]) / "B-only.txt").read_text() == "target remains recoverable"
    # The retained target packet independently restores its old SQL and files on C.
    c, _, ec = app_at(tmp_path / "C")
    review_c = c.post("/api/project-library/restore-plan", json={"folder":backup["folder_path"]}).json()
    recovered = confirm(c, "/api/project-library/restore", review_c)
    assert (Path(recovered["member"]["folder_path"]) / "B-only.txt").read_text() == "target remains recoverable"
    with Session(ec) as session:
        assert session.get(ProjectModel, member_b["project_id"]).product_name == "B historical test"
    for engine in (ea, eb, ec):
        engine.dispose()


def test_stale_coverage_reports_saved_snapshot_not_unprepared_changes(tmp_path):
    client, _, engine = app_at(tmp_path / "library")
    member = prepared_manual_project(client, tmp_path / "old")
    url = f"/api/project-library/members/{member['member_id']}/portable"
    saved = client.get(url).json()
    (Path(member["folder_path"]) / "unprepared.txt").write_text("new notes",encoding="utf-8")
    with Session(engine) as session:
        session.add(ReportAuthoringRevisionModel(project_id=member["project_id"], revision=1, payload_json=json.dumps({"sources":[],"mappings":[]})))
        session.commit()
    stale = client.get(url).json()
    assert stale["status"] == "stale"
    assert stale["created_at"] == saved["created_at"]
    assert stale["file_count"] == saved["file_count"]
    assert stale["record_counts"]["report_authoring_revisions"] == 0
    assert stale["warnings"]
    refreshed = client.post(url.rsplit("/",1)[0]+"/prepare", json={"fingerprint":stale["fingerprint"]}).json()
    assert refreshed["status"] == "current"
    assert refreshed["file_count"] == saved["file_count"]+1
    assert refreshed["record_counts"]["report_authoring_revisions"] == 1
    engine.dispose()


def test_self_consistent_portable_package_never_grants_target_public_write_permission(tmp_path):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from backend.infrastructure.project_library.files import fingerprint
    from sqlalchemy import text
    with seed_full_graph(tmp_path / "A") as graph:
        client=graph["lab"]["client"]
        source=str(graph["project_root"])
        plan=client.post("/api/project-library/intake-plan",json={"sources":[source]}).json()
        batch=client.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[source]}).json()
        member_id=batch["items"][0]["member_id"]
        status=client.get(f"/api/project-library/members/{member_id}/portable").json()
        response=client.post(f"/api/project-library/members/{member_id}/prepare",json={"fingerprint":status["fingerprint"]})
        assert response.status_code==200,response.text
        travelling=tmp_path/"USB"/Path(source).name
        shutil.copytree(source,travelling)
    packet_path=travelling/".connlab/portable/snapshot.json"
    packet=json.loads(packet_path.read_text(encoding="utf-8"))
    body=packet["body"]
    body["rows"]["project_public_folder_workflow_states"][0]["auto_sync_enabled"]=True
    content={key:body[key] for key in ("project_id","ltr_number","title","rows","files","assets","operation_history")}
    body["source_signature"]=fingerprint(content)
    packet_path.write_text(json.dumps({"body":body,"checksum":fingerprint(body)}),encoding="utf-8")
    b,_,engine=app_at(tmp_path/"B")
    plan=b.post("/api/project-library/restore-plan",json={"folder":str(travelling)}).json()
    assert plan["status"]=="ready",plan
    confirm(b,"/api/project-library/restore",plan)
    with Session(engine) as session:
        assert session.execute(text("SELECT auto_sync_enabled FROM project_public_folder_workflow_states")).scalar_one()==0
    engine.dispose()


def test_carried_test_data_cannot_be_bypassed_by_plain_intake_or_empty_takeover(tmp_path):
    a, _, ea = app_at(tmp_path/"A")
    b, _, eb = app_at(tmp_path/"B")
    source=prepared_manual_project(a,tmp_path/"old")
    for corrupt in (False,True):
        if corrupt:
            (Path(source["folder_path"])/".connlab/portable/snapshot.json").write_text("{}",encoding="utf-8")
        plan=b.post("/api/project-library/intake-plan",json={"sources":[source["folder_path"]]}).json()
        assert plan["items"][0]["status"]=="conflict",plan
        assert plan["items"][0]["can_keep_separate"] is False
        assert "restore" in " ".join(plan["items"][0]["warnings"]).lower()
        result=b.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[source["folder_path"]],"keep_separate":[source["folder_path"]]}).json()
        assert result["counts"]["added"]==0,result
    raw=tmp_path/"Raw old project"
    raw.mkdir()
    (raw/"note.txt").write_text("raw note",encoding="utf-8")
    plan=b.post("/api/project-library/intake-plan",json={"sources":[str(raw)]}).json()
    result=b.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[str(raw)]}).json()
    member=b.get("/api/project-library/members/"+result["items"][0]["member_id"]).json()
    packet=Path(member["folder_path"])/".connlab/portable/snapshot.json"
    packet.parent.mkdir(parents=True)
    packet.write_text("{}",encoding="utf-8")
    review=b.get(f"/api/project-library/members/{member['member_id']}/takeover-plan").json()
    assert review["blockers"]
    assert b.post(f"/api/project-library/members/{member['member_id']}/takeover",json={k:review[k] for k in ("plan_id","fingerprint")}).status_code==409
    with Session(eb) as session:
        assert list(session.scalars(select(ProjectModel)))==[]
    ea.dispose();eb.dispose()


@pytest.mark.parametrize("damage", ["version", "shape", "checksum", "unbound_path", "duplicate_record"])
def test_portable_inspection_rejects_unsupported_or_damaged_packet_before_target_write(tmp_path,damage):
    from backend.infrastructure.project_library.files import fingerprint
    a,_,ea=app_at(tmp_path/"A")
    b,settings_b,eb=app_at(tmp_path/"B")
    member=prepared_manual_project(a,tmp_path/"old")
    path=Path(member["folder_path"])/".connlab/portable/snapshot.json"
    packet=json.loads(path.read_text(encoding="utf-8"))
    body=packet["body"]
    if damage=="version":
        body["format_version"]=99
    elif damage=="shape":
        body["assets"]=[{"blob":"broken"}]
    elif damage=="unbound_path":
        body["rows"]["project_folder_records"]=[{"folder_id":"bad-folder", "project_id":member["project_id"], "folder_path":str(tmp_path/"untrusted-outside"), "created_on":"2026-10-27"}]
    elif damage=="duplicate_record":
        body["rows"]["projects"].append(dict(body["rows"]["projects"][0]))
    else:
        body["title"]="checksum changed"
    if damage!="checksum":
        content={key:body[key] for key in ("project_id","ltr_number","title","rows","files","assets","operation_history")}
        body["source_signature"]=fingerprint(content)
        packet["checksum"]=fingerprint(body)
    path.write_text(json.dumps(packet),encoding="utf-8")
    review=b.post("/api/project-library/restore-plan",json={"folder":member["folder_path"]})
    assert review.status_code==200,review.text
    assert review.json()["status"]=="unavailable",review.text
    assert review.json()["blockers"]
    assert not list(settings_b.projects_dir.iterdir())
    with Session(eb) as session:
        assert list(session.scalars(select(ProjectModel)))==[]
    ea.dispose();eb.dispose()


def test_replacement_blocks_cross_project_non_fk_references(tmp_path):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from backend.infrastructure.storage.models_result_dataset import ResultDatasetRevisionModel
    from dataclasses import replace
    from tests.unit.test_result_dataset_repository import _dataset
    from backend.infrastructure.storage.repositories.result_dataset import ResultDatasetRepository
    with seed_full_graph(tmp_path/"A") as graph:
        a=graph["lab"]["client"];source=str(graph["project_root"])
        plan=a.post("/api/project-library/intake-plan",json={"sources":[source]}).json()
        batch=a.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[source]}).json()
        member_id=batch["items"][0]["member_id"]
        info=a.get(f"/api/project-library/members/{member_id}/portable").json()
        assert a.post(f"/api/project-library/members/{member_id}/prepare",json={"fingerprint":info["fingerprint"]}).status_code==200
        travelling=tmp_path/"USB"/Path(source).name
        shutil.copytree(source,travelling)
    b,_,engine=app_at(tmp_path/"B")
    plan=b.post("/api/project-library/restore-plan",json={"folder":str(travelling)}).json()
    member=confirm(b,"/api/project-library/restore",plan)["member"]
    (Path(member["folder_path"])/"B-change.txt").write_text("target revision",encoding="utf-8")
    with Session(engine) as session:
        session.add(ProjectModel(project_id="other-project",product_name="Unrelated",requestor="QA",status="draft"))
        session.flush()
        dataset=replace(_dataset("other-dataset",1),project_id="other-project",confirmed_matrix_id="cmv-1",confirmed_matrix_revision=2)
        ResultDatasetRepository(session).create_dataset(dataset)
        session.commit()
    review=b.post("/api/project-library/restore-plan",json={"folder":str(travelling),"resolution":"replace_working_copy"}).json()
    assert review["status"]=="unavailable",review
    assert any("another project" in message.lower() for message in review["blockers"])
    with Session(engine) as session:
        assert session.get(ResultDatasetRevisionModel,"other-dataset") is not None
    engine.dispose()


def test_daily_title_and_location_remain_available_with_preserved_incomplete_history(tmp_path):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from sqlalchemy import text
    with seed_full_graph(tmp_path/"source") as graph:
        client=graph["lab"]["client"];source=str(graph["project_root"])
        plan=client.post("/api/project-library/intake-plan",json={"sources":[source]}).json()
        batch=client.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[source]}).json()
        member_id=batch["items"][0]["member_id"]
        with graph["lab"]["factory"]() as session:
            session.execute(text("UPDATE project_matrix_draft_rows SET source_row_snapshot_id='missing-historic-source' WHERE draft_row_id='pmdr-1'"))
            session.commit()
        blocked=client.get(f"/api/project-library/members/{member_id}/portable").json()
        assert blocked["blockers"]
        plan=client.post(f"/api/project-library/members/{member_id}/title-plan",json={"title":"Changed descriptive title"}).json()
        changed=confirm(client,f"/api/project-library/members/{member_id}/title",plan)
        assert changed["title"]=="Changed descriptive title"
        root=Path(changed["folder_path"])
        moved=root.with_name(root.name+" relocated")
        root.rename(moved)
        plan=client.post(f"/api/project-library/members/{member_id}/location-plan",json={"folder":str(moved)}).json()
        located=confirm(client,f"/api/project-library/members/{member_id}/location",plan)
        assert located["folder_path"]==str(moved)
        with graph["lab"]["factory"]() as session:
            assert session.execute(text("SELECT source_row_snapshot_id FROM project_matrix_draft_rows WHERE draft_row_id='pmdr-1'")).scalar_one()=="missing-historic-source"
        assert client.get(f"/api/project-library/members/{member_id}/portable").json()["blockers"]


@pytest.mark.parametrize("foreign_token", ["@data/local_identity.sqlite3", "@data/report_authoring_assets/" + "a"*64 + "/evidence.docx"])
def test_receiver_rejects_self_consistent_assets_outside_own_project_partition(tmp_path,foreign_token):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from backend.infrastructure.project_library.files import fingerprint
    from backend.infrastructure.project_library.ownership import transform_paths
    with seed_full_graph(tmp_path/"A") as graph:
        a=graph["lab"]["client"];source=str(graph["project_root"])
        plan=a.post("/api/project-library/intake-plan",json={"sources":[source]}).json()
        batch=a.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[source]}).json()
        member_id=batch["items"][0]["member_id"]
        info=a.get(f"/api/project-library/members/{member_id}/portable").json()
        assert a.post(f"/api/project-library/members/{member_id}/prepare",json={"fingerprint":info["fingerprint"]}).status_code==200
        travelling=tmp_path/"USB"/Path(source).name
        shutil.copytree(source,travelling)
    path=travelling/".connlab/portable/snapshot.json"
    packet=json.loads(path.read_text(encoding="utf-8"));body=packet["body"]
    original=body["assets"][0]["token"]
    body["assets"][0]["token"]=foreign_token
    body["rows"]=transform_paths(body["rows"],lambda value:foreign_token if value==original else value)
    content={key:body[key] for key in ("project_id","ltr_number","title","rows","files","assets","operation_history")}
    body["source_signature"]=fingerprint(content)
    path.write_text(json.dumps({"body":body,"checksum":fingerprint(body)}),encoding="utf-8")
    b,settings,engine=app_at(tmp_path/"B")
    sentinel=settings.data_dir/foreign_token[len("@data/"):]
    sentinel.parent.mkdir(parents=True,exist_ok=True)
    sentinel.write_bytes(b"target private or unrelated project sentinel")
    review=b.post("/api/project-library/restore-plan",json={"folder":str(travelling)}).json()
    assert review["status"]=="unavailable",review
    assert any("partition" in message.lower() for message in review["blockers"]),review
    assert sentinel.read_bytes()==b"target private or unrelated project sentinel"
    assert not list(settings.projects_dir.iterdir())
    engine.dispose()


@pytest.mark.parametrize("crash_after_commit,fresh_review,change", [
    (False,False,None),(False,True,None),(True,False,None),(True,True,None),
    (False,False,"file"),(False,True,"file"),(False,True,"database"),(False,True,"source"),
    (False,True,"parcel_only"),
])
def test_process_crash_during_working_copy_switch_recovers_review(tmp_path,crash_after_commit,fresh_review,change):
    import subprocess
    import sys
    a,_,ea=app_at(tmp_path/"A");b,settings,eb=app_at(tmp_path/"B")
    source=prepared_manual_project(a,tmp_path/"old")
    plan=b.post("/api/project-library/restore-plan",json={"folder":source["folder_path"]}).json()
    target=confirm(b,"/api/project-library/restore",plan)["member"]
    (Path(target["folder_path"])/"target-old.txt").write_text("old evidence",encoding="utf-8")
    with Session(eb) as session:
        session.get(ProjectModel,target["project_id"]).product_name="Original complete B SQL"
        session.commit()
    from backend.infrastructure.files.generation_journal import GenerationJournal
    generation=GenerationJournal(settings.data_dir/"project_folder_generation")
    previous=generation.create(target["project_id"],"update_in_place","previous-generation-context")
    previous["status"]="completed";generation.save(previous)
    plan=b.post("/api/project-library/restore-plan",json={"folder":source["folder_path"],"resolution":"replace_working_copy"}).json()
    assert plan["status"]=="ready",plan
    script = """
import os
from pathlib import Path
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from backend.shared.config import Settings
from backend.infrastructure.storage.repositories.project_library import ProjectLibraryRepository
from backend.infrastructure.project_library.files import LibraryFiles
from backend.infrastructure.project_library.portability import ProjectPortability
root=Path(ROOT)
settings=Settings(root/'Data',root/'Projects',root/'Templates',root/'Data'/'connlab.sqlite3',library_root=root)
engine=create_engine('sqlite:///'+str(settings.database_path))
from backend.infrastructure.storage.database import init_db
init_db(engine)
with Session(engine) as session:
    store=ProjectLibraryRepository(session)
    operation=ProjectPortability(store,LibraryFiles(settings),settings)
    original=operation._insert
    def crash(*args,**kwargs):
        if AFTER:
            original(*args,**kwargs)
        os._exit(42)
    operation._insert=crash
    operation.restore(PLAN,EXPECTED,'replace_working_copy')
""".replace("ROOT",repr(str(settings.library_root))).replace("AFTER",repr(crash_after_commit)).replace("PLAN",repr(plan["plan_id"])).replace("EXPECTED",repr(plan["fingerprint"]))
    result=subprocess.run([sys.executable,"-c",script],capture_output=True,text=True,timeout=20)
    assert result.returncode==42,result.stderr
    assert b.get("/api/project-library/members/"+target["member_id"]).json()["can_continue"] is False
    _assert_replacement_fences_public_writers(settings,eb,target["project_id"],plan["plan_id"])
    histories=[item for item in b.get("/api/project-library/members").json()["items"] if item["is_history"]]
    original=next(item for item in histories if (Path(item["folder_path"])/"target-old.txt").exists())
    packet=json.loads((Path(original["folder_path"])/".connlab/portable/snapshot.json").read_text(encoding="utf-8"))
    assert packet["body"]["rows"]["projects"][0]["product_name"]=="Original complete B SQL"
    review_folder=source["folder_path"]
    if change=="parcel_only":
        source_folder=Path(source["folder_path"])
        source_folder.rename(source_folder.with_name(source_folder.name+" unavailable-source"))
        review_folder=target["folder_path"]
    if change=="file":
        (Path(target["folder_path"])/"added-after-crash.txt").write_text("retain new operator note",encoding="utf-8")
    elif change=="database":
        with Session(eb) as session:
            session.get(ProjectModel,target["project_id"]).product_name="External edit after crash"
            session.commit()
    elif change=="source":
        (Path(source["folder_path"])/"new-source-note.txt").write_text("new source revision",encoding="utf-8")
        info=a.get(f"/api/project-library/members/{source['member_id']}/portable").json()
        assert a.post(f"/api/project-library/members/{source['member_id']}/prepare",json={"fingerprint":info["fingerprint"]}).status_code==200
    if change in {"file","database","source"}:
        refused=b.post("/api/project-library/restore",json={"plan_id":plan["plan_id"],"fingerprint":plan["fingerprint"],"resolution":"replace_working_copy"})
        assert refused.status_code==409,refused.text
        fresh=b.post("/api/project-library/restore-plan",json={"folder":review_folder}).json()
        assert fresh["status"]=="unavailable",fresh
        assert fresh["blockers"],fresh
        if change=="file":
            assert (Path(target["folder_path"])/"added-after-crash.txt").read_text()=="retain new operator note"
        assert (Path(original["folder_path"])/"target-old.txt").read_text()=="old evidence"
        assert LibraryFiles(settings).pending_replacement_gate(target["project_id"])
        ea.dispose();eb.dispose()
        return
    if fresh_review:
        fresh=b.post("/api/project-library/restore-plan",json={"folder":review_folder}).json()
        assert fresh["status"]=="conflict",fresh
        scene_before = (Path(target["folder_path"])/"result.txt").read_bytes()
        packet_before = (Path(original["folder_path"])/".connlab/portable/snapshot.json").read_bytes()
        reject_unselected_restore(b, fresh)
        assert (Path(target["folder_path"])/"result.txt").read_bytes() == scene_before
        assert (Path(original["folder_path"])/".connlab/portable/snapshot.json").read_bytes() == packet_before
        assert LibraryFiles(settings).pending_replacement_gate(target["project_id"])
        renewed=b.post("/api/project-library/restore-plan",json={"folder":review_folder,"resolution":"replace_working_copy"}).json()
        assert renewed["backup_path"]==plan["backup_path"]
        assert renewed["plan_id"]!=plan["plan_id"]
        plan=renewed
    response=b.post("/api/project-library/restore",json={"plan_id":plan["plan_id"],"fingerprint":plan["fingerprint"],"resolution":"replace_working_copy"})
    assert response.status_code==200,response.text
    restored=response.json()
    assert restored["member"]["can_continue"] is True
    backup=b.get("/api/project-library/members/"+restored["backup_member_id"]).json()
    assert (Path(backup["folder_path"])/"target-old.txt").read_text()=="old evidence"
    assert not (Path(restored["member"]["folder_path"])/"target-old.txt").exists()
    assert LibraryFiles(settings).pending_replacement_gate(target["project_id"]) is None
    assert generation.read(target["project_id"])==previous
    ea.dispose();eb.dispose()


def test_working_copy_switch_respects_existing_project_writer_slot(tmp_path):
    a,_,ea=app_at(tmp_path/"A");b,settings,eb=app_at(tmp_path/"B")
    source=prepared_manual_project(a,tmp_path/"old")
    plan=b.post("/api/project-library/restore-plan",json={"folder":source["folder_path"]}).json()
    member=confirm(b,"/api/project-library/restore",plan)["member"]
    (Path(member["folder_path"])/"target-note.txt").write_text("working data",encoding="utf-8")
    plan=b.post("/api/project-library/restore-plan",json={"folder":source["folder_path"],"resolution":"replace_working_copy"}).json()
    request={"plan_id":plan["plan_id"],"fingerprint":plan["fingerprint"],"resolution":"replace_working_copy"}
    with LibraryFiles(settings).project_lock(member["project_id"]):
        busy=b.post("/api/project-library/restore",json=request)
        assert busy.status_code==409,busy.text
        assert (Path(member["folder_path"])/"target-note.txt").read_text()=="working data"
    assert b.post("/api/project-library/restore",json=request).status_code==200
    ea.dispose();eb.dispose()


@pytest.mark.parametrize("state",["renamed","missing","ambiguous"])
def test_manual_inner_folder_change_requires_verified_explicit_reconnection(tmp_path,state):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from sqlalchemy import text
    with seed_full_graph(tmp_path/"A") as graph:
        client=graph["lab"]["client"];root=graph["project_root"];old=graph["business_folder"]
        plan=client.post("/api/project-library/intake-plan",json={"sources":[str(root)]}).json()
        batch=client.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[str(root)]}).json()
        member_id=batch["items"][0]["member_id"]
        assert client.get("/api/project-library/members/"+member_id).json()["can_continue"] is True
        renamed=old.with_name("DL-2026-05-001 Operator revised title")
        if state=="missing":
            old.rename(root.parent/"temporarily unavailable business folder")
        else:
            old.rename(renamed)
            if state=="ambiguous":
                (root/"DL-2026-05-001 Another folder").mkdir()
        refreshed=client.get("/api/project-library/members/"+member_id).json()
        assert refreshed["can_continue"] is False,refreshed
        assert refreshed["warnings"]
        listing=client.get("/api/project-library/members").json()["items"]
        assert next(item for item in listing if item["member_id"]==member_id)["can_continue"] is False
        review=client.post(f"/api/project-library/members/{member_id}/location-plan",json={"folder":str(root)})
        if state!="renamed":
            assert review.status_code==409,review.text
            return
        assert review.status_code==200,review.text
        assert review.json()["warnings"]
        repaired=confirm(client,f"/api/project-library/members/{member_id}/location",review.json())
        assert repaired["can_continue"] is True,repaired
        assert repaired["title"]=="Operator revised title"
        manifest=json.loads((root/".connlab/manifest.json").read_text(encoding="utf-8"))
        assert manifest["official_project_folder_path"]==str(renamed)
        with graph["lab"]["factory"]() as session:
            assert session.execute(text("SELECT official_folder_path FROM project_official_workspace_records WHERE project_id='P1'")).scalar_one()==str(renamed)
            assert session.execute(text("SELECT path FROM file_assets WHERE asset_id='qa-llcr-workbook'")).scalar_one().startswith(str(renamed))
            assert session.execute(text("SELECT COUNT(*) FROM ltr_records WHERE project_id='P1'")).scalar_one()==2
        assert (renamed/"Test results"/"DL-2026-05-001 LLCR Record.xlsx").exists()


@pytest.mark.parametrize("state",["files_received","sql_committed","foreign_marker","changed_file"])
def test_fresh_restore_review_recovers_owned_receive_without_adopting_foreign_changes(tmp_path,state):
    import subprocess,sys
    a,_,ea=app_at(tmp_path/"A");b,settings,eb=app_at(tmp_path/"B")
    source=prepared_manual_project(a,tmp_path/"old")
    plan=b.post("/api/project-library/restore-plan",json={"folder":source["folder_path"]}).json()
    script="""
import os
from pathlib import Path
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from backend.shared.config import Settings
from backend.infrastructure.storage.database import init_db
from backend.infrastructure.storage.repositories.project_library import ProjectLibraryRepository
from backend.infrastructure.project_library.files import LibraryFiles
from backend.infrastructure.project_library.portability import ProjectPortability
root=Path(ROOT)
settings=Settings(root/'Data',root/'Projects',root/'Templates',root/'Data'/'connlab.sqlite3',library_root=root)
engine=create_engine('sqlite:///'+str(settings.database_path));init_db(engine)
with Session(engine) as session:
    operation=ProjectPortability(ProjectLibraryRepository(session),LibraryFiles(settings),settings)
    original=operation._insert
    def crash(*args,**kwargs):
        if COMMIT:
            original(*args,**kwargs)
        os._exit(42)
    operation._insert=crash
    operation.restore(PLAN,EXPECTED)
""".replace("ROOT",repr(str(settings.library_root))).replace("COMMIT",repr(state=="sql_committed")).replace("PLAN",repr(plan["plan_id"])).replace("EXPECTED",repr(plan["fingerprint"]))
    crashed=subprocess.run([sys.executable,"-c",script],capture_output=True,text=True,timeout=20)
    assert crashed.returncode==42,crashed.stderr
    target=Path(plan["target_path"])
    if state=="foreign_marker":
        path=target/".connlab/library.json";data=json.loads(path.read_text(encoding="utf-8"));data["member_id"]="foreign-owner";path.write_text(json.dumps(data),encoding="utf-8")
    elif state=="changed_file":
        (target/"result.txt").write_text("unreviewed target change",encoding="utf-8")
    fresh=b.post("/api/project-library/restore-plan",json={"folder":source["folder_path"]}).json()
    assert fresh["plan_id"]!=plan["plan_id"]
    if state in {"foreign_marker","changed_file"}:
        assert fresh["status"]=="conflict",fresh
        with Session(eb) as session:
            assert session.get(ProjectModel,source["project_id"]) is None
    else:
        assert fresh["status"] in {"ready","already_added"},fresh
        restored=confirm(b,"/api/project-library/restore",fresh)
        assert restored["member"]["project_id"]==source["project_id"]
        assert restored["member"]["can_continue"] is True
        assert len(b.get("/api/project-library/members").json()["items"])==1
    ea.dispose();eb.dispose()


def test_link_existing_project_takeover_honors_writer_slot_and_stale_review(tmp_path):
    from backend.infrastructure.storage.models import ProjectFolderRecordModel
    from datetime import date
    client,settings,engine=app_at(tmp_path/"library")
    source=tmp_path/"DL-2026-09-032 Existing version";source.mkdir()
    (source/"result.txt").write_text("existing evidence",encoding="utf-8")
    (source/".connlab").mkdir()
    (source/".connlab/manifest.json").write_text(json.dumps({"project_id":"existing-P1","dl_number":"DL-2026-09-032"}),encoding="utf-8")
    with Session(engine) as session:
        session.add(ProjectModel(project_id="existing-P1",product_name="Existing",requestor="Lab",status="draft",registry_state="active"))
        session.flush()
        session.add(ProjectFolderRecordModel(folder_id="existing-folder",project_id="existing-P1",folder_path=str(source),created_on=date.today()))
        session.commit()
    plan=client.post("/api/project-library/intake-plan",json={"sources":[str(source)]}).json()
    batch=client.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[str(source)]}).json()
    member_id=batch["items"][0]["member_id"]
    review=client.get(f"/api/project-library/members/{member_id}/takeover-plan").json()
    with LibraryFiles(settings).project_lock("existing-P1"):
        response=client.post(f"/api/project-library/members/{member_id}/takeover",json={k:review[k] for k in ("plan_id","fingerprint")})
        assert response.status_code==409,response.text
        with Session(engine) as session:
            assert session.get(ProjectFolderRecordModel,"existing-folder").folder_path==str(source)
    with Session(engine) as session:
        session.get(ProjectModel,"existing-P1").registry_revision+=1
        session.commit()
    assert client.post(f"/api/project-library/members/{member_id}/takeover",json={k:review[k] for k in ("plan_id","fingerprint")}).status_code==409
    review=client.get(f"/api/project-library/members/{member_id}/takeover-plan").json()
    received=client.get("/api/project-library/members/"+member_id).json()
    manifest=Path(received["folder_path"])/".connlab/manifest.json"
    manifest.write_text(json.dumps({"project_id":"foreign-project","dl_number":"DL-2026-09-032"}),encoding="utf-8")
    assert client.post(f"/api/project-library/members/{member_id}/takeover",json={k:review[k] for k in ("plan_id","fingerprint")}).status_code==409
    with Session(engine) as session:
        assert session.get(ProjectFolderRecordModel,"existing-folder").folder_path==str(source)
    manifest.write_text(json.dumps({"project_id":"existing-P1","dl_number":"DL-2026-09-032"}),encoding="utf-8")
    changed=confirm(client,f"/api/project-library/members/{member_id}/takeover",review)["member"]
    assert changed["project_id"]=="existing-P1"
    with Session(engine) as session:
        assert session.get(ProjectFolderRecordModel,"existing-folder").folder_path==changed["folder_path"]
    engine.dispose()


def test_member_list_does_not_read_project_contents_and_reports_unchecked_snapshot(tmp_path,monkeypatch):
    from backend.infrastructure.project_library import files as files_module
    client,_,engine=app_at(tmp_path/"library")
    member=prepared_manual_project(client,tmp_path/"old")
    (Path(member["folder_path"])/"daily-measurement.bin").write_bytes(bytes(1024*1024))
    original=files_module.file_hash;reads=[]
    def track(path):
        reads.append(str(path));return original(path)
    monkeypatch.setattr(files_module,"file_hash",track)
    for query in ("", "History"):
        response=client.get("/api/project-library/members",params={"q":query})
        assert response.status_code==200,response.text
        assert response.json()["items"][0]["snapshot_status"]=="not_checked"
        assert response.json()["items"][0]["snapshot_created_at"] is None
    assert reads==[],reads
    detail=client.get("/api/project-library/members/"+member["member_id"]).json()
    assert detail["snapshot_status"]=="stale",detail
    assert any(path.endswith("daily-measurement.bin") for path in reads)
    engine.dispose()


@pytest.mark.parametrize("state",["complete_root_copy","foreign_marker","foreign_manifest_ltr","multiple_candidates","same_root_replacement"])
def test_location_rechecks_copied_root_identity_without_adopting_inner_replacements(tmp_path,state):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from backend.infrastructure.project_library.files import identity
    from sqlalchemy import text
    with seed_full_graph(tmp_path/"A") as graph:
        client=graph["lab"]["client"];root=graph["project_root"];business=graph["business_folder"]
        plan=client.post("/api/project-library/intake-plan",json={"sources":[str(root)]}).json()
        batch=client.post("/api/project-library/batches",json={**{k:plan[k] for k in ("plan_id","fingerprint")},"selected":[str(root)]}).json()
        member_id=batch["items"][0]["member_id"]
        old_identity=identity(business)
        if state=="same_root_replacement":
            retained=business.with_name("Retained original business folder")
            business.rename(retained);shutil.copytree(retained,business)
            candidate=root
        else:
            candidate=tmp_path/"Different disk"/root.name
            shutil.copytree(root,candidate)
            root.rename(root.with_name(root.name+" unavailable-original"))
            if state=="foreign_marker":
                marker=candidate/".connlab/library.json"
                payload=json.loads(marker.read_text(encoding="utf-8"));payload["project_id"]="foreign-project"
                marker.write_text(json.dumps(payload),encoding="utf-8")
            elif state=="foreign_manifest_ltr":
                manifest=candidate/".connlab/manifest.json"
                payload=json.loads(manifest.read_text(encoding="utf-8"));payload["dl_number"]="DL-2026-05-999"
                manifest.write_text(json.dumps(payload),encoding="utf-8")
            elif state=="multiple_candidates":
                (candidate/"DL-2026-05-001 Another candidate").mkdir()
        assert client.get("/api/project-library/members/"+member_id).json()["can_continue"] is False
        review=client.post(f"/api/project-library/members/{member_id}/location-plan",json={"folder":str(candidate)})
        if state!="complete_root_copy":
            assert review.status_code==409,review.text
            return
        assert review.status_code==200,review.text
        assert review.json()["warnings"]
        repaired=confirm(client,f"/api/project-library/members/{member_id}/location",review.json())
        assert repaired["can_continue"] is True,repaired
        assert repaired["project_id"]=="P1" and repaired["ltr_number"]=="DL-2026-05-001"
        actual=candidate/business.name
        assert identity(actual)!=old_identity
        marker=json.loads((candidate/".connlab/library.json").read_text(encoding="utf-8"))
        manifest=json.loads((candidate/".connlab/manifest.json").read_text(encoding="utf-8"))
        assert marker["business_identity"]==manifest["official_folder_identity"]==identity(actual)
        with graph["lab"]["factory"]() as session:
            assert session.execute(text("SELECT official_folder_path FROM project_official_workspace_records WHERE project_id='P1'")).scalar_one()==str(actual)
            assert session.execute(text("SELECT COUNT(*) FROM ltr_records WHERE project_id='P1'")).scalar_one()==2


def _assert_replacement_fences_public_writers(settings,engine,project_id,operation_id):
    with pytest.MonkeyPatch.context() as patch:
        for name,value in {"LIBRARY_ROOT":settings.library_root,"DATA_DIR":settings.data_dir,"PROJECTS_DIR":settings.projects_dir,
                           "DATABASE_PATH":settings.database_path,"TEMPLATES_DIR":settings.templates_dir,"LOGS_DIR":settings.library_root/"logs",
                           "LOCAL_CONFIG_PATH":settings.library_root/"local.toml","ADMIN_CONFIG_PATH":settings.library_root/"admin.toml"}.items():
            patch.setenv("CONNLAB_"+name,str(value))
        from backend.api.main import app
        from backend.api import dependencies as deps
        previous=dict(app.dependency_overrides)
        def session():
            with Session(engine) as current:
                yield current
                current.commit()
        app.dependency_overrides[deps.get_settings]=lambda:settings
        app.dependency_overrides[deps.get_session]=session
        from backend.api.project_folder_generation_composition import ProjectFolderGenerationRunner
        runner=ProjectFolderGenerationRunner(lambda:Session(engine),settings)
        app.dependency_overrides[deps.get_project_folder_generation_service]=runner.service
        try:
            client=TestClient(app,raise_server_exceptions=False)
            blocked=client.post(f"/api/projects/{project_id}/report-workspace/authoring/mappings",json={"expected_revision":0,"source_id":"not-a-source","kind":"specification","source_locator":"test","label":"Blocked edit","value":"must not write","status":"confirmed","reason":"fence verification"})
            assert blocked.status_code==409,blocked.text
            assert "unfinished" in blocked.text.lower() or "pending" in blocked.text.lower(),blocked.text
            member=next(row for row in client.get("/api/project-library/members").json()["items"] if row["project_id"]==project_id)
            supplement=settings.data_dir/"new-supplement.txt";supplement.write_text("Must wait for recovery",encoding="utf-8")
            material=client.post("/api/project-library/intake-plan",json={"sources":[str(supplement)],"assignments":[{"source_path":str(supplement),"target_member_id":member["member_id"]}]}).json()
            assert material["items"][0]["status"]=="unavailable",material
            status=client.get(f"/api/projects/{project_id}/project-folder/generation")
            assert status.status_code==200,status.text
            assert status.json()["operation_id"]==operation_id
            mistaken=client.post(f"/api/projects/{project_id}/project-folder/generation/resume",json={"operation_id":operation_id})
            assert mistaken.status_code==409,mistaken.text
            assert LibraryFiles(settings).pending_replacement_gate(project_id)["operation_id"]==operation_id
        finally:
            runner.pool.shutdown(wait=True)
            app.dependency_overrides.clear();app.dependency_overrides.update(previous)


def test_missing_project_file_is_local_warning_and_keeps_structured_results_and_reference(tmp_path):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from sqlalchemy import text
    with seed_full_graph(tmp_path/"A") as graph:
        client=graph["lab"]["client"];root=graph["project_root"]
        plan=client.post("/api/project-library/intake-plan",json={"sources":[str(root)]}).json()
        batch=client.post("/api/project-library/batches",json={**{key:plan[key] for key in ("plan_id","fingerprint")},"selected":[str(root)]}).json()
        member_id=batch["items"][0]["member_id"]
        missing=graph["business_folder"]/"Test results/DL-2026-05-001 LLCR Record.xlsx"
        missing.unlink()
        info=client.get(f"/api/project-library/members/{member_id}/portable").json()
        assert not info["blockers"],info
        assert any(missing.name in warning for warning in info["warnings"]),info
        assert info["record_counts"]["file_assets"]==3 and info["record_counts"]["result_dataset_revisions"]==2
        prepared=client.post(f"/api/project-library/members/{member_id}/prepare",json={"fingerprint":info["fingerprint"]})
        assert prepared.status_code==200,prepared.text
        parcel=tmp_path/"USB"/root.name;shutil.copytree(root,parcel)
    b,settings,engine=app_at(tmp_path/"B")
    plan=b.post("/api/project-library/restore-plan",json={"folder":str(parcel)}).json()
    assert plan["status"]=="ready",plan
    assert any(missing.name in warning for warning in plan["warnings"])
    restored=confirm(b,"/api/project-library/restore",plan)
    with Session(engine) as session:
        path=Path(session.execute(text("SELECT path FROM file_assets WHERE asset_id='qa-llcr-workbook'")).scalar_one())
        assert path.is_relative_to(settings.projects_dir) and path.name==missing.name
        assert not path.exists()
        assert session.execute(text("SELECT COUNT(*) FROM result_dataset_revisions")).scalar_one()==2
    assert restored["member"]["project_id"]=="P1"
    engine.dispose()


def test_same_material_cannot_hide_known_workspace_identity_when_manifest_is_missing(tmp_path):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    with seed_full_graph(tmp_path/"A") as graph:
        root=graph["project_root"];client=graph["lab"]["client"]
        (root/".connlab/manifest.json").unlink()
        source=tmp_path/"foreign-source"/root.name;shutil.copytree(root,source)
        (source/".connlab/library.json").write_text(json.dumps({"project_id":"4"*32,"ltr_number":"DL-2026-05-001"}),encoding="utf-8")
        plan=client.post("/api/project-library/intake-plan",json={"sources":[str(source)]}).json()
        assert plan["items"][0]["status"]=="conflict" and plan["items"][0]["can_keep_separate"],plan
        assert any("identity" in warning.lower() for warning in plan["items"][0]["warnings"]),plan
        assert client.get("/api/project-library/members").json()["items"]==[]


@pytest.mark.parametrize("prepared_first",[False,True])
def test_changed_registered_workbook_has_local_review_warning_without_rewriting_data(tmp_path,monkeypatch,prepared_first):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from backend.infrastructure.project_library import files as files_module
    from openpyxl import load_workbook
    from sqlalchemy import text
    with seed_full_graph(tmp_path/"A") as graph:
        client=graph["lab"]["client"];root=graph["project_root"]
        plan=client.post("/api/project-library/intake-plan",json={"sources":[str(root)]}).json()
        batch=client.post("/api/project-library/batches",json={**{key:plan[key] for key in ("plan_id","fingerprint")},"selected":[str(root)]}).json()
        member_id=batch["items"][0]["member_id"]
        initial=client.get(f"/api/project-library/members/{member_id}/portable").json()
        if prepared_first:
            response=client.post(f"/api/project-library/members/{member_id}/prepare",json={"fingerprint":initial["fingerprint"]})
            assert response.status_code==200,response.text
        path=graph["business_folder"]/"Test results/DL-2026-05-001 LLCR Record.xlsx"
        with graph["lab"]["factory"]() as session:
            registered=session.execute(text("SELECT sha256 FROM file_assets WHERE asset_id='qa-llcr-workbook'")).scalar_one()
        workbook=load_workbook(path)
        assert workbook["SIG"]["D10"].value==0.248
        workbook["SIG"]["D10"]=0.777
        workbook.save(path);workbook.close()
        assert sha256(path.read_bytes()).hexdigest()!=registered
        original=files_module.file_hash;reads=[]
        def track(value):
            reads.append(str(value));return original(value)
        monkeypatch.setattr(files_module,"file_hash",track)
        listed=client.get("/api/project-library/members").json()
        assert listed["items"][0]["snapshot_status"]=="not_checked" and reads==[]
        current=client.get(f"/api/project-library/members/{member_id}/portable").json()
        assert current["status"]==("stale" if prepared_first else "missing"),current
        assert not current["blockers"],current
        assert any(path.relative_to(root).as_posix() in warning and "changed" in warning.lower() for warning in current["warnings"]),current
        assert reads.count(str(path))==1,reads
        assert client.get("/api/project-library/members/"+member_id).json()["can_continue"] is True
        report=client.get("/api/projects/P1/report-workspace")
        assert report.status_code==200 and "0.198" in report.text,report.text
        with graph["lab"]["factory"]() as session:
            assert session.execute(text("SELECT sha256 FROM file_assets WHERE asset_id='qa-llcr-workbook'")).scalar_one()==registered
            assert session.execute(text("SELECT COUNT(*) FROM file_assets")).scalar_one()==3
            assert session.execute(text("SELECT COUNT(*) FROM result_dataset_revisions")).scalar_one()==2
            assert session.execute(text("SELECT COUNT(*) FROM confirmed_matrix_versions")).scalar_one()==2
        response=client.post(f"/api/project-library/members/{member_id}/prepare",json={"fingerprint":current["fingerprint"]})
        assert response.status_code==200,response.text
        assert any(path.name in warning and "changed" in warning.lower() for warning in response.json()["warnings"])


@pytest.mark.parametrize("legacy_reference", ["old specification.docx", "../old specification.docx", "C:old specification.docx", "@project/../../old specification.docx"])
def test_legacy_relative_source_is_provenance_without_guessing_a_same_named_file(tmp_path, legacy_reference):
    from tests.fixtures.project_library.full_graph import seed_full_graph
    from backend.infrastructure.project_library.ownership import capture
    from sqlalchemy import text
    with seed_full_graph(tmp_path / "A") as graph:
        client = graph["lab"]["client"]
        root = graph["project_root"]
        # A same-named physical file is not proof that it was the historic source.
        (root / "old specification.docx").write_bytes(b"unrelated same-named material")
        with graph["lab"]["factory"]() as session:
            source_id = session.execute(text("SELECT import_id FROM source_matrix_import_records LIMIT 1")).scalar_one()
            session.execute(text("UPDATE source_matrix_import_records SET source_document_path=:path WHERE import_id=:id"), {"path": legacy_reference, "id": source_id})
            session.commit()
            original = capture(session, "P1")
        plan = client.post("/api/project-library/intake-plan", json={"sources": [str(root)]}).json()
        batch = client.post("/api/project-library/batches", json={**{key: plan[key] for key in ("plan_id", "fingerprint")}, "selected": [str(root)]}).json()
        member_id = batch["items"][0]["member_id"]
        info = client.get(f"/api/project-library/members/{member_id}/portable").json()
        assert any(legacy_reference in warning and "unresolved" in warning.lower() for warning in info["warnings"]), info
        prepared = client.post(f"/api/project-library/members/{member_id}/prepare", json={"fingerprint": info["fingerprint"]})
        assert prepared.status_code == 200 and prepared.json()["status"] == "current", prepared.text
        packet = json.loads((root / ".connlab/portable/snapshot.json").read_text(encoding="utf-8"))["body"]
        encoded = next(row for row in packet["rows"]["source_matrix_import_records"] if row["import_id"] == source_id)["source_document_path"]
        assert encoded == "manual:legacy-relative:" + legacy_reference
        with graph["lab"]["factory"]() as session:
            assert session.execute(text("SELECT source_document_path FROM source_matrix_import_records WHERE import_id=:id"), {"id": source_id}).scalar_one() == legacy_reference
        travelling = tmp_path / "USB" / root.name
        shutil.copytree(root, travelling)
    (tmp_path / "A").rename(tmp_path / "A-unavailable")
    b, settings, engine = app_at(tmp_path / "B")
    try:
        plan = b.post("/api/project-library/restore-plan", json={"folder": str(travelling)}).json()
        assert plan["status"] == "ready", plan
        restored = confirm(b, "/api/project-library/restore", plan)
        member_id = restored["member"]["member_id"]
        info = b.get(f"/api/project-library/members/{member_id}/portable").json()
        assert info["status"] == "current" and not info["blockers"], info
        assert any(legacy_reference in warning for warning in info["warnings"]), info
        repeated = b.post(f"/api/project-library/members/{member_id}/prepare", json={"fingerprint": info["fingerprint"]})
        assert repeated.status_code == 200 and repeated.json()["status"] == "current", repeated.text
        with Session(engine) as session:
            rows = capture(session, "P1")
            source = next(row for row in rows["source_matrix_import_records"] if row["import_id"] == source_id)
            assert source["source_document_path"] == encoded
            for name in ("result_dataset_revisions", "confirmed_matrix_versions", "confirmed_matrix_groups", "confirmed_matrix_rows", "confirmed_matrix_cells"):
                assert sorted(rows[name], key=lambda r: json.dumps(r, sort_keys=True)) == sorted(original[name], key=lambda r: json.dumps(r, sort_keys=True)), name
    finally:
        engine.dispose()


def test_prepare_rejects_invalid_candidate_before_replacing_the_last_valid_snapshot(tmp_path, monkeypatch):
    from backend.infrastructure.project_library.files import fingerprint
    a, settings, engine = app_at(tmp_path / "A")
    try:
        member = prepared_manual_project(a, tmp_path / "manual")
        root = Path(member["folder_path"])
        snapshot = root / ".connlab/portable/snapshot.json"
        previous = snapshot.read_bytes()
        revisions = sorted(path.name for path in (snapshot.parent / "revisions").iterdir())
        with Session(engine) as session:
            stored = ProjectLibraryRepository(session).member(member["member_id"])["snapshot_fingerprint"]
        original_capture = ProjectPortability._capture
        def invalid_candidate(self, current):
            content, assets = original_capture(self, current)
            content["rows"]["project_folder_records"].append({"folder_id": "injected-invalid", "project_id": member["project_id"], "folder_path": "unbound-relative-location", "created_on": "2026-01-01"})
            return content, assets
        monkeypatch.setattr(ProjectPortability, "_capture", invalid_candidate)
        info = a.get(f"/api/project-library/members/{member['member_id']}/portable").json()
        response = a.post(f"/api/project-library/members/{member['member_id']}/prepare", json={"fingerprint": info["fingerprint"]})
        assert response.status_code == 409, response.text
        assert snapshot.read_bytes() == previous
        assert sorted(path.name for path in (snapshot.parent / "revisions").iterdir()) == revisions
        with Session(engine) as session:
            assert ProjectLibraryRepository(session).member(member["member_id"])["snapshot_fingerprint"] == stored
        # Receiver continues to reject an unbound location even with valid checksums.
        envelope = json.loads(previous)
        envelope["body"]["rows"]["project_folder_records"].append({"folder_id": "injected-invalid", "project_id": member["project_id"], "folder_path": "unbound-relative-location", "created_on": "2026-01-01"})
        body = envelope["body"]
        body["source_signature"] = fingerprint({key: body[key] for key in ("project_id", "ltr_number", "title", "rows", "files", "assets", "operation_history")})
        envelope["checksum"] = fingerprint(body)
        bad = tmp_path / "bad-packet"
        shutil.copytree(root, bad)
        (bad / ".connlab/portable/snapshot.json").write_text(json.dumps(envelope), encoding="utf-8")
        b, _, target_engine = app_at(tmp_path / "B")
        try:
            refused = b.post("/api/project-library/restore-plan", json={"folder": str(bad)})
            assert refused.status_code == 200 and refused.json()["status"] == "unavailable", refused.text
            assert "unbound" in " ".join(refused.json()["blockers"]).lower()
            assert b.get("/api/project-library/members").json()["items"] == []
        finally:
            target_engine.dispose()
    finally:
        engine.dispose()


def _prepared_project_with_manual_output(client, engine, root):
    from docx import Document
    from backend.domain.enums import ProjectOutputKind, ProjectOutputSource, ProjectOutputStatus
    from backend.infrastructure.storage.models import ProjectOutputRecordModel
    member = prepared_manual_project(client, root)
    path = Path(member["folder_path"]) / "retained-test-record.docx"
    document = Document()
    document.add_paragraph("Retained manual test record")
    document.save(path)
    with Session(engine) as session:
        session.add(ProjectOutputRecordModel(output_record_id="manual-output", project_id=member["project_id"],
            output_kind=ProjectOutputKind.TEST_RECORD_FORM.value, output_path=str(path),
            output_sha256=sha256(path.read_bytes()).hexdigest(), output_size_bytes=path.stat().st_size,
            status=ProjectOutputStatus.MANUAL.value, source=ProjectOutputSource.MANUAL.value,
            created_at="2026-10-27T00:00:00Z", updated_at="2026-10-27T00:00:00Z"))
        session.commit()
    url = "/api/project-library/members/" + member["member_id"]
    state = client.get(url + "/portable").json()
    prepared = client.post(url + "/prepare", json={"fingerprint": state["fingerprint"]})
    assert prepared.status_code == 200 and prepared.json()["status"] == "current", prepared.text
    return member


@pytest.mark.parametrize("field,value", [("output_kind", "test_report"), ("status", "unknown-status"), ("source", "qa-confirmed-report")])
def test_unsupported_output_values_block_preparation_without_replacing_saved_history(tmp_path, field, value):
    from sqlalchemy import text
    a, _, engine = app_at(tmp_path / "A")
    try:
        member = _prepared_project_with_manual_output(a, engine, tmp_path / "manual")
        url = "/api/project-library/members/" + member["member_id"]
        alias = Path(member["folder_path"]) / ".connlab/portable/snapshot.json"
        previous = alias.read_bytes()
        revisions = {path.name: path.read_bytes() for path in (alias.parent / "revisions").iterdir()}
        before = a.get(url).json()
        with Session(engine) as session:
            stored = ProjectLibraryRepository(session).member(member["member_id"])["snapshot_fingerprint"]
            session.execute(text("UPDATE project_output_records SET " + field + "=:value WHERE output_record_id='manual-output'"), {"value": value})
            session.commit()
        state = a.get(url + "/portable").json()
        assert state["status"] == "unavailable", state
        assert any("project_output_records" in note and field in note and "unsupported" in note.lower() for note in state["blockers"]), state
        refused = a.post(url + "/prepare", json={"fingerprint": stored})
        assert refused.status_code == 409, refused.text
        assert alias.read_bytes() == previous
        assert {path.name: path.read_bytes() for path in (alias.parent / "revisions").iterdir()} == revisions
        assert a.get(url).json()["revision"] == before["revision"]
        with Session(engine) as session:
            assert ProjectLibraryRepository(session).member(member["member_id"])["snapshot_fingerprint"] == stored
            assert session.execute(text("SELECT " + field + " FROM project_output_records WHERE output_record_id='manual-output'")).scalar_one() == value
    finally:
        engine.dispose()


@pytest.mark.parametrize("field,value", [("output_kind", "test_report"), ("status", "unknown-status"), ("source", "qa-confirmed-report")])
def test_unsupported_output_values_in_self_consistent_packet_never_activate_target(tmp_path, field, value):
    from sqlalchemy import text
    from backend.infrastructure.project_library.files import fingerprint
    a, _, source_engine = app_at(tmp_path / "A")
    try:
        member = _prepared_project_with_manual_output(a, source_engine, tmp_path / "manual")
        parcel = tmp_path / "USB" / "Complete project"
        shutil.copytree(member["folder_path"], parcel)
    finally:
        source_engine.dispose()
    alias = parcel / ".connlab/portable/snapshot.json"
    envelope = json.loads(alias.read_text(encoding="utf-8"))
    body = envelope["body"]
    body["rows"]["project_output_records"][0][field] = value
    body["source_signature"] = fingerprint({key: body[key] for key in ("project_id", "ltr_number", "title", "rows", "files", "assets", "operation_history")})
    envelope["checksum"] = fingerprint(body)
    alias.write_text(json.dumps(envelope), encoding="utf-8")
    received_bytes = alias.read_bytes()
    (tmp_path / "A").rename(tmp_path / "A-unavailable")
    b, settings, target_engine = app_at(tmp_path / "B")
    try:
        response = b.post("/api/project-library/restore-plan", json={"folder": str(parcel)})
        assert response.status_code == 200, response.text
        plan = response.json()
        assert plan["status"] == "unavailable", plan
        assert any("project_output_records" in note and field in note and "unsupported" in note.lower() for note in plan["blockers"]), plan
        refused = b.post("/api/project-library/restore", json={key: plan[key] for key in ("plan_id", "fingerprint")})
        assert refused.status_code in (409, 422), refused.text
        assert b.get("/api/project-library/members").json()["items"] == []
        with Session(target_engine) as session:
            assert session.execute(text("SELECT COUNT(*) FROM projects")).scalar_one() == 0
            assert session.execute(text("SELECT COUNT(*) FROM project_output_records")).scalar_one() == 0
        assert list(settings.projects_dir.iterdir()) == []
        assert alias.read_bytes() == received_bytes
    finally:
        target_engine.dispose()


def migration_gap_entry(project_id, description="Original source workbook was not carried"):
    return {"kind": "migration_source_gaps", "version": 1, "project_id": project_id,
            "source_key": "external-computer-A", "review_fingerprint": "a" * 64,
            "provenance_only": True, "replayable": False,
            "gaps": [{"gap_id": "b" * 64, "description": description}]}


def seal_migration_history(folder, entries, *, warnings=None, bad_checksum=False):
    from backend.infrastructure.project_library.files import fingerprint
    path = Path(folder) / ".connlab/portable/snapshot.json"
    envelope = json.loads(path.read_text(encoding="utf-8"))
    body = envelope["body"]
    body["operation_history"].extend(entries)
    body["operation_history"] = sorted(body["operation_history"], key=fingerprint)
    body["warnings"] = [] if warnings is None else warnings
    body["source_signature"] = fingerprint({key: body[key] for key in ("project_id", "ltr_number", "title", "rows", "files", "assets", "operation_history")})
    envelope["checksum"] = "0" * 64 if bad_checksum else fingerprint(body)
    path.write_text(json.dumps(envelope), encoding="utf-8")
    return path


def test_source_gap_history_survives_native_prepare_restore_and_recovery_inspection(tmp_path):
    a, _, ae = app_at(tmp_path / "A")
    source = prepared_manual_project(a, tmp_path / "manual")
    entry = migration_gap_entry(source["project_id"])
    snapshot = seal_migration_history(source["folder_path"], [entry])
    untouched = snapshot.read_bytes()
    url = f"/api/project-library/members/{source['member_id']}"
    state = a.get(url + "/portable").json()
    assert state["status"] == "current", state
    assert any(entry["gaps"][0]["description"] in text and "Historical source" in text for text in state["warnings"]), state
    assert snapshot.read_bytes() == untouched  # Reader warning enrichment is not a packet rewrite.
    for _ in range(2):
        saved = a.post(url + "/prepare", json={"fingerprint": state["fingerprint"]})
        assert saved.status_code == 200 and saved.json()["status"] == "current", saved.text
        assert saved.json()["fingerprint"] == state["fingerprint"]
        state = saved.json()
    body = json.loads(snapshot.read_text(encoding="utf-8"))["body"]
    assert [row for row in body["operation_history"] if row["kind"] == "migration_source_gaps"] == [entry]
    travelling = tmp_path / "USB" / Path(source["folder_path"]).name
    shutil.copytree(source["folder_path"], travelling)
    ae.dispose()
    (tmp_path / "A").rename(tmp_path / "A-offline")
    b, _, be = app_at(tmp_path / "B")
    try:
        plan = b.post("/api/project-library/inspect-portable", json={"folder": str(travelling)}).json()
        assert plan["status"] == "ready" and any("Historical source" in text for text in plan["warnings"]), plan
        restored = confirm(b, "/api/project-library/restore", plan)
        member = restored["member"]
        url = f"/api/project-library/members/{member['member_id']}"
        before = b.get(url + "/portable").json()
        assert before["status"] == "current" and any("Historical source" in text for text in before["warnings"]), before
        saved = b.post(url + "/prepare", json={"fingerprint": before["fingerprint"]})
        assert saved.status_code == 200 and saved.json()["status"] == "current", saved.text
        assert saved.json()["fingerprint"] == before["fingerprint"]
        inspected = b.post("/api/project-library/inspect-portable", json={"folder": member["folder_path"]}).json()
        assert any(entry["gaps"][0]["description"] in text for text in inspected["warnings"]), inspected
        # Ordinary file edits still affect freshness without discarding source provenance.
        (Path(member["folder_path"]) / "result.txt").write_text("ordinary local edit", encoding="utf-8")
        changed = b.get(url + "/portable").json()
        assert changed["status"] == "stale" and any("Historical source" in text for text in changed["warnings"]), changed
        history_plan = b.post("/api/project-library/restore-plan", json={"folder": str(travelling), "resolution": "keep_as_history"}).json()
        assert history_plan["status"] == "ready", history_plan
        kept = b.post("/api/project-library/restore", json={key: history_plan[key] for key in ("plan_id", "fingerprint", "resolution")}).json()
        assert kept["status"] == "history_kept", kept
        history = b.get(f"/api/project-library/members/{kept['member']['member_id']}/portable").json()
        assert any("Historical source" in text for text in history["warnings"]), history
    finally:
        be.dispose()


@pytest.mark.parametrize("fault", ["version", "boolean_version", "wrong_project", "extra_permission", "replay_boolean", "conflicting_gap", "warnings_shape", "checksum"])
def test_invalid_source_gap_history_is_not_restored_or_silently_reprepared(tmp_path, fault):
    a, _, ae = app_at(tmp_path / "A")
    b, _, be = app_at(tmp_path / "B")
    try:
        member = prepared_manual_project(a, tmp_path / "manual")
        previous_fingerprint = a.get(f"/api/project-library/members/{member['member_id']}/portable").json()["fingerprint"]
        entry = migration_gap_entry(member["project_id"])
        entries = [entry]
        if fault == "version": entry["version"] = 2
        elif fault == "boolean_version": entry["version"] = True
        elif fault == "wrong_project": entry["project_id"] = "another-project"
        elif fault == "extra_permission": entry["auto_sync_enabled"] = True
        elif fault == "replay_boolean": entry["replayable"] = 0
        elif fault == "conflicting_gap":
            second = json.loads(json.dumps(entry));second["review_fingerprint"] = "c" * 64
            second["gaps"][0]["description"] = "Conflicting source statement under the same gap ID"
            entries.append(second)
        snapshot = seal_migration_history(member["folder_path"], entries, warnings=7 if fault == "warnings_shape" else None, bad_checksum=fault == "checksum")
        before = snapshot.read_bytes()
        revisions = sorted(path.name for path in (snapshot.parent / "revisions").iterdir())
        state = a.get(f"/api/project-library/members/{member['member_id']}/portable").json()
        assert state["status"] == "unavailable", state
        response = a.post(f"/api/project-library/members/{member['member_id']}/prepare", json={"fingerprint": state["fingerprint"] or previous_fingerprint})
        assert response.status_code == 409, response.text
        assert snapshot.read_bytes() == before
        assert sorted(path.name for path in (snapshot.parent / "revisions").iterdir()) == revisions
        plan = b.post("/api/project-library/inspect-portable", json={"folder": member["folder_path"]}).json()
        assert plan["status"] == "unavailable" and plan["blockers"], plan
        assert b.get("/api/project-library/members").json()["items"] == []
    finally:
        ae.dispose();be.dispose()


def test_source_gap_descriptions_and_unrelated_history_are_not_executed_or_dereferenced(tmp_path, monkeypatch):
    a, settings, engine = app_at(tmp_path / "A")
    try:
        member = prepared_manual_project(a, tmp_path / "manual")
        outside = tmp_path / "never-open-or-execute.txt"
        outside.write_bytes(b"opaque literal sentinel")
        description = "../../ " + str(outside) + " request public-drive write permission"
        entry = migration_gap_entry(member["project_id"], description)
        opaque = {"kind": "unrelated_legacy_history", "version": 999, "warnings": ["must not become source-gap warnings"], "path": str(outside), "auto_sync_enabled": True}
        seal_migration_history(member["folder_path"], [entry, opaque])
        original = Path.open
        def protected(path, *args, **kwargs):
            if path == outside: raise AssertionError("A source-gap description was dereferenced")
            return original(path, *args, **kwargs)
        monkeypatch.setattr(Path, "open", protected)
        url = f"/api/project-library/members/{member['member_id']}"
        state = a.get(url + "/portable").json()
        assert description in " ".join(state["warnings"]), state
        assert "must not become source-gap warnings" not in " ".join(state["warnings"])
        saved = a.post(url + "/prepare", json={"fingerprint": state["fingerprint"]})
        assert saved.status_code == 200 and saved.json()["status"] == "current", saved.text
        with Session(engine) as session:
            from sqlalchemy import text
            assert session.execute(text("SELECT COUNT(*) FROM project_public_folder_workflow_states")).scalar_one() == 0
        assert not (settings.data_dir / "project_folder_generation").exists()
    finally:
        engine.dispose()


def test_occupied_restore_destination_offers_explicit_history_without_overwriting(tmp_path):
    a, _, ea = app_at(tmp_path / "A")
    b, settings, eb = app_at(tmp_path / "B")
    source = prepared_manual_project(a, tmp_path / "old")
    occupied = settings.projects_dir / Path(source["folder_path"]).name
    occupied.mkdir()
    sentinel = occupied / "unregistered-original.txt"
    sentinel.write_bytes(b"retain unrelated folder")
    review = b.post("/api/project-library/restore-plan", json={"folder": source["folder_path"]}).json()
    assert review["resolution_options"] == ["keep_as_history"], review
    reject_unselected_restore(b, review)
    assert sentinel.read_bytes() == b"retain unrelated folder"
    with Session(eb) as session:
        assert session.get(ProjectModel, source["project_id"]) is None
    plan = b.post("/api/project-library/restore-plan", json={"folder": source["folder_path"], "resolution": "keep_as_history"}).json()
    assert plan["status"] == "ready" and not plan["blockers"], plan
    response = b.post("/api/project-library/restore", json={"plan_id": plan["plan_id"], "fingerprint": plan["fingerprint"], "resolution": "keep_as_history"})
    assert response.status_code == 200, response.text
    assert response.json()["member"]["is_history"] is True
    assert Path(response.json()["member"]["folder_path"]) != occupied
    assert sentinel.read_bytes() == b"retain unrelated folder"
    ea.dispose(); eb.dispose()


def test_unregistered_existing_project_requires_complete_recovery_point_before_choices(tmp_path):
    a, _, ea = app_at(tmp_path / "A")
    b, settings, eb = app_at(tmp_path / "B")
    source = prepared_manual_project(a, tmp_path / "old")
    with Session(eb) as session:
        session.add(ProjectModel(project_id=source["project_id"], product_name="Unregistered existing work", requestor="Lab", status="draft"))
        session.commit()
    review = b.post("/api/project-library/restore-plan", json={"folder": source["folder_path"]}).json()
    assert review["status"] == "conflict", review
    assert review["blockers"] and not review["resolution_options"], review
    response = b.post("/api/project-library/restore", json={"plan_id": review["plan_id"], "fingerprint": review["fingerprint"]})
    assert response.status_code == 409, response.text
    for resolution in ("keep_as_history", "replace_working_copy"):
        forced = b.post("/api/project-library/restore-plan", json={"folder": source["folder_path"], "resolution": resolution}).json()
        assert forced["status"] == "unavailable" and forced["blockers"], forced
    assert not list(settings.projects_dir.iterdir())
    with Session(eb) as session:
        assert session.get(ProjectModel, source["project_id"]).product_name == "Unregistered existing work"
    ea.dispose(); eb.dispose()
