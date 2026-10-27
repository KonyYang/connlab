"""Independent, nonempty project graph for portable-project acceptance.

No product imports occur before explicit isolated environment binding. This file
contains synthetic laboratory records only; no operator data or source-machine paths.
"""
from __future__ import annotations

from contextlib import contextmanager
from dataclasses import replace
from datetime import date
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from itertools import count
import json
from pathlib import Path
import shutil

PROJECT_ID = "P1"
LTR = "DL-2026-05-001"
MATRIX_ID = "cmv-1"
GROUP_ID = "cmg-cmv-1-2"
ROW_ID = "cmr-cmv-1-2"
AT = "2026-10-27T00:00:00Z"


def _hash(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def _ok(response):
    assert response.status_code == 200, (response.status_code, response.text)
    return response.json()


def _isolated_environment(root, patch):
    root = Path(root).absolute()
    # A caller must deliberately pass a fresh isolated directory. Never accept an
    # existing business DB or use any machine/default path as an output root.
    if (root / "connlab.sqlite3").exists() or (root / "data" / "connlab.sqlite3").exists():
        raise ValueError("Seed requires a fresh isolated root, not an existing database.")
    root.mkdir(parents=True, exist_ok=True)
    for key, value in {
        "CONNLAB_LIBRARY_ROOT": root,
        "CONNLAB_DATA_DIR": root / "data",
        "CONNLAB_PROJECTS_DIR": root / "projects",
        "CONNLAB_TEMPLATES_DIR": root / "templates",
        "CONNLAB_DATABASE_PATH": root / "connlab.sqlite3",
        "CONNLAB_LOGS_DIR": root / "logs",
        "CONNLAB_LOCAL_CONFIG_PATH": root / "isolated.local.toml",
        "CONNLAB_ADMIN_CONFIG_PATH": root / "isolated.admin.toml",
    }.items():
        patch.setenv(key, str(value))
    (root / "isolated.local.toml").write_text("# Isolated QA configuration\n", encoding="utf-8")
    (root / "isolated.admin.toml").write_text("# No real workbook or credentials\n", encoding="utf-8")
    return root


def _seed_owned_sources(lab):
    from backend.infrastructure.storage.models import (
        ApplicationFormModel, SampleInfoModel, FileAssetModel, IntakePackageModel,
        IntakeAssetModel, IntakeCaseModel, IntakeDraftModel, ProjectTestPlanDraftModel,
        ProjectTemporaryContextModel, ProjectRequestMaterialCollectionModel,
        ProjectRequestMaterialCollectionItemModel, LtrRecordModel, LtrAssociationEventModel,
        ProjectLifecycleEventModel, ProjectRegistryEventModel,
    )
    from backend.infrastructure.storage.models_matrix_source import SourceMatrixImportRecordModel
    from tests.integration.test_report_authoring_api import docx_bytes

    settings = lab["settings"]
    local = settings.projects_dir / LTR
    business = local / f"{LTR} QA Connector Qualification Testing"
    source_book = local / "Source Book"
    for directory in (source_book, business / "Test results", business / "Photos", business / "Submitted Material", business / "E-mail"):
        directory.mkdir(parents=True, exist_ok=True)
    spec = source_book / "QA Specification.docx"
    spec.write_bytes(docx_bytes())
    intake_source = settings.data_dir / "intake" / "qa-package" / "source" / "request.msg"
    intake_source.parent.mkdir(parents=True)
    intake_source.write_bytes(b"QA laboratory request: retain original message and application lineage.\n")
    attachment = intake_source.parent.parent / "attachments" / "Application.docx"
    attachment.parent.mkdir()
    attachment.write_bytes(docx_bytes())
    placed = business / "Submitted Material" / "Application.docx"
    shutil.copyfile(attachment, placed)
    (business / "Photos" / "operator-note.txt").write_text("Operator-created photo note survives copying.\n", encoding="utf-8")

    with lab["factory"]() as session:
        session.add_all([
            ApplicationFormModel(form_id="qa-form", project_id=PROJECT_ID, form_no="QA-APPLICATION-01", revision="A", requester="QA Operator", request_date=date(2026, 10, 1), requested_testing="Connector LLCR qualification", lab_test_request_number=LTR),
            SampleInfoModel(sample_id="qa-sample", project_id=PROJECT_ID, product_name="QA Connector", part_number="QA-PART-032", revision="B", quantity=6, row_index=0, source_form_id="qa-form", material="Copper Alloy", plating="Gold"),
            IntakePackageModel(package_id="qa-package", source_type="outlook_msg", status="imported", source_original_name="request.msg", source_stored_path=str(intake_source), subject="QA connector qualification", body_text="Application and specification belong to P1", created_at=AT),
        ])
        session.flush()
        session.add(IntakeAssetModel(asset_id="qa-intake-application", package_id="qa-package", original_name="Application.docx", stored_path=str(attachment), extension=".docx", mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", size_bytes=attachment.stat().st_size, sha256=_hash(attachment), asset_role="application_form_candidate", candidate_score=100, created_at=AT))
        session.flush()
        session.add(IntakeCaseModel(case_id="case-1", package_id="qa-package", selected_form_asset_id="qa-intake-application", status="confirmed", confirmed_project_id=PROJECT_ID, created_at=AT, updated_at=AT, reviewer_notes="Explicitly accepted synthetic source"))
        session.flush()
        session.add(IntakeDraftModel(draft_id="draft-1", case_id="case-1", parsed_fields_json=json.dumps({"lab_test_request_number": LTR, "product_name": "QA Connector"}), sample_rows_json=json.dumps([{"part_number": "QA-PART-032", "quantity": 6}]), requested_testing_json=json.dumps(["LLCR"]), parser_warnings_json="[]", manual_overrides_json=json.dumps({"revision": "B"}), updated_at=AT))
        session.add_all([
            FileAssetModel(asset_id="asset-1", project_id=PROJECT_ID, asset_type="attachment", path=str(spec), original_name=spec.name, source_package_id="qa-package", source_intake_asset_id="qa-intake-application", source_role="specification", sha256=_hash(spec)),
            FileAssetModel(asset_id="qa-file-application", project_id=PROJECT_ID, asset_type="application_form", path=str(attachment), original_name=attachment.name, source_package_id="qa-package", source_intake_asset_id="qa-intake-application", source_role="application_form", sha256=_hash(attachment)),
            ProjectTemporaryContextModel(context_id="qa-temporary-context", project_id=PROJECT_ID, request_summary="Historic request context remains provenance", sample_description="6 QA Connector samples", test_item="LLCR", source_asset_ids_json=json.dumps(["asset-1", "qa-file-application"])),
            LtrRecordModel(ltr_id="qa-ltr-current", project_id=PROJECT_ID, ltr_number=LTR, status="registered", registered_on=date(2026, 10, 1), is_current_owner=True, owner_version=2),
            LtrRecordModel(ltr_id="qa-ltr-history", project_id=PROJECT_ID, ltr_number=LTR, status="registered", registered_on=date(2026, 9, 30), is_current_owner=False, superseded_at=AT, superseded_by_ltr_id="qa-ltr-current", superseded_reason="Preserved association history", owner_version=1),
            LtrAssociationEventModel(event_id="qa-ltr-event", ltr_number=LTR, event_type="association_reconfirmed", old_ltr_id="qa-ltr-history", old_project_id=PROJECT_ID, new_ltr_id="qa-ltr-current", new_project_id=PROJECT_ID, operator="QA Operator", reason="Keep history, no renumbering", created_at=AT),
            ProjectRegistryEventModel(event_id="qa-registry-event", operation_id="qa-registry-operation", project_id=PROJECT_ID, previous_state="active", new_state="active", revision=1, reason="Source identity verified", actor="QA Operator", changed_at=AT),
            ProjectLifecycleEventModel(event_id="qa-lifecycle-event", project_id=PROJECT_ID, event_type="resumed", previous_lifecycle_state="stopped", new_lifecycle_state="active", reason="QA continuing tests", operator="QA Operator", created_at=AT),
        ])
        for source in session.query(SourceMatrixImportRecordModel).filter_by(project_id=PROJECT_ID):
            source.source_document_path = str(spec)
            source.source_document_name = spec.name
            source.source_asset_id = "asset-1"
        session.add(ProjectTestPlanDraftModel(draft_id="ptpd-seed", project_id=PROJECT_ID, source_document_path=str(spec), source_document_name=spec.name, source_format=".docx", source_asset_id="asset-1", source_case_id="case-1", source_draft_id="draft-1", status="reviewed", version=1, payload_json=json.dumps({"groups": [{"group_key": "g2", "group_label": "G2", "sample_quantity_expression": "6"}], "rows": [{"test_item": "LLCR", "group_tokens": {"G2": "2"}}], "warnings": [], "blockers": []}), created_at=AT, updated_at=AT, reviewed_at=AT))
        session.add(ProjectRequestMaterialCollectionModel(collection_id="qa-collection", project_id=PROJECT_ID, workspace_id="qa-workspace", status="completed", item_count=1, copied_count=1, already_present_count=0, conflict_count=0, skipped_count=0, missing_source_count=0, created_at=AT, updated_at=AT, warnings_json="[]"))
        session.flush()
        session.add(ProjectRequestMaterialCollectionItemModel(item_id="qa-collected-application", collection_id="qa-collection", project_id=PROJECT_ID, source_asset_id="qa-file-application", source_asset_type="application_form", source_role="application_form", dedupe_key="qa-application-sha", source_path=str(attachment), original_name=attachment.name, target_area="submitted_material", target_path=str(placed), status="copied", action="copy", review_required=False, size_bytes=placed.stat().st_size, sha256=_hash(placed)))
        session.commit()
    return local, business, source_book, spec


def _seed_workspace(lab, local, business, source_book):
    from backend.application.official_project_workspace_service import OfficialWorkspaceRecord
    from backend.domain import ExternalResource, ExternalResourceType, ExternalResourceValidationStatus
    from backend.infrastructure.official_workspace_manifest import OfficialWorkspaceManifest, OfficialWorkspaceManifestGateway
    from backend.infrastructure.storage.repositories.official_workspace import ProjectOfficialWorkspaceRepository
    from backend.infrastructure.storage.repositories import ExternalResourceRepository
    from backend.infrastructure.storage.models import ProjectFolderRecordModel
    manifest = local / ".connlab" / "manifest.json"
    record = OfficialWorkspaceRecord("qa-workspace", PROJECT_ID, LTR, local, source_book, business, manifest, lab["settings"].templates_dir, AT)
    OfficialWorkspaceManifestGateway().write(manifest, OfficialWorkspaceManifest(1, PROJECT_ID, LTR, str(local), str(source_book), str(business), str(lab["settings"].templates_dir), AT))
    with lab["factory"]() as session:
        ProjectOfficialWorkspaceRepository(session).save(record)
        ExternalResourceRepository(session).upsert(ExternalResource("qa-project-root", ExternalResourceType.PROJECT_OUTPUT_ROOT, lab["settings"].projects_dir, True, ExternalResourceValidationStatus.VALID))
        session.add(ProjectFolderRecordModel(folder_id="qa-folder-record", project_id=PROJECT_ID, folder_path=str(business), created_on=date(2026, 10, 27)))
        session.commit()


def _seed_matrix_and_plan_history(lab):
    from backend.domain import ConfirmedMatrixStatus, MatrixStepContactPlan, MatrixStepContactFamily, contact_plan_to_json
    from backend.infrastructure.storage.models_confirmed_matrix_authority import ConfirmedMatrixVersionModel, ConfirmedMatrixStepQuantityModel
    from backend.infrastructure.storage.models_project_matrix_draft import ProjectMatrixDraftRecordModel
    from backend.infrastructure.storage.repositories import ConfirmedMatrixAuthorityRepository, ProjectMatrixDraftRepository, ContactMeasurementPlanAuthorityRepository
    from backend.application.contact_measurement_plan_bootstrap_service import ContactMeasurementPlanBootstrapService
    from backend.application.contact_measurement_plan_lifecycle_service import ContactMeasurementPlanLifecycleService
    from tests.unit.test_confirmed_matrix_authority_repository import _build_confirmed_snapshot

    with lab["factory"]() as session:
        repo = ConfirmedMatrixAuthorityRepository(session)
        current = repo.get(MATRIX_ID)
        session.get(ConfirmedMatrixVersionModel, MATRIX_ID).confirmed_revision = 2
        session.flush()
        old = replace(current, version=replace(current.version, confirmed_matrix_id="qa-matrix-history", confirmed_revision=1, is_active_authority=False, status=ConfirmedMatrixStatus.SUPERSEDED, superseded_by_confirmed_matrix_id=MATRIX_ID, superseded_at=AT, superseded_reason="Retain source matrix history"), groups=tuple(replace(g, confirmed_group_id="qa-history-"+g.confirmed_group_id, confirmed_matrix_id="qa-matrix-history") for g in current.groups), rows=tuple(replace(r, confirmed_row_id="qa-history-"+r.confirmed_row_id, confirmed_matrix_id="qa-matrix-history") for r in current.rows), cells=tuple(replace(c, confirmed_cell_id="qa-history-"+c.confirmed_cell_id, confirmed_matrix_id="qa-matrix-history", confirmed_group_id="qa-history-"+c.confirmed_group_id, confirmed_row_id="qa-history-"+c.confirmed_row_id) for c in current.cells), duration_authorities=(), step_quantities=(), step_text_overrides=())
        repo.create_snapshot(old)
        contact = MatrixStepContactPlan("llcr", "eligible", True, None, False, "1", (MatrixStepContactFamily("signal_pin", "Signal Pin", "1", "Signal Pin contact", "SIG", True, False),))
        session.add(ConfirmedMatrixStepQuantityModel(confirmed_step_quantity_id="qa-confirmed-quantity", confirmed_matrix_id=MATRIX_ID, confirmed_group_id=GROUP_ID, confirmed_row_id=ROW_ID, draft_group_id="pmdg-2", draft_row_id="pmdr-2", step_sequence=2, step_suffix_note="", raw_token="2", test_points_per_sample="1", readings_per_point="1", contact_points_per_sample="1", source="matrix_step_override", review_required=False, confirmed_at=AT, contact_plan_json=contact_plan_to_json(contact)))
        session.flush()
        store = ContactMeasurementPlanAuthorityRepository(session)
        serial = count(1)
        ids = lambda: f"qa-measurement-{next(serial)}"
        bootstrap = ContactMeasurementPlanBootstrapService(repository=store, clock=lambda: AT, id_factory=ids)
        bootstrap.bootstrap(repo.get(MATRIX_ID), actor="QA Operator")
        lifecycle = ContactMeasurementPlanLifecycleService(repository=store, confirmed_store=repo, bootstrap_service=bootstrap, clock=lambda: AT, id_factory=ids)
        editable = lifecycle.open_draft(PROJECT_ID, actor="QA Operator")
        session.commit()
        active = store.get_active_revision(PROJECT_ID)
        targets = store.targets(active.measurement_plan_revision_id)
        assert targets and targets[0].readings_per_sample == 1
        lab["measurement_target_key"] = targets[0].stable_target_key
        lab["measurement_draft_id"] = editable


def _seed_point_and_fee_history(lab):
    from backend.infrastructure.storage.models_contact_point_profile import ContactPointProfileRootModel, ContactPointProfileRevisionModel, ContactPointProfileCategoryModel, ContactPointProfileCrCategorySelectionModel
    from backend.infrastructure.storage.repositories.project_schedule import ProjectScheduleRepository
    from backend.infrastructure.storage.repositories.confirmed_fee_authority import ConfirmedFeeAuthorityRepository
    from backend.infrastructure.storage.repositories.fee_evaluation_pricing_draft_edit import FeeEvaluationPricingDraftEditRepository
    from tests.unit.test_project_schedule_repository import _revision
    from tests.unit.test_confirmed_fee_authority_repository import _version
    from tests.unit.test_fee_evaluation_pricing_draft_repository import _snapshot
    with lab["factory"]() as session:
        session.add(ContactPointProfileRootModel(contact_point_profile_root_id="qa-point-root", project_id=PROJECT_ID, created_at=AT, updated_at=AT))
        session.flush()
        session.add_all([
            ContactPointProfileRevisionModel(contact_point_profile_revision_id="qa-point-history", contact_point_profile_root_id="qa-point-root", revision_sequence=1, state="superseded", revision_fingerprint="qa-point-history-fingerprint", created_by="QA Operator", created_at=AT, updated_at=AT, confirmed_by="QA Operator", confirmed_at=AT, superseded_at=AT, superseded_reason="Explicit point revision"),
            ContactPointProfileRevisionModel(contact_point_profile_revision_id="qa-point-current", contact_point_profile_root_id="qa-point-root", revision_sequence=2, parent_revision_id="qa-point-history", state="confirmed", revision_fingerprint="qa-point-current-fingerprint", created_by="QA Operator", created_at=AT, updated_at=AT, confirmed_by="QA Operator", confirmed_at=AT),
        ])
        session.flush()
        for revision, size in (("qa-point-history", 2), ("qa-point-current", 4)):
            session.add(ContactPointProfileCategoryModel(contact_point_profile_category_snapshot_id=revision+"-signal", contact_point_profile_revision_id=revision, category_id="qa-signal", category_ordinal=0, label="Signal", normalized_label_key="signal", count_per_sample=size, record_prefix="SIG", normalized_prefix_key="sig", included=True, point_expression=f"1-{size}"))
        session.flush()
        session.add(ContactPointProfileCrCategorySelectionModel(contact_point_profile_cr_selection_id="qa-cr-selection", contact_point_profile_revision_id="qa-point-current", category_id="qa-signal", selection_ordinal=0))
        session.get(ContactPointProfileRootModel, "qa-point-root").active_confirmed_revision_id = "qa-point-current"
        for seq, state in ((1, "superseded"), (2, "confirmed")):
            schedule = replace(_revision(f"qa-schedule-{seq}", seq, state=state), based_on_confirmed_matrix_id=MATRIX_ID, based_on_confirmed_matrix_revision=2, based_on_basic_information_version=1, sample_received_date="2026-05-30", test_start_date="2026-06-01", test_complete_date="2026-06-10", estimated_completion_date="2026-06-12")
            ProjectScheduleRepository(session).add(schedule)
        pricing = _snapshot("fed-1", unit_price="25")
        row = replace(pricing.edited_values.rows[0], confirmed_group_id=GROUP_ID, confirmed_row_id=ROW_ID, step_token="2", testing_fee="25")
        pricing = replace(pricing, confirmed_revision=2, edited_values=replace(pricing.edited_values, rows=(row,)))
        FeeEvaluationPricingDraftEditRepository(session).upsert_current(pricing)
        for seq, fee in ((1, "41"), (2, "52")):
            ConfirmedFeeAuthorityRepository(session).create(replace(_version(f"qa-fee-{seq}", revision=seq, testing_fee_total=fee), confirmed_revision=2, pricing_snapshot_json=json.dumps({"rows": [{"confirmed_group_id": GROUP_ID, "confirmed_row_id": ROW_ID, "step_token": "2", "testing_fee": fee}]})))
        session.commit()


def _seed_results_and_reports(lab, business):
    from openpyxl import Workbook
    from docx import Document
    from backend.domain.enums import ProjectOutputKind, ProjectOutputSource, ProjectOutputStatus
    from backend.domain.result_dataset_models import ResultDatasetSourceIdentity
    from backend.infrastructure.storage.repositories.result_dataset import ResultDatasetRepository
    from backend.infrastructure.storage.models import FileAssetModel, ProjectOutputRecordModel
    from backend.infrastructure.files.report_draft_workspace import ReportDraftWorkspace
    from tests.unit.test_result_dataset_repository import _dataset, _report
    from tests.integration.test_report_authoring_api import docx_bytes, upload, BASE

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SIG"
    sheet["D10"] = 0.248
    sheet["K10"] = 0.198
    original = business / "Test results" / f"{LTR} LLCR Record.xlsx"
    workbook.save(original)
    workbook.close()
    source = ResultDatasetSourceIdentity(original.name, _hash(original), original.stat().st_size)
    report_dir = lab["settings"].data_dir / "generated_test_reports"
    report_dir.mkdir(parents=True)
    record_dir = lab["settings"].data_dir / "generated_test_records"
    record_dir.mkdir(parents=True)
    with lab["factory"]() as session:
        repository = ResultDatasetRepository(session)
        for seq in (1, 2):
            dataset = _dataset(f"qa-dataset-{seq}", seq)
            entry = replace(dataset.payload.entries[0], result_id=f"{GROUP_ID}:{ROW_ID}:2", confirmed_group_id=GROUP_ID, group_label="G2", confirmed_row_id=ROW_ID)
            summary = replace(dataset.payload.summary_rows[0], group_label="G2")
            dataset = replace(dataset, confirmed_matrix_id=MATRIX_ID, confirmed_matrix_revision=2, source=source, payload=replace(dataset.payload, entries=(entry,), summary_rows=(summary,)))
            repository.create_dataset(dataset)
            report_path = report_dir / f"{LTR} Report Rev {seq}.docx"
            report_path.write_bytes(docx_bytes())
            report = replace(_report(f"qa-report-{seq}", seq, dataset.dataset_id), file_name=report_path.name, file_path=str(report_path), file_sha256=_hash(report_path), size_bytes=report_path.stat().st_size, confirmed_matrix_id=MATRIX_ID, base_report_revision_id="qa-report-1" if seq == 2 else None)
            repository.create_report_revision(report)
            record_path = record_dir / f"{LTR} Test Record Rev {seq}.docx"
            record = Document()
            record.add_heading(f"{LTR} Test Record", level=1)
            record.add_paragraph(f"Retained test-plan output revision {seq}")
            record.save(record_path)
            session.add(ProjectOutputRecordModel(output_record_id=f"qa-output-{seq}", project_id=PROJECT_ID, draft_id="ptpd-seed", draft_version=1, output_kind=ProjectOutputKind.TEST_RECORD_FORM.value, output_path=str(record_path), output_sha256=_hash(record_path), output_size_bytes=record_path.stat().st_size, source_context_signature="qa-source-context", status=ProjectOutputStatus.CURRENT.value if seq == 2 else ProjectOutputStatus.STALE.value, source=ProjectOutputSource.SYSTEM_GENERATED.value, created_at=AT, updated_at=AT))
        session.add(FileAssetModel(asset_id="qa-llcr-workbook", project_id=PROJECT_ID, asset_type="other", path=str(original), original_name=original.name, source_role="llcr_raw_workbook", sha256=source.sha256))
        session.commit()
    state = _ok(upload(lab, content=docx_bytes()))
    original_source_id = state["sources"][0]["source_id"]
    state = _ok(lab["client"].post(BASE + "/mappings", json={"expected_revision": state["revision"], "source_id": original_source_id, "kind": "specification", "source_locator": "QA revision table page 27", "label": "QA revision date", "value": "2026-06-18", "status": "confirmed", "reason": "Explicit QA decision"}))
    state = _ok(lab["client"].post(BASE + "/base", json={"source_id": original_source_id, "expected_revision": state["revision"]}))
    state = _ok(upload(lab, role="cr", name="QA-CR.csv", revision=state["revision"], content=b"SN,Voltage_mV,Current_A\nSN1,38.48,450\n"))
    cr_source = state["sources"][-1]
    step = next(step for step in state["context"]["steps"] if step["test_item"].casefold() == "llcr")
    request = {"expected_revision": state["revision"], "source_id": cr_source["source_id"], "kind": "cr", "sheet_name": "", "start_row": 2, "end_row": 2, "columns": {"sample": 1, "voltage": 2, "current": 3}, "unit": "mV", **{key: step[key] for key in ("group_id", "row_id", "step_token")}}
    preview = _ok(lab["client"].post(BASE + "/results/preview", json=request))
    state = _ok(lab["client"].post(BASE + "/results/confirm", json={**request, "preview_sha256": preview["preview_sha256"]}))
    working = ReportDraftWorkspace(lab["settings"].data_dir / "report_working_drafts")
    report_path = report_dir / f"{LTR} Report Rev 2.docx"
    draft1 = working.update(PROJECT_ID, LTR, report_path, lambda src, dst: shutil.copyfile(src, dst), lambda: None, matrix_id=MATRIX_ID)
    draft2 = working.update(PROJECT_ID, LTR, draft1.file_path, lambda src, dst: shutil.copyfile(src, dst), lambda: None, matrix_id=MATRIX_ID)
    return {"llcr_workbook": original, "authoring_state": state, "working_draft_revision": draft2.report_revision_id, "authoring_original_sha256": sha256(docx_bytes()).hexdigest()}


def _inspect(lab, business, seed):
    from sqlalchemy import inspect, text
    from backend.infrastructure.storage.repositories.result_dataset import ResultDatasetRepository
    from backend.infrastructure.storage.repositories.report_authoring import ReportAuthoringRepository
    counts = {}
    with lab["factory"]() as session:
        for name in inspect(session.bind).get_table_names():
            counts[name] = session.execute(text('SELECT COUNT(*) FROM "' + name + '"')).scalar_one()
        violations = [list(row) for row in session.execute(text("PRAGMA foreign_key_check"))]
        assert violations == [], violations
        dataset = ResultDatasetRepository(session).get_dataset("qa-dataset-2")
        assert str(dataset.payload.entries[0].measurements[0].value) == "0.198"
        assert dataset.payload.entries[0].measurements[0].unit == "mΩ"
        authoring = ReportAuthoringRepository(session).get(PROJECT_ID)
        assert authoring.mappings[0].value == "2026-06-18"
    files = {str(path.relative_to(lab["root"])): _hash(path) for path in lab["root"].rglob("*") if path.is_file() and path.suffix.lower() in {".docx", ".xlsx", ".csv", ".msg", ".txt"}}
    return {"project_id": PROJECT_ID, "ltr_number": LTR, "confirmed_matrix_id": MATRIX_ID, "confirmed_matrix_revision": 2, "llcr_group_id": GROUP_ID, "llcr_row_id": ROW_ID, "llcr_measurement": "0.198", "llcr_unit": "mΩ", "llcr_original_measurement": "0.248", "llcr_summary_fill": "FFFACD", "confirmed_fee_testing_total": "52", "point_profile_count": 4, "authoring_mapping_value": "2026-06-18", "authoring_cr_resistance_mohm": "0.0855111111111111", "business_folder": str(business), "source_counts": counts, "nonempty_tables": {name: n for name, n in counts.items() if n}, "file_sha256": files, "working_draft_revision": seed["working_draft_revision"], "foreign_key_violations": violations}


def _seed_remaining_families(lab, business):
    from backend.domain import ProjectMatrixDraftStatus
    from backend.infrastructure.storage.models import (
        MatrixFeePendingRebaseModel, ProjectCleanupAuditRecordModel,
        PublicDriveUploadFileRecordModel, ProjectPublicFolderWorkflowStateModel,
        ProjectPublicFolderWorkflowOperationModel, ProjectPublicFolderWorkflowFileRecordModel,
        LtrDuplicateResolutionTokenModel, FrozenFieldRevisionRequestModel,
    )
    from backend.infrastructure.storage.models_matrix_source import SourceMatrixDurationAuthorityModel
    from backend.infrastructure.storage.models_project_matrix_draft import (
        ProjectMatrixDraftRecordModel, ProjectMatrixDraftStepQuantityModel,
        ProjectMatrixDraftStepTextOverrideModel, ProjectMatrixDraftDurationAuthorityModel,
    )
    from backend.infrastructure.storage.models_confirmed_matrix_authority import ConfirmedMatrixStepTextOverrideModel
    from backend.infrastructure.storage.models_contact_measurement_plan_authority import MeasurementPlanImpactModel, MeasurementPlanRevisionModel
    from backend.infrastructure.storage.repositories import ProjectMatrixDraftRepository
    archive = lab["root"] / "public-history" / LTR
    archive.mkdir(parents=True)
    published = archive / f"{LTR} Report Rev 2.docx"
    local_report = lab["settings"].data_dir / "generated_test_reports" / published.name
    shutil.copyfile(local_report, published)
    with lab["factory"]() as session:
        draft_repo = ProjectMatrixDraftRepository(session)
        old = draft_repo.get("pmd-1")
        session.get(ProjectMatrixDraftRecordModel, "pmd-1").status = "superseded"
        session.flush()
        live = replace(old,
            record=replace(old.record, project_matrix_draft_id="qa-working-matrix", source_import_id=None, base_confirmed_matrix_id=MATRIX_ID, status=ProjectMatrixDraftStatus.DRAFT, updated_at=AT),
            groups=tuple(replace(g, draft_group_id="qa-live-"+g.draft_group_id, project_matrix_draft_id="qa-working-matrix") for g in old.groups),
            rows=tuple(replace(r, draft_row_id="qa-live-"+r.draft_row_id, project_matrix_draft_id="qa-working-matrix") for r in old.rows),
            cells=tuple(replace(c, draft_cell_id="qa-live-"+c.draft_cell_id, project_matrix_draft_id="qa-working-matrix", draft_group_id="qa-live-"+c.draft_group_id, draft_row_id="qa-live-"+c.draft_row_id) for c in old.cells))
        draft_repo.create_snapshot(live)
        session.add(ProjectMatrixDraftStepQuantityModel(draft_step_quantity_id="qa-draft-quantity", project_matrix_draft_id="qa-working-matrix", draft_group_id="qa-live-pmdg-2", draft_row_id="qa-live-pmdr-2", step_sequence=2, step_suffix_note="", raw_token="2", test_points_per_sample="1", readings_per_point="1", contact_points_per_sample="1", source="matrix_step_override", review_required=False, updated_at=AT))
        session.add(ProjectMatrixDraftStepTextOverrideModel(project_matrix_draft_id="qa-working-matrix", draft_group_id="qa-live-pmdg-2", draft_row_id="qa-live-pmdr-2", step_sequence=2, step_suffix_note="", description="QA draft LLCR description", requirement="≤0.25mΩ"))
        session.add(ConfirmedMatrixStepTextOverrideModel(confirmed_matrix_id=MATRIX_ID, confirmed_group_id=GROUP_ID, confirmed_row_id=ROW_ID, step_sequence=2, step_suffix_note="", description="QA confirmed LLCR description", requirement="≤0.25mΩ"))
        duration = dict(step_sequence=2, step_suffix_note="", duration_value=Decimal("1.5"), duration_unit="hours", normalized_hours=Decimal("1.5"), source_kind="manual_edit", source_field="duration", source_import_id=old.record.source_import_id, source_fingerprint="qa-duration-source", lineage_fingerprint="qa-duration-lineage", authority_revision="1", status="usable", created_at=AT, updated_at=AT)
        session.add(SourceMatrixDurationAuthorityModel(source_duration_authority_id="qa-source-duration", source_snapshot_id=old.record.source_snapshot_id, source_group_snapshot_id=old.groups[1].source_group_snapshot_id, source_row_snapshot_id=old.rows[1].source_row_snapshot_id, **duration))
        session.add(ProjectMatrixDraftDurationAuthorityModel(draft_duration_authority_id="qa-draft-duration", project_matrix_draft_id="qa-working-matrix", draft_group_id="qa-live-pmdg-2", draft_row_id="qa-live-pmdr-2", **duration))
        session.add(MatrixFeePendingRebaseModel(pending_rebase_id="qa-fee-rebase", project_id=PROJECT_ID, project_matrix_draft_id="qa-working-matrix", base_confirmed_matrix_id=MATRIX_ID, base_confirmed_revision=2, fee_rule_version_id="fee_rules_v2026_06_03", matrix_draft_payload_signature="qa-editable-matrix-fingerprint", generation=1, payload_json=json.dumps({"rows": [{"confirmed_group_id": GROUP_ID, "confirmed_row_id": ROW_ID, "draft_group_id": "qa-live-pmdg-2", "draft_row_id": "qa-live-pmdr-2", "step_token": "2", "testing_fee": "25"}]}), created_at=AT, updated_at=AT))
        plan_draft = session.get(MeasurementPlanRevisionModel, lab["measurement_draft_id"])
        session.add(MeasurementPlanImpactModel(measurement_plan_impact_id="qa-resolved-impact", measurement_plan_root_id=plan_draft.measurement_plan_root_id, editable_revision_id=plan_draft.measurement_plan_revision_id, stable_target_key=lab["measurement_target_key"], impact_subject_key=lab["measurement_target_key"], impact_identity_key="cmp-impact:v1|qa-reviewed-mapping", category="coverage_changed", severity="review", before_evidence_fingerprint="qa-before-coverage", after_evidence_fingerprint="qa-after-coverage", resolution_state="resolved", reason="Operator retained the reviewed signal contact", created_at=AT))
        session.add(ProjectCleanupAuditRecordModel(cleanup_id="qa-cleanup-history", project_id=PROJECT_ID, cleanup_type="legacy_registry_review", previous_status="draft", new_status="ltr_registered", reason="Historic duplicate reviewed; source retained", operator="QA Operator", created_at=AT, details_json=json.dumps({"source_asset_ids": ["asset-1"], "deleted_files": []})))
        session.add(ProjectPublicFolderWorkflowStateModel(project_id=PROJECT_ID, auto_sync_enabled=True, sync_locked=False, submitted_at=AT, submit_operation_id="qa-public-completed", last_sync_operation_id="qa-public-completed", created_at=AT, updated_at=AT))
        session.add(ProjectPublicFolderWorkflowOperationModel(operation_id="qa-public-completed", project_id=PROJECT_ID, operation_type="submit", status="completed", preview_hash="qa-public-review", requested_at=AT, started_at=AT, completed_at=AT, operator="QA Operator", public_root=str(archive.parent), public_root_class="qa-isolated-public-history", public_folder_year=2026, year_source="ltr", local_official_folder_path=str(business), public_open_path=str(archive), target_path=str(archive), counts_json=json.dumps({"copied": 1}), blockers_json="[]", warnings_json="[]", conflicts_json="[]", snapshot_json=json.dumps({"files": [{"relative_path": published.name, "sha256": _hash(published)}]}), metadata_json=json.dumps({"historic_only": True, "no_target_publish_permission": True})))
        session.add(ProjectPublicFolderWorkflowFileRecordModel(project_id=PROJECT_ID, relative_path=published.name, public_path=str(published), local_fingerprint=_hash(local_report), public_fingerprint=_hash(published), updated_at=AT, operation_id="qa-public-completed"))
        session.add(PublicDriveUploadFileRecordModel(project_id=PROJECT_ID, relative_path=published.name, public_path=str(published), local_fingerprint=_hash(local_report), public_fingerprint=_hash(published), uploaded_at=AT, operation_id="qa-public-completed"))
        session.add(LtrDuplicateResolutionTokenModel(token_id="qa-used-historic-token", ltr_number=LTR, existing_ltr_id="qa-ltr-history", existing_project_id=PROJECT_ID, current_case_id="case-1", current_project_id=PROJECT_ID, conflict_fingerprint="qa-reviewed-ownership", expires_at="2026-01-01T00:00:00Z", used_at=AT, created_at="2025-12-31T00:00:00Z", created_by="QA Operator", metadata_json=json.dumps({"historical": True})))
        session.add(FrozenFieldRevisionRequestModel(request_id="qa-completed-frozen-request", intake_case_id="case-1", project_id=PROJECT_ID, ltr_record_id="qa-ltr-current", ltr_number=LTR, status="completed", requested_by="QA Operator", reason="Historic user-confirmed revision", field_changes_json=json.dumps({"revision": {"before": "A", "after": "B"}}), created_at=AT, updated_at=AT))
        session.commit()
    # A recognizable fake credential in a local-only database is negative-test
    # data: project export must exclude this machine identity file altogether.
    import sqlite3
    identity = sqlite3.connect(lab["settings"].data_dir / "local_identity.sqlite3")
    identity.execute("CREATE TABLE qa_credentials (secret TEXT NOT NULL)")
    identity.execute("INSERT INTO qa_credentials VALUES (?)", ("QA-NOT-A-REAL-CREDENTIAL-SENTINEL",))
    identity.commit()
    identity.close()


@contextmanager
def seed_full_graph(root):
    """Yield {lab, expected, project_root}; caller must close before rebinding B."""
    from pytest import MonkeyPatch
    patch = MonkeyPatch()
    root = _isolated_environment(root, patch)
    # Imports now occur only after explicit sandbox binding.
    from tests.integration.test_internal_report_generation_api import lab as lab_fixture
    generator = lab_fixture.__wrapped__(root, patch)
    try:
        lab = next(generator)
        from backend.api import dependencies as deps
        from backend.api.main import app
        lab["settings"] = replace(lab["settings"], library_root=root)
        app.dependency_overrides[deps.get_settings] = lambda: lab["settings"]
        local, business, source_book, _ = _seed_owned_sources(lab)
        _seed_workspace(lab, local, business, source_book)
        _seed_matrix_and_plan_history(lab)
        _seed_point_and_fee_history(lab)
        seed = _seed_results_and_reports(lab, business)
        _seed_remaining_families(lab, business)
        result = _ok(lab["client"].post("/api/application-forms/qa-form/precheck/run"))
        assert result["issues"], "The deliberately incomplete application must produce review issues"
        workspace_preview = lab["client"].get("/api/projects/P1/official-workspace/preview")
        assert workspace_preview.status_code == 200, workspace_preview.text
        expected = _inspect(lab, business, seed)
        yield {"lab": lab, "expected": expected, "project_root": local, "business_folder": business}
    finally:
        try:
            next(generator)
        except StopIteration:
            pass
        finally:
            generator.close()
            patch.undo()
