from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path

from docx import Document

from backend.infrastructure.office.equipment_id_document_reader import (
    EquipmentIdDocumentReader,
)


def test_reads_paragraphs_and_table_cells_in_order_and_deduplicates(tmp_path: Path) -> None:
    source = tmp_path / "EquipmentID.docx"
    document = Document()
    document.add_paragraph("DG-Q-0033")
    document.add_paragraph("Customer fixture L-0002")
    document.add_paragraph("dg-q-0033")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "DG-Q-0103"
    table.cell(0, 1).text = "Customer fixture L-0002"
    document.save(source)

    result = EquipmentIdDocumentReader().read(source)

    assert result.source_path == source
    assert result.references == (
        "DG-Q-0033",
        "Customer fixture L-0002",
        "DG-Q-0103",
    )


def test_rejects_an_empty_equipment_id_document(tmp_path: Path) -> None:
    source = tmp_path / "EquipmentID.docx"
    Document().save(source)

    try:
        EquipmentIdDocumentReader().read(source)
    except ValueError as exc:
        assert "does not contain equipment references" in str(exc)
    else:
        raise AssertionError("Expected an empty EquipmentID.docx to be rejected")


def test_reads_a_password_protected_equipment_document_through_readable_copy(
    tmp_path: Path,
) -> None:
    protected = tmp_path / "EquipmentID.docx"
    protected.write_bytes(b"encrypted")
    readable = tmp_path / "readable.docx"
    document = Document()
    document.add_paragraph("DG-Q-0033")
    document.save(readable)
    calls: list[Path] = []

    class _PackageGateway:
        @contextmanager
        def readable_copy(self, source: Path):
            calls.append(source)
            yield readable

    result = EquipmentIdDocumentReader(
        protected_package_gateway=_PackageGateway(),
    ).read(protected)

    assert calls == [protected]
    assert result.references == ("DG-Q-0033",)
