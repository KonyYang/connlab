from io import BytesIO

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment

from backend.infrastructure.office.fee_form_import_gateway import read_fee_form


def workbook_bytes():
    book = Workbook()
    sheet = book.active
    sheet.append(["Testing Prices"])
    sheet.append([])
    sheet.append([])
    sheet.append(["Group", "Man-hour", "Description", "Unit price", "Unit Type", "Units", "Base fee", "Discount", "Testing fee"])
    sheet.append(["Aa", 0.5, "Sample preparation", 50, "per sample", 5, 0, 0, "=D5*F5"])
    sheet.append([None, 1, "Visual Inspection", 10, "per photo", 3, 0, 0.1, "=D6*F6"])
    sheet["I6"].comment = Comment("Reuse this note", "Lab")
    sheet.append([None, 1, "Report preparation", 50, "per report", 1, 0, 0])
    sheet.append(["条件确认", 2])
    sheet.append(["Total"])
    sheet.append(["Grand Cost"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_read_exported_fields_without_importing_totals_or_executing_formulas():
    result = read_fee_form(workbook_bytes(), "fee.xlsx")
    assert len(result["rows"]) == 3
    assert result["rows"][1] == {
        "group": "Aa", "description": "Visual Inspection", "rowKind": "matrix_step",
        "values": {"spendTime": "1", "unitPrice": "10", "unitType": "per photo", "units": "3",
                   "baseFee": "0", "discount": "10", "notes": "Reuse this note"},
    }
    assert result["rows"][2]["rowKind"] == "manual_trailing"


def test_reject_unrelated_workbook_and_macro_extension():
    book = Workbook()
    stream = BytesIO()
    book.save(stream)
    book.close()
    with pytest.raises(ValueError, match="Fee Form"):
        read_fee_form(stream.getvalue(), "other.xlsx")
    with pytest.raises(ValueError, match="xls"):
        read_fee_form(workbook_bytes(), "fee.xlsm")


def test_accept_actual_export_header_labels_and_numeric_group_labels():
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(workbook_bytes()))
    headers = ['Group', 'Time\n(Unit: Hour)', 'Description', 'Unit Price(A)', '', 'Units(B)',
               'Base Fee(C)', 'Price \nPercent Off (d)', 'Testing Fee(RMB)\n=A*B*(1-d)+C']
    for column, value in enumerate(headers, 1):
        book.active.cell(4, column).value = value
    book.active["A5"] = 2.0
    stream = BytesIO()
    book.save(stream)
    book.close()
    result = read_fee_form(stream.getvalue(), "fee.xlsx")
    assert result["rows"][0]["group"] == "2"
    assert result["rows"][1]["group"] == "2"


@pytest.mark.parametrize("value", ["=1+1", "-1", "NaN", "1e-999999", "1e999999"])
def test_reject_invalid_editable_price_with_row_context(value):
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(workbook_bytes()))
    book.active["D6"] = value
    stream = BytesIO()
    book.save(stream)
    book.close()
    with pytest.raises(ValueError, match="row 6: unitPrice"):
        read_fee_form(stream.getvalue(), "fee.xlsx")


def test_reject_ambiguous_export_sheets_instead_of_selecting_arbitrarily():
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(workbook_bytes()))
    book.copy_worksheet(book.active)
    stream = BytesIO()
    book.save(stream)
    book.close()
    with pytest.raises(ValueError, match="exactly one"):
        read_fee_form(stream.getvalue(), "fee.xlsx")
