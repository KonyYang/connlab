from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from hashlib import sha256
from types import SimpleNamespace
import shutil
import msoffcrypto
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from backend.infrastructure.office import customer_report_document_gateway as gateway_module
from backend.infrastructure.office.customer_report_document_gateway import (
    CustomerReportDocumentGateway,
)
from backend.infrastructure.office.office_protected_document_gateway import (
    WordPackageProtectionState,
)
from backend.shared.office_document_password import OFFICE_DOCUMENT_PASSWORD
from backend.application.customer_report_projection_service import CustomerReportReadError


def test_reads_lineage_from_encrypted_report_without_changing_original(tmp_path):
    plain = tmp_path / "plain.docx"
    Document().save(plain)
    with ZipFile(plain, "a") as package:
        package.writestr("docProps/custom.xml", (
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
            'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            '<property name="ConnLabSourceReportSHA256"><vt:lpwstr>'
            + "a" * 64 + '</vt:lpwstr></property></Properties>'
        ))
    protected = tmp_path / "protected.docx"
    with plain.open("rb") as source, protected.open("wb") as output:
        msoffcrypto.OfficeFile(source).encrypt(OFFICE_DOCUMENT_PASSWORD, output)
    before = protected.read_bytes(), protected.stat().st_mtime_ns
    assert CustomerReportDocumentGateway().read_source_report_sha256(protected) == "a" * 64
    assert (protected.read_bytes(), protected.stat().st_mtime_ns) == before


def test_invalid_lineage_xml_is_not_reported_as_missing_fingerprint(tmp_path):
    report = tmp_path / "report.docx"
    Document().save(report)
    with ZipFile(report, "a") as package:
        package.writestr("docProps/custom.xml", b"<broken")
    with pytest.raises(CustomerReportReadError, match="source information"):
        CustomerReportDocumentGateway().read_source_report_sha256(report)


def test_protected_report_with_unknown_password_is_preserved(tmp_path):
    plain = tmp_path / "plain.docx"
    Document().save(plain)
    report = tmp_path / "protected.docx"
    with plain.open("rb") as source, report.open("wb") as output:
        msoffcrypto.OfficeFile(source).encrypt("isolated-test-password", output)
    before = report.read_bytes(), report.stat().st_mtime_ns
    with pytest.raises(CustomerReportReadError, match="password protection"):
        CustomerReportDocumentGateway().read_source_report_sha256(report)
    assert (report.read_bytes(), report.stat().st_mtime_ns) == before


def test_customer_report_gateway_rejects_missing_internal_source(tmp_path: Path) -> None:
    template = tmp_path / "E-4515.docx"
    template.write_bytes(b"template")

    with pytest.raises(FileNotFoundError, match="Internal report source"):
        CustomerReportDocumentGateway().generate_customer_report(
            source_path=tmp_path / "missing.docx",
            template_path=template,
            output_path=tmp_path / "customer.docx",
        )


def test_customer_report_gateway_never_overwrites_source(tmp_path: Path) -> None:
    source = tmp_path / "internal.docx"
    template = tmp_path / "E-4515.docx"
    source.write_bytes(b"internal")
    template.write_bytes(b"template")

    with pytest.raises(ValueError, match="must not replace"):
        CustomerReportDocumentGateway().generate_customer_report(
            source_path=source,
            template_path=template,
            output_path=source,
        )
    assert source.read_bytes() == b"internal"


def test_generated_customer_report_records_current_internal_fingerprint(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    _minimal_document(source, customer=False)
    template = tmp_path / "E-4515_F.docx"
    _minimal_document(template, customer=True)
    output = tmp_path / "customer.docx"
    monkeypatch.setattr(
        "backend.infrastructure.office.customer_report_document_gateway._generate_with_word",
        lambda source_path, target_path: None,
    )

    gateway = CustomerReportDocumentGateway()
    gateway.generate_customer_report(
        source_path=source,
        template_path=template,
        output_path=output,
    )

    assert gateway.read_source_report_sha256(output) == sha256(source.read_bytes()).hexdigest()
    assert gateway.read_source_report_sha256(template) is None


def test_customer_report_gateway_retries_transient_word_file_lock(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    _minimal_document(source, customer=False)
    template = tmp_path / "E-4515_F.docx"
    _minimal_document(template, customer=True)
    output = tmp_path / "customer.docx"
    attempts = 0
    original_write = gateway_module._write_source_report_sha256

    def write_after_word_releases_file(path: Path, fingerprint: str) -> None:
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            lock_error = PermissionError(13, "file is being used", str(path))
            lock_error.winerror = 32
            raise lock_error
        original_write(path, fingerprint)

    monkeypatch.setattr(gateway_module, "_generate_with_word", lambda *_args: None)
    monkeypatch.setattr(
        gateway_module,
        "_write_source_report_sha256",
        write_after_word_releases_file,
    )

    CustomerReportDocumentGateway(
        transient_retry_delay=lambda _seconds: None,
    ).generate_customer_report(
        source_path=source,
        template_path=template,
        output_path=output,
    )

    assert attempts == 2
    assert output.is_file()


def test_customer_report_gateway_does_not_retry_regular_permission_error(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    _minimal_document(source, customer=False)
    template = tmp_path / "E-4515_F.docx"
    _minimal_document(template, customer=True)
    attempts = 0

    def reject_write(_path: Path, _fingerprint: str) -> None:
        nonlocal attempts
        attempts += 1
        raise PermissionError(13, "access denied")

    monkeypatch.setattr(gateway_module, "_generate_with_word", lambda *_args: None)
    monkeypatch.setattr(gateway_module, "_write_source_report_sha256", reject_write)

    with pytest.raises(PermissionError, match="access denied"):
        CustomerReportDocumentGateway(
            transient_retry_delay=lambda _seconds: None,
        ).generate_customer_report(
            source_path=source,
            template_path=template,
            output_path=tmp_path / "customer.docx",
        )

    assert attempts == 1


def test_customer_report_restores_template_password_protection_after_audit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    _minimal_document(source, customer=False)
    template = tmp_path / "protected-E-4515_F.docx"
    _minimal_document(template, customer=True)
    output = tmp_path / "customer.docx"
    calls: list[tuple[str, Path]] = []

    class _ProtectedPackageGateway:
        def stage_editable_copy(self, source_path: Path, output_path: Path):
            calls.append(("stage", source_path))
            shutil.copy2(source_path, output_path)
            return WordPackageProtectionState(was_password_protected=True)

        def restore_password_protection(self, editable_path: Path, state) -> None:
            assert state.was_password_protected is True
            assert editable_path.is_file()
            calls.append(("restore", editable_path))

    monkeypatch.setattr(
        "backend.infrastructure.office.customer_report_document_gateway._generate_with_word",
        lambda source_path, target_path: None,
    )

    CustomerReportDocumentGateway(
        protected_package_gateway=_ProtectedPackageGateway(),
    ).generate_customer_report(
        source_path=source,
        template_path=template,
        output_path=output,
    )

    assert calls[0] == ("stage", template)
    assert calls[1][0] == "restore"
    assert output.is_file()


def test_customer_report_gateway_uses_short_word_working_copy_names(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    _minimal_document(source, customer=False)
    template = tmp_path / "E-4515_F.docx"
    _minimal_document(template, customer=True)
    output = tmp_path / ("DL-001-CR " + "Long qualification report name " * 5 + ".docx")
    opened_targets: list[Path] = []

    def capture_word_target(_source_path: Path, target_path: Path) -> None:
        opened_targets.append(Path(target_path))

    monkeypatch.setattr(
        "backend.infrastructure.office.customer_report_document_gateway._generate_with_word",
        capture_word_target,
    )

    CustomerReportDocumentGateway().generate_customer_report(
        source_path=source,
        template_path=template,
        output_path=output,
    )

    assert len(opened_targets) == 1
    assert opened_targets[0].parent == output.parent
    assert len(opened_targets[0].name) <= 64


def test_customer_report_gateway_reports_word_and_verification_stages(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    _minimal_document(source, customer=False)
    template = tmp_path / "E-4515_F.docx"
    _minimal_document(template, customer=True)
    stages: list[str] = []

    def generate_in_word(_source_path, _target_path, *, progress) -> None:
        progress("opening_word")
        progress("copying_content")

    monkeypatch.setattr(gateway_module, "_generate_with_word", generate_in_word)

    CustomerReportDocumentGateway().generate_customer_report(
        source_path=source,
        template_path=template,
        output_path=tmp_path / "customer.docx",
        progress=stages.append,
    )

    assert stages == [
        "preparing_template",
        "opening_word",
        "copying_content",
        "verifying_output",
        "protecting_output",
    ]


def test_customer_report_gateway_rejects_internal_disclosure_in_body(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    _minimal_document(source, customer=False)
    template = tmp_path / "E-4515_F.docx"
    _minimal_document(template, customer=True)

    def retain_internal_disclosure(_source_path: Path, target_path: Path) -> None:
        document = Document(target_path)
        document.add_paragraph(
            "This Laboratory Test Report shall not be reproduced except in full "
            "unless written permission is received from the Laboratory Manager."
        )
        document.save(target_path)

    monkeypatch.setattr(
        "backend.infrastructure.office.customer_report_document_gateway._generate_with_word",
        retain_internal_disclosure,
    )

    with pytest.raises(ValueError, match="internal-only disclosures"):
        CustomerReportDocumentGateway().generate_customer_report(
            source_path=source,
            template_path=template,
            output_path=tmp_path / "customer.docx",
        )


def test_customer_report_gateway_rejects_template_report_number_placeholder(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    _minimal_document(source, customer=False)
    template = tmp_path / "E-4515_F.docx"
    _minimal_document(template, customer=True)
    template_document = Document(template)
    template_document.sections[0].first_page_header.tables[0].cell(2, 0).text = (
        "WW-XXXX-YY-ZZZ-CR"
    )
    template_document.save(template)
    monkeypatch.setattr(
        "backend.infrastructure.office.customer_report_document_gateway._generate_with_word",
        lambda source_path, target_path: None,
    )

    with pytest.raises(ValueError, match="placeholder"):
        CustomerReportDocumentGateway().generate_customer_report(
            source_path=source,
            template_path=template,
            output_path=tmp_path / "customer.docx",
        )


@pytest.mark.parametrize("headers,values,expected", [
    (
        ["Description", "Part #", "Lot", "Base Material", "Contact Plating",
         "Lubricant", "Housing Material"],
        ["Connector", "P-1", "L-1", "C7025R-H", "Gold PE123", "PE456", "LCP E130i"],
        ["Connector", "P-1", "L-1", "Copper Alloy", "Gold Lubricant", "PE456",
         "High Temperature Resin"],
    ),
    (
        ["Description", "Part #", "Base\nMaterial", "Contact\nPlating", "Lubricant",
         "Housing\nMaterial"],
        ["Connector", "P-1", "CuNiSi C7025", "Gold PE123", "PE456", "PA9T G1300"],
        ["Connector", "P-1", "Copper Alloy", "Gold Lubricant", "PE456",
         "High Temperature Resin"],
    ),
    (
        ["Description", "Part #", "Housing Material", "Lot", "Contact Plating", "Base Mat."],
        ["Connector", "P-1", "PPS 40% GF", "L-1", "Gold", "Phosphor Bronze C5191"],
        ["Connector", "P-1", "High Temperature Resin", "L-1", "Gold", "Copper Alloy"],
    ),
    (
        ["Description", "Part #", "Lot#", "Base Mat’l", "Contact\n Plating",
         "Contact Lubricant", "Housing Material"],
        ["Connector", "P-1", "L-1", "HP: C18070;", "Gold", "PE456", "HTN"],
        ["Connector", "P-1", "L-1", "Copper Alloy", "Gold", "PE456", "High Temperature Resin"],
    ),
    (
        ["Description", "Part #", "Base Material", "Housing Material"],
        ["Connector", "P-1", " Copper\nAlloy ", " High Temperature Resin "],
        ["Connector", "P-1", "Copper Alloy", "High Temperature Resin"],
    ),
])
def test_customer_report_generalizes_sample_materials_without_changing_sources(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, headers, values, expected,
) -> None:
    source = tmp_path / "internal.docx"
    template = tmp_path / "template.docx"
    output = tmp_path / "customer.docx"
    _minimal_document(source, customer=False)
    _minimal_document(template, customer=True)
    document = Document(source)
    sample = document.add_table(rows=2, cols=len(headers))
    for column, (header, value) in enumerate(zip(headers, values)):
        sample.cell(0, column).text = header
        sample.cell(1, column).text = value
    document.paragraphs[2]._p.addnext(sample._tbl)
    document.save(source)
    before = source.read_bytes(), template.read_bytes()

    monkeypatch.setattr(gateway_module, "_generate_with_word", _generate_samples_with_fake_word)
    CustomerReportDocumentGateway().generate_customer_report(
        source_path=source, template_path=template, output_path=output,
    )

    assert [cell.text for cell in Document(output).tables[0].rows[1].cells] == expected
    assert (source.read_bytes(), template.read_bytes()) == before


@pytest.mark.parametrize("base,housing", [
    ("C7025R-H", "High Temperature Resin"),
    ("Copper alloy", "High Temperature Resin"),
    ("CuSn6", "High Temperature Resin"),
    ("Copper Alloy", "LCP E130i"),
    (" Copper Alloy ", "High Temperature Resin"),
])
def test_customer_report_rejects_unredacted_sample_materials_before_publication(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, base, housing,
) -> None:
    source = tmp_path / "internal.docx"
    template = tmp_path / "template.docx"
    output = tmp_path / "customer.docx"
    _minimal_document(source, customer=False)
    _minimal_document(template, customer=True)
    document = Document(template)
    sample = document.add_table(rows=2, cols=4)
    for column, (header, value) in enumerate(zip(
        ["Description", "Part #", "Base Material", "Housing Material"],
        ["Connector", "P-1", base, housing],
    )):
        sample.cell(0, column).text = header
        sample.cell(1, column).text = value
    document.paragraphs[2]._p.addnext(sample._tbl)
    document.save(template)
    before = source.read_bytes(), template.read_bytes()
    # Simulate Word saving a legacy/unprocessed sample table; the real package audit must reject it.
    monkeypatch.setattr(gateway_module, "_generate_with_word", lambda *_args: None)

    with pytest.raises(ValueError, match="sample material"):
        CustomerReportDocumentGateway().generate_customer_report(
            source_path=source, template_path=template, output_path=output,
        )

    assert not output.exists()
    assert (source.read_bytes(), template.read_bytes()) == before
    assert not list(tmp_path.glob(".customer-report.*"))


@pytest.mark.parametrize("base,housing", [
    ("Stainless Steel SUS304", "N/A"),
    ("Aluminum 6061", "NA"),
    ("", ""),
    ("N/A", "Not Applicable"),
    ("-", "—"),
    ("Copper Alloy", "High Temperature Resin"),
])
def test_customer_report_preserves_non_copper_material_and_absent_housing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, base, housing,
) -> None:
    source = tmp_path / "internal.docx"
    template = tmp_path / "template.docx"
    output = tmp_path / "customer.docx"
    _minimal_document(source, customer=False)
    _minimal_document(template, customer=True)
    document = Document(source)
    sample = document.add_table(rows=2, cols=4)
    for column, (header, value) in enumerate(zip(
        ["Description", "Part #", "Base Material", "Housing Material"],
        ["Connector", "P-1", base, housing],
    )):
        sample.cell(0, column).text = header
        sample.cell(1, column).text = value
    document.paragraphs[2]._p.addnext(sample._tbl)
    document.save(source)
    monkeypatch.setattr(gateway_module, "_generate_with_word", _generate_samples_with_fake_word)

    CustomerReportDocumentGateway().generate_customer_report(
        source_path=source, template_path=template, output_path=output,
    )

    assert [cell.text for cell in Document(output).tables[0].rows[1].cells] == [
        "Connector", "P-1", base, housing,
    ]


def test_customer_report_filters_non_product_samples_and_keeps_other_sections(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "internal.docx"
    template = tmp_path / "template.docx"
    output = tmp_path / "customer.docx"
    _minimal_document(source, customer=False)
    _minimal_document(template, customer=True)
    document = Document(source)
    sample = document.add_table(rows=6, cols=7)
    for column, header in enumerate([
        "Description", "Part #", "Lot#", "Base Mat’l", "Contact\n Plating",
        "Contact Lubricant", "Housing Material",
    ]):
        sample.cell(0, column).text = header
    rows = [
        ["Plug", "P-1", "L-1", "HP: C18070;", "Gold PE123", "PE456", "PA9T"],
        ["Receptacle", "P-2", "L-2", "HP: C19210", "Gold", "PE456", "HTN"],
        ["Test PCB", "Board-1", "L-3", "2oz. Cu", "Gold", "N/A", "FR4"],
        ["Test Fixture", "Fixture-1", "L-4", "Steel", "N/A", "N/A", "N/A"],
        ["Adapter", "Adapter-1", "L-5", "C7025", "N/A", "N/A", "PPS"],
    ]
    for row_index, values in enumerate(rows, start=1):
        for column, value in enumerate(values):
            sample.cell(row_index, column).text = value
    document.paragraphs[2]._p.addnext(sample._tbl)
    unrelated = deepcopy(sample._tbl)
    document.paragraphs[5]._p.addnext(unrelated)
    document.save(source)
    before = source.read_bytes(), template.read_bytes()
    monkeypatch.setattr(gateway_module, "_generate_with_word", _generate_samples_with_fake_word)

    CustomerReportDocumentGateway().generate_customer_report(
        source_path=source, template_path=template, output_path=output,
    )

    result = Document(output)
    assert [[cell.text for cell in row.cells] for row in result.tables[0].rows[1:]] == [
        ["Plug", "P-1", "L-1", "Copper Alloy", "Gold Lubricant", "PE456", "High Temperature Resin"],
        ["Receptacle", "P-2", "L-2", "Copper Alloy", "Gold", "PE456", "High Temperature Resin"],
    ]
    assert [[cell.text for cell in row.cells] for row in result.tables[1].rows[1:]] == rows
    assert (source.read_bytes(), template.read_bytes()) == before


@pytest.mark.parametrize("housing,description_width,housing_width,should_widen", [
    ("PA9T", 122, 53, True),
    ("HTN", 88, 53, False),
    ("HTN", 122, 95, False),
    ("N/A", 122, 53, False),
])
def test_customer_report_allocates_space_for_generic_housing_without_widening_table(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
    housing, description_width, housing_width, should_widen,
) -> None:
    source = tmp_path / "internal.docx"
    template = tmp_path / "template.docx"
    output = tmp_path / "customer.docx"
    _minimal_document(source, customer=False)
    _minimal_document(template, customer=True)
    document = Document(source)
    sample = document.add_table(rows=2, cols=7)
    widths = [description_width, 65, 43, 65, 65, 65, housing_width]
    for column, (header, value, width) in enumerate(zip(
        ["Description", "Part #", "Lot#", "Base Mat’l", "Contact Plating",
         "Contact Lubricant", "Housing Material"],
        ["Connector", "P-1", "L-1", "HP: C18070;", "Gold", "PE456", housing],
        widths,
    )):
        sample.cell(0, column).text = header
        sample.cell(1, column).text = value
        sample.columns[column].width = Pt(width)
        for row in sample.rows:
            row.cells[column].width = Pt(width)
    document.paragraphs[2]._p.addnext(sample._tbl)
    unrelated = deepcopy(sample._tbl)
    document.paragraphs[5]._p.addnext(unrelated)
    document.save(source)
    before = source.read_bytes(), template.read_bytes()
    monkeypatch.setattr(gateway_module, "_generate_with_word", _generate_samples_with_fake_word)

    CustomerReportDocumentGateway().generate_customer_report(
        source_path=source, template_path=template, output_path=output,
    )

    result = Document(output)
    after = [column.width.pt for column in result.tables[0].columns]
    assert sum(after) == pytest.approx(sum(widths))
    assert after[1:6] == pytest.approx(widths[1:6])
    if should_widen:
        assert after[6] > widths[6]
        assert after[0] < widths[0]
    else:
        assert after == pytest.approx(widths)
    assert [column.width.pt for column in result.tables[1].columns] == pytest.approx(widths)
    assert (source.read_bytes(), template.read_bytes()) == before


def _generate_samples_with_fake_word(source_path, target_path):
    """Copy raw source body as Word would, then use production sanitization against mutable cells."""
    source = Document(source_path)
    target = Document(target_path)
    for element in list(target.element.body):
        if element.tag != qn("w:sectPr"):
            target.element.body.remove(element)
    for element in source.element.body:
        if element.tag != qn("w:sectPr"):
            target.element.body.sectPr.addprevious(deepcopy(element))
    gateway_module._sanitize_customer_body(_WordDocument(target))
    target.save(target_path)


class _WordCollection:
    def __init__(self, values):
        self.values = values

    @property
    def Count(self):
        return len(self.values)

    def __call__(self, index):
        return self.values[index - 1]


class _WordCell:
    def __init__(self, cell):
        self.cell = cell
        self.Range = self

    @property
    def Text(self):
        return self.cell.text + "\r\x07"

    @Text.setter
    def Text(self, value):
        self.cell.text = value


class _WordRow:
    def __init__(self, row):
        self.row = row
        self.Cells = _WordCollection([_WordCell(cell) for cell in row.cells])

    def Delete(self):
        self.row._tr.getparent().remove(self.row._tr)


class _WordColumn:
    def __init__(self, table, index):
        self.table = table
        self.index = index

    @property
    def Width(self):
        return self.table.columns[self.index].width.pt

    def SetWidth(self, width, ruler_style):
        assert ruler_style == 0  # Word's wdAdjustNone keeps the other column widths unchanged.
        self.table.columns[self.index].width = Pt(width)
        for row in self.table.rows:
            row.cells[self.index].width = Pt(width)


class _WordTable:
    def __init__(self, table, document):
        self.table = table
        self.document = document
        self.Range = self

    @property
    def Rows(self):
        return _WordCollection([_WordRow(row) for row in self.table.rows])

    @property
    def Columns(self):
        return _WordCollection([
            _WordColumn(self.table, index) for index in range(len(self.table.columns))
        ])

    @property
    def Duplicate(self):
        return self

    @property
    def Text(self):
        return "\r".join(cell.text for row in self.table.rows for cell in row.cells)

    @property
    def Start(self):
        return self.document.positions()[self.table._tbl][0]

    @property
    def Cells(self):
        return _WordCollection([
            _WordCell(cell) for row in self.table.rows for cell in row.cells
        ])


class _WordRange:
    def __init__(self, document, start=0, end=None):
        self.document = document
        self.Start = start
        self.End = len(document.text()) if end is None else end
        self.Find = SimpleNamespace(ClearFormatting=lambda: None, Execute=self.find)

    def find(self):
        found = self.document.text().find(self.Find.Text, self.Start, self.End)
        if found == -1:
            return False
        self.Start, self.End = found, found + len(self.Find.Text)
        return True

    @property
    def Text(self):
        return self.document.text()[self.Start:self.End]

    @property
    def Duplicate(self):
        return _WordRange(self.document, self.Start, self.End)

    def Paragraphs(self, _index):
        text = self.document.text()
        start = text.rfind("\r", 0, self.Start) + 1
        end = text.find("\r", self.End)
        return SimpleNamespace(Range=_WordRange(self.document, start, end + 1))

    @property
    def Tables(self):
        positions = self.document.positions()
        return _WordCollection([
            _WordTable(table, self.document) for table in self.document.document.tables
            if self.Start <= positions[table._tbl][0] < self.End
        ])


class _WordDocument:
    def __init__(self, document):
        self.document = document
        self.Shapes = _WordCollection([])

    def positions(self):
        result = {}
        position = 0
        for element in self.document.element.body:
            text = (
                "".join(node.text or "" for node in element.iter(qn("w:t")))
                if element.tag == qn("w:p") else "table"
            )
            result[element] = position, text + "\r"
            position += len(text) + 1
        return result

    def text(self):
        return "".join(text for _position, text in self.positions().values())

    @property
    def Content(self):
        return _WordRange(self)

    @property
    def Tables(self):
        return _WordCollection([_WordTable(table, self) for table in self.document.tables])

    def Range(self, start, end):
        return _WordRange(self, start, end)


def test_heading_index_performs_one_lookup_per_heading_and_reuses_matches() -> None:
    class _ParagraphRange:
        def __init__(self, text: str, start: int, end: int) -> None:
            self.Text = text
            self.Start = start
            self.End = end

        @property
        def Duplicate(self):
            return _ParagraphRange(self.Text, self.Start, self.End)

    class _Paragraph:
        def __init__(self, text: str, start: int, end: int) -> None:
            self.Range = _ParagraphRange(text, start, end)

    class _Paragraphs:
        def __init__(self, paragraph: _Paragraph) -> None:
            self._paragraph = paragraph

        def __call__(self, _index: int):
            return self._paragraph

    class _Finder:
        def __init__(self, search) -> None:
            self._search = search
            self.Text = ""

        def ClearFormatting(self) -> None:
            return None

        def Execute(self) -> bool:
            document = self._search._document
            document.lookup_count += 1
            position = document.text.find(
                self.Text,
                self._search.Start,
                self._search.End,
            )
            if position < 0:
                return False
            self._search.Start = position
            self._search.End = position + len(self.Text)
            self._search._paragraph = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.Range.Start <= position < paragraph.Range.End
            )
            return True

    class _SearchRange:
        def __init__(self, document, start: int, end: int) -> None:
            self._document = document
            self._paragraph = None
            self.Start = start
            self.End = end

        @property
        def Duplicate(self):
            return _SearchRange(self._document, self.Start, self.End)

        @property
        def Find(self):
            return _Finder(self)

        @property
        def Paragraphs(self):
            return _Paragraphs(self._paragraph)

    class _Document:
        def __init__(self, text: str) -> None:
            self.text = text
            self.lookup_count = 0
            self.paragraphs = []
            start = 0
            for paragraph_text in text.split("\r")[:-1]:
                end = start + len(paragraph_text) + 1
                self.paragraphs.append(
                    _Paragraph(f"{paragraph_text}\r", start, end)
                )
                start = end
            self.Content = _SearchRange(self, 0, len(text))

        def Range(self, start: int, end: int):
            return _SearchRange(self, start, end)

    document = _Document(
        "cover\rPURPOSE\rbody\r2. CONCLUSIONS\rmore body\r"
        "REVISION RECORD\rtail\r"
    )

    headings = gateway_module._NumberedHeadingIndex.from_document(
        document,
        ((1, "PURPOSE"), (2, "CONCLUSIONS"), (8, "REVISION RECORD")),
    )
    purpose, purpose_text = headings.get(1, "PURPOSE")
    conclusions, conclusions_text = headings.get(2, "CONCLUSIONS")
    revision, revision_text = headings.get(8, "REVISION RECORD")
    headings.get(1, "PURPOSE")
    headings.get(8, "REVISION RECORD")

    assert document.lookup_count == 3
    assert (purpose.Start, purpose_text) == (6, "PURPOSE")
    assert (conclusions.Start, conclusions_text) == (19, "2. CONCLUSIONS")
    assert (revision.Start, revision_text) == (44, "REVISION RECORD")


def test_disclosure_scan_returns_only_matching_paragraph_offsets() -> None:
    text = (
        "Customer-visible introduction.\r"
        "This Laboratory Test Report shall not be reproduced except in full.\r"
        "Visible result paragraph.\r"
        "*.: internal note\r"
    )

    spans = gateway_module._internal_disclosure_spans(text)

    assert [text[start:end] for start, end in spans] == [
        "This Laboratory Test Report shall not be reproduced except in full.\r",
        "*.: internal note\r",
    ]


def test_word_application_never_updates_external_links() -> None:
    class _Options:
        UpdateLinksAtOpen = True

    class _Word:
        Visible = True
        DisplayAlerts = 1
        ScreenUpdating = True
        AutomationSecurity = 0
        Options = _Options()

    word = _Word()

    gateway_module._configure_word_application(word)

    assert word.Visible is False
    assert word.DisplayAlerts == 0
    assert word.ScreenUpdating is False
    assert word.AutomationSecurity == 3
    assert word.Options.UpdateLinksAtOpen is False


def test_customer_body_destination_clears_only_second_section_placeholder() -> None:
    class _Range:
        def __init__(self, start: int, end: int) -> None:
            self.Start = start
            self.End = end
            self.deleted = False

        @property
        def Duplicate(self):
            return self

        def Delete(self) -> None:
            self.deleted = True

    class _Section:
        def __init__(self, start: int, end: int) -> None:
            self.Range = _Range(start, end)

    class _Sections:
        Count = 2

        def __init__(self) -> None:
            self.values = {1: _Section(0, 1), 2: _Section(1, 40)}

        def __call__(self, index: int):
            return self.values[index]

    class _Document:
        def __init__(self) -> None:
            self.Sections = _Sections()
            self.requested_range: tuple[int, int] | None = None

        def Range(self, start: int, end: int):
            self.requested_range = (start, end)
            return _Range(start, end)

    document = _Document()

    destination = gateway_module._prepare_customer_body_destination(document)

    assert document.Sections(1).Range.deleted is False
    assert document.Sections(2).Range.deleted is True
    assert document.Sections(2).Range.End == 39
    assert document.requested_range == (1, 1)
    assert (destination.Start, destination.End) == (1, 1)


def test_section_normalization_preserves_template_headers_and_two_section_geometry() -> None:
    class _HeaderOrFooterRange:
        @property
        def FormattedText(self):
            return None

        @FormattedText.setter
        def FormattedText(self, _value) -> None:
            raise AssertionError("template header/footer content must not be replaced")

    class _HeaderOrFooter:
        def __init__(self) -> None:
            self.LinkToPrevious = False
            self.Range = _HeaderOrFooterRange()

    class _Collection:
        def __init__(self) -> None:
            self.values = {kind: _HeaderOrFooter() for kind in (1, 2, 3)}

        def __call__(self, kind: int):
            return self.values[kind]

    class _PageSetup:
        def __init__(self, *, left: int, right: int, title: bool) -> None:
            self.SectionStart = 0
            self.OddAndEvenPagesHeaderFooter = False
            self.DifferentFirstPageHeaderFooter = title
            self.TopMargin = 100
            self.BottomMargin = 100
            self.LeftMargin = left
            self.RightMargin = right
            self.HeaderDistance = 20
            self.FooterDistance = 20
            self.PageHeight = 1000
            self.PageWidth = 800

    class _Section:
        def __init__(self, *, left: int, right: int, title: bool) -> None:
            self.PageSetup = _PageSetup(left=left, right=right, title=title)
            self.Headers = _Collection()
            self.Footers = _Collection()

    class _Sections:
        Count = 2

        def __init__(self, values: dict[int, _Section]) -> None:
            self.values = values

        def __call__(self, index: int):
            return self.values[index]

    target = type("Document", (), {})()
    target.Sections = _Sections(
        {
            1: _Section(left=999, right=999, title=False),
            2: _Section(left=999, right=999, title=False),
        }
    )
    template = type("Document", (), {})()
    template.Sections = _Sections(
        {
            1: _Section(left=1008, right=1008, title=True),
            2: _Section(left=720, right=720, title=True),
        }
    )

    gateway_module._restore_customer_section_geometry(target, template)

    assert target.Sections(1).PageSetup.LeftMargin == 1008
    assert target.Sections(2).PageSetup.LeftMargin == 720
    assert target.Sections(2).PageSetup.RightMargin == 720
    assert target.Sections(1).PageSetup.DifferentFirstPageHeaderFooter is True
    assert target.Sections(2).PageSetup.DifferentFirstPageHeaderFooter is True


def test_field_refresh_skips_body_fields_and_updates_headers_and_footers() -> None:
    class _Fields:
        def __init__(self, *, fail: bool = False) -> None:
            self.fail = fail
            self.updates = 0

        def Update(self) -> None:
            if self.fail:
                raise AssertionError("body fields may contain external links")
            self.updates += 1

    class _Range:
        def __init__(self) -> None:
            self.Fields = _Fields()

    class _HeaderOrFooter:
        def __init__(self) -> None:
            self.Range = _Range()

    class _Collection:
        def __init__(self) -> None:
            self.values = {kind: _HeaderOrFooter() for kind in (1, 2, 3)}

        def __call__(self, kind: int):
            return self.values[kind]

    class _Section:
        def __init__(self) -> None:
            self.Headers = _Collection()
            self.Footers = _Collection()

    class _Sections:
        Count = 2

        def __init__(self) -> None:
            self.values = {1: _Section(), 2: _Section()}

        def __call__(self, index: int):
            return self.values[index]

    document = type("Document", (), {})()
    document.Fields = _Fields(fail=True)
    document.Sections = _Sections()
    document.repaginate_count = 0
    document.Repaginate = lambda: setattr(
        document,
        "repaginate_count",
        document.repaginate_count + 1,
    )

    gateway_module._refresh_customer_fields(document)

    assert document.repaginate_count == 1
    for section_index in (1, 2):
        section = document.Sections(section_index)
        for kind in (1, 2, 3):
            assert section.Headers(kind).Range.Fields.updates == 1
            assert section.Footers(kind).Range.Fields.updates == 1


def test_continuation_header_report_number_uses_bold_arial_10_point_font() -> None:
    class _Font:
        Name = None
        Size = None
        Bold = None

    class _Range:
        def __init__(self, start: int, end: int) -> None:
            self.Start = start
            self.End = end
            self.Text = ""
            self.Font = _Font()
            self.duplicates: list[_Range] = []

        @property
        def Duplicate(self):
            duplicate = _Range(self.Start, self.End)
            self.duplicates.append(duplicate)
            return duplicate

    class _Cell:
        def __init__(self) -> None:
            self.Range = _Range(start=400, end=430)

    class _Table:
        def __init__(self) -> None:
            self.cell = _Cell()

        def Cell(self, row: int, column: int):
            assert (row, column) == (1, 1)
            return self.cell

    table = _Table()

    gateway_module._set_continuation_report_number(
        table,
        "DL-2026-07-013-CR",
    )

    assert table.cell.Range.Text == "Report No. DL-2026-07-013-CR"
    report_number = table.cell.Range.duplicates[-1]
    assert (report_number.Start, report_number.End) == (411, 428)
    assert report_number.Font.Name == "Arial"
    assert report_number.Font.Size == 10
    assert report_number.Font.Bold is True


def test_customer_revision_note_is_left_aligned(monkeypatch: pytest.MonkeyPatch) -> None:
    class _Format:
        Alignment = 1

    class _Paragraph:
        Format = _Format()

    class _Paragraphs:
        def __call__(self, index: int):
            assert index == 1
            return _Paragraph()

    note = type("Range", (), {"Paragraphs": _Paragraphs()})()
    monkeypatch.setattr(gateway_module, "_find_text", lambda *_args, **_kwargs: note)

    gateway_module._left_align_customer_revision_note(object())

    assert note.Paragraphs(1).Format.Alignment == 0


def _minimal_document(path: Path, *, customer: bool) -> None:
    document = Document()
    document.sections[0].different_first_page_header_footer = True
    header = document.sections[0].first_page_header
    table = header.add_table(rows=5, cols=3, width=1)
    table.cell(0, 1).text = "CUSTOMER TEST REPORT" if customer else "LABORATORY TEST REPORT"
    table.cell(2, 0).text = "DL-001-CR" if customer else "DL-001"
    for heading in (
        "PURPOSE",
        "CONCLUSIONS",
        "SAMPLE DESCRIPTION",
        "TEST DESCRIPTION",
        "TEST METHODS/REQUIREMENTS",
        "TEST RESULTS",
        "REVISION RECORD",
    ):
        document.add_paragraph(heading)
    document.save(path)
