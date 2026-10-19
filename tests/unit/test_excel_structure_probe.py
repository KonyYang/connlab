from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH

from backend.infrastructure.office import (
    ExcelWorkbookGateway,
    UnsupportedLtrWorkbookError,
)
from backend.infrastructure.office.excel_tabular_layout import ExcelTabularLayout


def test_tabular_read_decodes_excel_calibration_date_cells(tmp_path: Path) -> None:
    workbook = tmp_path / "calibration.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Equipment ID", "Last Cal.", "Cal. Due"])
    sheet.append(["DG-Q-0134", 46020, 46384])
    sheet["B2"].number_format = "mm-dd-yy"
    sheet["C2"].number_format = "dd-mmm-yyyy"
    book.save(workbook)
    book.close()
    original = workbook.read_bytes()

    result = ExcelWorkbookGateway().read_tabular_rows(
        workbook, expected_headers=("Equipment ID", "Last Cal.", "Cal. Due")
    )

    assert result.rows[0]["Last Cal."] == "2025-12-29"
    assert result.rows[0]["Cal. Due"] == "2026-12-28"
    assert workbook.read_bytes() == original


@pytest.mark.parametrize("explicit_layout", [False, True])
@pytest.mark.parametrize(
    ("epoch", "serial", "expected"),
    [
        (WINDOWS_EPOCH, 46020, "2025-12-29"),
        (MAC_EPOCH, 44558, "2025-12-29"),
        (MAC_EPOCH, 0, "1904-01-01"),
        (MAC_EPOCH, 0.5, "1904-01-01T12:00:00"),
        (WINDOWS_EPOCH, 46020.5, "2025-12-29T12:00:00"),
        (WINDOWS_EPOCH, 59, "1900-02-28"),
        (WINDOWS_EPOCH, 60, "1900-02-28"),
        (WINDOWS_EPOCH, 61, "1900-03-01"),
    ],
)
def test_tabular_dates_follow_workbook_epoch_and_preserve_time(
    tmp_path: Path, explicit_layout: bool, epoch, serial: float, expected: str,
) -> None:
    workbook = tmp_path / "dates.xlsx"
    book = Workbook()
    book.epoch = epoch
    sheet = book.active
    sheet.append(["Value"])
    sheet.append([serial])
    sheet["A2"].number_format = '[$-409]dd-mmm-yyyy h:mm:ss'
    book.save(workbook)
    book.close()
    layout = (
        ExcelTabularLayout(header_row_number=1, required_header_columns=(("Value", 1),))
        if explicit_layout else None
    )

    result = ExcelWorkbookGateway().read_tabular_rows(
        workbook, expected_headers=("Value",), layout=layout,
    )

    assert result.rows[0]["Value"] == expected


@pytest.mark.parametrize(
    ("value", "number_format", "expected"),
    [
        (46020, "General", "46020"),
        (46020, "0.00", "46020"),
        (46020, '0 "days"', "46020"),
        (46020, '0 \\d', "46020"),
        (46020, "0%", "46020"),
        (46020, "0.00E+00", "46020"),
        (46020, "@", "46020"),
        (0.5, "h:mm", "0.5"),
        (0.5, "h:mm AM/PM", "0.5"),
        (0.5, "h:mm:ss AM/PM", "0.5"),
        (0.5, "h:mm A/P", "0.5"),
        (0.5, "mm:ss", "0.5"),
        (1.5, "[h]:mm:ss", "1.5"),
        (1.5, "[mm]:ss", "1.5"),
        ("46020", "yyyy-mm-dd", "46020"),
        ("N/A", "yyyy-mm-dd", "N/A"),
        ("Not calibrated", "yyyy-mm-dd", "Not calibrated"),
        ("29 Dec 2025", "yyyy-mm-dd", "29 Dec 2025"),
        (True, "yyyy-mm-dd", "1"),
        ("#N/A", "yyyy-mm-dd", "#N/A"),
        (46020, "mmm", "2025-12-29"),
        (46020, 'yyyy-mm-dd "hours"', "2025-12-29"),
    ],
)
def test_tabular_date_decoding_preserves_other_cell_semantics(
    tmp_path: Path, value, number_format: str, expected: str,
) -> None:
    workbook = tmp_path / "values.xlsx"
    book = Workbook()
    sheet = book.active
    # Deliberately use a date column header: it must not trigger numeric guessing.
    sheet.append(["Last Calibration Date"])
    sheet.append([value])
    sheet["A2"].number_format = number_format
    book.save(workbook)
    book.close()

    result = ExcelWorkbookGateway().read_tabular_rows(
        workbook, expected_headers=("Last Calibration Date",),
    )

    assert result.rows[0]["Last Calibration Date"] == expected


@pytest.mark.parametrize(
    ("cell_xml", "expected"),
    [
        ('<c r="B2" s="0"><v>46020</v></c>', "2025-12-29"),
        ('<c r="B2" s="0" t="n"><v>46020.5</v></c>', "2025-12-29T12:00:00"),
        ('<c r="B2" s="0"><f>DATE(2025,12,29)</f><v>46020</v></c>', "2025-12-29"),
        ('<c r="B2" s="0"><f>DATE(2025,12,29)</f></c>', ""),
        ('<c r="B2" s="0" t="str"><v>46020</v></c>', "46020"),
        ('<c r="B2" s="0" t="s"><v>0</v></c>', "46020"),
        ('<c r="B2" s="0" t="d"><v>2025-12-29T12:00:00</v></c>', "2025-12-29T12:00:00"),
        ('<c r="B2" s="0" t="e"><v>#N/A</v></c>', "#N/A"),
        ('<c r="B2" s="0" t="b"><v>1</v></c>', "1"),
        ('<c r="B2" s="4"><v>46020</v></c>', "46020"),
    ],
)
def test_tabular_date_decoding_uses_only_numeric_values_and_formula_cache(
    tmp_path: Path, cell_xml: str, expected: str,
) -> None:
    workbook = tmp_path / "cached.xlsx"
    _write_minimal_xlsx(workbook, sheet_name="Data", rows=[["ID", "Value"]])
    with ZipFile(workbook) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = (
        _sheet_xml([["ID", "Value"]]).replace(
            "</sheetData>",
            '<row r="2"><c r="A2" t="inlineStr"><is><t>EQ-1</t></is></c>'
            f'{cell_xml}</row></sheetData>',
        ).encode("utf-8")
    )
    parts["xl/styles.xml"] = (
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<cellXfs count="1"><xf numFmtId="14"/></cellXfs></styleSheet>'
    ).encode("utf-8")
    parts["xl/sharedStrings.xml"] = (
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<si><t>46020</t></si></sst>'
    ).encode("utf-8")
    with ZipFile(workbook, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    result = ExcelWorkbookGateway().read_tabular_rows(
        workbook, expected_headers=("ID", "Value"),
    )

    assert result.rows[0]["Value"] == expected


@pytest.mark.parametrize(
    ("num_fmt_id", "explicit_code", "expected"),
    [
        *[(identifier, None, "2025-12-29") for identifier in (27, 28, 29, 30, 31, 36, 50, 51, 54, 57, 58)],
        *[(identifier, None, "46020") for identifier in (32, 33, 34, 35, 52, 53, 55, 56)],
        (31, "0.00", "46020"),
        (58, "h:mm", "46020"),
        (34, "yyyy-mm-dd", "2025-12-29"),
        (52, "0", "46020"),
    ],
)
def test_tabular_dates_resolve_only_unambiguous_implied_calendar_formats(
    tmp_path: Path, num_fmt_id: int, explicit_code: str | None, expected: str,
) -> None:
    workbook = tmp_path / "implied-formats.xlsx"
    _write_minimal_xlsx(workbook, sheet_name="Data", rows=[["Value"], ["46020"]])
    with ZipFile(workbook) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b'<c r="A2" t="inlineStr"><is><t>46020</t></is></c>',
        b'<c r="A2" s="0"><v>46020</v></c>',
    )
    custom_format = (
        f'<numFmts count="1"><numFmt numFmtId="{num_fmt_id}" formatCode="{explicit_code}"/></numFmts>'
        if explicit_code is not None else ""
    )
    parts["xl/styles.xml"] = (
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'{custom_format}<cellXfs count="1"><xf numFmtId="{num_fmt_id}"/></cellXfs></styleSheet>'
    ).encode("utf-8")
    with ZipFile(workbook, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    result = ExcelWorkbookGateway().read_tabular_rows(workbook, expected_headers=("Value",))

    assert result.rows[0]["Value"] == expected


def test_excel_structure_probe_accepts_matching_sheet_headers_and_date(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "standard.xlsx"
    _write_minimal_xlsx(
        workbook,
        sheet_name="Standard Records",
        rows=[
            ["LTR Number", "Date", "Test Item", "Sample Description"],
            ["DL-2026-05-001", "2026-05-07", "Qualification", "Samples"],
        ],
    )

    result = ExcelWorkbookGateway().probe_structure(
        workbook,
        expected_headers=("LTR Number", "Test Item", "Sample Description"),
        expected_date_headers=("Date",),
        expected_sheet_name_patterns=(r".*record.*",),
    )

    assert result.valid is True
    assert result.matched_sheet_names == ("Standard Records",)
    assert result.missing_headers == ()
    assert result.missing_date_headers == ()


def test_excel_structure_probe_reports_missing_header(tmp_path: Path) -> None:
    workbook = tmp_path / "standard.xlsx"
    _write_minimal_xlsx(
        workbook,
        sheet_name="Standard Records",
        rows=[["LTR Number", "Date"]],
    )

    result = ExcelWorkbookGateway().probe_structure(
        workbook,
        expected_headers=("LTR Number", "Test Item"),
        expected_date_headers=("Date",),
        expected_sheet_name_patterns=(r".*record.*",),
    )

    assert result.valid is False
    assert result.missing_headers == ("Test Item",)
    assert result.failure_reason == "Missing required headers: Test Item"


def test_excel_structure_probe_reports_missing_sheet(tmp_path: Path) -> None:
    workbook = tmp_path / "standard.xlsx"
    _write_minimal_xlsx(workbook, sheet_name="Other", rows=[["LTR Number"]])

    result = ExcelWorkbookGateway().probe_structure(
        workbook,
        expected_headers=("LTR Number",),
        expected_sheet_name_patterns=(r".*record.*",),
    )

    assert result.valid is False
    assert result.matched_sheet_names == ()
    assert result.failure_reason == "No worksheet matched the expected sheet rules."


def test_excel_structure_probe_is_xlsx_only(tmp_path: Path) -> None:
    workbook = tmp_path / "legacy.xls"
    workbook.write_bytes(b"legacy")

    with pytest.raises(UnsupportedLtrWorkbookError, match="supports .xlsx only"):
        ExcelWorkbookGateway().probe_structure(
            workbook,
            expected_headers=("LTR Number",),
        )


def _write_minimal_xlsx(
    path: Path,
    *,
    sheet_name: str,
    rows: list[list[str]],
) -> None:
    """Write a minimal XLSX package that the gateway can read."""
    with ZipFile(path, "w") as archive:
        archive.writestr(
            "xl/workbook.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
  xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="{sheet_name}" sheetId="1" r:id="rId1"/>
  </sheets>
</workbook>""",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
    Target="worksheets/sheet1.xml"/>
</Relationships>""",
        )
        archive.writestr("xl/worksheets/sheet1.xml", _sheet_xml(rows))


def _sheet_xml(rows: list[list[str]]) -> str:
    """Return worksheet XML using inline string cells."""
    row_xml = []
    for row_index, row in enumerate(rows, start=1):
        cells = []
        for column_index, value in enumerate(row):
            reference = f"{chr(65 + column_index)}{row_index}"
            cells.append(
                f'<c r="{reference}" t="inlineStr"><is><t>{value}</t></is></c>'
            )
        row_xml.append(f'<row r="{row_index}">{"".join(cells)}</row>')
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{"".join(row_xml)}</sheetData>'
        "</worksheet>"
    )
