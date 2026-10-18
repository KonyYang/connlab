from hashlib import sha256
from io import BytesIO, StringIO
import csv
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
import pytest

from backend.application.temperature_data_preparation import prepare_data, suggest_selection
from backend.domain.temperature_data import DataSelection, WorkbookTable
from backend.domain.temperature_rise import Coefficients, analyze_temperature_rise, generate_derating, calculate_current
from backend.infrastructure.office.temperature_workbook_gateway import TemperatureWorkbookGateway
from tests.unit.test_temperature_rise_calculations import baseline_rows


@pytest.mark.parametrize('decimals,expected,number_format', [(0, [0, 10, 20], '0'), (1, [0, 10.1, 20.1], '0.0')])
def test_export_quantizes_current_formulas_caches_and_axes_but_keeps_original_values(tmp_path, decimals, expected, number_format):
    table = WorkbookTable('precision.csv', ('Data',), 'Data', (
        ('Scan', 'Time', '101 <1_TC> (C)', 'Ambient', 'Current', 'Aux (A)'),
        *((index + 1, '2026/10/17 10:00', temperature, 23, current, 125.07)
          for index, (current, temperature) in enumerate([(0.0087, 23)] * 2 + [(10.05, 25)] * 2 + [(20.05, 30)] * 2)),
    ))
    selection = DataSelection(1, 2, 7, 4, 5, (3,), 1, current_decimals=decimals)
    prepared = prepare_data(table, selection, acknowledge_warnings=True)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=1, zero_intercept=True)
    output = tmp_path / 'precision.xlsx'
    TemperatureWorkbookGateway().write(output, table=table, selection=selection, prepared=prepared, analysis=analysis,
        maximum_coefficients=analysis.maximum_fit.coefficients, average_coefficients=analysis.average_fit.coefficients,
        target_rise=30, derating=generate_derating(analysis.average_fit.coefficients, max_temperature=105, step=2.5, ambient_point=75))
    with BytesIO(output.read_bytes()) as stream:
        book = load_workbook(stream, data_only=True)
        try:
            assert book['Initial Data']['E2'].value == 0.0087
            assert book['Initial Data']['E4'].value == 10.05
            assert book['Initial Data']['F4'].value == 125.07
            target, area = next(book.defined_names['TemperatureRiseSummary'].destinations)
            summary = list(book[target][area])
            assert [cell.value for cell in summary[1][1:]] == expected
            assert [cell.value for cell in summary[0][1:]] == [None, *([125] * 2 if decimals == 0 else [125.1] * 2)]
        finally:
            book.close()
    book = load_workbook(output)
    try:
        rise, derating = book['T-riseChart'], book['Derating']
        assert rise._charts[0].x_axis.numFmt.formatCode == number_format
        assert rise._charts[0].y_axis.numFmt.formatCode == '0'
        assert derating._charts[0].x_axis.numFmt.formatCode == '0'
        assert derating._charts[0].y_axis.numFmt.formatCode == number_format
        assert derating['G2'].number_format == number_format
        assert derating._charts[0].series[2].dLbls.numFmt == number_format
        target, area = next(book.defined_names['CurrentCalculator'].destinations)
        assert list(book[target][area])[-1][-1].number_format == number_format
        target, area = next(book.defined_names['StageTemperatureRise'].destinations)
        for row in list(book[target][area])[1:]:
            assert row[-1].value.startswith('=ROUND(') and row[-1].value.endswith(f',{decimals})')
            assert row[-1].number_format == number_format
    finally:
        book.close()


def test_reader_preserves_bytes_and_accepts_explicit_sheet(tmp_path):
    source = tmp_path / 'scanner.xlsx'
    book = Workbook()
    book.active.title = 'Notes'
    data = book.create_sheet('Initial Data')
    data.append(['Scan', 'Time', 'TC1', 'Ambient', 'Current'])
    data.append([1, '17:17', 28, 25, 10])
    data.append([2, '17:18', 29, 25, 10])
    book.save(source)
    digest = sha256(source.read_bytes()).hexdigest()
    table = TemperatureWorkbookGateway().read(source)
    assert table.sheet_name == 'Initial Data'
    assert table.sheet_names == ('Notes', 'Initial Data')
    assert table.rows[-1] == (2, '17:18', 29, 25, 10)
    choice = suggest_selection(table)
    assert choice.temperature_columns == (3,)
    assert choice.ambient_column == 4
    assert choice.current_column == 5
    assert sha256(source.read_bytes()).hexdigest() == digest
    with pytest.raises(ValueError, match='sheet'):
        TemperatureWorkbookGateway().read(source, sheet_name='Absent')


@pytest.mark.parametrize('encoding', ['utf-8-sig', 'utf-16', 'gb18030'])
@pytest.mark.parametrize('delimiter', [',', '\t', ';'])
def test_csv_reader_preserves_scanner_metadata_quotes_decimal_values_and_row_positions(tmp_path, encoding, delimiter):
    stream = StringIO(newline='')
    writer = csv.writer(stream, delimiter=delimiter)
    writer.writerows([
        ['名称:', '扫描仪, "原始数据"'],
        ['通道', '名称', '功能', '增益', '偏移'],
        ['101', '1_H1', '温度', '1', '0'],
        ['扫描', '时间', '101 <1_H1> (C)', '313 <ambient> (C)', '317 <Current> (VDC)'],
        ['1', '2024/10/28 17:10:16:391', '29.328', '25.465', '3.0010263'],
        [],
        ['2', '2024/10/28 17:11:16:375', '', '25.459', '0'],
    ])
    source = tmp_path / 'scanner.CSV'
    content = stream.getvalue().encode(encoding)
    source.write_bytes(content)
    table = TemperatureWorkbookGateway().read(source)
    assert source.read_bytes() == content
    assert table.sheet_names == ('Initial Data',)
    assert table.rows[0] == ('名称:', '扫描仪, "原始数据"')
    assert table.rows[4][-1] == '3.0010263'
    assert table.rows[5] == ()
    assert table.rows[6][2] is None
    choice = suggest_selection(table)
    assert (choice.header_row, choice.start_row, choice.end_row) == (4, 5, 7)
    assert choice.excluded_rows == ()
    review = prepare_data(table, choice)
    assert not review.ready
    assert review.measurements[0].current == 3
    assert review.measurements[0].ambient == 25.465
    with pytest.raises(ValueError, match='sheet'):
        TemperatureWorkbookGateway().read(source, sheet_name='Absent')


@pytest.mark.parametrize('content', [b'\xff', b'Scan,Time,TC,Ambient,Current\n1,"unterminated', b'a\x00,b\n1,2'])
def test_csv_reader_reports_unreadable_encoding_or_malformed_data(content):
    with pytest.raises(ValueError, match='CSV'):
        TemperatureWorkbookGateway().read_upload(content, 'bad.csv')


@pytest.mark.parametrize(('limit', 'value', 'content'), [
    ('MAX_ROWS', 2, b'a,b\n1,2\n3,4'),
    ('MAX_COLUMNS', 2, b'a,b,c\n1,2,3'),
    ('MAX_CELLS', 5, b'a,b\n1,2\n3,4'),
    ('MAX_FILE_BYTES', 3, b'a,b\n1,2'),
])
def test_csv_reader_obeys_existing_import_resource_limits(monkeypatch, limit, value, content):
    monkeypatch.setattr('backend.infrastructure.office.temperature_workbook_gateway.' + limit, value)
    with pytest.raises(ValueError, match='Limit|25 MB'):
        TemperatureWorkbookGateway().read_upload(content, 'large.csv')


def test_export_contains_native_xy_charts_formulas_cached_values_and_source_mapping(tmp_path):
    rows = baseline_rows()
    headers = ('Original Scan', *(f'TC{i}' for i in range(1, 21)), 'Ambient', 'Current')
    values = tuple((row.source_row, *row.temperatures, row.ambient, row.current) for row in rows)
    table = WorkbookTable('scanner.xlsx', ('Initial Data',), 'Initial Data', (headers, *values))
    selection = DataSelection(1, 2, 301, 22, 23, tuple(range(2, 22)), 4)
    prepared = prepare_data(table, selection)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=4, zero_intercept=True)
    derating = generate_derating(analysis.average_fit.coefficients, max_temperature=105, step=2.5, ambient_point=75)
    output = tmp_path / 'result.xlsx'
    TemperatureWorkbookGateway().write(output, table=table, selection=selection, prepared=prepared,
        analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=30, derating=derating)
    book = load_workbook(output, data_only=False)
    assert book.sheetnames == ['Setup', 'Initial Data', 'T-riseChart', 'Derating']
    assert book.active.title == 'T-riseChart'
    assert book['T-riseChart'].sheet_view.tabSelected
    assert not any(sheet.sheet_view.tabSelected for sheet in book if sheet.title != 'T-riseChart')
    assert len(book['T-riseChart']._charts) == 1
    assert len(book['Derating']._charts) == 1
    assert book['T-riseChart']._charts[0].__class__.__name__ == 'ScatterChart'
    assert [series.trendline.order for series in book['T-riseChart']._charts[0].series] == [2, 2]
    assert all(series.trendline.intercept == 0 for series in book['T-riseChart']._charts[0].series)
    assert 'SQRT' in book['Derating']['B6'].value
    assert book['Derating'].freeze_panes is None
    assert book['Initial Data']['B1'].value == 'TC1'
    book.close()
    cached = load_workbook(output, data_only=True)
    assert cached['Derating']['B36'].value == pytest.approx(derating.annotation.basic)
    assert cached['Derating']['C36'].value == pytest.approx(derating.annotation.derated)
    assert cached['Derating']['B48'].value == 0
    cached.close()
    with ZipFile(output) as archive:
        assert not any('vba' in name.lower() for name in archive.namelist())
        charts = [name for name in archive.namelist() if name.startswith('xl/charts/chart') and name.endswith('.xml')]
        assert len(charts) == 2
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        for name in charts:
            tree = ET.fromstring(archive.read(name))
            assert tree.findall('.//c:numCache/c:pt', ns), 'Charts must render before Excel recalculates.'
            drawing = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
            for line in tree.iter(f'{drawing}ln'):
                fills = [child for child in line if child.tag in {
                    f'{drawing}{kind}' for kind in ('noFill', 'solidFill', 'gradFill', 'pattFill')
                }]
                assert len(fills) <= 1, 'Excel rejects mutually exclusive line-fill elements.'


def test_export_refuses_existing_file(tmp_path):
    source = tmp_path / 'existing.xlsx'
    source.write_bytes(b'keep')
    with pytest.raises(ValueError, match='new'):
        TemperatureWorkbookGateway().write(source, table=None, selection=None, prepared=None,
            analysis=None, maximum_coefficients=None, average_coefficients=None, target_rise=None, derating=None)
    assert source.read_bytes() == b'keep'


def test_export_keeps_original_data_positions_and_marks_only_real_stage_endpoints(tmp_path):
    headers = ('Scan', 'Time', '101 <1_HS> (C)', 'Ambient (C)', 'Current (VDC)', 'Unused')
    readings = tuple((i, '18:00', 25 + .01 * current ** 2, 25, current, '=literal')
                     for i, current in enumerate((10, 10, 10, 20, 20, 30, 30), 1))
    table = WorkbookTable('original.csv', ('Data',), 'Data', (
        ('Instrument', 'Scanner'), (), headers, *readings, ('Footer',)))
    selection = DataSelection(3, 4, 10, 4, 5, (3,), 1, excluded_rows=(4,))
    prepared = prepare_data(table, selection)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=1, zero_intercept=True)
    content = TemperatureWorkbookGateway().render('output.xlsx', table=table, selection=selection,
        prepared=prepared, analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=30, derating=None)
    book = load_workbook(BytesIO(content))
    initial = book['Initial Data']
    assert initial.max_row == 11
    assert initial.max_column == 6
    for row, values in enumerate(table.rows, 1):
        assert tuple(initial.cell(row, col).value for col in range(1, len(values) + 1)) == values
    assert initial['F4'].data_type == 's'
    yellow = {cell.row for row in initial for cell in row if cell.fill.fgColor.rgb == '00FFFF00'}
    assert yellow == {6, 8, 10}
    assert all(initial.cell(row, col).fill.fgColor.rgb == '00FFFF00'
               for row in yellow for col in range(1, 7))
    assert book.sheetnames == ['Setup', 'Initial Data', 'T-riseChart', 'Derating']
    assert book.active.title == 'T-riseChart'
    assert all(cell.value is None for row in book['Setup'] for cell in row)
    book.close()


def _report_export(tmp_path, maximum=None, average=None, *, include_derating=False, zero_intercept=True):
    rows = baseline_rows()
    table = WorkbookTable('scanner.xlsx', ('Initial Data',), 'Initial Data', (
        ('Scan', *(f'TC{i}' for i in range(1, 21)), 'Ambient', 'Current'),
        *((row.source_row, *row.temperatures, row.ambient, row.current) for row in rows)))
    selection = DataSelection(1, 2, 301, 22, 23, tuple(range(2, 22)), 4)
    prepared = prepare_data(table, selection)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=4, zero_intercept=zero_intercept)
    output = tmp_path / 'report-ready.xlsx'
    TemperatureWorkbookGateway().write(output, table=table, selection=selection, prepared=prepared,
        analysis=analysis, maximum_coefficients=maximum or analysis.maximum_fit.coefficients,
        average_coefficients=average or analysis.average_fit.coefficients, target_rise=30,
        derating=generate_derating(average or analysis.average_fit.coefficients,
            max_temperature=105, step=2.5, ambient_point=75) if include_derating else None)
    return output, analysis


def test_exported_current_axes_and_plot_data_use_default_integer_precision(tmp_path):
    output, _ = _report_export(tmp_path, include_derating=True)
    book = load_workbook(output)
    try:
        for name in ('T-riseChart', 'Derating'):
            chart = book[name]._charts[0]
            for axis in (chart.x_axis, chart.y_axis):
                assert axis.numFmt.formatCode == '0'
                assert axis.numFmt.sourceLinked is False
        rise = book['T-riseChart']._charts[0]
        assert float(rise.series[0].xVal.numRef.numCache.pt[1].v) == 18
        assert float(rise.series[0].yVal.numRef.numCache.pt[-1].v) == pytest.approx(36.545)
        derating = book['Derating']._charts[0]
        assert derating.series[2].dLbls.numFmt == '0'
    finally:
        book.close()
    with ZipFile(output) as archive:
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        for name in ('xl/charts/chart1.xml', 'xl/charts/chart2.xml'):
            root = ET.fromstring(archive.read(name))
            formats = root.findall('.//c:valAx/c:numFmt', ns)
            assert len(formats) == 2
            assert all(item.get('formatCode') == '0' and item.get('sourceLinked') == '0' for item in formats)


def multi_current_source(*, measured_zero=True):
    currents = ([0, 0, 10, 10, 10] if measured_zero else [10, 10]) + [20, 20, 30, 30]
    headers = ('Scan', 'Time', '101 <1_HS> (C)', '102 <1_C> (C)', '313 <Ambient T> (C)',
               '316 <Low Power> (A)', '315 <High Power> (A)', '317 <A4_B4 Signal> (A)',
               '318 <Others Signal> (A)', '319 <Varying> (A)', '320 <Off> (A)')
    # CSV numeric strings must stay literal in Initial Data, but remain numeric in report formulas.
    rows = tuple((index, '2026-10-14 18:00', str(22 + .01 * current ** 2 + .2 * current),
                  str(21 + .005 * current ** 2 + .1 * current), '20', str(current),
                  str(125 + index * .001), '3.001', '1.02', str(5 + index), '0')
                 for index, current in enumerate(currents, 1))
    return WorkbookTable('multi.csv', ('Data',), 'Data', (headers, *rows))


@pytest.mark.parametrize('measured_zero', [False, True])
def test_multi_current_report_has_named_auxiliary_stage_rows_without_changing_curve_inputs(tmp_path, measured_zero):
    table = multi_current_source(measured_zero=measured_zero)
    choice = DataSelection(1, 2, len(table.rows), 5, 6, (3, 4), 2,
                           excluded_rows=(6,) if measured_zero else ())
    prepared = prepare_data(table, choice, acknowledge_warnings=True)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=2,
                                        zero_intercept=not measured_zero)
    output = tmp_path / 'multi.xlsx'
    TemperatureWorkbookGateway().write(output, table=table, selection=choice, prepared=prepared,
        analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=30, derating=None)
    book = load_workbook(output)
    try:
        sheet = book['T-riseChart']
        summary = sheet[next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
        assert [row[0].value for row in summary[:6]] == [
            'High Power(A)', 'A4_B4 Signal(A)', 'Others Signal(A)', 'Low Power(A)', 'HS', 'C']
        assert all(cell.font.bold and cell.fill.fgColor.rgb == '00D9D9D9'
                   and cell.alignment.horizontal == 'center' for row in summary[:4] for cell in row)
        assert all(cell.number_format == '0' for row in summary[:4] for cell in row[1:])
        assert summary[0][2].data_type == 'f'
        assert summary[4][2].number_format == '0.0'
        main_row = summary[3][0].row
        for series in sheet._charts[0].series:
            assert series.xVal.numRef.f.endswith(f'$B${main_row}:$E${main_row}')
        raw = sheet[next(book.defined_names['KeyStageReadings'].destinations)[1].replace('$', '')]
        if measured_zero:
            assert summary[0][1].value == f'=ROUND(G{raw[1][0].row}*1,0)'
        else:
            assert all(row[1].value is None for row in summary[:3])
        assert not any(cell.comment for row in sheet for cell in row)
        assert list(book['Initial Data'].values) == list(table.rows)
    finally:
        book.close()
    cached = load_workbook(output, data_only=True)
    try:
        sheet = cached['T-riseChart']
        summary = sheet[next(cached.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
        assert [cell.value for cell in summary[0][1:]] == (
            [125] * 4 if measured_zero else [None, 125, 125, 125])
        assert [cell.value for cell in summary[1][1:]] == ([3] * 4 if measured_zero else [None, 3, 3, 3])
        assert [cell.value for cell in summary[2][1:]] == ([1] * 4 if measured_zero else [None, 1, 1, 1])
        assert [cell.value for cell in summary[3][1:]] == [0, 10, 20, 30]
        assert [cell.value for cell in summary[-2][1:]] == ([2, 5, 10, 17] if measured_zero else [0, 5, 10, 17])
    finally:
        cached.close()


def test_multi_current_export_uses_literal_edited_current_names_with_safe_blank_fallback(tmp_path):
    table = multi_current_source()
    choice = DataSelection(1, 2, len(table.rows), 5, 6, (3, 4), 2,
                           channel_names=((6, '01 Custom Main'), (7, ''), (8, '=1+1')))
    prepared = prepare_data(table, choice, acknowledge_warnings=True)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=2, zero_intercept=False)
    content = TemperatureWorkbookGateway().render('edited-currents.xlsx', table=table, selection=choice,
        prepared=prepared, analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=30, derating=None)
    book = load_workbook(BytesIO(content))
    try:
        sheet = book['T-riseChart']
        summary = sheet[next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
        assert [row[0].value for row in summary[:4]] == [
            'Current 7(A)', '=1+1(A)', 'Others Signal(A)', 'Custom Main(A)']
        assert all(row[0].data_type == 's' for row in summary[:4])
    finally:
        book.close()


@pytest.mark.parametrize('zero_intercept', [False, True])
@pytest.mark.parametrize('measured_zero', [False, True])
def test_export_dual_mode_points_warnings_and_raw_preservation(tmp_path, zero_intercept, measured_zero):
    table = multi_current_source(measured_zero=measured_zero)
    selection = DataSelection(1, 2, len(table.rows), 5, 6, (3, 4), 2)
    prepared = prepare_data(table, selection, acknowledge_warnings=True)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=2,
                                        zero_intercept=zero_intercept)
    content = TemperatureWorkbookGateway().render('dual.xlsx', table=table, selection=selection,
        prepared=prepared, analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=30, derating=None)
    book = load_workbook(BytesIO(content), data_only=True)
    try:
        sheet = book['T-riseChart']
        area = next(book.defined_names['TemperatureRiseSummary'].destinations)[1]
        summary = sheet[area]
        expected_currents = ([0] if zero_intercept or measured_zero else []) + [10, 20, 30]
        expected_maxima = ([0 if zero_intercept else 2] if zero_intercept or measured_zero else []) + [5, 10, 17]
        assert [cell.value for cell in summary[3][1:]] == expected_currents
        assert [cell.value for cell in summary[-2][1:]] == expected_maxima
        assert [point.current for point in analysis.points] == expected_currents
        assert list(book['Initial Data'].values) == list(table.rows)
        warnings = [cell.value for row in sheet for cell in row
                    if isinstance(cell.value, str) and 'extrapolated' in cell.value]
        assert warnings == list(analysis.warnings)
        assert bool(warnings) is (not zero_intercept and not measured_zero)
        raw_area = next(book.defined_names['KeyStageReadings'].destinations)[1]
        assert sheet[raw_area][1][0].value == (2 if measured_zero and not zero_intercept else 5 if measured_zero else 2)
    finally:
        book.close()
    book = load_workbook(BytesIO(content))
    try:
        chart = book['T-riseChart']._charts[0]
        assert chart.series[0].trendline.intercept == (0 if zero_intercept else None)
        assert [point.v for point in chart.series[0].xVal.numRef.numCache.pt] == expected_currents
        assert [point.v for point in chart.series[0].yVal.numRef.numCache.pt] == expected_maxima
    finally:
        book.close()


@pytest.mark.parametrize('include_derating', [False, True])
def test_export_uses_unfrozen_source_table_alignment_and_integer_current_display_without_comments(tmp_path, include_derating):
    output, _ = _report_export(tmp_path, include_derating=include_derating)
    with ZipFile(output) as archive:
        assert not any('comments' in name.lower() for name in archive.namelist())
    book = load_workbook(output)
    try:
        assert book['Initial Data'].freeze_panes is None
        assert all(cell.comment is None for sheet in book for row in sheet for cell in row)
        sheet = book['T-riseChart']
        for name in ('TemperatureRiseSummary', 'CurrentCalculator', 'KeyStageReadings',
                     'StageTemperatureRise', 'AverageCoefficients'):
            block = sheet[next(book.defined_names[name].destinations)[1].replace('$', '')]
            for row_index, row in enumerate(block):
                for column_index, cell in enumerate(row):
                    is_time_data = (name in ('KeyStageReadings', 'StageTemperatureRise')
                                    and row_index > 0 and column_index == 1)
                    assert cell.alignment.horizontal == ('left' if is_time_data else 'center'), (name, cell.coordinate)
        summary = sheet[next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
        assert all(cell.number_format == '0' for cell in summary[0][1:])
        assert summary[0][2].data_type == 'f'
        assert summary[1][2].number_format == '0.0'
    finally:
        book.close()
    cached = load_workbook(output, data_only=True)
    try:
        summary = cached['T-riseChart'][next(cached.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
        assert [cell.value for cell in summary[0][1:]] == pytest.approx(
            [0, 18, 31, 40, 44, 66, 75])
        assert summary[-2][-1].value == pytest.approx(36.545)
    finally:
        cached.close()


def test_export_formats_measurement_labels_and_keeps_only_max_calculator_visible(tmp_path):
    headers = ('Scan', 'Time', '101 <1_HS> (C)', '102 <1_H> (C)', '103 <1_C> (C)',
               '104 <1_RS> (C)', 'Ambient (C)', 'Current (VDC)')
    table = WorkbookTable('names.csv', ('Data',), 'Data', (headers, *(
        (i, '18:00', 25 + current ** 2 / 100, 26, 27, 28, 25, current)
        for i, current in enumerate((10, 10, 20, 20, 30, 30), 1))))
    selection = DataSelection(1, 2, 7, 7, 8, (3, 4, 5, 6), 4)
    prepared = prepare_data(table, selection)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=4, zero_intercept=True)
    content = TemperatureWorkbookGateway().render('names.xlsx', table=table, selection=selection,
        prepared=prepared, analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=30, derating=None)
    book = load_workbook(BytesIO(content))
    sheet = book['T-riseChart']
    summary = sheet[next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
    assert [row[0].value for row in summary] == [
        'Current(A)', 'HS', 'H', 'C', 'RS',
        '1# Single Max', 'Max T-Rise', 'Avg of max T-Rise on each sample']
    assert sheet.column_dimensions['A'].width == 18
    assert sheet.column_dimensions['B'].width == sheet.column_dimensions['C'].width
    calculator = sheet[next(book.defined_names['CurrentCalculator'].destinations)[1].replace('$', '')]
    assert not sheet.row_dimensions[calculator[1][0].row].hidden
    average = sheet[next(book.defined_names['AverageCoefficients'].destinations)[1].replace('$', '')]
    assert all(sheet.row_dimensions[row[0].row].hidden for row in average)
    for series in sheet._charts[0].series:
        equation_row = sheet[series.trendline.trendlineLbl.tx.strRef.f.split('!')[1].replace('$', '')].row
        assert sheet.row_dimensions[equation_row].hidden
    chart = sheet._charts[0]
    assert chart.layout.manualLayout.w >= .85
    book.close()
    with ZipFile(BytesIO(content)) as archive:
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        root = ET.fromstring(archive.read('xl/charts/chart1.xml'))
        assert root.find('c:roundedCorners', ns).get('val') == '1'


@pytest.mark.parametrize('sheet_name,width_emu', [('T-riseChart', 6768000), ('Derating', 6912000)])
def test_exported_charts_use_eighty_percent_of_the_previous_dimensions(tmp_path, sheet_name, width_emu):
    output, _ = _report_export(tmp_path, include_derating=True)
    book = load_workbook(output)
    try:
        extent = book[sheet_name]._charts[0].anchor.ext
        assert extent.cx == width_emu
        assert extent.cy == pytest.approx(4032000, abs=1, rel=0)
    finally:
        book.close()


def test_derating_layout_has_horizontal_bordered_inputs_and_chart_beside_data(tmp_path):
    output, _ = _report_export(tmp_path, include_derating=True)
    book = load_workbook(output)
    try:
        sheet = book['Derating']
        assert [cell.value for cell in sheet[1][:4]] == ['AVG a', 'AVG b', 'AVG c', 'Max Temp (°C)']
        assert [sheet.cell(1, col).value for col in (6, 7, 8)] == [
            'Ambient Point (°C)', 'Basic (A)', '80% Derating (A)']
        assert not any(isinstance(cell.value, str) and
                       (cell.value == 'Current vs Ambient Temperature' or cell.value.startswith('Max working temperature:'))
                       for row in sheet for cell in row)
        for area in ('A1:D2', 'F1:H2', 'J1:K2', f'A5:C{sheet.max_row}'):
            for row in sheet[area]:
                for cell in row:
                    assert cell.alignment.horizontal == 'center', cell.coordinate
                    assert all(getattr(cell.border, side).style == 'thin'
                               for side in ('left', 'right', 'top', 'bottom')), cell.coordinate
        assert sheet['A2'].number_format == '0.000000'
        assert sheet['B6'].value == '=(-$B$2+SQRT($B$2^2+4*$A$2*($D$2-A6)))/(2*$A$2)'
        assert sheet['G2'].value == '=(-$B$2+SQRT($B$2^2+4*$A$2*($D$2-F2)))/(2*$A$2)'
        assert sheet['H2'].value == '=G2*0.8'
        assert sheet._charts[0].anchor._from.col == 4
        assert sheet._charts[0].anchor._from.row == 4
    finally:
        book.close()


def test_derating_chart_matches_reference_and_guide_follows_editable_ambient_point(tmp_path):
    output, analysis = _report_export(tmp_path, include_derating=True)
    book = load_workbook(output)
    sheet = book['Derating']
    chart = sheet._charts[0]
    assert sheet.freeze_panes is None
    assert chart.graphical_properties.solidFill.srgbClr == 'FFFFE0'
    assert [series.tx.v for series in chart.series[:2]] == ['Basic', '80% Derating']
    colors = ['8B0000', 'FFA500']
    assert [series.graphicalProperties.line.solidFill.srgbClr for series in chart.series[:2]] == colors
    for series, color in zip(chart.series[2:4], colors):
        assert series.marker.symbol == 'circle'
        assert series.marker.graphicalProperties.line.solidFill.srgbClr == color
        assert series.dLbls.numFmt == '0'
        assert series.dLbls.dLblPos == 'r'
        font = series.dLbls.txPr.p[0].pPr.defRPr
        assert font.b and font.solidFill.srgbClr == color
    for axis in (chart.x_axis, chart.y_axis):
        assert axis.majorGridlines.spPr.line.prstDash == 'dash'
    assert chart.x_axis.majorUnit == 5
    assert chart.x_axis.scaling.max == 105
    assert [entry.idx for entry in chart.legend.legendEntry if entry.delete] == [2, 3, 4]
    assert chart.legend.layout.manualLayout.x < .15
    assert sheet['J1'].value == '=$F$2'
    assert sheet['J2'].value == '=$F$2'
    assert sheet['K2'].value == '=$G$2'
    assert sheet['G2'].number_format == sheet['H2'].number_format == '0'
    guide = chart.series[4]
    assert guide.graphicalProperties.line.solidFill.srgbClr == '0000FF'
    assert guide.graphicalProperties.line.prstDash == 'dash'
    assert [float(point.v) for point in guide.xVal.numRef.numCache.pt] == [75, 75]
    assert [float(point.v) for point in guide.yVal.numRef.numCache.pt] == pytest.approx([0, calculate_current(analysis.average_fit.coefficients, 30)])
    book.close()
    cached = load_workbook(output, data_only=True)
    assert cached['Derating']['K2'].value == pytest.approx(calculate_current(analysis.average_fit.coefficients, 30))
    assert cached['Derating']['J1'].value == cached['Derating']['J2'].value == 75
    cached.close()
    with ZipFile(output) as archive:
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
              'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        root = ET.fromstring(archive.read('xl/charts/chart2.xml'))
        assert root.find('.//c:plotArea/c:spPr/a:solidFill/a:srgbClr', ns).get('val') == 'FFFFFF'
        assert root.findall('.//c:scatterChart/c:ser', ns)[4].find('c:marker/c:symbol', ns).get('val') == 'none'
        titles = root.findall('.//c:title', ns)
        assert [''.join(t.itertext()) for t in titles] == [
            'Current Carrying Capacity De-rating Curve', 'Ambient Temperature [°C]', 'Current [Amp]']
        assert all(t.find('.//a:defRPr', ns).get('b') == '1' for t in titles)


def test_report_export_highlights_copied_readings_without_recoloring_calculations(tmp_path):
    output, _ = _report_export(tmp_path)
    book = load_workbook(output)
    try:
        sheet = book['T-riseChart']
        raw = sheet[next(book.defined_names['KeyStageReadings'].destinations)[1].replace('$', '')]
        rises = sheet[next(book.defined_names['StageTemperatureRise'].destinations)[1].replace('$', '')]
        assert all(cell.fill.patternType == 'solid' and cell.fill.fgColor.rgb == '00FFFF00'
                   for row in raw[1:] for cell in row)
        assert all(cell.fill.fgColor.rgb == '00D9D9D9' for cell in raw[0])
        assert all(cell.fill.patternType is None for row in rises[1:] for cell in row)
        assert all(cell.fill.patternType is None for cell in sheet[raw[-1][0].row + 1])
    finally:
        book.close()


def test_report_export_has_traceable_key_readings_and_copyable_formula_summary(tmp_path):
    output, analysis = _report_export(tmp_path)
    book = load_workbook(output)
    sheet = book['T-riseChart']
    summary = next(book.defined_names['TemperatureRiseSummary'].destinations)[1]
    block = sheet[summary.replace('$', '')]
    assert block[0][0].value == 'Current(A)'
    labels = [row[0].value for row in block]
    assert labels[1:6] == ['TC1', 'TC2', 'TC3', 'TC4', '1# Single Max']
    assert labels[-2:] == ['Max T-Rise', 'Avg of max T-Rise on each sample']
    assert block[1][2].data_type == 'f'
    assert block[5][2].value.startswith('=MAX(')
    assert block[-1][2].value.startswith('=AVERAGE(')
    assert block[1][2].number_format == '0.0'
    assert block[5][2].fill.fgColor != block[1][2].fill.fgColor
    assert block[-2][0].font.bold
    assert block[-2][2].border.bottom.style is not None
    raw_range = next(book.defined_names['KeyStageReadings'].destinations)[1]
    raw = sheet[raw_range.replace('$', '')]
    assert len(raw) == 7  # six genuine endpoints, no fabricated raw row for the inserted origin
    assert [cell.value for cell in raw[0][:2]] == ['扫描', '时间']
    assert raw[1][0].value == 81  # original scanner scan, distinct from source row 51
    assert raw[1][1].value is None  # This fixture has no scanner time; never invent one.
    assert raw[0][0].comment is None  # Details stay in the data, not cell annotations.
    assert not sheet.protection.sheet
    book.close()
    cached = load_workbook(output, data_only=True)
    values = cached['T-riseChart'][summary.replace('$', '')]
    assert [cell.value for cell in values[-2][1:]] == pytest.approx([p.maximum for p in analysis.points])
    assert [cell.value for cell in values[-1][1:]] == pytest.approx([p.average for p in analysis.points])
    cached.close()


@pytest.mark.parametrize('identifier_headers', [('扫描', '时间'), ('Scan', 'Time')])
@pytest.mark.parametrize('boolean_identifier', [False, True])
def test_endpoint_tables_share_scan_time_columns_without_changing_readings(tmp_path, identifier_headers, boolean_identifier):
    headers = (*identifier_headers, 'TC1', 'Spare', 'Ambient', 'Current')
    rows = ((10, '8/19 18:06:10', 29, 27, 25, 10),
            (50, '8/19 18:56:10', 30, 28, 25, 10),
            (100, '8/19 19:46:10', 39, 37, 25, 20),
            (150, '8/19 20:36:10', 40, 38, 25, 20),
            (200, '8/19 21:26:10', 54, 52, 25, 30),
            (300, '8/19 22:16:10', 55, 53, 25, 30))
    if boolean_identifier:
        rows = tuple((False, *row[1:]) for row in rows)
    table = WorkbookTable('scanner.xlsx', ('Data',), 'Data', (headers, *rows))
    selection = DataSelection(1, 2, 7, 5, 6, (4, 3), 2)
    prepared = prepare_data(table, selection)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=2, zero_intercept=True)
    output = tmp_path / 'scan-time.xlsx'
    TemperatureWorkbookGateway().write(output, table=table, selection=selection, prepared=prepared,
        analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=30, derating=None)
    book = load_workbook(output)
    sheet = book['T-riseChart']
    raw = sheet[next(book.defined_names['KeyStageReadings'].destinations)[1].replace('$', '')]
    rises = sheet[next(book.defined_names['StageTemperatureRise'].destinations)[1].replace('$', '')]
    for block in (raw, rises):
        assert [cell.value for cell in block[0][:2]] == ['扫描', '时间']
        assert all(cell.value != 'Original Row' for cell in block[0])
        assert block[0][1].alignment.horizontal == 'center'
        assert all(row[1].alignment.horizontal == 'left' for row in block[1:])
        assert all(row[1].alignment.vertical == 'center' for row in block[1:])
        assert all(row[0].alignment.horizontal == 'center' for row in block[1:])
    assert [tuple(cell.value for cell in row[:2]) for row in raw[1:]] == [rows[i][:2] for i in (1, 3, 5)]
    raw_row = raw[1][0].row
    assert rises[1][0].value is False if boolean_identifier else rises[1][0].value == f'=A{raw_row}'
    assert rises[1][1].value == f'=B{raw_row}'
    assert rises[1][2].value == f'=D{raw_row}-$E{raw_row}'  # Confirmed spare replaces the first channel.
    assert book['Initial Data']['A2'].value == rows[0][0]  # Original scan stays at its source position.
    book.close()
    cached = load_workbook(output, data_only=True)
    sheet = cached['T-riseChart']
    rises = sheet[next(cached.defined_names['StageTemperatureRise'].destinations)[1].replace('$', '')]
    assert [tuple(cell.value for cell in row[:2]) for row in rises[1:]] == [rows[i][:2] for i in (1, 3, 5)]
    assert [row[2].value for row in rises[1:]] == [3, 13, 28]
    summary = sheet[next(cached.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
    assert [cell.value for cell in summary[-2][1:]] == pytest.approx([0, 5, 15, 30])
    cached.close()


def test_report_export_has_editable_y_calculator_and_named_bold_colored_equations(tmp_path):
    output, analysis = _report_export(tmp_path)
    book = load_workbook(output)
    sheet = book['T-riseChart']
    area = next(book.defined_names['CurrentCalculator'].destinations)[1]
    calculator = sheet[area.replace('$', '')]
    assert [c.value for c in calculator[0]] == ['y (°C)', 'a', 'b', 'c', 'x (A)']
    assert calculator[1][0].value == 30
    assert all(cell.data_type == 'f' and 'ROUND(' in cell.value for cell in calculator[1][1:4])
    assert calculator[1][-1].data_type == 'f'
    assert 'SQRT' in calculator[1][-1].value
    assert calculator[1][-1].font.bold
    assert calculator[1][-1].font.color.rgb == '00FFA500'
    assert calculator[1][-1].number_format == '0'
    book.close()
    with ZipFile(output) as archive:
        ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
              'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        root = ET.fromstring(archive.read('xl/charts/chart1.xml'))
        labels = root.findall('.//c:trendline/c:trendlineLbl', ns)
        assert len(labels) == 2
        equations = [''.join(e.itertext()) for e in labels]
        assert 'Max T-Rise: y =' in equations[0]
        assert 'Avg of max T-Rise on each sample: y =' in equations[1]
        colors = []
        for label in labels:
            assert label.find('.//a:defRPr', ns).get('b') == '1'
            colors.append(label.find('.//a:srgbClr', ns).get('val'))
            assert label.find('c:tx/c:strRef/c:f', ns) is not None, 'Equation label stays linked when data is edited in Excel.'
        assert colors[0] != colors[1]
    cached = load_workbook(output, data_only=True)
    assert cached['T-riseChart'][calculator[1][-1].coordinate].value == pytest.approx(calculate_current(analysis.maximum_fit.coefficients, 30))
    cached.close()


@pytest.mark.parametrize('zero_intercept', [True, False])
def test_report_chart_names_follow_summary_cells_and_labels_stack_at_upper_left(tmp_path, zero_intercept):
    output, _ = _report_export(tmp_path, zero_intercept=zero_intercept)
    book = load_workbook(output)
    sheet = book['T-riseChart']
    assert sheet.freeze_panes is None
    assert book['Initial Data'].freeze_panes is None
    assert book['Derating'].freeze_panes is None  # No Derating was requested for this export.
    summary = sheet[next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')]
    labels = [row[0] for row in summary[-2:]]
    chart = sheet._charts[0]
    plot = chart.layout.manualLayout
    positions = []
    for series, cell in zip(chart.series, labels):
        title = series.tx.strRef
        assert title.f == f"'T-riseChart'!$A${cell.row}"
        assert title.strCache.pt[0].v == cell.value
        label = series.trendline.trendlineLbl
        equation_cell = sheet[label.tx.strRef.f.split('!')[1].replace('$', '')]
        name_cell = sheet[f'A{equation_cell.row}']
        assert name_cell.value == f'=A{cell.row}'
        assert label.tx.strRef.strCache.pt[0].v.startswith(cell.value + ': y = ')
        position = label.layout.manualLayout
        assert plot.x <= position.x <= plot.x + .02
        assert plot.y <= position.y <= plot.y + .11
        assert all(paragraph.pPr.algn == 'l' for paragraph in label.txPr.p)
        positions.append(position)
    assert positions[0].x == positions[1].x
    assert positions[1].y > positions[0].y
    book.close()


@pytest.mark.parametrize('constant_rise', [False, True])
def test_report_formulas_honor_spare_channel_order_current_scale_and_unconstrained_fit(tmp_path, constant_rise):
    headers = ('Current (mA)', 'Scanner Note', 'Spare', 'Ambient', 'TC1', 'Unused', 'Time')
    rows = []
    for current, rise in ((10000, 4), (20000, 9), (30000, 16)):
        value = 0 if constant_rise else rise
        rows.extend([(current, '=1+1', 25 + value, 25, 25 + value / 2, 999, '10:00')] * 2)
    table = WorkbookTable('reordered.xlsx', ('Data',), 'Data', (headers, *rows))
    selection = DataSelection(1, 2, 7, 4, 1, (5, 3), 2, current_multiplier=.001)
    prepared = prepare_data(table, selection)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=2, zero_intercept=False)
    output = tmp_path / 'mapped-report.xlsx'
    TemperatureWorkbookGateway().write(output, table=table, selection=selection, prepared=prepared,
        analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=None, derating=None)
    book = load_workbook(output)
    sheet = book['T-riseChart']
    raw = sheet[next(book.defined_names['KeyStageReadings'].destinations)[1].replace('$', '')]
    assert raw[1][0].value is None  # Scanner Note is not a scan counter.
    assert raw[1][1].value == '10:00'
    assert raw[1][3].value == '=1+1'
    assert raw[1][3].data_type == 's'
    rises = sheet[next(book.defined_names['StageTemperatureRise'].destinations)[1].replace('$', '')]
    raw_row = raw[1][0].row
    assert rises[1][2].value == f'=G{raw_row}-$F{raw_row}'  # TC1 first, replacement Spare second
    assert rises[1][3].value == f'=E{raw_row}-$F{raw_row}'
    assert rises[1][-1].value == f'=ROUND(C{raw_row}*0.001,0)'
    summary = next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')
    assert sheet._charts[0].series[0].trendline.intercept is None
    assert any('LINEST(' in cell.value and 'TRUE' in cell.value
               for row in sheet for cell in row if cell.data_type == 'f')
    book.close()
    cached = load_workbook(output, data_only=True)
    block = cached['T-riseChart'][summary]
    assert [cell.value for cell in block[0][1:]] == [10, 20, 30]
    assert [cell.value for cell in block[-2][1:]] == pytest.approx([p.maximum for p in analysis.points])
    assert [cell.value for cell in block[-1][1:]] == pytest.approx([p.average for p in analysis.points])
    cached.close()


def test_report_preserves_operator_coefficient_overrides_instead_of_replacing_them_with_fits(tmp_path):
    maximum, average = Coefficients(.005, .1, 0), Coefficients(.004, .2, 0)
    output, _ = _report_export(tmp_path, maximum, average)
    book = load_workbook(output)
    sheet = book['T-riseChart']
    maximum_area = next(book.defined_names['CurrentCalculator'].destinations)[1].replace('$', '')
    average_area = next(book.defined_names['AverageCoefficients'].destinations)[1].replace('$', '')
    assert [cell.value for cell in sheet[maximum_area][1][1:4]] == [.005, .1, 0]
    assert [cell.value for cell in sheet[average_area][1][1:4]] == [.004, .2, 0]
    book.close()


def test_unconstrained_export_keeps_measured_zero_baseline_and_normalizes_formula(tmp_path):
    headers = ('Scan', 'Time', 'TC1', 'TC2', 'Ambient', 'Current')
    rows = [(1, '10:00', 29, 28, 20, .009999), (2, '10:01', 24, 23, 22, .011427)]
    for current in (10, 20, 30):
        for _ in range(2):
            rows.append((len(rows) + 1, '10:02', 25 + .01 * current ** 2 + .2 * current + 2,
                         25 + .005 * current ** 2 + .1 * current + 1, 25, current))
    table = WorkbookTable('background.xlsx', ('Data',), 'Data', (headers, *rows))
    selection = DataSelection(1, 2, 9, 5, 6, (3, 4), 1)
    prepared = prepare_data(table, selection, acknowledge_warnings=True)
    analysis = analyze_temperature_rise(prepared.measurements, thermocouples_per_sample=1, zero_intercept=False)
    output = tmp_path / 'background-report.xlsx'
    TemperatureWorkbookGateway().write(output, table=table, selection=selection, prepared=prepared,
        analysis=analysis, maximum_coefficients=analysis.maximum_fit.coefficients,
        average_coefficients=analysis.average_fit.coefficients, target_rise=10, derating=None)
    book = load_workbook(output)
    sheet = book['T-riseChart']
    raw_area = next(book.defined_names['KeyStageReadings'].destinations)[1].replace('$', '')
    rise_area = next(book.defined_names['StageTemperatureRise'].destinations)[1].replace('$', '')
    summary_area = next(book.defined_names['TemperatureRiseSummary'].destinations)[1].replace('$', '')
    assert sheet[raw_area][1][0].value == 2
    assert sheet[raw_area][1][-1].value == .011427  # The scanner reading stays intact.
    formula = sheet[rise_area][1][-1].value
    assert formula.startswith('=ROUND(') and formula.endswith(',0)')
    for series in sheet._charts[0].series:
        assert series.trendline.intercept is None
        assert series.xVal.numRef.numCache.pt[0].v == 0
        assert series.yVal.numRef.numCache.pt[0].v in (2, 1.5)
    book.close()
    cached = load_workbook(output, data_only=True)
    summary = cached['T-riseChart'][summary_area]
    assert [cell.value for cell in summary[0][1:]] == [0, 10, 20, 30]
    assert [cell.value for cell in summary[-2][1:]] == [2, 5, 10, 17]
    assert [cell.value for cell in summary[-1][1:]] == [1.5, 3.75, 7.5, 12.75]
    assert cached['T-riseChart'][rise_area][1][-1].value == 0
    calculator = next(cached.defined_names['CurrentCalculator'].destinations)[1].replace('$', '')
    assert [cell.value for cell in cached['T-riseChart'][calculator][1]] == pytest.approx([10, .01, .2, 2, 20])
    cached.close()
