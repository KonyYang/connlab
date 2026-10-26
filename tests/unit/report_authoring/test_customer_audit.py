from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

from backend.infrastructure.office.report_authoring.customer_audit import audit_customer_package


def document(tmp_path, text="Copper Alloy and stainless steel housing"):
    path = tmp_path / "customer.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    return path


def add_parts(path, additions):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts.update(additions)
    with ZipFile(path, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)


def reviewed_thumbnail(path):
    # The test's python-docx starter includes a known thumbnail. Real callers
    # must obtain this decision from an explicit visual review, not this helper.
    with ZipFile(path) as archive:
        name = "docProps/thumbnail.jpeg"
        return {name: sha256(archive.read(name)).hexdigest()}


def test_safe_customer_package_passes_without_changing_source(tmp_path):
    path = document(tmp_path)
    original = path.read_bytes()
    result = audit_customer_package(path, sensitive_terms=["C18150", "PA66"], forbidden_project_ids=["DL-2026-09-021"], reviewed_parts=reviewed_thumbnail(path))
    assert result.passed and result.issues == ()
    assert path.read_bytes() == original


def test_sensitive_material_split_across_runs_and_header_cannot_pass(tmp_path):
    path = document(tmp_path)
    doc = Document(path)
    para = doc.add_paragraph()
    para.add_run("C18")
    para.add_run("150")
    doc.sections[0].header.paragraphs[0].text = "PA66 grade"
    doc.save(path)
    result = audit_customer_package(path, sensitive_terms=["C18150", "PA66"])
    assert not result.passed
    assert {issue.location for issue in result.issues if issue.code == "sensitive_text"} >= {"word/document.xml", "word/header1.xml"}


def test_revision_comment_and_properties_are_checked_not_just_visible_text(tmp_path):
    path = document(tmp_path)
    doc = Document(path)
    para = doc.add_paragraph()
    deletion = OxmlElement("w:del")
    para._p.append(deletion)
    run = OxmlElement("w:r")
    text = OxmlElement("w:delText")
    text.text = "C18150"
    run.append(text)
    deletion.append(run)
    doc.core_properties.comments = "DL-2026-09-021"
    doc.save(path)
    add_parts(path, {"word/comments.xml": b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:comment><w:p><w:r><w:t>PA66</w:t></w:r></w:p></w:comment></w:comments>'})
    result = audit_customer_package(path, sensitive_terms=["C18150", "PA66"], forbidden_project_ids=["DL-2026-09-021"])
    assert {issue.code for issue in result.issues} >= {"tracked_changes", "comments", "sensitive_text", "old_project_identity"}
    assert not result.passed


def test_external_target_and_unreadable_embedding_require_review(tmp_path):
    path = document(tmp_path)
    add_parts(path, {
        "word/embeddings/oleObject1.bin": b"opaque binary content",
        "word/_rels/header1.xml.rels": b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="image" TargetMode="External" Target="file:///D:/private/C18150.png"/></Relationships>',
    })
    result = audit_customer_package(path, sensitive_terms=["C18150"])
    assert {issue.code for issue in result.issues} >= {"unreviewed_embedding", "external_relationship", "sensitive_text"}
    assert not result.passed


def test_unused_old_image_is_reported_even_when_no_text_leaks(tmp_path):
    path = document(tmp_path)
    stream = BytesIO()
    Image.new("RGB", (12, 12), "red").save(stream, format="PNG")
    add_parts(path, {"word/media/retired.png": stream.getvalue()})
    result = audit_customer_package(path)
    assert any(issue.code == "unreferenced_media" for issue in result.issues)
    assert not result.passed


def test_reviewed_opaque_object_is_bound_to_exact_bytes(tmp_path):
    path = document(tmp_path)
    name, content = "word/embeddings/logo.bin", b"reviewed fixed PBrush logo"
    add_parts(path, {name: content})
    reviews = {**reviewed_thumbnail(path), name: sha256(content).hexdigest()}
    assert audit_customer_package(path, reviewed_parts=reviews).passed
    add_parts(path, {name: b"different object"})
    assert not audit_customer_package(path, reviewed_parts=reviews).passed


def test_visible_images_require_visual_check_but_never_count_as_unused(tmp_path):
    path = document(tmp_path)
    doc = Document(path)
    stream = BytesIO()
    Image.new("RGB", (12, 12), "blue").save(stream, format="PNG")
    doc.add_picture(BytesIO(stream.getvalue()))
    doc.save(path)
    result = audit_customer_package(path)
    assert not result.passed
    assert any(issue.code == "visual_review_required" for issue in result.issues)
    assert not any(issue.code == "unreferenced_media" for issue in result.issues)
    with ZipFile(path) as archive:
        reviews = {name: sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("word/media/")}
    reviews.update(reviewed_thumbnail(path))
    assert audit_customer_package(path, reviewed_parts=reviews).passed


def test_thumbnail_and_html_alternate_content_are_not_silently_ignored(tmp_path):
    path = document(tmp_path)
    add_parts(path, {"word/afchunk.html": b"<html><body>Private C18150</body></html>"})
    result = audit_customer_package(path, sensitive_terms=["C18150"])
    locations = {issue.location for issue in result.issues}
    assert "docProps/thumbnail.jpeg" in locations
    assert "word/afchunk.html" in locations
    assert not result.passed


def test_simple_field_external_content_requires_review(tmp_path):
    path = document(tmp_path)
    doc = Document(path)
    field = OxmlElement("w:fldSimple")
    from docx.oxml.ns import qn
    field.set(qn("w:instr"), 'INCLUDETEXT "D:/private/internal.docx"')
    doc.add_paragraph()._p.append(field)
    doc.save(path)
    result = audit_customer_package(path)
    assert any(issue.code == "external_field" for issue in result.issues)
