from dataclasses import asdict
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from openpyxl import load_workbook
import pytest

from backend.domain.temperature_rise import Measurement, analyze_temperature_rise
from backend.infrastructure.office.report_authoring.temperature_chart import insert_temperature_chart


def source_report(tmp_path):
    path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Fixed approved wording")
    document.sections[0].header.paragraphs[0].text = "Original Release"
    region = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "connlab-report-result:temperature-1")
    properties.append(tag)
    region.append(properties)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    content.append(paragraph)
    region.append(content)
    document._element.body.insert(1, region)
    document.save(path)
    return path


def analysis():
    rows = [Measurement(row, current, 20, temperatures) for row, current, temperatures in (
        (1, 10, (22, 21)), (2, 10, (22, 21)),
        (3, 20, (28, 24)), (4, 20, (28, 24)),
        (5, 30, (38, 29)), (6, 30, (38, 29)),
    )]
    return asdict(analyze_temperature_rise(rows, thermocouples_per_sample=1, zero_intercept=True))


def test_native_chart_uses_domain_results_and_preserves_other_report_parts(tmp_path):
    source = source_report(tmp_path)
    original = source.read_bytes()
    target = tmp_path / "updated.docx"
    insert_temperature_chart(source, target, analysis(), "connlab-report-result:temperature-1")
    assert source.read_bytes() == original
    with ZipFile(BytesIO(original)) as before, ZipFile(target) as after:
        for name in before.namelist():
            if name not in {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}:
                assert after.read(name) == before.read(name)
        charts = [name for name in after.namelist() if name.startswith("word/charts/") and name.endswith(".xml")]
        assert len(charts) == 1
        chart = etree.fromstring(after.read(charts[0]))
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        assert chart.find("c:externalData", ns) is not None
        assert len(chart.findall(".//c:ser", ns)) == 4
        for line in chart.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ln"):
            fills = [node for node in line if etree.QName(node).localname in {"noFill", "solidFill", "gradFill", "pattFill"}]
            assert len(fills) <= 1  # OOXML fill is a choice, not simultaneous values.
        embed = next(name for name in after.namelist() if name.endswith(".xlsx"))
        book = load_workbook(BytesIO(after.read(embed)), data_only=True)
        sheet = book.active
        assert sheet["A3"].value == 10
        assert sheet["B3"].value == 2
        assert sheet["C3"].value == 1.5
        assert sheet["F2"].value == "Model origin (not measured)"
        assert sheet["D4"].value == pytest.approx(8)
        book.close()
        text = after.read("word/document.xml").decode()
        assert "Fixed approved wording" in text
        assert "Figure" not in text


def test_chart_requires_unique_region_and_new_output(tmp_path):
    source = source_report(tmp_path)
    with pytest.raises(ValueError, match="region"):
        insert_temperature_chart(source, tmp_path / "missing.docx", analysis(), "missing")
    assert not (tmp_path / "missing.docx").exists()
    with pytest.raises(ValueError, match="existing"):
        insert_temperature_chart(source, source, analysis(), "connlab-report-result:temperature-1")


def test_negative_measured_baseline_remains_visible(tmp_path):
    rows = [Measurement(i, current, 20, (temperature,)) for i, (current, temperature) in enumerate(
        [(0, 19), (10, 21), (10, 21), (20, 27), (20, 27), (30, 37), (30, 37)], 1)]
    result = asdict(analyze_temperature_rise(rows, thermocouples_per_sample=1, zero_intercept=False))
    target = tmp_path / "negative-baseline.docx"
    insert_temperature_chart(source_report(tmp_path), target, result, "connlab-report-result:temperature-1")
    with ZipFile(target) as archive:
        name = next(name for name in archive.namelist() if name.startswith("word/charts/") and name.endswith(".xml"))
        chart = etree.fromstring(archive.read(name))
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        y_axis = chart.findall(".//c:valAx", ns)[1]
        assert float(y_axis.find("c:scaling/c:min", ns).get("val")) <= -1
