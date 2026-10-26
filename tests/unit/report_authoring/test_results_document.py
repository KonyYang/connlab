from io import BytesIO
from zipfile import ZipFile
import pytest
from docx import Document
from backend.infrastructure.office.report_authoring.results_document import ReportResultsDocument


def fixture_report(path):
    doc = Document()
    doc.add_paragraph('Purpose: approved manual wording').runs[0].bold = True
    doc.add_paragraph('Group 1 Test Results')
    table = doc.add_table(rows=2, cols=4)
    for cell, value in zip(table.rows[0].cells, ['Step', 'Test Item', 'Result', 'Comment']):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ['2', 'Retention Force', 'Old value', 'Pass']):
        cell.text = value
    doc.add_paragraph('Appendix B: Old Orange Mating Force')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Old data'
    doc.add_paragraph('Fixed ending')
    doc.save(path)


def mechanical():
    return {'result_id': 'r1', 'kind': 'mechanical', 'target': {'group_label': '1', 'step_token': '2', 'group_id': 'g', 'row_id': 'r'},
        'outcome': 'pending', 'data': {'test_type': 'retention', 'quantity': 'hold', 'unit': 'N',
        'rows': [{'sample_id': 'SN1', 'record_id': 'rec9', 'color': 'Orange', 'value': '410', 'hold_seconds': '6'},
                 {'sample_id': 'SN2', 'record_id': 'rec10', 'color': 'Black', 'value': '420', 'hold_seconds': '6'}],
        'summary': {'min': '410', 'max': '420', 'average': '415', 'stdev': '7.071'}}}


def test_mechanical_replaces_selected_appendix_and_result_without_fake_pass(tmp_path):
    source, output = tmp_path/'source.docx', tmp_path/'output.docx'
    fixture_report(source)
    before = source.read_bytes()
    writer = ReportResultsDocument()
    targets = writer.inspect(source, mechanical())
    assert targets['result_rows'] == [{'table_index': 0, 'row_index': 1, 'step': '2', 'test_item': 'Retention Force'}]
    assert targets['appendix_tables'][0]['table_index'] == 1
    writer.update(source, output, mechanical(), result_table=0, result_row=1, appendix_table=1,
                  appendix_title='Appendix B: Retention Force — Orange and Black')
    after = Document(output)
    assert after.tables[0].cell(1, 3).text == ''
    assert '415' in after.tables[0].cell(1, 2).text
    assert 'hold' in after.tables[0].cell(1, 2).text
    assert 'Pending' not in '\n'.join(p.text for p in after.paragraphs)
    assert after.paragraphs[0].runs[0].bold
    assert after.paragraphs[-1].text == 'Fixed ending'
    with ZipFile(output) as archive:
        xml = archive.read('word/document.xml').decode()
        assert 'SN1' in xml and 'rec9' in xml and 'Black' in xml and 'Orange' in xml
        assert 'Old Orange Mating Force' not in xml
        assert 'Retention Force' in xml and 'Orange and Black' in xml
    assert source.read_bytes() == before
    with ZipFile(BytesIO(before)) as old, ZipFile(output) as new:
        assert all(old.read(name) == new.read(name) for name in old.namelist() if name != 'word/document.xml')


def test_wrong_result_target_and_existing_output_fail_without_mutation(tmp_path):
    source = tmp_path/'source.docx'
    fixture_report(source)
    writer = ReportResultsDocument()
    with pytest.raises(ValueError, match='Group and Step'):
        writer.update(source, tmp_path/'no.docx', mechanical(), result_table=1, result_row=1, appendix_table=None, appendix_title='Appendix C')
    with pytest.raises(ValueError, match='exists'):
        writer.update(source, source, mechanical(), result_table=0, result_row=1, appendix_table=None, appendix_title='Appendix C')


def test_appendix_title_is_derived_from_confirmed_measurement_not_old_mating_label(tmp_path):
    source, output = tmp_path/'source.docx', tmp_path/'updated.docx'
    fixture_report(source)
    ReportResultsDocument().update(source, output, mechanical(), result_table=0, result_row=1,
        appendix_table=1, appendix_title='Appendix B: Old Orange Mating Force')
    paragraphs = [p.text for p in Document(output).paragraphs]
    assert 'Appendix B: Retention Force — Hold (N); Orange and Black' in paragraphs
    assert not any('Mating' in paragraph for paragraph in paragraphs)


def test_new_appendix_precedes_end_marker(tmp_path):
    source, output = tmp_path/'source.docx', tmp_path/'updated.docx'
    fixture_report(source)
    doc = Document(source); doc.add_paragraph('*** End of Report ***'); doc.save(source)
    ReportResultsDocument().update(source, output, mechanical(), result_table=0, result_row=1,
        appendix_table=None, appendix_title='Appendix D')
    with ZipFile(output) as archive:
        xml = archive.read('word/document.xml').decode()
        assert xml.index('Appendix D: Retention') < xml.index('End of Report')


def test_result_update_protects_identified_group_starts_without_reformatting_other_blocks(tmp_path):
    from lxml import etree
    from backend.infrastructure.office.report_authoring.word_copy import NS
    source, output = tmp_path/'source.docx', tmp_path/'updated.docx'
    fixture_report(source)
    document = Document(source)
    heading = document.add_paragraph('Group B Test Results')
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = False
    document.add_paragraph('')
    table = document.add_table(rows=4, cols=3)
    for cell, label in zip(table.rows[0].cells, ['No.', 'Result', 'Comment']): cell.text = label
    for row in table.rows[1:]:
        for cell in row.cells: cell.text = 'Unchanged downstream result'
    document.add_paragraph('Manual ending').paragraph_format.keep_with_next = False
    document.save(source)
    evidence = ReportResultsDocument().update(source, output, mechanical(), result_table=0, result_row=1,
        appendix_table=1, appendix_title='Appendix B')
    result = Document(output)
    group = next(p for p in result.paragraphs if p.text == 'Group B Test Results')
    assert group.paragraph_format.keep_with_next is True
    assert group.paragraph_format.page_break_before is True
    assert result.paragraphs[-1].paragraph_format.keep_with_next is False
    assert 'Group B Test Results' in evidence['pagination_groups']
    with ZipFile(source) as before, ZipFile(output) as after:
        old, new = [etree.fromstring(z.read('word/document.xml')) for z in (before, after)]
        old_rows = old.findall('w:body/w:tbl', NS)[-1].findall('w:tr', NS)
        new_rows = new.findall('w:body/w:tbl', NS)[-1].findall('w:tr', NS)
        assert all(etree.tostring(a) == etree.tostring(b) for a, b in zip(old_rows[2:], new_rows[2:]))


def test_temperature_explicit_region_replaces_old_table_chart_and_conclusion_only(tmp_path):
    from dataclasses import asdict
    from backend.infrastructure.office.temperature_workbook_gateway import TemperatureWorkbookGateway
    from backend.application.tools_temperature_service import ToolsTemperatureService
    from backend.domain.temperature_data import DataSelection
    from backend.infrastructure.office.report_authoring.temperature_chart import insert_temperature_chart
    from PIL import Image
    source, output = tmp_path/'temperature.docx', tmp_path/'updated.docx'
    image = BytesIO(); Image.new('RGB', (40, 20), 'red').save(image, 'PNG')
    doc = Document(); doc.add_paragraph('Group C Test Results')
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['Step', 'Result', 'Comment']): cell.text = value
    table.cell(1, 0).text = '2'; table.cell(1, 1).text = 'Old rise'; table.cell(1, 2).text = 'Pass'
    doc.add_picture(BytesIO(image.getvalue())); doc.add_paragraph('Figure 14: Setup of Temperature Rise Testing')
    doc.add_paragraph('Table-1: Temperature Rise Measurements (Unit: °C)')
    old = doc.add_table(rows=1, cols=2); old.cell(0, 0).text = 'Applied current'; old.cell(0, 1).text = 'OLD MEASUREMENT'
    doc.add_picture(BytesIO(image.getvalue()))
    doc.add_paragraph('The maximum rise reached 45°C at 496.9A.')
    doc.add_paragraph('Group D Test Results'); doc.add_paragraph('KEEP UNRELATED NEXT GROUP'); doc.save(source)
    workbook = TemperatureWorkbookGateway()
    data = workbook.read_upload(b'Ambient,I,T1\n20,100,26\n20,100,26\n20,200,45\n20,200,45\n', 'data.csv')
    selection = DataSelection(header_row=1, start_row=2, end_row=5, ambient_column=1, current_column=2,
        temperature_columns=(3,), thermocouples_per_sample=1, sample_count=1)
    analysis = ToolsTemperatureService(workbook).analyze(data, selection, zero_intercept=True)
    result = {'result_id': 'new-temperature', 'kind': 'temperature', 'target': {'group_label': 'C', 'step_token': '2'},
        'outcome': 'pending', 'data': {'analysis': asdict(analysis), 'selection': {'ambient_column': 1, 'temperature_columns': [3]},
            'raw_rows': [{'source_row': i+1, 'values': list(row)} for i, row in enumerate(data.rows)]}}
    writer = ReportResultsDocument(temperature_chart=insert_temperature_chart)
    preview = writer.inspect(source, result)
    blocks = preview['evidence_blocks']
    start = next(item['block_index'] for item in blocks if item['text'].startswith('Table-1'))
    end = next(item['block_index'] for item in blocks if '496.9A' in item['text'])
    assert preview['unreconciled_temperature_evidence']
    with pytest.raises(ValueError, match='competing appendix'):
        writer.update(source, tmp_path/'competing.docx', result, result_table=0, result_row=1,
            appendix_table=None, appendix_title='Appendix D')
    writer.update(source, output, result, result_table=0, result_row=1, appendix_table=None,
        appendix_title='', evidence_start=start, evidence_end=end)
    with ZipFile(output) as package:
        xml = package.read('word/document.xml').decode()
        assert 'OLD MEASUREMENT' not in xml and '496.9A' not in xml
        assert 'Figure 14' in xml and 'KEEP UNRELATED NEXT GROUP' in xml
        assert 'Avg of Max' in xml and 'Ambient (°C)' in xml
    assert not writer.inspect(output, result)['unreconciled_temperature_evidence']
    with pytest.raises(ValueError, match='region'):
        writer.update(source, tmp_path/'bad.docx', result, result_table=0, result_row=1, appendix_table=None,
            appendix_title='', evidence_start=start, evidence_end=end+1)
    with pytest.raises(ValueError, match='complete old temperature'):
        writer.update(source, tmp_path/'partial.docx', result, result_table=0, result_row=1, appendix_table=None,
            appendix_title='', evidence_start=start, evidence_end=start+1)
    with pytest.raises(ValueError, match='complete old temperature'):
        writer.update(source, tmp_path/'conclusion-only.docx', result, result_table=0, result_row=1, appendix_table=None,
            appendix_title='', evidence_start=end, evidence_end=end)
    # The old chart shares its image with the retained setup: removing the chart must keep that image.
    with ZipFile(source) as before, ZipFile(output) as after:
        assert after.read('word/media/image1.png') == before.read('word/media/image1.png')
    updated_blocks = writer.inspect(output, result)['evidence_blocks']
    owned = next(block['block_index'] for block in updated_blocks if block['owned_current_result'])
    repeated = tmp_path/'repeated.docx'
    writer.update(output, repeated, result, result_table=0, result_row=1, appendix_table=None,
        appendix_title='', evidence_start=owned, evidence_end=owned)
    with ZipFile(repeated) as package:
        xml = package.read('word/document.xml').decode()
        assert xml.count('connlab-report-result:new-temperature') == 1
        assert len([name for name in package.namelist() if name.startswith('word/charts/') and name.endswith('.xml') and '/_rels/' not in name]) == 1
        assert '496.9A' not in xml
