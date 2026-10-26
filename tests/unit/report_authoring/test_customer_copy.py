from io import BytesIO
from zipfile import ZipFile
from docx import Document
from backend.infrastructure.office.report_authoring.customer_copy import CustomerReportCopy
from backend.infrastructure.office.report_authoring.temperature_chart import insert_temperature_chart
from backend.application.tools_temperature_service import ToolsTemperatureService
from backend.infrastructure.office.temperature_workbook_gateway import TemperatureWorkbookGateway
from backend.domain.temperature_data import DataSelection
from dataclasses import asdict


def test_customer_copy_keeps_complete_temperature_table_and_editable_curve(tmp_path):
    source, template, output = (tmp_path/name for name in ('source.docx', 'template.docx', 'customer.docx'))
    for path in (source, template):
        doc = Document(); doc.add_paragraph('Fixed customer wording'); doc.add_paragraph('*** End of Report ***'); doc.save(path)
    originals = source.read_bytes(), template.read_bytes()
    class Converter:
        def generate_customer_report(self, *, source_path, template_path, output_path):
            output_path.write_bytes(template_path.read_bytes())
    workbook = TemperatureWorkbookGateway()
    table = workbook.read_upload(b'Ambient,I,T1,T2\n20,100,26,28\n20,100,26,28\n20,200,41,45\n20,200,41,45\n', 'data.csv')
    analysis = ToolsTemperatureService(workbook).analyze(table,
        DataSelection(header_row=1, start_row=2, end_row=5, ambient_column=1, current_column=2,
            temperature_columns=(3, 4), thermocouples_per_sample=1, sample_count=2), zero_intercept=True)
    result = {'kind': 'temperature', 'result_id': 'r1', 'target': {'group_label': '1', 'step_token': '2'},
        'data': {'analysis': asdict(analysis), 'selection': {'ambient_column': 1, 'temperature_columns': [3, 4]},
            'raw_rows': [{'source_row': index + 1, 'values': list(values)} for index, values in enumerate(table.rows)]}}
    adapter = CustomerReportCopy(converter=Converter(), template=lambda: template, temperature_chart=insert_temperature_chart)
    adapter.generate(source, output, [result])
    with ZipFile(output) as archive:
        body = archive.read('word/document.xml').decode()
        assert 'Channel 1' in body and 'Avg of Max' in body and '25' in body
        assert 'Ambient (°C)' in body and 'Channel 1 measured' in body and '41' in body
        assert 'zero-intercept: yes' in body and 'R²=' in body
        assert 'tblBorders' in body and 'rFonts' in body
        assert body.index('Temperature rise') < body.index('End of Report')
        assert 'Figure' not in body
        assert any(name.startswith('word/charts/') for name in archive.namelist())
        assert any(name.startswith('word/embeddings/') and name.endswith('.xlsx') for name in archive.namelist())
    audit = adapter.audit(output, sensitive_terms=[], forbidden_project_ids=[], reviewed_parts={})
    assert not audit['passed']
    assert any(item['code'] == 'unreviewed_embedding' for item in audit['issues'])
    assert (source.read_bytes(), template.read_bytes()) == originals


def test_customer_pagination_keeps_group_heading_header_and_first_row_without_chaining_table(tmp_path):
    from lxml import etree
    from backend.infrastructure.office.report_authoring.word_copy import NS
    source, output = tmp_path/'source.docx', tmp_path/'customer.docx'
    document = Document()
    document.add_paragraph('Fixed text before the result group')
    heading = document.add_paragraph('Group Ad Test Results')
    heading.paragraph_format.keep_with_next = False
    table = document.add_table(rows=4, cols=3)
    for cell, label in zip(table.rows[0].cells, ['No.', 'Result', 'Comment']): cell.text = label
    for index, row in enumerate(table.rows[1:], 1):
        for column, cell in enumerate(row.cells):
            cell.text = f'value-{index}-{column}'
            cell.paragraphs[0].paragraph_format.keep_with_next = True if index == 1 else False
    document.add_paragraph('Fixed ending'); document.save(source)
    before = source.read_bytes()
    class Converter:
        def generate_customer_report(self, *, source_path, template_path, output_path):
            output_path.write_bytes(source_path.read_bytes())
    CustomerReportCopy(converter=Converter(), template=lambda: source, temperature_chart=None).generate(source, output, [])
    with ZipFile(source) as old, ZipFile(output) as new:
        old_root = etree.fromstring(old.read('word/document.xml'))
        root = etree.fromstring(new.read('word/document.xml'))
        assert root.xpath('//w:t/text()', namespaces=NS) == old_root.xpath('//w:t/text()', namespaces=NS)
        heading = root.find('w:body', NS)[1]
        assert heading.find('w:pPr/w:keepNext', NS).get('{'+NS['w']+'}val', '1') == '1'
        rows = root.find('w:body/w:tbl', NS).findall('w:tr', NS)
        assert rows[0].find('w:trPr/w:tblHeader', NS) is not None
        for row in rows[:2]: assert row.find('w:trPr/w:cantSplit', NS) is not None
        assert all(p.find('w:pPr/w:keepNext', NS).get('{'+NS['w']+'}val', '1') == '1' for p in rows[0].xpath('.//w:p', namespaces=NS))
        assert all(p.find('w:pPr/w:keepNext', NS).get('{'+NS['w']+'}val') == '0' for p in rows[1].xpath('.//w:p', namespaces=NS))
        assert etree.tostring(rows[2]) == etree.tostring(old_root.find('w:body/w:tbl', NS).findall('w:tr', NS)[2])
        assert all(old.read(name) == new.read(name) for name in old.namelist() if name != 'word/document.xml')
    assert source.read_bytes() == before


def test_customer_audit_requires_actual_current_identity_not_a_review_checkbox(tmp_path):
    source = tmp_path/'customer.docx'
    document = Document(); document.add_paragraph('Fixed approved wording'); document.save(source)
    adapter = CustomerReportCopy(converter=None, template=None, temperature_chart=None)
    audit = adapter.audit(source, sensitive_terms=[], forbidden_project_ids=[], reviewed_parts={},
        expected_project_dl='DL-2026-10-099')
    assert not audit['passed']
    assert any(item['code'] == 'missing_current_project_identity' for item in audit['issues'])
    header = document.sections[0].header.paragraphs[0]
    header.add_run('DL-2026-'); header.add_run('10-099'); document.save(source)
    updated = adapter.audit(source, sensitive_terms=[], forbidden_project_ids=[], reviewed_parts={},
        expected_project_dl='DL-2026-10-099')
    assert not any(item['code'] == 'missing_current_project_identity' for item in updated['issues'])
