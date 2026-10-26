import pytest
from docx import Document
from backend.infrastructure.office.report_authoring.word_copy import ReportCopyWriter


def test_cascading_fact_replacements_reject_without_changing_output(tmp_path):
    source, target = tmp_path/'source.docx', tmp_path/'output.docx'
    doc = Document(); doc.add_paragraph('SourceA;SourceB'); doc.save(source)
    original = source.read_bytes()
    with pytest.raises(ValueError, match='cascading'):
        ReportCopyWriter().apply_facts(source, target, [('SourceA', 'SourceB'), ('SourceB', 'SourceC')])
    assert not target.exists() and source.read_bytes() == original


def test_reference_identity_split_across_runs_is_replaced_without_changing_format(tmp_path):
    source, target = tmp_path/'source.docx', tmp_path/'copy.docx'
    document = Document()
    header = document.sections[0].header.paragraphs[0]
    header.add_run('DL-2026-').bold = True
    header.add_run('09-031').italic = True
    document.add_paragraph('Approved fixed wording. DL-2026-09-0310 is a different identifier.')
    document.save(source)
    ReportCopyWriter().prepare(source, target, source_dl='DL-2026-09-031', project_dl='DL-2026-10-099', reference=True)
    result = Document(target)
    assert result.sections[0].header.paragraphs[0].text == 'DL-2026-10-099'
    assert any(run.bold for run in result.sections[0].header.paragraphs[0].runs)
    assert result.paragraphs[0].text == 'Approved fixed wording. DL-2026-09-0310 is a different identifier.'


def test_template_without_old_identity_does_not_run_an_empty_replacement(tmp_path):
    source, target = tmp_path/'template.docx', tmp_path/'copy.docx'
    document = Document(); document.add_paragraph('Fixed approved wording.'); document.save(source)
    warnings = ReportCopyWriter().prepare(source, target, source_dl='', project_dl='DL-2026-10-099', reference=True)
    assert Document(target).paragraphs[0].text == 'Fixed approved wording.'
    assert any('identity' in message.casefold() for message in warnings)
