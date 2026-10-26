from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from lxml import etree
from PIL import Image
import pytest

from backend.infrastructure.office.report_authoring.photo_slots import (
    PhotoReplacement, inspect_photo_slots, read_photo_slot, replace_photo_slots,
)


def picture(colour, size=(80, 60)):
    stream = BytesIO()
    Image.new("RGB", size, colour).save(stream, format="PNG")
    return stream.getvalue()


def report(tmp_path):
    path = tmp_path / "approved.docx"
    document = Document()
    document.add_paragraph("Approved fixed wording")
    original = picture("red")
    for _ in range(2):
        document.add_picture(BytesIO(original), width=Inches(2))
        document.add_paragraph("Figure: approved specimen")
    document.add_paragraph("*** End of Report ***")
    document.save(path)
    return path, original


def test_only_selected_occurrence_changes_when_two_slots_share_media(tmp_path):
    source, original = report(tmp_path)
    before = source.read_bytes()
    slots = inspect_photo_slots(source)
    assert len(slots) == 2
    assert slots[0].slot_id != slots[1].slot_id
    assert read_photo_slot(source, slots[0].slot_id) == (original, "image/png")
    output = tmp_path / "copy.docx"
    replacement = picture("blue")
    result = replace_photo_slots(source, output,
        [PhotoReplacement(slots[0].slot_id, replacement)], expected_source_sha256=sha256(before).hexdigest())
    assert result.changed
    assert source.read_bytes() == before
    updated = inspect_photo_slots(output)
    assert read_photo_slot(output, updated[0].slot_id)[0] == replacement
    assert read_photo_slot(output, updated[1].slot_id)[0] == original
    with ZipFile(source) as a, ZipFile(output) as b:
        for name in a.namelist():
            if name not in {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}:
                assert a.read(name) == b.read(name), name
    assert Document(output).paragraphs[-1].text == "*** End of Report ***"


def test_identical_replacement_is_byte_identical_and_never_overwrites(tmp_path):
    source, original = report(tmp_path)
    slot = inspect_photo_slots(source)[0]
    target = tmp_path / "copy.docx"
    result = replace_photo_slots(source, target, [PhotoReplacement(slot.slot_id, original)],
        expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    assert not result.changed
    assert target.read_bytes() == source.read_bytes()
    with pytest.raises(FileExistsError):
        replace_photo_slots(source, target, [PhotoReplacement(slot.slot_id, original)],
            expected_source_sha256=sha256(source.read_bytes()).hexdigest())


def test_stale_source_or_unknown_slot_writes_nothing(tmp_path):
    source, _ = report(tmp_path)
    target = tmp_path / "copy.docx"
    with pytest.raises(ValueError, match="changed"):
        replace_photo_slots(source, target, [], expected_source_sha256="0" * 64)
    with pytest.raises(ValueError, match="slot"):
        replace_photo_slots(source, target, [PhotoReplacement("missing", picture("blue"))],
            expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    assert not target.exists()


def test_crop_and_shape_geometry_remain_unchanged_for_compatible_image(tmp_path):
    source, _ = report(tmp_path)
    doc = Document(source)
    shape = doc.inline_shapes[0]
    crop = OxmlElement("a:srcRect")
    crop.set("l", "10000")
    shape._inline.graphic.graphicData.pic.blipFill.blip.addnext(crop)
    doc.save(source)
    slot = inspect_photo_slots(source)[0]
    output = tmp_path / "copy.docx"
    replace_photo_slots(source, output, [PhotoReplacement(slot.slot_id, picture("blue"))],
        expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    left = Document(source).inline_shapes[0]._inline
    right = Document(output).inline_shapes[0]._inline
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    # Ignore only the intentionally replaced relationship, not crop/extent/anchor formatting.
    left = etree.fromstring(etree.tostring(left))
    right = etree.fromstring(etree.tostring(right))
    for node in (left, right):
        for blip in node.xpath(".//a:blip", namespaces=ns):
            blip.set("{" + ns["r"] + "}embed", "selected-photo")
    assert etree.tostring(left) == etree.tostring(right)


def test_different_aspect_requires_crop_choice_instead_of_stretching(tmp_path):
    source, _ = report(tmp_path)
    slot = inspect_photo_slots(source)[0]
    target = tmp_path / "copy.docx"
    with pytest.raises(ValueError, match="crop"):
        replace_photo_slots(source, target, [PhotoReplacement(slot.slot_id, picture("blue", (60, 80)))],
            expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    assert not target.exists()
    replace_photo_slots(source, target,
        [PhotoReplacement(slot.slot_id, picture("blue", (60, 80)), crop_mode="centre_crop")],
        expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    assert Document(target).inline_shapes[0].width == Document(source).inline_shapes[0].width


def test_deleted_picture_is_not_a_visible_slot(tmp_path):
    source, _ = report(tmp_path)
    doc = Document(source)
    drawing = doc.inline_shapes[0]._inline.getparent()
    run = drawing.getparent()
    parent = run.getparent()
    deletion = OxmlElement("w:del")
    parent.replace(run, deletion)
    deletion.append(run)
    doc.save(source)
    assert len(inspect_photo_slots(source)) == 1


def test_transformed_slot_requires_explicit_acceptance_and_preserves_transform(tmp_path):
    source, _ = report(tmp_path)
    doc = Document(source)
    transform = doc.inline_shapes[0]._inline.graphic.graphicData.pic.spPr.xfrm
    transform.set("rot", "5400000")
    transform.set("flipH", "1")
    doc.save(source)
    slot = inspect_photo_slots(source)[0]
    assert slot.rotation == 5400000 and slot.flip_horizontal
    output = tmp_path / "copy.docx"
    with pytest.raises(ValueError, match="rotation"):
        replace_photo_slots(source, output, [PhotoReplacement(slot.slot_id, picture("blue"))],
            expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    assert not output.exists()
    replace_photo_slots(source, output,
        [PhotoReplacement(slot.slot_id, picture("blue"), accept_slot_transform=True)],
        expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    updated = inspect_photo_slots(output)[0]
    assert updated.rotation == slot.rotation and updated.flip_horizontal


def test_external_image_is_not_followed_or_replaced(tmp_path):
    source, _ = report(tmp_path)
    with ZipFile(source) as archive:
        contents = {item.filename: archive.read(item) for item in archive.infolist()}
    relations = etree.fromstring(contents["word/_rels/document.xml.rels"])
    for relation in relations:
        if relation.get("Type", "").endswith("/image"):
            relation.set("TargetMode", "External")
            relation.set("Target", "https://example.invalid/private-image.png")
    contents["word/_rels/document.xml.rels"] = etree.tostring(relations)
    with ZipFile(source, "w") as archive:
        for name, content in contents.items():
            archive.writestr(name, content)
    slots = inspect_photo_slots(source)
    assert all(not slot.supported for slot in slots)
    output = tmp_path / "copy.docx"
    with pytest.raises(ValueError, match="Linked"):
        replace_photo_slots(source, output, [PhotoReplacement(slots[0].slot_id, picture("blue"))],
            expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    assert not output.exists()


def test_exif_rotated_same_bytes_is_noop_without_reinterpreting_image(tmp_path):
    stream = BytesIO()
    photo = Image.new("RGB", (80, 60), "red")
    exif = Image.Exif()
    exif[274] = 6
    photo.save(stream, format="JPEG", exif=exif)
    source = tmp_path / "original.docx"
    doc = Document()
    doc.add_picture(BytesIO(stream.getvalue()), width=Inches(2))
    doc.save(source)
    output = tmp_path / "copy.docx"
    result = replace_photo_slots(source, output, [PhotoReplacement(inspect_photo_slots(source)[0].slot_id, stream.getvalue())],
        expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    assert not result.changed
    assert output.read_bytes() == source.read_bytes()


def test_same_centre_crop_replay_is_noop(tmp_path):
    source, _ = report(tmp_path)
    change = PhotoReplacement(inspect_photo_slots(source)[0].slot_id, picture("blue", (60, 80)), crop_mode="centre_crop")
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    replace_photo_slots(source, first, [change], expected_source_sha256=sha256(source.read_bytes()).hexdigest())
    result = replace_photo_slots(first, second, [change], expected_source_sha256=sha256(first.read_bytes()).hexdigest())
    assert not result.changed
    assert first.read_bytes() == second.read_bytes()


def test_moved_from_image_is_not_a_current_photo(tmp_path):
    source, _ = report(tmp_path)
    doc = Document(source)
    drawing = doc.inline_shapes[0]._inline.getparent()
    run = drawing.getparent()
    parent = run.getparent()
    moved = OxmlElement("w:moveFrom")
    parent.replace(run, moved)
    moved.append(run)
    doc.save(source)
    assert len(inspect_photo_slots(source)) == 1
