from __future__ import annotations

from pathlib import Path

import msoffcrypto
import pytest
from docx import Document

from backend.infrastructure.office.office_file_password_gateway import (
    OfficeFilePasswordGateway,
)


class _ComRuntime:
    def __init__(self) -> None:
        self.initialized = 0
        self.uninitialized = 0

    def CoInitialize(self) -> None:
        self.initialized += 1

    def CoUninitialize(self) -> None:
        self.uninitialized += 1


class _OfficeDocument:
    def __init__(self, output_writer=None) -> None:
        self.SaveFormat = 16
        self.FileFormat = 51
        self.Password = ""
        self.output_writer = output_writer
        self.closed = False

    def SaveAs2(self, **kwargs) -> None:
        Path(kwargs["FileName"]).write_bytes(b"word-encrypted")
        kwargs["Password"] = self.Password
        self.output_writer.append(kwargs)

    def SaveAs(self, *args, **kwargs) -> None:
        if kwargs:
            Path(kwargs["Filename"]).write_bytes(b"excel-encrypted")
            self.output_writer.append(kwargs)
        else:
            Path(args[0]).write_bytes(b"powerpoint-encrypted")
            self.output_writer.append({"path": args[0], "password": self.Password})

    def Close(self, *_args, **_kwargs) -> None:
        self.closed = True


class _Documents:
    def __init__(self, saves: list[dict], opens: list[dict]) -> None:
        self.saves = saves
        self.opens = opens

    def Open(self, **kwargs):
        self.opens.append(kwargs)
        return _OfficeDocument(self.saves)


class _Presentations:
    def __init__(self, saves: list[dict], opens: list[dict]) -> None:
        self.saves = saves
        self.opens = opens

    def Open(self, **kwargs):
        self.opens.append(kwargs)
        return _OfficeDocument(self.saves)


class _App:
    def __init__(self, kind: str, saves: list[dict], opens: list[dict]) -> None:
        self.quit = False
        self.Documents = _Documents(saves, opens)
        self.Workbooks = _Documents(saves, opens)
        self.Presentations = _Presentations(saves, opens)

    def Quit(self) -> None:
        self.quit = True


def test_docx_password_roundtrip_does_not_start_word_com(tmp_path: Path) -> None:
    source = tmp_path / "report.docx"
    Document().save(source)
    protected = tmp_path / "protected.docx"
    decrypted = tmp_path / "decrypted.docx"
    gateway = OfficeFilePasswordGateway(
        dispatch=lambda _kind: pytest.fail("DOCX encryption must not start Word COM"),
        com_runtime=_ComRuntime(),
    )

    gateway.encrypt_and_verify(
        source_path=source,
        output_path=protected,
        password="DGLAB",
        office_kind="word",
    )
    with protected.open("rb") as protected_stream:
        package = msoffcrypto.OfficeFile(protected_stream)
        package.load_key(password="DGLAB", verify_password=True)
        with decrypted.open("wb") as decrypted_stream:
            package.decrypt(decrypted_stream)

    assert decrypted.read_bytes() == source.read_bytes()


def test_legacy_doc_and_excel_save_with_open_password_then_verify(tmp_path: Path) -> None:
    runtime = _ComRuntime()
    saves: list[dict] = []
    opens: list[dict] = []
    apps: list[_App] = []

    def dispatch(kind: str) -> _App:
        app = _App(kind, saves, opens)
        apps.append(app)
        return app

    verified: list[tuple[Path, str]] = []
    gateway = OfficeFilePasswordGateway(
        dispatch=dispatch,
        com_runtime=runtime,
        output_verifier=lambda path, kind: verified.append((path, kind)),
    )
    word = tmp_path / "report.doc"
    excel = tmp_path / "data.xlsx"
    word.write_bytes(b"word")
    excel.write_bytes(b"excel")

    gateway.encrypt_and_verify(
        source_path=word,
        output_path=tmp_path / "word-stage.doc",
        password="DGLAB",
        office_kind="word",
    )
    gateway.encrypt_and_verify(
        source_path=excel,
        output_path=tmp_path / "excel-stage.xlsx",
        password="202608007",
        office_kind="excel",
    )

    assert saves[0]["Password"] == "DGLAB"
    assert saves[1]["Password"] == "202608007"
    assert opens[0]["PasswordDocument"] == "DGLAB"
    assert opens[0]["WritePasswordDocument"] == "DGLAB"
    assert verified == [
        (tmp_path / "word-stage.doc", "word"),
        (tmp_path / "excel-stage.xlsx", "excel"),
    ]
    assert runtime.initialized == runtime.uninitialized == 2
    assert all(app.quit for app in apps)


def test_powerpoint_sets_open_password_and_runs_encrypted_container_verifier(tmp_path: Path) -> None:
    runtime = _ComRuntime()
    saves: list[dict] = []
    opens: list[dict] = []
    verified: list[Path] = []
    gateway = OfficeFilePasswordGateway(
        dispatch=lambda kind: _App(kind, saves, opens),
        com_runtime=runtime,
        output_verifier=lambda path, _kind: verified.append(path),
    )
    source = tmp_path / "slides.pptx"
    output = tmp_path / "slides-stage.pptx"
    source.write_bytes(b"powerpoint")

    gateway.encrypt_and_verify(
        source_path=source,
        output_path=output,
        password="DGLAB",
        office_kind="powerpoint",
    )

    assert saves == [{"path": str(output), "password": "DGLAB"}]
    assert opens == [
        {
            "FileName": f"{source.resolve()}::DGLAB",
            "ReadOnly": False,
            "Untitled": False,
            "WithWindow": False,
        }
    ]
    assert verified == [output]
    assert runtime.initialized == runtime.uninitialized == 1


def test_macro_enabled_files_are_rejected_before_com_initialization(tmp_path: Path) -> None:
    runtime = _ComRuntime()
    source = tmp_path / "macro.docm"
    source.write_bytes(b"macro")
    gateway = OfficeFilePasswordGateway(
        dispatch=lambda _kind: pytest.fail("COM must not start for unsupported files"),
        com_runtime=runtime,
    )

    with pytest.raises(ValueError, match="Unsupported or mismatched"):
        gateway.encrypt_and_verify(
            source_path=source,
            output_path=tmp_path / "stage.docm",
            password="DGLAB",
            office_kind="word",
        )

    assert runtime.initialized == runtime.uninitialized == 0


def test_office_automation_error_never_exposes_document_password(tmp_path: Path) -> None:
    runtime = _ComRuntime()
    source = tmp_path / "slides.pptx"
    source.write_bytes(b"powerpoint")

    class _FailingPresentations:
        def Open(self, **kwargs):
            raise RuntimeError(f"Unable to open {kwargs['FileName']}")

    class _FailingApp:
        Presentations = _FailingPresentations()

        def Quit(self) -> None:
            pass

    gateway = OfficeFilePasswordGateway(
        dispatch=lambda _kind: _FailingApp(),
        com_runtime=runtime,
    )

    with pytest.raises(Exception) as error:
        gateway.encrypt_and_verify(
            source_path=source,
            output_path=tmp_path / "stage.pptx",
            password="DGLAB",
            office_kind="powerpoint",
        )

    assert "DGLAB" not in str(error.value)
    assert str(source.resolve()) not in str(error.value)
    assert error.value.__cause__ is None
