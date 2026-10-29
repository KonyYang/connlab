from datetime import date
from pathlib import Path

from docx import Document
import pytest
from openpyxl import Workbook

from backend.application.external_excel_read_service import (
    EquipmentCalibrationReadResult, EquipmentCalibrationRow,
    ExternalExcelReadService,
)
from backend.application.tools_equipment_report_service import ToolsEquipmentReportService
from backend.infrastructure.office.equipment_id_document_reader import EquipmentIdDocumentReader
from backend.domain import ExternalResource, ExternalResourceType
from backend.infrastructure.office.test_report_document_gateway import (
    TestReportDocumentGateway as ReportDocumentGateway,
)


class Catalog:
    def __init__(self, path):
        self.path = path

    def read_equipment_calibrations(self):
        return EquipmentCalibrationReadResult(str(self.path), ("All Equip.",), (
            EquipmentCalibrationRow("DG-Q-0033", "Multimeter", "Keysight",
                                    "2026-01-01", "2027-01-01", "All Equip."),
        ))


class OfficeWriter:
    def synchronize_equipment_list(self, *, source_path, output_path, rows):
        self.rows = rows
        output_path.write_bytes(b"updated-copy")
        return output_path


def service(tmp_path):
    catalog_path = tmp_path / "catalog.xlsx"
    catalog_path.write_bytes(b"stable catalog")
    writer = OfficeWriter()
    instance = ToolsEquipmentReportService(
        source_reader=EquipmentIdDocumentReader(), catalog_reader=Catalog(catalog_path),
        report_writer=writer, today=lambda: date(2026, 10, 3),
    )
    return instance, writer


def test_updates_uploaded_report_copy_from_equipment_document_without_project(tmp_path: Path):
    report = tmp_path / "Internal Report.docx"
    report.write_bytes(b"original report")
    equipment = tmp_path / "equipment.docx"
    document = Document()
    document.add_paragraph("Q-0033, DG-Q-0033, DG-Q-9999")
    document.save(equipment)
    before = equipment.read_bytes()
    instance, writer = service(tmp_path)

    result = instance.update(source_path=report, output_path=tmp_path / "updated.docx",
                             equipment_path=equipment)

    assert result.output_path.read_bytes() == b"updated-copy"
    assert report.read_bytes() == b"original report"
    assert equipment.read_bytes() == before
    assert len(writer.rows) == 2
    assert writer.rows[0].item == "Multimeter"
    assert writer.rows[0].last_calibration == "01-Jan-2026"
    assert writer.rows[1].item == ""
    assert result.statistics == {"filled": 2, "unmatched": ["DG-Q-9999"],
                                 "incomplete": [], "expired": []}


def test_pasted_ids_are_normalized_and_do_not_replace_original_report(tmp_path):
    instance, writer = service(tmp_path)
    report = tmp_path / "source.docx"
    report.write_bytes(b"source")
    result = instance.update(source_path=report, output_path=tmp_path / "copy.docx",
                             references_text="DG-Q-0033\nQ-0033")
    assert result.statistics["filled"] == 1
    assert writer.rows[0].id_number == "DG-Q-0033"
    assert report.read_bytes() == b"source"


def test_real_xlsx_calibration_dates_reach_new_word_report_copy(tmp_path: Path):
    workbook = tmp_path / "catalog.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "All Equip."
    sheet.append(["Catalog"])
    sheet.append([])
    sheet.append([])
    sheet.append(["Item (Equipment Name)", "", "Manufacturer", "ID Number", "Last Cal.", "Cal. Due"])
    sheet.append(["Load Frame", "", "Instron", "DG-Q-0134", 46020, 46384])
    sheet.append(["Meter", "", "Keysight", "DG-Q-0002", "29 Dec 2025", "19 Oct 2026"])
    sheet.append(["Fixture", "", "Lab", "DG-L-0001", "Not calibrated", "N/A"])
    sheet.append(["Ordinary numbers", "", "Lab", "DG-Q-0003", 46020, 46384])
    sheet.append(["Expired meter", "", "Lab", "DG-Q-0004", "2025-10-18", "2026-10-18"])
    sheet["E5"].number_format = "mm-dd-yy"
    sheet["F5"].number_format = "dd-mmm-yyyy"
    book.save(workbook)
    book.close()
    report = tmp_path / "Internal Report.docx"
    document = Document()
    document.add_paragraph("Keep reviewed test results")
    document.add_paragraph("7. EQUIPMENTS")
    table = document.add_table(rows=2, cols=5)
    for cell, heading in zip(table.rows[0].cells,
                             ("Item", "Manufacturer", "ID Number", "Last Cal.", "Cal. Due")):
        cell.text = heading
    document.save(report)
    selection = tmp_path / "selection.docx"
    document = Document()
    document.add_paragraph("Q-0134, Q-0002, L-0001, Q-0003, Q-0004")
    document.save(selection)
    originals = {path: path.read_bytes() for path in (workbook, report, selection)}
    resource = ExternalResource(
        resource_id="equipment", resource_type=ExternalResourceType.EQUIPMENT_CALIBRATION_EXCEL,
        path=workbook, active=True,
    )

    class ResourceStore:
        def get_by_type(self, resource_type):
            return resource

    instance = ToolsEquipmentReportService(
        source_reader=EquipmentIdDocumentReader(),
        catalog_reader=ExternalExcelReadService(ResourceStore()),
        report_writer=ReportDocumentGateway(), today=lambda: date(2026, 10, 19),
    )

    result = instance.update(source_path=report, output_path=tmp_path / "updated-copy.docx",
                             equipment_path=selection)

    updated = Document(result.output_path)
    assert [tuple(cell.text for cell in row.cells[2:]) for row in updated.tables[0].rows[1:]] == [
        ("DG-Q-0134", "29-Dec-2025", "28-Dec-2026"),
        ("DG-Q-0002", "29-Dec-2025", "19-Oct-2026"),
        ("DG-L-0001", "Not calibrated", "N/A"),
        ("DG-Q-0003", "", ""),
        ("DG-Q-0004", "18-Oct-2025", "18-Oct-2026"),
    ]
    assert result.statistics == {"filled": 5, "unmatched": [],
                                 "incomplete": ["Q-0003"], "expired": ["Q-0004"]}
    assert updated.paragraphs[0].text == "Keep reviewed test results"
    due = updated.tables[0].rows[5].cells[4]
    assert all(str(run.font.color.rgb) == "FF0000" for p in due.paragraphs for run in p.runs if run.text)
    assert all(path.read_bytes() == original for path, original in originals.items())


@pytest.mark.parametrize("selection", [{}, {"references_text": " "},
                                        {"equipment_path": Path("id.docx"), "references_text": "Q-0033"}])
def test_requires_exactly_one_equipment_source(tmp_path, selection):
    instance, _writer = service(tmp_path)
    report = tmp_path / "source.docx"
    report.write_bytes(b"source")
    output = tmp_path / "copy.docx"
    with pytest.raises(ValueError, match="Choose EquipmentID.docx"):
        instance.update(source_path=report, output_path=output, **selection)
    assert not output.exists()


def test_existing_output_is_never_replaced_or_deleted(tmp_path):
    instance, _writer = service(tmp_path)
    report = tmp_path / "source.docx"
    report.write_bytes(b"source")
    output = tmp_path / "copy.docx"
    output.write_bytes(b"keep")
    with pytest.raises(ValueError, match="new .docx copy"):
        instance.update(source_path=report, output_path=output, references_text="Q-0033")
    assert output.read_bytes() == b"keep"


def test_calibration_change_during_writing_discards_only_updated_copy(tmp_path):
    instance, writer = service(tmp_path)
    report = tmp_path / "source.docx"
    report.write_bytes(b"source")
    write = writer.synchronize_equipment_list
    def changing_writer(**kwargs):
        result = write(**kwargs)
        (tmp_path / "catalog.xlsx").write_bytes(b"changed")
        return result
    writer.synchronize_equipment_list = changing_writer
    output = tmp_path / "copy.docx"
    with pytest.raises(ValueError, match="changed"):
        instance.update(source_path=report, output_path=output, references_text="Q-0033")
    assert not output.exists()
    assert report.read_bytes() == b"source"


def test_catalog_configuration_switch_during_writing_stops_download(tmp_path):
    workbook = tmp_path / "catalog.xlsx"
    workbook.write_bytes(b"catalog")
    catalog = Catalog(workbook)
    writer = OfficeWriter()
    instance = ToolsEquipmentReportService(source_reader=EquipmentIdDocumentReader(),
        catalog_reader=catalog, report_writer=writer)
    report = tmp_path / "source.docx"
    report.write_bytes(b"source")
    replacement = tmp_path / "other.xlsx"
    replacement.write_bytes(b"other catalog")
    write = writer.synchronize_equipment_list
    def switching_writer(**kwargs):
        result = write(**kwargs)
        catalog.path = replacement
        return result
    writer.synchronize_equipment_list = switching_writer
    output = tmp_path / "copy.docx"
    with pytest.raises(ValueError, match="changed"):
        instance.update(source_path=report, output_path=output, references_text="Q-0033")
    assert not output.exists()
