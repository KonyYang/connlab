from __future__ import annotations

from dataclasses import replace
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from backend.infrastructure.files.ir_dwv_record_artifact_store import (
    IrDwvRecordArtifactStore,
)

from backend.infrastructure.office.ir_dwv_record_workbook_gateway import (
    IrDwvEquipment,
    IrDwvRecordGroup,
    IrDwvRecordProjection,
    IrDwvRecordStep,
    IrDwvRecordWorkbookGateway,
    copy_block,
)
from backend.infrastructure.office.ir_dwv_record_workbook_layout import (
    block_layout,
    plan_block_placements,
)
from scripts.inventory_ir_dwv_template import inventory_workbook


FIXTURE = Path(__file__).parents[1] / "fixtures" / "ir_dwv" / "IR_DWV_Template.xlsx"


@pytest.mark.parametrize("samples", [1, 3])
def test_small_sample_forms_keep_five_slots_but_only_number_actual_samples(tmp_path, samples):
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=tmp_path / "slots.xlsx", projection=_projection(samples=samples))
    workbook = load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        assert [sheet.cell(21, col).value for col in range(3, 8)] == [f"{index + 1}#" if index < samples else None for index in range(5)]
        assert [sheet.cell(21, col).value for col in range(8, 13)] == [f"{index + 1}#" if index < samples else None for index in range(5)]
        for start in (3, 8):
            assert all(sheet.cell(row, col).value is None for row in range(19, 34) for col in range(start + samples, start + 5))
        assert all(sheet.cell(row, col).value is None for row in range(14 + samples, 19) for col in (2, 8))
        assert sheet["H34"].value == ("=MIN(H22:H33)" if samples == 1 else "=MIN(H22:J33)")
        assert sheet["H19"].value == samples + 1
    finally:
        workbook.close()


@pytest.fixture
def template_with_separate_ir_statistics(tmp_path):
    workbook = load_workbook(FIXTURE)
    sheet = workbook.worksheets[0]
    sheet.unmerge_cells("C36:G37")
    sheet.merge_cells("C36:G36")
    sheet.merge_cells("C37:G37")
    for row, function in zip(range(34, 38), ("MIN", "MAX", "AVERAGE", "STDEV"), strict=True):
        sheet.cell(row, 3).value = f"={function}(C22:G33)"
    template = tmp_path / "separate-ir-statistics.xlsx"
    workbook.save(template)
    workbook.close()
    return template


@pytest.mark.parametrize("samples", [3, 5, 6, 7])
@pytest.mark.parametrize("pair_count", [1, 13])
def test_user_template_statistics_keep_both_sides_formulas_and_actual_sample_ranges(tmp_path, template_with_separate_ir_statistics, samples, pair_count):
    from hashlib import sha256
    template = template_with_separate_ir_statistics
    before = sha256(template.read_bytes()).hexdigest()
    writer = IrDwvRecordWorkbookGateway(template)
    writer.template_fingerprint()
    path = writer.write(output_path=tmp_path / "new-stats.xlsx", projection=_projection(samples=samples, steps=3, pairs=tuple(f"Pair {i}" for i in range(pair_count))))
    assert sha256(template.read_bytes()).hexdigest() == before
    workbook = load_workbook(path)
    try:
        for placement in plan_block_placements(sample_count=samples, pair_count=pair_count, step_count=3):
            sheet = workbook.worksheets[placement.sheet_index]
            layout = block_layout(origin_column=placement.origin_column, sample_count=samples, pair_count=pair_count)
            merges = {str(merged) for merged in sheet.merged_cells.ranges}
            for offset, function in enumerate(("MIN", "MAX", "AVERAGE", "STDEV")):
                row = layout.stats_min_row + offset
                for start, last in ((layout.ir_start_column, layout.ir_end_column), (layout.dwv_start_column, layout.dwv_end_column)):
                    actual_range = f"{get_column_letter(start)}{layout.data_start_row}:{get_column_letter(start + samples - 1)}{layout.data_end_row}"
                    assert sheet.cell(row, start).value == f"={function}({actual_range})"
                    assert f"{get_column_letter(start)}{row}:{get_column_letter(last)}{row}" in merges
            assert all(sheet.cell(row, col).value is None for row in range(layout.data_start_row, layout.data_end_row + 1) for col in range(layout.ir_start_column, layout.dwv_end_column + 1))
    finally:
        workbook.close()


@pytest.mark.parametrize("samples", [3, 7])
@pytest.mark.parametrize("kind", ["ir", "dwv"])
def test_separate_template_statistics_leave_only_the_unselected_side_blank(tmp_path, template_with_separate_ir_statistics, samples, kind):
    projection = _projection(samples=samples)
    step = replace(projection.groups[0].steps[0], ir_enabled=kind == "ir", dwv_enabled=kind == "dwv")
    projection = replace(projection, groups=(replace(projection.groups[0], steps=(step,)),))
    path = IrDwvRecordWorkbookGateway(template_with_separate_ir_statistics).write(output_path=tmp_path / "one-side-stats.xlsx", projection=projection)
    workbook = load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        layout = block_layout(origin_column=2, sample_count=samples, pair_count=1)
        active, inactive = (layout.ir_start_column, layout.dwv_start_column) if kind == "ir" else (layout.dwv_start_column, layout.ir_start_column)
        assert all(sheet.cell(row, inactive).value is None for row in range(layout.stats_min_row, layout.stats_stdev_row + 1))
        assert all(str(sheet.cell(row, active).value).startswith("=") for row in range(layout.stats_min_row, layout.stats_stdev_row + 1))
    finally:
        workbook.close()


@pytest.mark.parametrize("samples", [3, 5, 7])
def test_legacy_template_retains_ir_slashes_and_joined_average_stdev_area(tmp_path, samples):
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=tmp_path / "legacy-stats.xlsx", projection=_projection(samples=samples))
    workbook = load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        layout = block_layout(origin_column=2, sample_count=samples, pair_count=1)
        assert [sheet.cell(row, layout.ir_start_column).value for row in range(layout.stats_min_row, layout.stats_stdev_row + 1)] == ["/", "/", "/", None]
        assert f"C{layout.stats_average_row}:{get_column_letter(layout.ir_end_column)}{layout.stats_stdev_row}" in {str(merged) for merged in sheet.merged_cells.ranges}
    finally:
        workbook.close()


@pytest.mark.parametrize("samples", [3, 7])
@pytest.mark.parametrize("kind", ["both", "ir", "dwv"])
def test_record_writes_matrix_requirements_and_step_overrides_without_changing_units(tmp_path, samples, kind):
    from backend.application.matrix_editor_ir_dwv_record_projection import build_matrix_editor_ir_dwv_record_projection
    from backend.application.matrix_editor_llcr_cr_record_projection import MatrixEditorLlcrCrRecordGroupInput as Group, MatrixEditorLlcrCrRecordRowInput as Row
    from backend.application.matrix_step_text_output import MatrixStepTextOutputOverride
    from backend.domain.matrix_contact_measurement_models import MatrixPointProfile
    projection = build_matrix_editor_ir_dwv_record_projection(
        project_id="requirements", groups=(Group("g1", "1", str(samples)),),
        rows=(Row("IR", condition="500 VDC 2 min", requirement="1000 MΩ minimum", group_values={"g1": "2"}),
              Row("DWV", condition="1500 VDC 1 min", requirement="1 mA maximum; no breakdown", group_values={"g1": "3"})),
        point_profile=MatrixPointProfile((), electrical_point_pairs="Odd&Even"),
        step_text_overrides=(MatrixStepTextOutputOverride("g1", 1, 2, requirement="IR override: 2000 MΩ minimum"),),
    ).workbook
    step = replace(projection.groups[0].steps[0], ir_enabled=kind != "dwv", dwv_enabled=kind != "ir")
    projection = replace(projection, groups=(replace(projection.groups[0], steps=(step,)),))
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=tmp_path / "requirements.xlsx", projection=projection)
    workbook = load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        remarks = sheet["B11"].value
        assert ("IR Requirement: IR override: 2000 MΩ minimum" in remarks) == (kind != "dwv")
        assert ("DWV Requirement: 1 mA maximum; no breakdown" in remarks) == (kind != "ir")
        assert "1000 MΩ minimum" not in remarks
        layout = block_layout(origin_column=2, sample_count=samples, pair_count=1)
        assert sheet.cell(layout.units_row, layout.ir_start_column).value == ("GΩ" if kind != "dwv" else None)
        assert sheet.cell(layout.units_row, layout.dwv_start_column).value == ("nA" if kind != "ir" else None)
        assert sheet["B11"].alignment.wrap_text
        assert len(sheet._images) == 1
    finally:
        workbook.close()


def test_multiline_requirements_remain_complete_and_visible_after_a_shorter_round(tmp_path):
    from backend.application.matrix_editor_ir_dwv_record_projection import IrDwvRecordSourceStep
    base = _projection(steps=2)
    requirement = "\n".join(f"Acceptance criterion {index}" for index in range(1, 7))
    source = IrDwvRecordSourceStep("q", "r", 2, "", "2", "IR", "500 VDC", requirement)
    first = replace(base.groups[0].steps[0], ir_source_step=source)
    second = replace(base.groups[0].steps[1], ir_source_step=replace(source, step_sequence=6, source_step="6", requirement="Short final criterion"))
    projection = replace(base, groups=(replace(base.groups[0], steps=(first, second)),))
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=tmp_path / "long-requirements.xlsx", projection=projection)
    workbook = load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        assert requirement in sheet["B11"].value
        assert "Short final criterion" in sheet["O11"].value
        assert sheet.row_dimensions[11].height >= 7 * sheet["B11"].font.sz
        assert sheet["B11"].alignment.wrap_text and sheet["O11"].alignment.wrap_text
    finally:
        workbook.close()


@pytest.mark.parametrize("normal_font", ["Calibri", "宋体"])
def test_logo_font_metric_does_not_reduce_existing_remarks_height_budget(tmp_path, normal_font):
    from copy import copy
    from backend.application.matrix_editor_ir_dwv_record_projection import IrDwvRecordSourceStep
    template_book = load_workbook(FIXTURE)
    style = next(style for style in template_book._named_styles if style.builtinId == 0)
    font = copy(style.font)
    font.name, font.sz = normal_font, 11
    style.font = font
    template = tmp_path / "font-template.xlsx"
    template_book.save(template)
    template_book.close()
    base = _projection()
    # This complete IR label is 36 Latin characters. The established conservative
    # Remarks budget fits 35, hence two requirement lines plus the Remarks label.
    source = IrDwvRecordSourceStep("q", "r", 2, "", "2", "IR", "500 VDC", "A" * 20)
    step = replace(base.groups[0].steps[0], ir_source_step=source)
    path = IrDwvRecordWorkbookGateway(template).write(
        output_path=tmp_path / "font-output.xlsx",
        projection=replace(base, groups=(replace(base.groups[0], steps=(step,)),)),
    )
    workbook = load_workbook(path)
    try:
        assert workbook.active["B11"].value == "Remarks:\nIR Requirement: " + "A" * 20
        assert workbook.active.row_dimensions[11].height == pytest.approx(70.3)
    finally:
        workbook.close()


@pytest.mark.parametrize("requirement", ["X" * 32768, "\n".join(["Acceptance criterion"] * 30)], ids=["cell-text-overflow", "row-height-overflow"])
def test_requirements_fail_explicitly_instead_of_truncating_or_clipping(tmp_path, requirement):
    from backend.application.matrix_editor_ir_dwv_record_projection import IrDwvRecordSourceStep
    base = _projection()
    source = IrDwvRecordSourceStep("q", "r", 2, "", "2", "IR", "500 VDC", requirement)
    step = replace(base.groups[0].steps[0], ir_source_step=source)
    projection = replace(base, groups=(replace(base.groups[0], steps=(step,)),))
    output = tmp_path / "overflow.xlsx"
    with pytest.raises(ValueError, match="Matrix requirements.*Remarks"):
        IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=output, projection=projection)
    assert not output.exists()


def test_explicitly_empty_requirement_preserves_blank_remarks_template_layout(tmp_path):
    from copy import copy
    from backend.application.matrix_editor_ir_dwv_record_projection import IrDwvRecordSourceStep
    base = _projection()
    source = IrDwvRecordSourceStep("q", "r", 2, "", "2", "IR", "500 VDC", "")
    step = replace(base.groups[0].steps[0], ir_source_step=source)
    projection = replace(base, groups=(replace(base.groups[0], steps=(step,)),))
    output = IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=tmp_path / "empty-requirement.xlsx", projection=projection)
    workbook, source_book = load_workbook(output), load_workbook(FIXTURE)
    try:
        sheet, prototype = workbook.worksheets[0], source_book.worksheets[0]
        assert sheet["B11"].value == "Remarks:"
        assert copy(sheet["B11"].alignment) == copy(prototype["B11"].alignment)
        assert sheet.row_dimensions[11].height == prototype.row_dimensions[11].height
    finally:
        workbook.close()
        source_book.close()


@pytest.mark.parametrize("samples,pair_count", [(3, 1), (5, 13), (6, 1), (7, 1)])
def test_each_form_clones_first_template_logo_defaults_and_result_units(tmp_path, samples, pair_count):
    from hashlib import sha256
    source_book = load_workbook(FIXTURE)
    source = source_book.worksheets[0]
    source._images = source._images[:1]
    source["E3"], source["H3"] = "Template instrument default", "Template-gage"
    source["C20"], source["H20"] = "MΩ", "µA"
    template = tmp_path / "user-template.xlsx"
    source_book.save(template)
    source_book.close()
    before = sha256(template.read_bytes()).hexdigest()
    projection = replace(_projection(samples=samples, steps=4, pairs=tuple(f"Pair {i}" for i in range(pair_count))), equipment=())
    path = IrDwvRecordWorkbookGateway(template).write(output_path=tmp_path / "logos.xlsx", projection=projection)
    assert sha256(template.read_bytes()).hexdigest() == before
    workbook, source_book = load_workbook(path), load_workbook(template)
    try:
        prototype = source_book.worksheets[0]._images[0]
        prototype_bytes = prototype._data()
        assert sum(len(sheet._images) for sheet in workbook.worksheets) == 4
        for placement in plan_block_placements(sample_count=samples, pair_count=pair_count, step_count=4):
            sheet = workbook.worksheets[placement.sheet_index]
            layout = block_layout(origin_column=placement.origin_column, sample_count=samples, pair_count=pair_count)
            assert sheet.cell(3, layout.origin_column + 3).value == "Template instrument default"
            assert sheet.cell(3, layout.origin_column + 6).value == "Template-gage"
            assert sheet.cell(3, layout.origin_column + 7).value is None
            assert sheet.cell(3, layout.origin_column + 8).value is None
            assert sheet.cell(layout.units_row, layout.ir_start_column).value == "MΩ"
            assert sheet.cell(layout.units_row, layout.dwv_start_column).value == "µA"
            logo = sheet._images[placement.block_index]
            assert logo._data() == prototype_bytes
            shift = layout.origin_column - 2 + layout.width - 11
            for marker_name in ("_from", "to"):
                expected, actual = getattr(prototype.anchor, marker_name), getattr(logo.anchor, marker_name)
                assert (actual.row, actual.rowOff) == (expected.row, expected.rowOff)
                if samples <= 5 or marker_name == "to":
                    assert (actual.col, actual.colOff) == (expected.col + shift, expected.colOff)
    finally:
        workbook.close()
        source_book.close()


@pytest.mark.parametrize("samples", [3, 5, 6, 7])
@pytest.mark.parametrize("normal_font,pixel_widths,expected_width_emu", [
    ("Calibri", {9.0: 63, 16.46484375: 115, 17.86328125: 125, 14.46484375: 101}, 1556550),
    ("宋体", {9.0: 72, 16.46484375: 132, 17.86328125: 143, 14.46484375: 116}, 1728000),
], ids=["standard-7px-digit", "template-Song-11pt-8px-digit"])
def test_logo_preserves_displayed_rectangle_and_right_edge_on_every_form(tmp_path, samples, normal_font, pixel_widths, expected_width_emu):
    from copy import copy
    from hashlib import sha256
    source_book = load_workbook(FIXTURE)
    style = next(style for style in source_book._named_styles if style.builtinId == 0)
    font = copy(style.font)
    font.name, font.sz = normal_font, 11
    style.font = font
    source = source_book.worksheets[0]
    source._images = source._images[:1]
    anchor = source._images[0].anchor
    # The user's actual saved anchor. Independent OpenXML reference values
    # cover both the standard seven-pixel metric and this template's Song 11pt.
    anchor._from.col, anchor._from.colOff = 10, 202407
    anchor._from.row, anchor._from.rowOff = 0, 202407
    anchor.to.col, anchor.to.colOff = 11, 568332
    anchor.to.row, anchor.to.rowOff = 5, 99578
    template = tmp_path / "logo-geometry.xlsx"
    source_book.save(template)
    source_book.close()
    before = sha256(template.read_bytes()).hexdigest()
    path = IrDwvRecordWorkbookGateway(template).write(output_path=tmp_path / "logo-geometry-output.xlsx", projection=_projection(samples=samples, steps=3))
    assert sha256(template.read_bytes()).hexdigest() == before
    workbook, source_book = load_workbook(path), load_workbook(template)
    try:
        source = source_book.worksheets[0]
        assert workbook._named_styles[0].font == source_book._named_styles[0].font
        # K = 125/143 pixels at 7/8-pixel digit metrics, respectively.
        # Native PNG width and cached shape-transform extents are not the frame.
        expected_height_emu = sum(source.row_dimensions[row].height * 12700 for row in range(1, 6)) + 99578 - 202407
        for placement in plan_block_placements(sample_count=samples, pair_count=1, step_count=3):
            sheet = workbook.worksheets[placement.sheet_index]
            logo = sheet._images[placement.block_index]
            first, last = logo.anchor._from, logo.anchor.to
            width_emu = sum(pixel_widths[sheet.column_dimensions[get_column_letter(col + 1)].width] * 9525 for col in range(first.col, last.col)) + last.colOff - first.colOff
            height_emu = sum(sheet.row_dimensions[row].height * 12700 for row in range(first.row + 1, last.row + 1)) + last.rowOff - first.rowOff
            assert width_emu == expected_width_emu
            assert height_emu == expected_height_emu
            layout = block_layout(origin_column=placement.origin_column, sample_count=samples, pair_count=1)
            assert (last.col, last.colOff) == (layout.last_column - 1, 568332)
            if samples <= 5:
                assert (first.col, first.colOff) == (layout.last_column - 2, 202407)
            if samples == 7:
                assert [sheet.column_dimensions[get_column_letter(col)].width for col in range(layout.last_column - 2, layout.last_column + 1)] == [9.0, 9.0, 9.0]
    finally:
        workbook.close()
        source_book.close()


@pytest.mark.parametrize("samples", [1, 3, 7])
def test_dynamic_blank_context_does_not_leak_historical_headers(tmp_path, samples):
    projection = replace(_projection(samples=samples), request_number="NEW-LTR", product_name="New product",
                         equipment=(), tested_by="", checked_by="", approved_by="", requestor="",
                         start_date=None, finish_date=None, ambient_temperature="", relative_humidity="")
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=tmp_path / "blank.xlsx", projection=projection)
    workbook = load_workbook(path)
    try:
        values = [str(cell.value) for row in workbook.worksheets[0].iter_rows() for cell in row if cell.value is not None]
        for business_value in ("2024-12-30", "2025-12-29", "25.7", "48.5", "Peter Qiu", "Even Yang", "Gentle Zeng", "David Tao", "DL-2026-07-115", "Custom Pwr"):
            assert not any(business_value in value for value in values), (samples, business_value)
        assert values.count("FDQF-E-033") == 1
        assert values.count("Rev:") == 1
        assert values.count("Equipment Used") == 1
    finally:
        workbook.close()


@pytest.mark.parametrize("samples", [1, 3, 7])
def test_dynamic_headers_preserve_semantic_template_styles(tmp_path, samples):
    from copy import copy
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=tmp_path / "styles.xlsx", projection=_projection(samples=samples, steps=2))
    source_book, output_book = load_workbook(FIXTURE), load_workbook(path)
    try:
        source, sheet = source_book.worksheets[0], output_book.worksheets[0]
        fields = {"Instrument": "E2", "Gage ID": "H2", "Last Cal.": "I2", "Cal Due.": "J2",
                  "Start Date:2026/08/18": "B8", "Finish Date: 2026/08/18": "B9",
                  "Amb Temp: 25.7°C": "K8", "Rel. Hum.:48.5%RH": "K9",
                  "Request  No.:  DL-2026-07-115": "I10", "Product Name: Custom Pwr 14P VH": "I11", "Remarks:": "B11",
                  "Tested By: Even Yang": "B12", "Checked By:Peter Qiu": "E12", "Approved By:Gentle Zeng": "H12", "Requestor:David Tao": "K12"}
        for text, anchor in fields.items():
            targets = [cell for row in sheet.iter_rows(max_row=12) for cell in row if cell.value == text]
            assert len(targets) == 2
            for cell in targets:
                for attribute in ("font", "alignment", "fill", "border", "number_format", "protection"):
                    assert copy(getattr(cell, attribute)) == copy(getattr(source[anchor], attribute)), (samples, cell.coordinate, attribute)
        equipment_columns = {"DG-Q-0624": "H3", "2024/12/30": "I3", "2025/12/29": "J3"}
        for text, anchor in equipment_columns.items():
            for cell in (cell for row in sheet.iter_rows(min_row=3, max_row=7) for cell in row if cell.value == text):
                assert copy(cell.font) == copy(source[anchor].font)
                assert copy(cell.alignment) == copy(source[anchor].alignment)
        for text, anchor in {"Form No:": "K39", "FDQF-E-033": "L39", "Rev:": "K40", "G": "L40"}.items():
            targets = [cell for row in sheet.iter_rows(min_row=13) for cell in row if cell.value == text]
            assert len(targets) == 2, (samples, text, [cell.coordinate for cell in targets])
            for cell in targets:
                for attribute in ("font", "alignment", "fill", "border", "number_format", "protection"):
                    assert copy(getattr(cell, attribute)) == copy(getattr(source[anchor], attribute)), (samples, cell.coordinate, attribute)
    finally:
        source_book.close()
        output_book.close()


@pytest.mark.parametrize("samples", [1, 2])
@pytest.mark.parametrize("with_equipment", [True, False])
def test_small_sample_forms_preserve_all_headers_without_overlap(tmp_path, samples, with_equipment):
    projection = _projection(samples=samples)
    if not with_equipment:
        projection = replace(projection, equipment=())
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(output_path=tmp_path / "small.xlsx", projection=projection)
    workbook = load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        texts = [str(cell.value) for row in sheet.iter_rows(max_row=12) for cell in row if cell.value is not None]
        for label in ("Tested By: Even Yang", "Checked By:Peter Qiu", "Approved By:Gentle Zeng", "Requestor:David Tao", "Request  No.:  DL-2026-07-115", "Product Name: Custom Pwr 14P VH"):
            assert label in texts
        assert "Amb Temp: 25.7°C" in texts and "Start Date:2026/08/18" in texts
        if with_equipment:
            assert "DG-Q-0624" in texts and "2024/12/30" in texts and "2025/12/29" in texts
        layout = block_layout(origin_column=2, sample_count=samples, pair_count=1)
        assert all(sheet.cell(row, column).value is None
                   for row in range(layout.sequence_row, layout.data_end_row + 1)
                   for column in range(layout.dwv_end_column + 1, layout.last_column + 1))
    finally:
        workbook.close()


def test_template_cleaner_preserves_prototype_but_removes_historical_business_data(tmp_path):
    from copy import copy
    from scripts.inventory_ir_dwv_template import purify_template
    from hashlib import sha256
    before = sha256(FIXTURE.read_bytes()).hexdigest()
    cleaned = purify_template(FIXTURE, tmp_path / "clean.xlsx")
    assert sha256(FIXTURE.read_bytes()).hexdigest() == before
    source_book, clean_book = load_workbook(FIXTURE), load_workbook(cleaned)
    try:
        prototype, sheet = source_book.worksheets[0], clean_book.worksheets[0]
        assert len(clean_book.worksheets) == 1 and sheet.max_column <= 14
        for row in range(1, 41):
            for column in range(2, 13):
                for attribute in ("font", "fill", "border", "alignment", "number_format", "protection"):
                    assert copy(getattr(sheet.cell(row, column), attribute)) == copy(getattr(prototype.cell(row, column), attribute))
        assert sheet["H34"].value == prototype["H34"].value
        assert all(sheet.cell(row, col).value is None for row in range(22, 34) for col in range(2, 13))
        assert sheet["E3"].value == prototype["E3"].value
        assert sheet["H3"].value == prototype["H3"].value
        assert all(sheet.cell(row, col).value is None for row in range(3, 8) for col in (9, 10))
        assert len(sheet._images) == 1
        assert sheet._images[0]._data() == prototype._images[0]._data()
        assert sheet._images[0].anchor == prototype._images[0].anchor
        assert "Peter" not in str(sheet["E12"].value)
    finally:
        source_book.close(); clean_book.close()
    with pytest.raises(ValueError, match="different"):
        purify_template(FIXTURE, FIXTURE)
    with pytest.raises(FileExistsError):
        purify_template(FIXTURE, cleaned)
    IrDwvRecordWorkbookGateway(cleaned).write(output_path=tmp_path / "from-clean.xlsx", projection=_projection())


def _projection(
    *,
    samples: int = 5,
    steps: int = 1,
    pairs: tuple[str, ...] = ("Odd&Even",),
) -> IrDwvRecordProjection:
    return IrDwvRecordProjection(
        request_number="DL-2026-07-115",
        product_name="Custom Pwr 14P VH",
        tested_by="Even Yang",
        checked_by="Peter Qiu",
        approved_by="Gentle Zeng",
        requestor="David Tao",
        start_date="2026/08/18",
        finish_date="2026/08/18",
        ambient_temperature="25.7°C",
        relative_humidity="48.5%RH",
        equipment=(
            IrDwvEquipment(
                instrument="Eleactrical Safety Compliance Analyzer",
                gage_id="DG-Q-0624",
                last_calibration="2024/12/30",
                calibration_due="2025/12/29",
            ),
        ),
        groups=(
            IrDwvRecordGroup(
                label="Group 1",
                sample_count=samples,
                steps=tuple(
                    IrDwvRecordStep(
                        label=(
                            "Initial"
                            if index == 0
                            else "Final"
                            if index == steps - 1
                            else f"After Step {index}"
                        ),
                        ir_condition="500VDC 2min. mated",
                        dwv_condition="1500VDC 1min. Mated",
                        measurement_pairs=pairs,
                        expected_fee_quantity=samples * len(pairs),
                    )
                    for index in range(steps)
                ),
            ),
        ),
    )


def test_copy_block_preserves_formatting_merges_dimensions_and_translates_formulas():
    source_book = load_workbook(FIXTURE)
    source = source_book["Group 1"]
    target = source_book.create_sheet("Copied")

    copy_block(
        source,
        target,
        source_min_row=1,
        source_max_row=40,
        source_min_column=2,
        source_max_column=12,
        target_min_row=1,
        target_min_column=15,
    )

    assert target["O1"].value == source["B1"].value
    assert target["O1"]._style == source["B1"]._style
    assert target.row_dimensions[1].height == source.row_dimensions[1].height
    assert target.column_dimensions["O"].width == source.column_dimensions["B"].width
    assert "O1:Q2" in {str(item) for item in target.merged_cells.ranges}
    assert target["U34"].value == "=MIN(U22:Y33)"
    source_book.close()


def test_copy_block_preserves_grouped_widths_overrides_and_implicit_columns(tmp_path: Path):
    workbook = Workbook()
    source = workbook.active
    source.sheet_format.defaultColWidth = 9
    source.column_dimensions.group("B", "E", hidden=True)
    source.column_dimensions["B"].width = 18
    source.column_dimensions["C"].width = 7
    target = workbook.create_sheet("Copied")
    target.sheet_format.defaultColWidth = 9

    copy_block(
        source,
        target,
        source_min_row=1,
        source_max_row=1,
        source_min_column=2,
        source_max_column=6,
        target_min_row=1,
        target_min_column=8,
    )
    output = tmp_path / "column-dimensions.xlsx"
    workbook.save(output)
    workbook.close()

    reopened = load_workbook(output)
    try:
        copied = reopened["Copied"]
        assert [copied.column_dimensions[letter].width for letter in "HIJK"] == [
            18, 7, 18, 18,
        ]
        assert [copied.column_dimensions[letter].hidden for letter in "HIJK"] == [
            True, False, True, True,
        ]
        assert "L" not in copied.column_dimensions
        assert copied.sheet_format.defaultColWidth == 9
        assert set(reopened.active.column_dimensions) == {"B", "C"}
    finally:
        reopened.close()


@pytest.mark.parametrize("samples", [3, 5, 7])
def test_writer_preserves_template_effective_widths_for_resized_blocks(
    tmp_path: Path,
    samples: int,
):
    output = tmp_path / f"widths-{samples}.xlsx"
    IrDwvRecordWorkbookGateway(FIXTURE).write(
        output_path=output,
        projection=_projection(samples=samples, steps=2),
    )
    workbook = load_workbook(output)
    try:
        sheet = workbook["Group 1"]
        expected = {
            3: [16.46484375] * 9 + [17.86328125, 9],
            5: [16.46484375] * 9 + [17.86328125, 9],
            7: [16.46484375] * 11 + [17.86328125, 9, 9, 9],
        }[samples]
        placements = plan_block_placements(
            sample_count=samples, pair_count=1, step_count=2,
        )
        for placement in placements:
            assert [
                sheet.column_dimensions[get_column_letter(column)].width
                for column in range(
                    placement.origin_column,
                    placement.origin_column + len(expected),
                )
            ] == expected
        assert sheet.column_dimensions["A"].width == 14.46484375
        first_end = placements[0].origin_column + len(expected) - 1
        assert [
            sheet.column_dimensions[get_column_letter(column)].width
            for column in range(first_end + 1, placements[1].origin_column)
        ] == [14.46484375, 16.46484375]
        assert not any(
            dimension.min > placements[-1].origin_column + len(expected) - 1
            for dimension in sheet.column_dimensions.values()
        )
    finally:
        workbook.close()


@pytest.mark.parametrize("samples", [3, 5, 7])
def test_writer_resizes_samples_and_rebuilds_conditions_headers_and_statistics(
    tmp_path: Path,
    samples: int,
):
    target = tmp_path / f"record-{samples}.xlsx"
    IrDwvRecordWorkbookGateway(FIXTURE).write(
        output_path=target,
        projection=_projection(samples=samples),
    )

    workbook = load_workbook(target, data_only=False)
    sheet = workbook["Group 1"]
    layout = block_layout(
        origin_column=2,
        sample_count=samples,
        pair_count=1,
    )
    assert sheet.cell(layout.conditions_start_row, layout.origin_column).value == (
        "1. IR testing/ 1# 500VDC 2min. mated"
    )
    assert sheet.cell(layout.conditions_start_row, layout.dwv_start_column).value == (
        f"{samples + 1}. DWV testing/ 1# 1500VDC 1min. Mated"
    )
    assert [
        sheet.cell(layout.sample_id_row, column).value
        for column in range(layout.ir_start_column, layout.ir_end_column + 1)
    ] == [f"{index + 1}#" if index < samples else None for index in range(max(5, samples))]
    assert [
        sheet.cell(layout.sequence_row, column).value
        for column in range(layout.ir_start_column, layout.dwv_end_column + 1)
    ] == list(range(1, samples + 1)) + [None] * max(0, 5 - samples) + list(range(samples + 1, 2 * samples + 1)) + [None] * max(0, 5 - samples)
    assert sheet.cell(layout.data_start_row, layout.origin_column).value == "Odd&Even"
    assert all(
        sheet.cell(layout.data_start_row, column).value is None
        for column in range(layout.ir_start_column, layout.dwv_end_column + 1)
    )
    dwv_range = (
        f"{sheet.cell(layout.data_start_row, layout.dwv_start_column).coordinate}:"
        f"{sheet.cell(layout.data_end_row, layout.dwv_start_column + samples - 1).coordinate}"
    )
    assert sheet.cell(layout.stats_min_row, layout.dwv_start_column).value == f"=MIN({dwv_range})"
    assert sheet.cell(layout.stats_max_row, layout.dwv_start_column).value == f"=MAX({dwv_range})"
    assert sheet.cell(layout.stats_average_row, layout.dwv_start_column).value == (
        f"=AVERAGE({dwv_range})"
    )
    assert sheet.cell(layout.stats_stdev_row, layout.dwv_start_column).value == (
        f"=STDEV({dwv_range})"
    )
    workbook.close()


@pytest.mark.parametrize(
    ("step_count", "expected_sheets", "expected_origins"),
    [
        (2, ["Group 1"], [2, 15]),
        (3, ["Group 1"], [2, 15, 28]),
        (4, ["Group 1", "Group 1 (2)"], [2, 15, 28, 2]),
    ],
)
def test_writer_places_steps_across_sheets(
    tmp_path: Path,
    step_count: int,
    expected_sheets: list[str],
    expected_origins: list[int],
):
    placements = plan_block_placements(
        sample_count=5,
        pair_count=1,
        step_count=step_count,
    )
    assert [item.origin_column for item in placements] == expected_origins

    target = tmp_path / f"steps-{step_count}.xlsx"
    IrDwvRecordWorkbookGateway(FIXTURE).write(
        output_path=target,
        projection=_projection(steps=step_count),
    )
    workbook = load_workbook(target)
    assert workbook.sheetnames == expected_sheets
    workbook.close()


@pytest.mark.parametrize("samples,expected_origins", [
    (3, [2, 15, 28]), (5, [2, 15, 28]), (6, [2, 17, 32]), (7, [2, 19, 36]),
])
def test_three_rounds_stay_on_the_group_sheet_for_every_sample_width(tmp_path, samples, expected_origins):
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(
        output_path=tmp_path / "three-rounds.xlsx", projection=_projection(samples=samples, steps=3),
    )
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Group 1"]
        assert [cell.column for cell in workbook.active[10] if str(cell.value).startswith("Item/Process:")] == expected_origins
        assert len(workbook.active._images) == 3
    finally:
        workbook.close()


def test_three_round_pages_and_fourth_round_continuation_remain_independent_between_groups(tmp_path):
    base = _projection(samples=6, steps=4)
    first = replace(base.groups[0], label="Group 2")
    second = replace(_projection(samples=7, steps=3).groups[0], label="Group 6b")
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(
        output_path=tmp_path / "groups.xlsx", projection=replace(base, groups=(first, second)),
    )
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Group 2", "Group 2 (2)", "Group 6b"]
        for sheet, origins, labels, group in (
            (workbook["Group 2"], [2, 17, 32], ["Initial", "After Step 1", "After Step 2"], "Group 2"),
            (workbook["Group 2 (2)"], [2], ["Final"], "Group 2"),
            (workbook["Group 6b"], [2, 19, 36], ["Initial", "After Step 1", "Final"], "Group 6b"),
        ):
            headers = [cell for cell in sheet[10] if str(cell.value).startswith("Item/Process:")]
            assert [cell.column for cell in headers] == origins
            assert [cell.value for cell in headers] == [f"Item/Process: {group}  IR&DWV-{label}" for label in labels]
    finally:
        workbook.close()


@pytest.mark.parametrize("samples", [3, 5, 6, 7])
def test_every_adjacent_record_table_has_two_empty_columns_without_changing_capacity(tmp_path, samples):
    path = IrDwvRecordWorkbookGateway(FIXTURE).write(
        output_path=tmp_path / "gutters.xlsx", projection=_projection(samples=samples, steps=4),
    )
    workbook = load_workbook(path)
    try:
        first_origins = {3: [2, 15, 28], 5: [2, 15, 28], 6: [2, 17, 32], 7: [2, 19, 36]}[samples]
        expected_origins = (first_origins, [2])
        expected_widths = {
            3: [16.46484375] * 9 + [17.86328125, 9.0],
            5: [16.46484375] * 9 + [17.86328125, 9.0],
            6: [16.46484375] * 10 + [17.86328125, 9.0, 9.0],
            7: [16.46484375] * 11 + [17.86328125, 9.0, 9.0, 9.0],
        }[samples]
        footer_row = {3: 39, 5: 39, 6: 40, 7: 41}[samples]
        assert len(workbook.worksheets) == 2
        for sheet, expected in zip(workbook.worksheets, expected_origins, strict=True):
            origins = [cell.column for cell in sheet[10] if str(cell.value).startswith("Item/Process:")]
            ends = [cell.column for cell in sheet[footer_row] if cell.value == "FDQF-E-033"]
            for previous_end, following_origin in zip(ends[:-1], origins[1:], strict=True):
                assert following_origin - previous_end - 1 == 2
                gutters = range(previous_end + 1, following_origin)
                for row in range(1, sheet.max_row + 1):
                    for column in gutters:
                        cell = sheet.cell(row, column)
                        assert cell.value is None
                        assert all(not side or not side.style for side in (
                            cell.border.left, cell.border.right, cell.border.top, cell.border.bottom,
                        ))
                assert not any(
                    merged.min_col <= previous_end + 2 and merged.max_col >= previous_end + 1
                    for merged in sheet.merged_cells.ranges
                )
                assert [sheet.column_dimensions[get_column_letter(column)].width for column in gutters] == [14.46484375, 16.46484375]
            assert origins == expected
            assert sheet.max_column == ends[-1]
            # No fixed print area clips a moved form: the output uses its full
            # used range, which reaches AL for the third five-slot form.
            assert not sheet.print_area
            assert len(sheet._images) == len(origins)
            for origin, end, image in zip(origins, ends, sheet._images, strict=True):
                assert [sheet.column_dimensions[get_column_letter(column)].width for column in range(origin, end + 1)] == expected_widths
                assert image.anchor._from.col >= origin - 1
                assert image.anchor.to.col == end - 1
                # The template's verified Song-11pt metric makes width9 =72px.
                assert image.anchor.to.colOff <= 72 * 9525
    finally:
        workbook.close()


def test_writer_clears_all_template_measurements_when_pairs_are_empty(tmp_path: Path):
    target = tmp_path / "empty-pairs.xlsx"
    IrDwvRecordWorkbookGateway(FIXTURE).write(
        output_path=target,
        projection=_projection(pairs=()),
    )

    workbook = load_workbook(target, data_only=False)
    sheet = workbook["Group 1"]
    layout = block_layout(origin_column=2, sample_count=5, pair_count=0)
    for row in range(layout.data_start_row, layout.data_end_row + 1):
        assert sheet.cell(row, layout.origin_column).value is None
        for column in range(layout.ir_start_column, layout.dwv_end_column + 1):
            assert sheet.cell(row, column).value is None
    workbook.close()


def test_writer_blocks_a_template_with_a_changed_fingerprint(tmp_path: Path):
    changed_template = tmp_path / "changed.xlsx"
    workbook = load_workbook(FIXTURE)
    workbook["Group 1"]["B1"] = "Unexpected template"
    workbook.save(changed_template)
    workbook.close()

    with pytest.raises(ValueError, match="fingerprint"):
        IrDwvRecordWorkbookGateway(changed_template).write(
            output_path=tmp_path / "record.xlsx",
            projection=_projection(),
        )


@pytest.mark.parametrize("logo_state", ["missing", "unsupported_anchor"])
def test_template_validation_blocks_missing_or_unlocatable_first_block_logo(tmp_path, logo_state):
    from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
    template = tmp_path / "without-supported-logo.xlsx"
    workbook = load_workbook(FIXTURE)
    if logo_state == "missing":
        workbook.worksheets[0]._images = []
    else:
        workbook.worksheets[0]._images[0].anchor = AbsoluteAnchor()
    workbook.save(template)
    workbook.close()
    writer = IrDwvRecordWorkbookGateway(template)
    with pytest.raises(ValueError, match="first-block LOGO"):
        writer.template_fingerprint()
    output = tmp_path / "blocked.xlsx"
    with pytest.raises(ValueError, match="first-block LOGO"):
        writer.write(output_path=output, projection=_projection())
    assert not output.exists()


def test_writer_blocks_fee_quantity_mismatch(tmp_path: Path):
    projection = _projection()
    bad_step = replace(projection.groups[0].steps[0], expected_fee_quantity=99)
    bad_group = replace(projection.groups[0], steps=(bad_step,))

    with pytest.raises(ValueError, match="Fee quantity"):
        IrDwvRecordWorkbookGateway(FIXTURE).write(
            output_path=tmp_path / "record.xlsx",
            projection=replace(projection, groups=(bad_group,)),
        )


@pytest.mark.parametrize("kind", ["ir", "dwv"])
def test_writer_leaves_the_unselected_test_side_blank(tmp_path: Path, kind: str):
    base = _projection()
    step = replace(base.groups[0].steps[0], ir_enabled=kind == "ir", dwv_enabled=kind == "dwv")
    output = tmp_path / f"{kind}-only.xlsx"
    IrDwvRecordWorkbookGateway(FIXTURE).write(
        output_path=output,
        projection=replace(base, groups=(replace(base.groups[0], steps=(step,)),)),
    )
    workbook = load_workbook(output)
    try:
        sheet = workbook["Group 1"]
        inactive_start, inactive_end = (8, 12) if kind == "ir" else (3, 7)
        assert all(sheet.cell(row, column).value is None
                   for row in range(19, 22)
                   for column in range(inactive_start, inactive_end + 1))
        assert sheet["H14" if kind == "ir" else "B14"].value is None
        assert sheet["H34"].value is None if kind == "ir" else sheet["H34"].value == "=MIN(H22:L33)"
    finally:
        workbook.close()



def test_ir_dwv_artifact_store_uses_contained_uuid_paths(tmp_path: Path):
    store = IrDwvRecordArtifactStore(tmp_path)
    artifact = store.prepare_draft(project_id="project/one", record_type="ir_dwv")

    assert artifact.output_path.parent == tmp_path / "project_one"
    assert artifact.output_path.suffix == ".xlsx"
    artifact.output_path.touch()
    resolved = store.resolve(
        project_id="project/one",
        artifact_id=artifact.artifact_id,
    )
    assert resolved.output_path == artifact.output_path.resolve()

    with pytest.raises(ValueError, match="Record type"):
        store.prepare_draft(project_id="project/one", record_type="llcr")

def test_template_inventory_matches_the_frozen_fixture():
    expected_path = FIXTURE.with_name("template_inventory.json")
    expected = json.loads(expected_path.read_text(encoding="utf-8"))

    assert inventory_workbook(FIXTURE) == expected

@pytest.mark.parametrize(
    ("samples", "step_count", "expected_sheet_indexes"),
    [
        (3, 5, [0, 0, 0, 1, 1]),
        (6, 4, [0, 0, 0, 1]),
        (7, 4, [0, 0, 0, 1]),
    ],
)
def test_round_capacity_is_independent_of_block_width(
    samples: int,
    step_count: int,
    expected_sheet_indexes: list[int],
):
    placements = plan_block_placements(
        sample_count=samples,
        pair_count=1,
        step_count=step_count,
    )

    assert [item.sheet_index for item in placements] == expected_sheet_indexes


@pytest.mark.parametrize("samples", [0, -1, 3.5, True])
def test_writer_blocks_non_positive_or_non_integer_sample_counts(
    tmp_path: Path,
    samples,
):
    with pytest.raises(ValueError, match="positive integer"):
        IrDwvRecordWorkbookGateway(FIXTURE).write(
            output_path=tmp_path / "record.xlsx",
            projection=_projection(samples=samples),
        )


def test_writer_keeps_seventeen_sample_capacity_and_rejects_eighteen(tmp_path):
    gateway = IrDwvRecordWorkbookGateway(FIXTURE)
    path = gateway.write(output_path=tmp_path / "seventeen.xlsx", projection=_projection(samples=17, steps=3))
    workbook = load_workbook(path)
    try:
        assert len(workbook.worksheets) == 1
        assert [cell.column for cell in workbook.active[10] if str(cell.value).startswith("Item/Process:")] == [2, 39, 76]
        assert workbook.active.max_column == 110
    finally:
        workbook.close()
    rejected = tmp_path / "eighteen.xlsx"
    with pytest.raises(ValueError, match="cannot fit"):
        gateway.write(output_path=rejected, projection=_projection(samples=18))
    assert not rejected.exists()


def test_writer_blocks_when_registered_template_has_an_excel_lock(tmp_path: Path):
    template = tmp_path / "IR&DWV Template.xlsx"
    template.write_bytes(FIXTURE.read_bytes())
    template.with_name("~$IR&DWV Template.xlsx").write_bytes(b"locked")

    with pytest.raises(
        ValueError,
        match="请关闭 Excel 中的 IR&DWV Template.xlsx",
    ):
        IrDwvRecordWorkbookGateway(template).write(
            output_path=tmp_path / "record.xlsx",
            projection=_projection(),
        )
