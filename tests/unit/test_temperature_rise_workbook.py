from pathlib import Path
import struct
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from backend.application.tools_temperature_rise_service import ToolsTemperatureRiseService
from backend.infrastructure.office.temperature_rise_workbook_gateway import TemperatureRiseWorkbookGateway


def source_book(path: Path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Measurements"
    sheet.append(["Instrument setup"])
    for _ in range(29):
        sheet.append([])
    sheet.append(["Scan", "Time", "A HS (C)", "A C (C)", "B only (C)", "Ambient (C)", "Current (VDC)"])
    for index, current in enumerate([10, 10, 20, 20, 30, 30], 1):
        rise = current ** 2 * .01 + current * .1
        sheet.append([index, "2026/8/19 17:17:12:352", 20 + rise, 20 + rise - 1, 20 + rise - 2, 20, current])
    book.save(path)
    book.close()


def test_preview_finds_content_block_and_exports_only_confirmed_rows(tmp_path):
    source = tmp_path / "source.xlsx"
    source_book(source)
    original = source.read_bytes()
    service = ToolsTemperatureRiseService(TemperatureRiseWorkbookGateway())
    preview = service.preview(source)
    block = preview["blocks"][0]
    assert block["sheet"] == "Measurements"
    assert block["header_row"] == 31
    assert block["row_count"] == 6
    options = {"block_id": block["id"], "mapping": {"ambient": 5, "current": 6,
        "channels": [{"column": 2, "sample": "A", "point": "HS"},
                     {"column": 3, "sample": "A", "point": "C"},
                     {"column": 4, "sample": "B", "point": "only"}]},
        "selected_rows": [33, 35, 37], "current_mode": "amperes", "current_gain": 6666.67,
        "zero_intercept": True, "include_origin": True, "target_rise": 30,
        "max_temperature": 125, "ambient_temperatures": [20, 125], "derating_factor": .8}
    analysis = service.analyze(source, options)
    assert analysis["candidate_rows"] == [33, 35, 37]
    assert analysis["points"][0]["current"] == 10  # Already scaled: no gain applied.
    assert analysis["points"][0]["maximum"] == pytest.approx(2)
    assert analysis["points"][0]["average_of_max"] == pytest.approx(1)
    output = tmp_path / "new.xlsx"
    service.export(source, output, options)
    assert source.read_bytes() == original
    book = load_workbook(output)
    assert book.sheetnames == ["Setup", "T-riseChart", "Derating", "Initial Data"]
    assert len(book["T-riseChart"]._charts) == 1
    assert len(book["Derating"]._charts) == 1
    assert book["T-riseChart"]["A2"].value == 2
    assert book["Initial Data"]["G2"].value == 10
    assert "Non-measured" in str(book["Setup"]["B8"].value)
    assert book["T-riseChart"]["B13"].value == "2026/8/19 17:17:12:352"
    assert book["T-riseChart"]["C13"].value == pytest.approx(2)
    coefficient_row = next(row[0].row for row in book["Setup"] if row[0].value == "Max")
    assert book["Setup"].cell(coefficient_row, 2).value == pytest.approx(analysis["max_curve"]["a"])
    book.close()


def test_plain_csv_voltage_conversion_and_invalid_mapping(tmp_path):
    source = tmp_path / "source.csv"
    source.write_text("Scan,Time,Sample HS,Ambient,Current\n1,12:00,22,20,1\n2,12:01,26,20,2\n3,12:02,32,20,3\n", encoding="utf-8")
    service = ToolsTemperatureRiseService(TemperatureRiseWorkbookGateway())
    block = service.preview(source)["blocks"][0]
    options = {"block_id": block["id"], "mapping": {"ambient": 3, "current": 4,
        "channels": [{"column": 2, "sample": "Sample", "point": "HS"}]},
        "selected_rows": [2, 3, 4], "current_mode": "voltage", "current_gain": 10,
        "ambient_temperatures": [20, 125]}
    result = service.analyze(source, options)
    assert result["points"][0]["current"] == 10
    options["mapping"]["channels"][0]["column"] = 3
    with pytest.raises(ValueError, match="distinct"):
        service.analyze(source, options)


def test_corrupt_middle_measurement_is_not_silently_truncated(tmp_path):
    source = tmp_path / "broken.csv"
    source.write_text("Scan,Time,1_HS,Ambient,Current\n1,12:00,22,20,10\n2,12:01,26,,20\n3,12:02,32,20,30\n")
    with pytest.raises(ValueError, match="row 3"):
        ToolsTemperatureRiseService(TemperatureRiseWorkbookGateway()).preview(source)


def test_real_biff_xls_is_read_offline_without_office_or_new_dependencies(tmp_path):
    # A valid Excel BIFF2 worksheet stream, parsed by the real installed xlrd reader.
    def record(code, data):
        return struct.pack("<HH", code, len(data)) + data
    content = record(0x0009, struct.pack("<HH", 0, 0x10))
    rows = [[320, "Current", "VDC", 6666.67, "True"],
            ["Scan", "Time", "1_HS", "Ambient", "320 <Current> (VDC)"],
            [1, "12:00", 22, 20, 10], [2, "12:01", 26, 20, 20], [3, "12:02", 32, 20, 30]]
    for row, values in enumerate(rows):
        for column, value in enumerate(values):
            prefix = struct.pack("<HH", row, column) + b"\0\0\0"
            content += record(4, prefix + bytes([len(value)]) + value.encode("ascii")) if isinstance(value, str) else record(3, prefix + struct.pack("<d", value))
    content += record(0xA, b"")
    source = tmp_path / "readings.xls"
    source.write_bytes(content)
    service = ToolsTemperatureRiseService(TemperatureRiseWorkbookGateway())
    block = service.preview(source)["blocks"][0]
    assert block["row_count"] == 3
    assert block["current_metadata"][0] == 320.0
    result = service.analyze(source, {"block_id": block["id"], "mapping": block["suggested_mapping"], "selected_rows": block["candidate_rows"]})
    assert result["max_curve"]["a"] == pytest.approx(.01)
    assert source.read_bytes() == content


def test_excel_date_arbitrary_sheet_order_and_explicit_derating_contract(tmp_path):
    source = tmp_path / "dates.xlsm"
    book = Workbook()
    book.active.title = "Notes"
    book.active.append(["No measurement data here"])
    sheet = book.create_sheet("Unspecified measurement name")
    sheet.append(["Current", "Ambient", "1_HS", "Scan", "Time"])
    for scan, (current, temperature) in enumerate([(10, 22), (20, 26), (30, 32)], 1):
        sheet.append([current, 20, temperature, scan, datetime(2026, 8, 19, 17, 17, 12, 352000)])
    book.save(source)
    book.close()
    service = ToolsTemperatureRiseService(TemperatureRiseWorkbookGateway())
    block = service.preview(source)["blocks"][0]
    assert block["sheet"] == "Unspecified measurement name"
    assert block["records"][0]["values"][4] == "2026-08-19 17:17:12.352000"
    result = service.analyze(source, {"block_id": block["id"], "mapping": block["suggested_mapping"], "selected_rows": block["candidate_rows"], "ambient_temperatures": [95, 125]})
    assert result["target_current"] == pytest.approx(50)
    assert result["derating"][0] == pytest.approx({"ambient": 95, "allowable_rise": 30, "basic_current": 50, "derated_current": 40})
    assert result["derating"][1]["derated_current"] == pytest.approx(0)
    with pytest.raises(ValueError, match="source"):
        service.export(source, source, {"block_id": block["id"], "mapping": block["suggested_mapping"], "selected_rows": block["candidate_rows"]})


@pytest.mark.parametrize("change,message", [
    ({"selected_rows": [2, 2]}, "distinct source rows"),
    ({"selected_rows": [2]}, "distinct measured current"),
    ({"derating_factor": 1.1}, "Derating factor"),
    ({"current_mode": "voltage", "current_gain": 0}, "gain"),
    ({"ambient_temperatures": [20, 126]}, "non-negative"),
])
def test_invalid_selections_and_parameters_do_not_generate_workbooks(tmp_path, change, message):
    source = tmp_path / "input.csv"
    source.write_text("Scan,Time,1_HS,Ambient,Current\n1,12:00,22,20,10\n2,12:01,26,20,20\n3,12:02,32,20,30\n")
    service = ToolsTemperatureRiseService(TemperatureRiseWorkbookGateway())
    block = service.preview(source)["blocks"][0]
    options = {"block_id": block["id"], "mapping": block["suggested_mapping"], "selected_rows": block["candidate_rows"], **change}
    output = tmp_path / "new.xlsx"
    with pytest.raises(ValueError, match=message):
        service.export(source, output, options)
    assert not output.exists()


def test_derating_zero_allowable_rise_exports_zero_current_with_negative_linear_coefficient(tmp_path):
    source = tmp_path / "negative-b.csv"
    source.write_text("Scan,Time,1_HS,Ambient,Current\n1,12:00,21,20,10\n2,12:01,28,20,20\n3,12:02,41,20,30\n")
    service = ToolsTemperatureRiseService(TemperatureRiseWorkbookGateway())
    block = service.preview(source)["blocks"][0]
    options = {"block_id": block["id"], "mapping": block["suggested_mapping"], "selected_rows": block["candidate_rows"], "ambient_temperatures": [20, 125]}
    result = service.analyze(source, options)
    assert result["derating"][-1]["basic_current"] == 0
    assert result["derating"][-1]["derated_current"] == 0
    output = tmp_path / "new.xlsx"
    service.export(source, output, options)
    book = load_workbook(output)
    assert book["Derating"]["C3"].value == 0
    assert book["Derating"]["D3"].value == 0
    book.close()


def test_blank_preamble_rows_before_numbered_current_channel_do_not_break_preview(tmp_path):
    source = tmp_path / "blank-metadata.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Instrument setup"])
    sheet.append([])
    sheet.append([])
    sheet.append(["Scan", "Time", "1_HS", "Ambient", "320 <Current> (VDC)"])
    for scan, current, temperature in [(1, 10, 22), (2, 20, 26), (3, 30, 32)]:
        sheet.append([scan, "12:00", temperature, 20, current])
    book.save(source)
    book.close()
    block = ToolsTemperatureRiseService(TemperatureRiseWorkbookGateway()).preview(source)["blocks"][0]
    assert block["header_row"] == 4
    assert block["row_count"] == 3
    assert block["current_metadata"] == []
