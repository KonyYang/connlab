import os
from hashlib import sha256

from docx import Document
import pytest

from backend.application.customer_report_source_resolver import DownloadedInternalReportUnavailable
from backend.infrastructure.files.downloaded_internal_report_source import DownloadedInternalReportDrafts


DL = "DL-2026-09-002"


def draft(root, suffix="", *, dl=DL, header_dl=None, form="E-3707", modified=100):
    path = root / f"{dl} Connector Report_Rev_A_Draft{suffix}.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "LABORATORY TEST REPORT " + (header_dl or dl)
    document.sections[0].footer.paragraphs[0].text = f"Form {form} Rev. H"
    document.add_paragraph("DRAFT test content")
    document.save(path)
    os.utime(path, (modified, modified))
    return path


def test_selects_latest_matching_internal_draft_not_another_project_or_customer(tmp_path):
    draft(tmp_path)
    latest = draft(tmp_path, " (1)", modified=200)
    draft(tmp_path, dl=DL + "A", modified=300)
    draft(tmp_path, dl=DL + "-CR", form="E-4515", modified=400)
    source = DownloadedInternalReportDrafts(tmp_path).latest(DL)
    assert source.file_path == latest
    assert source.mode == "managed_draft"
    assert source.file_name == latest.name


@pytest.mark.parametrize("problem", ["wrong_project", "customer", "similar_form", "corrupt"])
def test_rejects_invalid_latest_instead_of_silently_using_older_draft(tmp_path, problem):
    draft(tmp_path)
    latest = draft(tmp_path, " (1)", header_dl="DL-2026-09-999" if problem == "wrong_project" else DL,
                   form={"customer": "E-4515", "similar_form": "E-37070"}.get(problem, "E-3707"), modified=200)
    if problem == "corrupt":
        latest.write_bytes(b"unfinished download")
    original = latest.read_bytes()
    with pytest.raises(DownloadedInternalReportUnavailable):
        DownloadedInternalReportDrafts(tmp_path).latest(DL)
    assert latest.read_bytes() == original


def test_ignores_non_drafts_temporary_files_and_nested_downloads(tmp_path):
    valid = draft(tmp_path)
    for name in [valid.name.replace("_Draft", ""), "~$" + valid.name, valid.name + ".crdownload"]:
        (tmp_path / name).write_bytes(b"not a source")
    nested = tmp_path / "nested"
    nested.mkdir()
    draft(nested, modified=300)
    original = valid.read_bytes()
    source = DownloadedInternalReportDrafts(tmp_path).latest(DL)
    assert source.file_path == valid
    assert source.file_sha256 == sha256(original).hexdigest()
    assert valid.read_bytes() == original


def test_browser_copy_number_breaks_equal_timestamp_tie(tmp_path):
    draft(tmp_path)
    latest = draft(tmp_path, " (2)")
    draft(tmp_path, " (1)")
    assert DownloadedInternalReportDrafts(tmp_path).latest(DL).file_path == latest


def test_different_latest_sources_without_ordering_are_not_guessed(tmp_path):
    first = draft(tmp_path)
    other = tmp_path / first.name.replace("Connector", "Other Connector")
    other.write_bytes(first.read_bytes() + b"different")
    os.utime(other, (100, 100))
    with pytest.raises(DownloadedInternalReportUnavailable, match="Several latest"):
        DownloadedInternalReportDrafts(tmp_path).latest(DL)


def test_missing_downloads_or_matching_file_returns_none(tmp_path):
    assert DownloadedInternalReportDrafts(tmp_path / "absent").latest(DL) is None
    draft(tmp_path, dl="DL-2026-09-003")
    assert DownloadedInternalReportDrafts(tmp_path).latest(DL) is None


def test_reads_protected_internal_report_without_changing_it(tmp_path):
    import msoffcrypto
    from backend.shared.office_document_password import OFFICE_DOCUMENT_PASSWORD
    plain = draft(tmp_path)
    protected = tmp_path / plain.name.replace("Connector", "Protected Connector")
    with plain.open("rb") as source, protected.open("wb") as target:
        msoffcrypto.OfficeFile(source).encrypt(OFFICE_DOCUMENT_PASSWORD, target)
    original = protected.read_bytes()
    selected = DownloadedInternalReportDrafts(tmp_path).latest(DL)
    assert selected.file_path == protected
    assert protected.read_bytes() == original
