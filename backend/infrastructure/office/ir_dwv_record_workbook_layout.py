"""Central layout calculator for template-driven IR/DWV record workbooks."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter

TEMPLATE_SAMPLE_COUNT = 5
TEMPLATE_PAIR_ROW_COUNT = 12
TEMPLATE_FIRST_COLUMN = 2
TEMPLATE_LAST_COLUMN = 37
TEMPLATE_FIRST_BLOCK_ORIGINS = (2, 15, 26)


@dataclass(frozen=True, slots=True)
class IrDwvBlockLayout:
    origin_column: int
    sample_count: int
    pair_count: int
    ir_start_column: int
    ir_end_column: int
    dwv_start_column: int
    dwv_end_column: int
    last_column: int
    conditions_start_row: int
    conditions_end_row: int
    sequence_row: int
    units_row: int
    sample_id_row: int
    data_start_row: int
    data_end_row: int
    stats_min_row: int
    stats_max_row: int
    stats_average_row: int
    stats_stdev_row: int
    spacer_row: int
    footer_form_row: int
    footer_revision_row: int

    @property
    def width(self) -> int:
        return self.last_column - self.origin_column + 1

    @property
    def height(self) -> int:
        return self.footer_revision_row


@dataclass(frozen=True, slots=True)
class IrDwvBlockPlacement:
    sheet_index: int
    block_index: int
    origin_column: int


def block_layout(
    *,
    origin_column: int,
    sample_count: int,
    pair_count: int,
) -> IrDwvBlockLayout:
    if sample_count < 1:
        raise ValueError("IR/DWV sample count must be a positive integer.")
    if pair_count < 0:
        raise ValueError("IR/DWV measurement pair count cannot be negative.")
    ir_start = origin_column + 1
    ir_end = ir_start + sample_count - 1
    dwv_start = ir_end + 1
    dwv_end = dwv_start + sample_count - 1
    conditions_start = 14
    conditions_end = conditions_start + sample_count - 1
    sequence_row = conditions_end + 1
    units_row = sequence_row + 1
    sample_id_row = units_row + 1
    data_start = sample_id_row + 1
    data_end = data_start + max(TEMPLATE_PAIR_ROW_COUNT, pair_count) - 1
    stats_min = data_end + 1
    return IrDwvBlockLayout(
        origin_column=origin_column,
        sample_count=sample_count,
        pair_count=pair_count,
        ir_start_column=ir_start,
        ir_end_column=ir_end,
        dwv_start_column=dwv_start,
        dwv_end_column=dwv_end,
        last_column=dwv_end,
        conditions_start_row=conditions_start,
        conditions_end_row=conditions_end,
        sequence_row=sequence_row,
        units_row=units_row,
        sample_id_row=sample_id_row,
        data_start_row=data_start,
        data_end_row=data_end,
        stats_min_row=stats_min,
        stats_max_row=stats_min + 1,
        stats_average_row=stats_min + 2,
        stats_stdev_row=stats_min + 3,
        spacer_row=stats_min + 4,
        footer_form_row=stats_min + 5,
        footer_revision_row=stats_min + 6,
    )


def block_origins(*, sample_count: int, pair_count: int) -> tuple[int, ...]:
    layout = block_layout(
        origin_column=TEMPLATE_FIRST_COLUMN,
        sample_count=sample_count,
        pair_count=pair_count,
    )
    if sample_count == TEMPLATE_SAMPLE_COUNT:
        return TEMPLATE_FIRST_BLOCK_ORIGINS
    origins: list[int] = []
    origin = TEMPLATE_FIRST_COLUMN
    while origin + layout.width - 1 <= TEMPLATE_LAST_COLUMN:
        origins.append(origin)
        origin += layout.width + 2
    if not origins:
        raise ValueError("IR/DWV block cannot fit within the template sheet.")
    return tuple(origins)


def plan_block_placements(
    *,
    sample_count: int,
    pair_count: int,
    step_count: int,
) -> tuple[IrDwvBlockPlacement, ...]:
    if step_count < 0:
        raise ValueError("IR/DWV step count cannot be negative.")
    origins = block_origins(sample_count=sample_count, pair_count=pair_count)
    return tuple(
        IrDwvBlockPlacement(
            sheet_index=index // len(origins),
            block_index=index % len(origins),
            origin_column=origins[index % len(origins)],
        )
        for index in range(step_count)
    )


def dynamic_merge_ranges(layout: IrDwvBlockLayout) -> tuple[str, ...]:
    ranges: list[tuple[int, int, int, int]] = []
    origin = layout.origin_column
    last = layout.last_column

    def add(min_row: int, min_col: int, max_row: int, max_col: int) -> None:
        if min_col <= max_col and min_row <= max_row and (min_col != max_col or min_row != max_row):
            ranges.append((min_row, min_col, max_row, max_col))

    add(1, origin, 2, min(origin + 2, last))
    add(1, min(origin + 3, last), 1, min(origin + 8, last))
    if layout.sample_count >= TEMPLATE_SAMPLE_COUNT:
        add(2, min(origin + 3, last), 2, min(origin + 5, last))
    add(3, origin, 3, min(origin + 2, last))
    if layout.sample_count >= TEMPLATE_SAMPLE_COUNT:
        add(3, min(origin + 3, last), 3, min(origin + 5, last))
    for row in range(4, 8):
        add(row, origin, row, min(origin + 2, last))
    for row in (8, 9):
        add(row, origin, row, min(origin + 2, last))
        add(row, max(origin, last - 1), row, last)

    first_header_end = min(last, layout.ir_end_column + (1 if layout.sample_count > 1 else 0))
    for row in (10, 11):
        add(row, origin, row, first_header_end)
        add(row, first_header_end + 1, row, last)

    width = layout.width
    base, remainder = divmod(width, 4)
    sizes = [base + (1 if index < remainder else 0) for index in range(4)]
    cursor = origin
    for size in sizes:
        if size > 1:
            add(12, cursor, 12, cursor + size - 1)
        cursor += size

    for row in range(layout.conditions_start_row, layout.conditions_end_row + 1):
        add(row, origin, row, layout.ir_end_column)
        add(row, layout.dwv_start_column, row, last)
    add(13, origin, 13, layout.ir_end_column)
    add(13, layout.dwv_start_column, 13, last)

    add(layout.stats_min_row, layout.ir_start_column, layout.stats_min_row, layout.ir_end_column)
    add(layout.stats_min_row, layout.dwv_start_column, layout.stats_min_row, last)
    add(layout.stats_max_row, layout.ir_start_column, layout.stats_max_row, layout.ir_end_column)
    add(layout.stats_max_row, layout.dwv_start_column, layout.stats_max_row, last)
    add(
        layout.stats_average_row,
        layout.ir_start_column,
        layout.stats_stdev_row,
        layout.ir_end_column,
    )
    add(layout.stats_average_row, layout.dwv_start_column, layout.stats_average_row, last)
    add(layout.stats_stdev_row, layout.dwv_start_column, layout.stats_stdev_row, last)

    return tuple(
        f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        for min_row, min_col, max_row, max_col in ranges
    )
