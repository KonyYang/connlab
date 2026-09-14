"""Microsoft Word adapter for deriving E-4515 customer-report drafts."""

from __future__ import annotations

from dataclasses import dataclass
import gc
from hashlib import sha256
import logging
import os
from pathlib import Path
import re
import shutil
import time
from typing import Callable, Iterable, TypeVar
from uuid import uuid4
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from backend.infrastructure.office.office_lifecycle import OfficeAutomationUnavailable
from backend.infrastructure.office.office_protected_document_gateway import (
    ProtectedWordPackageGateway,
)
from backend.shared.office_document_password import OFFICE_DOCUMENT_PASSWORD


_INTERNAL_REPORT_LABEL = "LABORATORY TEST REPORT"
_CUSTOMER_REPORT_LABEL = "CUSTOMER TEST REPORT"
_REVISION_NOTE = "Note: Each new revision replaces/supersedes all previous revisions."
_CONTINUATION_REPORT_PREFIX = "Report No. "
_END_MARKER = "*** End of Report ***"
_ACCEPTANCE_TEXT = (
    "Unless otherwise specified, assessment of conformity to requirements is based on "
    "simple acceptance."
)
_SAMPLE_SCOPE_TEXT = "The results of testing only apply to the sample"
_KEEP_WITH_NEXT_CUSTOMER_HEADINGS = {
    "TEST DESCRIPTION",
    "TEST METHODS/REQUIREMENTS",
    "TEST RESULTS",
}
_CUSTOMER_HEADINGS = (
    "PURPOSE",
    "CONCLUSIONS",
    "SAMPLE DESCRIPTION",
    "TEST DESCRIPTION",
    "TEST METHODS/REQUIREMENTS",
    "TEST RESULTS",
    "REVISION RECORD",
)
_SOURCE_SHA256_PROPERTY = "ConnLabSourceReportSHA256"
_CUSTOM_PROPERTY_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
)
_VALUE_TYPE_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
)
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_TYPE_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_CUSTOM_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "custom-properties"
)
_CUSTOM_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.custom-properties+xml"
)
_INTERNAL_DISCLOSURE_PATTERNS = (
    re.compile(r"^\s*[#*]\s*[.:]", flags=re.IGNORECASE),
    re.compile(r"performed\s+and\s+results\s+reported\s+under", flags=re.IGNORECASE),
    re.compile(r"\bgiven\s+in\s+Appendix\s+[A-Z]\b", flags=re.IGNORECASE),
    re.compile(
        r"^This Laboratory Test Report shall not be reproduced except in full",
        flags=re.IGNORECASE,
    ),
    re.compile(
        r"^The contents are guaranteed to be originals \(not modified\)",
        flags=re.IGNORECASE,
    ),
)
_NON_PRODUCT_SAMPLE_PATTERN = re.compile(
    r"\b(?:PCB|BUS\s*BAR|BUSBAR|TEST\s+FIXTURE|FIXTURE|ADAPTER)\b",
    flags=re.IGNORECASE,
)
_TRANSIENT_WORD_FILE_LOCK_WINERRORS = {32, 33}
_TRANSIENT_WORD_FILE_LOCK_DELAYS_SECONDS = (
    0.05,
    0.1,
    0.2,
    0.4,
    0.8,
    1.0,
    1.0,
    1.0,
)
_T = TypeVar("_T")
_LOGGER = logging.getLogger(__name__)


def _is_transient_word_file_lock(exc: OSError) -> bool:
    return getattr(exc, "winerror", None) in _TRANSIENT_WORD_FILE_LOCK_WINERRORS


@dataclass(frozen=True, slots=True)
class _InternalHeaderValues:
    report_number: str
    report_date: str
    dates_tested: str
    title: str
    prepared_by: str
    approved_by: str


@dataclass(slots=True)
class _NumberedHeadingIndex:
    """Build once with Word Find and retain live ranges through later edits."""

    _matches: dict[tuple[int, str], tuple[object, str]]

    @classmethod
    def from_document(
        cls,
        document,
        headings: Iterable[tuple[int, str]],
    ) -> _NumberedHeadingIndex:
        ordered_headings = tuple(headings)
        matches: dict[tuple[int, str], tuple[object, str]] = {}
        for number, heading in ordered_headings:
            matches[(number, heading)] = _find_heading_paragraph(
                document,
                number,
                heading,
            )
        return cls(matches)

    def get(self, number: int, heading: str) -> tuple[object, str]:
        try:
            return self._matches[(number, heading)]
        except KeyError:
            raise ValueError(
                "Internal report is missing required content: "
                f"{number}. {heading}"
            ) from None


class CustomerReportDocumentGateway:
    """Create one customer draft without changing its internal-report source."""

    def __init__(
        self,
        *,
        protected_package_gateway=None,
        transient_retry_delay: Callable[[float], None] | None = None,
    ) -> None:
        self._protected_package_gateway = (
            protected_package_gateway or ProtectedWordPackageGateway()
        )
        self._transient_retry_delay = transient_retry_delay or time.sleep

    def generate_customer_report(
        self,
        *,
        source_path: Path,
        template_path: Path,
        output_path: Path,
        progress: Callable[[str], None] | None = None,
    ) -> Path:
        source = Path(source_path)
        template = Path(template_path)
        output = Path(output_path)
        _validate_paths(source, template, output)
        source_hash = _file_hash(source)
        temporary = output.with_name(
            f".customer-report.{uuid4().hex}.tmp{output.suffix}"
        )
        try:
            _emit_progress(progress, "preparing_template")
            protection_state = self._protected_package_gateway.stage_editable_copy(
                template,
                temporary,
            )
            if progress is None:
                _generate_with_word(source, temporary)
            else:
                _generate_with_word(source, temporary, progress=progress)
            _emit_progress(progress, "verifying_output")
            self._after_word_file_operation(
                lambda: _write_source_report_sha256(temporary, source_hash),
                operation="record source fingerprint",
                path=temporary,
            )
            self._after_word_file_operation(
                lambda: _audit_customer_report(temporary),
                operation="audit generated report",
                path=temporary,
            )
            current_source_hash = self._after_word_file_operation(
                lambda: _file_hash(source),
                operation="verify source fingerprint",
                path=source,
            )
            if current_source_hash != source_hash:
                raise ValueError("The internal report source changed during customer generation.")
            _emit_progress(progress, "protecting_output")
            self._after_word_file_operation(
                lambda: self._protected_package_gateway.restore_password_protection(
                    temporary,
                    protection_state,
                ),
                operation="restore report protection",
                path=temporary,
            )
            self._after_word_file_operation(
                lambda: os.replace(temporary, output),
                operation="publish generated report",
                path=temporary,
            )
        finally:
            try:
                temporary.unlink(missing_ok=True)
            except OSError as exc:
                if not _is_transient_word_file_lock(exc):
                    raise
                _LOGGER.warning(
                    "Word still holds customer-report working copy during cleanup: file=%s",
                    temporary.name,
                )
        return output

    def _after_word_file_operation(
        self,
        action: Callable[[], _T],
        *,
        operation: str,
        path: Path,
    ) -> _T:
        for attempt, delay_seconds in enumerate(
            _TRANSIENT_WORD_FILE_LOCK_DELAYS_SECONDS,
            start=1,
        ):
            try:
                return action()
            except OSError as exc:
                if not _is_transient_word_file_lock(exc):
                    raise
                _LOGGER.warning(
                    "Waiting for Microsoft Word to release customer-report working copy: "
                    "operation=%s file=%s attempt=%s winerror=%s",
                    operation,
                    path.name,
                    attempt,
                    getattr(exc, "winerror", None),
                )
                self._transient_retry_delay(delay_seconds)
        try:
            return action()
        except OSError as exc:
            if not _is_transient_word_file_lock(exc):
                raise
            raise RuntimeError(
                "Microsoft Word did not release the generated working copy. "
                "Wait a moment and try again."
            ) from exc

    def read_source_report_sha256(self, path: Path) -> str | None:
        """Read the invisible lineage fingerprint embedded in a customer DOCX."""
        report = Path(path)
        if not report.is_file() or report.suffix.casefold() != ".docx":
            return None
        try:
            with ZipFile(report) as package:
                content = package.read("docProps/custom.xml")
        except (KeyError, OSError):
            return None
        try:
            root = ET.fromstring(content)
        except ET.ParseError:
            return None
        for prop in root.findall(f"{{{_CUSTOM_PROPERTY_NS}}}property"):
            if prop.attrib.get("name") != _SOURCE_SHA256_PROPERTY or len(prop) != 1:
                continue
            value = (prop[0].text or "").strip().casefold()
            if re.fullmatch(r"[0-9a-f]{64}", value):
                return value
        return None


def _validate_paths(source: Path, template: Path, output: Path) -> None:
    if source.suffix.lower() != ".docx" or template.suffix.lower() != ".docx":
        raise ValueError("Customer report generation requires .docx source and template files.")
    if not source.is_file():
        raise FileNotFoundError("Internal report source file is missing.")
    if not template.is_file():
        raise FileNotFoundError("E-4515 customer report template file is missing.")
    if source.resolve() == output.resolve() or template.resolve() == output.resolve():
        raise ValueError("Customer report output must not replace its source or template.")
    if not output.parent.is_dir():
        raise FileNotFoundError("Customer report output folder is missing.")
    if output.exists():
        raise FileExistsError("Customer report output already exists and will not be replaced.")


def _generate_with_word(
    source_path: Path,
    target_path: Path,
    *,
    progress: Callable[[str], None] | None = None,
) -> None:
    try:
        import pythoncom  # type: ignore[import-not-found]
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - depends on Windows host
        raise OfficeAutomationUnavailable(
            "Customer report generation requires Microsoft Word automation on Windows."
        ) from exc

    pythoncom.CoInitialize()
    word = None
    source = None
    target = None
    template_reference = None
    update_links_at_open = None
    reference_path = target_path.with_name(
        f".customer-template.{uuid4().hex}.reference{target_path.suffix}"
    )
    try:
        shutil.copy2(target_path, reference_path)
        _emit_progress(progress, "opening_word")
        word = win32com.client.DispatchEx("Word.Application")
        update_links_at_open = bool(word.Options.UpdateLinksAtOpen)
        _configure_word_application(word)
        source = word.Documents.Open(
            str(source_path.resolve()),
            PasswordDocument=OFFICE_DOCUMENT_PASSWORD,
            WritePasswordDocument=OFFICE_DOCUMENT_PASSWORD,
            ReadOnly=True,
            AddToRecentFiles=False,
            ConfirmConversions=False,
            NoEncodingDialog=True,
        )
        if source is None:
            raise ValueError(
                "Microsoft Word did not return the internal report document."
            )
        target = word.Documents.Open(
            str(target_path.resolve()),
            PasswordDocument=OFFICE_DOCUMENT_PASSWORD,
            WritePasswordDocument=OFFICE_DOCUMENT_PASSWORD,
            ReadOnly=False,
            AddToRecentFiles=False,
            ConfirmConversions=False,
            NoEncodingDialog=True,
        )
        if target is None:
            raise ValueError(
                "Microsoft Word did not return the E-4515 customer template document."
            )
        try:
            _validate_document_types(source, target)
        except Exception as exc:
            raise ValueError(
                f"Customer report document validation failed: {_error_summary(exc)}"
            ) from exc
        template_reference = word.Documents.Open(
            str(reference_path.resolve()),
            PasswordDocument=OFFICE_DOCUMENT_PASSWORD,
            WritePasswordDocument=OFFICE_DOCUMENT_PASSWORD,
            ReadOnly=True,
            AddToRecentFiles=False,
            ConfirmConversions=False,
            NoEncodingDialog=True,
        )
        try:
            _emit_progress(progress, "copying_content")
            target_headings = _copy_customer_body(
                source,
                target,
                progress=progress,
            )
        except Exception as exc:
            raise ValueError(
                f"Customer report body generation failed: {_error_summary(exc)}"
            ) from exc
        try:
            _emit_progress(progress, "formatting_document")
            _normalize_customer_sections(
                target,
                template_reference,
                target_headings,
            )
            _copy_customer_header(source, target)
            _refresh_customer_fields(target)
        except Exception as exc:
            raise ValueError(
                f"Customer report header generation failed: {_error_summary(exc)}"
            ) from exc
        _emit_progress(progress, "saving_document")
        target.Save()
    except Exception as exc:
        summary = _error_summary(exc)
        raise ValueError(f"Unable to generate customer report: {summary[:320]}") from None
    finally:
        if template_reference is not None:
            try:
                template_reference.Close(SaveChanges=False)
            except Exception as exc:
                _LOGGER.warning(
                    "Unable to close customer-template reference in Word: %s",
                    _error_summary(exc),
                )
        if target is not None:
            try:
                target.Close(SaveChanges=False)
            except Exception as exc:
                _LOGGER.warning(
                    "Unable to close customer-report working copy in Word: %s",
                    _error_summary(exc),
                )
        if source is not None:
            try:
                source.Close(SaveChanges=False)
            except Exception as exc:
                _LOGGER.warning(
                    "Unable to close internal-report working copy in Word: %s",
                    _error_summary(exc),
                )
        if word is not None:
            if update_links_at_open is not None:
                try:
                    word.Options.UpdateLinksAtOpen = update_links_at_open
                except Exception as exc:
                    _LOGGER.warning(
                        "Unable to restore Microsoft Word link-update preference: %s",
                        _error_summary(exc),
                    )
            try:
                word.Quit()
            except Exception as exc:
                _LOGGER.warning(
                    "Unable to quit customer-report Word automation: %s",
                    _error_summary(exc),
                )
        template_reference = None
        target = None
        source = None
        word = None
        gc.collect()
        pythoncom.CoUninitialize()
        try:
            reference_path.unlink(missing_ok=True)
        except OSError as exc:
            if not _is_transient_word_file_lock(exc):
                raise
            _LOGGER.warning(
                "Word still holds customer-template reference during cleanup: file=%s",
                reference_path.name,
            )


def _validate_document_types(source, target) -> None:
    target_header = _first_page_header_table(target, "E-4515 template")
    try:
        source_header = _first_page_header_table(source, "internal report")
    except ValueError:
        source_header = None
    header_is_internal = (
        source_header is not None
        and _INTERNAL_REPORT_LABEL in _clean_text(source_header.Range.Text)
    )
    if (
        not header_is_internal
        and _find_text(source, _INTERNAL_REPORT_LABEL, required=False) is None
    ):
        raise ValueError("Selected source is not a Laboratory Test Report.")
    if _CUSTOMER_REPORT_LABEL not in _clean_text(target_header.Range.Text):
        raise ValueError("Configured E-4515 template is not a Customer Test Report.")


def _copy_customer_header(source, target) -> None:
    values = _extract_internal_header(source)
    target_table = _first_page_header_table(target, "E-4515 template")
    report_number = _customer_report_number(values.report_number)
    _set_cell_text(target_table, 3, 1, report_number)
    _set_cell_text(target_table, 3, 2, values.report_date)
    _set_cell_text(target_table, 3, 3, values.dates_tested)
    _set_cell_text(target_table, 5, 1, values.title)
    _set_cell_text(target_table, 5, 2, values.prepared_by)
    _set_cell_text(target_table, 5, 3, values.approved_by)

    for index in range(1, int(target.Sections.Count) + 1):
        continuation_header = target.Sections(index).Headers(1).Range
        if int(continuation_header.Tables.Count) < 1:
            raise ValueError("E-4515 continuation header table is missing.")
        _set_continuation_report_number(
            continuation_header.Tables(1),
            report_number,
        )


def _copy_customer_body(
    source,
    target,
    *,
    progress: Callable[[str], None] | None = None,
) -> _NumberedHeadingIndex:
    source_headings = _NumberedHeadingIndex.from_document(
        source,
        (
            (1, "PURPOSE"),
            (7, "EQUIPMENTS"),
            (8, "REVISION RECORD"),
        ),
    )
    purpose, _ = source_headings.get(1, "PURPOSE")
    equipment, _ = source_headings.get(7, "EQUIPMENTS")
    if int(equipment.Start) <= int(purpose.Start):
        raise ValueError("Internal report section order is invalid.")
    core = source.Range(int(purpose.Start), int(equipment.Start))
    destination = _prepare_customer_body_destination(target)
    destination.FormattedText = core.FormattedText

    revision, _ = source_headings.get(8, "REVISION RECORD")
    revision_note = _find_text(source, _REVISION_NOTE)
    if int(revision_note.End) <= int(revision.Start):
        raise ValueError("Internal report revision record is incomplete.")
    revision_content = source.Range(int(revision.Start), int(revision_note.End))
    end = target.Content.Duplicate
    end.End = max(int(end.Start), int(end.End) - 1)
    end.Collapse(0)
    end.FormattedText = revision_content.FormattedText

    marker = target.Content.Duplicate
    marker.End = max(int(marker.Start), int(marker.End) - 1)
    marker.Collapse(0)
    marker.InsertAfter(f"\r\r\r{_END_MARKER}\r")
    marker.Font.Name = "Times New Roman"
    marker.Font.Size = 12
    marker.Font.Bold = True
    marker.Font.Italic = True
    marker.ParagraphFormat.Alignment = 1

    target_heading_pairs = tuple(
        (8 if heading == "REVISION RECORD" else index, heading)
        for index, heading in enumerate(_CUSTOMER_HEADINGS, start=1)
    )
    target_headings = _NumberedHeadingIndex.from_document(target, target_heading_pairs)
    for number, heading in target_heading_pairs:
        found, matched_text = target_headings.get(number, heading)
        if matched_text == heading:
            found.ListFormat.RemoveNumbers()
            continue
        delete_prefix = target.Range(
            int(found.Start),
            int(found.Start) + len(matched_text) - len(heading),
        )
        delete_prefix.Text = ""
    _emit_progress(progress, "cleaning_content")
    _sanitize_customer_body(target)
    _normalize_customer_body_typography(target, target_headings)
    _format_customer_headings(target, target_headings)
    _merge_customer_disclaimer(target)
    return target_headings


def _emit_progress(progress: Callable[[str], None] | None, stage: str) -> None:
    if progress is not None:
        progress(stage)


def _configure_word_application(word) -> None:
    """Keep isolated automation silent and preserve linked content without refreshing it."""
    word.Visible = False
    word.DisplayAlerts = 0
    word.ScreenUpdating = False
    word.AutomationSecurity = 3
    word.Options.UpdateLinksAtOpen = False


def _prepare_customer_body_destination(document):
    """Clear the template's continuation body while retaining its section boundary."""
    if int(document.Sections.Count) != 2:
        raise ValueError("E-4515 customer report must retain exactly two sections.")
    continuation = document.Sections(2).Range.Duplicate
    start = int(continuation.Start)
    continuation.End = max(start, int(continuation.End) - 1)
    if int(continuation.End) > start:
        continuation.Delete()
    return document.Range(start, start)


def _normalize_customer_sections(
    document,
    template_reference,
    headings: _NumberedHeadingIndex,
) -> None:
    _normalize_customer_revision_page(document, headings)
    _left_align_customer_revision_note(document)
    _restore_customer_section_geometry(document, template_reference)


def _restore_customer_section_geometry(document, template_reference) -> None:
    template_section_count = int(template_reference.Sections.Count)
    document_section_count = int(document.Sections.Count)
    if template_section_count != 2 or document_section_count != template_section_count:
        raise ValueError("E-4515 customer report must retain exactly two sections.")
    for index in range(1, int(document.Sections.Count) + 1):
        section = document.Sections(index)
        donor = template_reference.Sections(index)
        section.PageSetup.SectionStart = donor.PageSetup.SectionStart
        section.PageSetup.OddAndEvenPagesHeaderFooter = (
            donor.PageSetup.OddAndEvenPagesHeaderFooter
        )
        section.PageSetup.DifferentFirstPageHeaderFooter = (
            donor.PageSetup.DifferentFirstPageHeaderFooter
        )
        for attribute in (
            "TopMargin",
            "BottomMargin",
            "LeftMargin",
            "RightMargin",
            "HeaderDistance",
            "FooterDistance",
            "PageHeight",
            "PageWidth",
        ):
            setattr(
                section.PageSetup,
                attribute,
                getattr(donor.PageSetup, attribute),
            )


def _normalize_customer_revision_page(
    document,
    headings: _NumberedHeadingIndex,
) -> None:
    revision, _ = headings.get(8, "REVISION RECORD")
    protected_anchor_starts = {
        int(document.Shapes(index).Anchor.Paragraphs(1).Range.Start)
        for index in range(1, int(document.Shapes.Count) + 1)
    }
    protected_anchor_starts.update(
        int(document.InlineShapes(index).Range.Paragraphs(1).Range.Start)
        for index in range(1, int(document.InlineShapes.Count) + 1)
    )
    preceding = document.Range(0, int(revision.Start))
    for paragraph_index in range(int(preceding.Paragraphs.Count), 0, -1):
        paragraph = preceding.Paragraphs(paragraph_index)
        start = int(paragraph.Range.Start)
        if _clean_text(paragraph.Range.Text) or start in protected_anchor_starts:
            break
        paragraph.Range.Delete()
    revision.Paragraphs(1).Format.PageBreakBefore = True


def _left_align_customer_revision_note(document) -> None:
    revision_note = _find_text(document, _REVISION_NOTE, required=False)
    if revision_note is not None:
        revision_note.Paragraphs(1).Format.Alignment = 0


def _refresh_customer_fields(document) -> None:
    document.Repaginate()
    for index in range(1, int(document.Sections.Count) + 1):
        section = document.Sections(index)
        for kind in (1, 2, 3):
            section.Headers(kind).Range.Fields.Update()
            section.Footers(kind).Range.Fields.Update()


def _extract_internal_header(document) -> _InternalHeaderValues:
    try:
        table = _first_page_header_table(document, "internal report")
    except ValueError:
        table = None
    if table is not None:
        return _InternalHeaderValues(
            report_number=_cell_text(table, 3, 1),
            report_date=_cell_text(table, 3, 2),
            dates_tested=_cell_text(table, 3, 3),
            title=_cell_text(table, 5, 2, preserve_paragraphs=True),
            prepared_by=_cell_text(table, 5, 3, preserve_paragraphs=True),
            approved_by=_cell_text(table, 5, 4, preserve_paragraphs=True),
        )

    for index in range(1, int(document.Tables.Count) + 1):
        candidate = document.Tables(index)
        text = _clean_text(candidate.Range.Text)
        if not all(
            label in text
            for label in (
                "REPORT NO.",
                "DATE OF REPORT",
                "DATES TESTED",
                "TITLE",
                "PREPARED",
                "APPROVED",
            )
        ):
            continue
        return _InternalHeaderValues(
            report_number=_labeled_cell_value(candidate, "REPORT NO."),
            report_date=_labeled_cell_value(candidate, "DATE OF REPORT"),
            dates_tested=_labeled_cell_value(candidate, "DATES TESTED"),
            title=_labeled_cell_value(candidate, "TITLE", preserve_paragraphs=True),
            prepared_by=_labeled_cell_value(
                candidate,
                "PREPARED BY/TITLE",
                preserve_paragraphs=True,
            ),
            approved_by=_labeled_cell_value(
                candidate,
                "APPROVED BY/TITLE",
                preserve_paragraphs=True,
            ),
        )
    raise ValueError("internal report first-page information table is missing.")


def _labeled_cell_value(table, label: str, *, preserve_paragraphs: bool = False) -> str:
    compact_label = _clean_text(label).casefold()
    for index in range(1, int(table.Range.Cells.Count) + 1):
        raw = str(table.Range.Cells(index).Range.Text).replace("\x07", "").rstrip("\r")
        compact = _clean_text(raw)
        if not compact.casefold().startswith(compact_label):
            continue
        value = raw[len(label) :].lstrip("\r :")
        result = value.rstrip("\r") if preserve_paragraphs else _clean_text(value)
        if result:
            return result
    raise ValueError(f"Internal report header value is missing: {label}")


def _sanitize_customer_body(document) -> None:
    protected_anchor_starts = {
        int(document.Shapes(index).Anchor.Paragraphs(1).Range.Start)
        for index in range(1, int(document.Shapes.Count) + 1)
        if not (
            int(document.Shapes(index).Type) == 5
            and float(document.Shapes(index).Width) > 520
            and float(document.Shapes(index).Height) < 1
        )
    }
    content = document.Content.Duplicate
    disclosure_paragraphs = _find_disclosure_paragraphs(
        document,
        str(content.Text),
    )
    for paragraph_range in reversed(disclosure_paragraphs):
        paragraph = paragraph_range.Paragraphs(1)
        text = _clean_text(paragraph.Range.Text)
        if not any(pattern.search(text) for pattern in _INTERNAL_DISCLOSURE_PATTERNS):
            continue
        start = int(paragraph.Range.Start)
        if start not in protected_anchor_starts:
            paragraph.Range.Delete()
            continue
        paragraph_content = paragraph.Range.Duplicate
        paragraph_content.End = max(
            int(paragraph_content.Start),
            int(paragraph_content.End) - 1,
        )
        if int(paragraph_content.End) <= int(paragraph_content.Start):
            continue
        tail = document.Range(
            int(paragraph_content.Start) + 1,
            int(paragraph_content.End),
        )
        tail.Text = ""
        anchor_character = document.Range(
            int(paragraph_content.Start),
            int(paragraph_content.Start) + 1,
        )
        anchor_character.Font.Hidden = False
        anchor_character.Font.Color = 16777215
        anchor_character.Font.Size = 1

    _remove_body_footer_artifacts(document)

    for index in range(1, int(document.Tables.Count) + 1):
        table = document.Tables(index)
        table_text = _clean_text(table.Range.Text)
        if all(
            label in table_text
            for label in ("Description", "Part #", "Lot", "Base Mat", "Contact Plating")
        ):
            _sanitize_sample_table(table)
            continue
        if "#" not in table_text and "MFG*" not in table_text.upper():
            continue
        for cell_index in range(1, int(table.Range.Cells.Count) + 1):
            cell = table.Range.Cells(cell_index)
            value = _clean_text(cell.Range.Text)
            updated = value
            if re.fullmatch(r"\d+#", value):
                updated = value[:-1]
            updated = re.sub(r"\bMFG\*", "MFG", updated, flags=re.IGNORECASE)
            if updated != value:
                cell.Range.Text = updated


def _internal_disclosure_spans(text: str) -> list[tuple[int, int]]:
    """Locate disclosure paragraphs in one bulk Word text read."""
    spans: list[tuple[int, int]] = []
    for match in re.finditer(r"[^\r]*(?:\r|$)", text):
        if match.start() == match.end():
            continue
        paragraph_text = _clean_text(match.group())
        if any(
            pattern.search(paragraph_text)
            for pattern in _INTERNAL_DISCLOSURE_PATTERNS
        ):
            spans.append((match.start(), match.end()))
    return spans


def _find_disclosure_paragraphs(document, content_text: str) -> list[object]:
    """Use bulk text to choose a few Word Find queries instead of scanning every range."""
    needles = {
        content_text[start:end].strip("\r\x07 ")[:180]
        for start, end in _internal_disclosure_spans(content_text)
    }
    matches: dict[int, object] = {}
    content_end = int(document.Content.End)
    for needle in needles:
        if not needle:
            continue
        search = document.Content.Duplicate
        while True:
            finder = search.Find
            finder.ClearFormatting()
            finder.Text = needle
            finder.Forward = True
            finder.Wrap = 0
            finder.MatchCase = False
            finder.MatchWholeWord = False
            finder.MatchWildcards = False
            if not finder.Execute():
                break
            paragraph = search.Paragraphs(1).Range.Duplicate
            paragraph_text = _clean_text(paragraph.Text)
            if any(
                pattern.search(paragraph_text)
                for pattern in _INTERNAL_DISCLOSURE_PATTERNS
            ):
                matches[int(paragraph.Start)] = paragraph
            next_start = int(search.End)
            if next_start >= content_end:
                break
            search = document.Range(next_start, content_end)
    return [matches[start] for start in sorted(matches)]


def _remove_body_footer_artifacts(document) -> None:
    occupied_anchors = {
        int(document.Shapes(index).Anchor.Paragraphs(1).Range.Start)
        for index in range(1, int(document.Shapes.Count) + 1)
        if not (
            int(document.Shapes(index).Type) == 5
            and float(document.Shapes(index).Width) > 520
            and float(document.Shapes(index).Height) < 1
        )
    }
    for index in range(int(document.Shapes.Count), 0, -1):
        shape = document.Shapes(index)
        if not (
            int(shape.Type) == 5
            and float(shape.Width) > 520
            and float(shape.Height) < 1
            and not _clean_text(shape.Anchor.Text)
        ):
            continue
        anchor = shape.Anchor.Paragraphs(1).Range.Duplicate
        anchor_start = int(anchor.Start)
        shape.Delete()
        if anchor_start not in occupied_anchors:
            anchor.Delete()

def _format_customer_headings(
    document,
    headings: _NumberedHeadingIndex,
) -> None:
    for index, heading in enumerate(_CUSTOMER_HEADINGS, start=1):
        number = 8 if heading == "REVISION RECORD" else index
        found, _ = headings.get(number, heading)
        paragraph = found.Paragraphs(1)
        paragraph.Range.Style = document.Styles(-1)
        paragraph.Format.SpaceBefore = 0
        paragraph.Format.SpaceAfter = 0
        paragraph.Format.LineSpacingRule = 1
        paragraph.Format.PageBreakBefore = False
        paragraph.Format.KeepWithNext = heading in _KEEP_WITH_NEXT_CUSTOMER_HEADINGS
        paragraph.Range.Font.Name = "Times New Roman"
        paragraph.Range.Font.Size = 12
        paragraph.Range.Font.Bold = True
        paragraph.Range.Font.Underline = 1


def _normalize_customer_body_typography(
    document,
    headings: _NumberedHeadingIndex,
) -> None:
    purpose, _ = headings.get(1, "PURPOSE")
    description, _ = headings.get(4, "TEST DESCRIPTION")
    introduction = document.Range(int(purpose.Start), int(description.Start))
    introduction.Font.Name = "Arial"
    for index in range(1, int(document.Tables.Count) + 1):
        table = document.Tables(index)
        table.Range.Font.Name = "Arial"
        table.Range.ParagraphFormat.SpaceBefore = 0
        table.Range.ParagraphFormat.SpaceAfter = 0
        table.Range.ParagraphFormat.LineSpacingRule = 0


def _sanitize_sample_table(table) -> None:
    for row_index in range(int(table.Rows.Count), 1, -1):
        row = table.Rows(row_index)
        description = _clean_text(row.Cells(1).Range.Text)
        if _NON_PRODUCT_SAMPLE_PATTERN.search(description):
            row.Delete()
            continue
        material = _clean_text(row.Cells(4).Range.Text)
        plating = _clean_text(row.Cells(5).Range.Text)
        customer_material = re.sub(
            r"Copper\s+alloy\s+and\s+SUS\d*",
            "Copper alloy/SUS",
            material,
            flags=re.IGNORECASE,
        )
        customer_material = re.sub(
            r"\bC\d{3,}[A-Z0-9R\-/]*\b",
            "Copper alloy",
            customer_material,
            flags=re.IGNORECASE,
        )
        customer_plating = re.sub(
            r"\bPE\d+\b",
            "Lubricant",
            plating,
            flags=re.IGNORECASE,
        )
        if customer_material != material:
            row.Cells(4).Range.Text = customer_material
        if customer_plating != plating:
            row.Cells(5).Range.Text = customer_plating


def _merge_customer_disclaimer(document) -> None:
    scope = _find_text(document, _SAMPLE_SCOPE_TEXT, required=False)
    acceptance = _find_text(document, _ACCEPTANCE_TEXT, required=False)
    scope_text: str | None = None
    if scope is not None:
        sentence = scope.Duplicate
        sentence.End = sentence.Start
        for _ in range(1000):
            sentence.MoveEnd(1, 1)
            if str(sentence.Characters.Last.Text) == ".":
                break
        else:
            raise ValueError("Customer-report sample-scope sentence is incomplete.")
        scope_text = _clean_text(sentence.Text)
        sentence.Delete()
        acceptance = _find_text(document, _ACCEPTANCE_TEXT, required=False)
    if acceptance is None:
        end = document.Content.Duplicate
        end.Collapse(0)
        end.InsertAfter(_ACCEPTANCE_TEXT)
        acceptance = _find_text(document, _ACCEPTANCE_TEXT)
    if scope_text:
        insert = acceptance.Duplicate
        insert.Collapse(0)
        insert.Text = f" {scope_text}"
        insert.Font.Name = acceptance.Font.Name
        insert.Font.Size = acceptance.Font.Size
        insert.Font.Bold = acceptance.Font.Bold
        insert.Font.Italic = acceptance.Font.Italic


def _find_text(document, text: str, *, required: bool = True):
    search = document.Content.Duplicate
    finder = search.Find
    finder.ClearFormatting()
    finder.Text = text
    finder.Forward = True
    finder.Wrap = 0
    finder.MatchCase = True
    finder.MatchWholeWord = False
    if finder.Execute():
        return search
    if required:
        raise ValueError(f"Internal report is missing required content: {text}")
    return None


def _find_heading_paragraph(document, number: int, heading: str):
    content_end = int(document.Content.End)
    search = document.Content.Duplicate
    accepted_text = {
        heading: heading,
        f"{number}. {heading}": f"{number}. {heading}",
        f"{number}.{heading}": f"{number}.{heading}",
    }
    while True:
        finder = search.Find
        finder.ClearFormatting()
        finder.Text = heading
        finder.Forward = True
        finder.Wrap = 0
        finder.MatchCase = True
        finder.MatchWholeWord = False
        if not finder.Execute():
            break
        paragraph = search.Paragraphs(1).Range
        matched_text = accepted_text.get(_clean_text(paragraph.Text))
        if matched_text is not None:
            found = paragraph.Duplicate
            found.End = max(int(found.Start), int(found.End) - 1)
            return found, matched_text
        next_start = int(search.End)
        if next_start >= content_end:
            break
        search = document.Range(next_start, content_end)
    raise ValueError(
        f"Internal report is missing required content: {number}. {heading}"
    )


def _first_page_header_table(document, label: str):
    if int(document.Sections.Count) < 1:
        raise ValueError(f"{label} has no document sections.")
    tables = document.Sections(1).Headers(2).Range.Tables
    for index in range(1, int(tables.Count) + 1):
        table = tables(index)
        if int(table.Rows.Count) >= 5:
            try:
                table.Cell(5, 3)
            except Exception:
                continue
            return table
    raise ValueError(f"{label} first-page header table is missing or incompatible.")


def _cell_text(table, row: int, column: int, *, preserve_paragraphs: bool = False) -> str:
    value = str(table.Cell(row, column).Range.Text).replace("\x07", "")
    if preserve_paragraphs:
        return value.rstrip("\r")
    return _clean_text(value)


def _set_cell_text(table, row: int, column: int, value: str) -> None:
    table.Cell(row, column).Range.Text = value


def _set_continuation_report_number(table, report_number: str) -> None:
    """Keep the continuation-header label distinct from the report identifier."""
    cell = table.Cell(1, 1)
    cell.Range.Text = f"{_CONTINUATION_REPORT_PREFIX}{report_number}"
    identifier = cell.Range.Duplicate
    identifier.Start = int(cell.Range.Start) + len(_CONTINUATION_REPORT_PREFIX)
    identifier.End = int(identifier.Start) + len(report_number)
    identifier.Font.Name = "Arial"
    identifier.Font.Size = 10
    identifier.Font.Bold = True


def _clean_text(value: str) -> str:
    return " ".join(value.replace("\r", " ").replace("\x07", " ").split())


def _customer_report_number(value: str) -> str:
    report_number = _clean_text(value)
    if not report_number:
        raise ValueError("Internal report number is blank.")
    match = re.search(r"\s+Rev\.[A-Z]+$", report_number, flags=re.IGNORECASE)
    if match:
        base = report_number[: match.start()]
        suffix = report_number[match.start() :]
    else:
        base = report_number
        suffix = ""
    if not base.upper().endswith("-CR"):
        base = f"{base}-CR"
    return f"{base}{suffix}"


def _audit_customer_report(path: Path) -> None:
    document = Document(path)
    body_paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    body_text = "\n".join(body_paragraphs)
    for heading in _CUSTOMER_HEADINGS:
        if heading not in body_text:
            raise ValueError(f"Generated customer report is missing {heading}.")
    if re.search(
        r"(?im)^\s*(?:7\s*\.\s*)?EQUIPMENTS\s*$",
        body_text,
    ) or re.search(r"\bAppendix\s+[A-Z]:", body_text, flags=re.IGNORECASE):
        raise ValueError("Generated customer report retained internal-only sections.")
    if any(
        pattern.search(paragraph)
        for paragraph in body_paragraphs
        for pattern in _INTERNAL_DISCLOSURE_PATTERNS
    ):
        raise ValueError("Generated customer report retained internal-only disclosures.")
    if (
        not document.sections
        or not document.sections[0].different_first_page_header_footer
    ):
        raise ValueError("Generated customer report is missing its first-page header.")
    first_page_header_text = " ".join(
        cell.text
        for section in document.sections
        for table in section.first_page_header.tables
        for row in table.rows
        for cell in row.cells
    )
    continuation_header_text = " ".join(
        cell.text
        for section in document.sections
        for table in section.header.tables
        for row in table.rows
        for cell in row.cells
    )
    header_text = _clean_text(
        f"{first_page_header_text} {continuation_header_text}"
    )
    if _CUSTOMER_REPORT_LABEL not in header_text:
        raise ValueError("Generated customer report header is invalid.")
    if "-CR" not in header_text:
        raise ValueError("Generated customer report number is missing -CR.")
    if re.search(r"(?:WW-XXXX-YY-ZZZ|XX-YY-ZZZ)", header_text, flags=re.IGNORECASE):
        raise ValueError("Generated customer report retained a template report-number placeholder.")


def _file_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _error_summary(exc: Exception) -> str:
    if isinstance(exc, OSError) and _is_transient_word_file_lock(exc):
        return (
            "Microsoft Word has not released the generated working copy yet. "
            "Wait a moment and try again."
        )
    summary = " ".join(str(exc).split()) or exc.__class__.__name__
    return re.sub(
        re.escape(OFFICE_DOCUMENT_PASSWORD),
        "[redacted]",
        summary,
        flags=re.IGNORECASE,
    )[:240]


def _write_source_report_sha256(path: Path, fingerprint: str) -> None:
    value = fingerprint.strip().casefold()
    if not re.fullmatch(r"[0-9a-f]{64}", value):
        raise ValueError("Customer report source fingerprint is invalid.")
    report = Path(path)
    replacement = report.with_name(f".{report.stem}.{uuid4().hex}.props{report.suffix}")
    try:
        with ZipFile(report, "r") as source:
            entries = [(item, source.read(item.filename)) for item in source.infolist()]
        payload = {item.filename: content for item, content in entries}
        payload["docProps/custom.xml"] = _custom_properties_xml(
            payload.get("docProps/custom.xml"),
            value,
        )
        payload["_rels/.rels"] = _package_relationships_xml(payload["_rels/.rels"])
        payload["[Content_Types].xml"] = _content_types_xml(
            payload["[Content_Types].xml"]
        )
        seen: set[str] = set()
        with ZipFile(replacement, "w", compression=ZIP_DEFLATED) as target:
            for item, _content in entries:
                target.writestr(item, payload[item.filename])
                seen.add(item.filename)
            for name in ("docProps/custom.xml",):
                if name not in seen:
                    target.writestr(name, payload[name])
        os.replace(replacement, report)
    finally:
        replacement.unlink(missing_ok=True)


def _custom_properties_xml(existing: bytes | None, fingerprint: str) -> bytes:
    ET.register_namespace("", _CUSTOM_PROPERTY_NS)
    ET.register_namespace("vt", _VALUE_TYPE_NS)
    if existing:
        root = ET.fromstring(existing)
    else:
        root = ET.Element(f"{{{_CUSTOM_PROPERTY_NS}}}Properties")
    properties = root.findall(f"{{{_CUSTOM_PROPERTY_NS}}}property")
    target = next(
        (item for item in properties if item.attrib.get("name") == _SOURCE_SHA256_PROPERTY),
        None,
    )
    if target is None:
        used = [int(item.attrib.get("pid", "1")) for item in properties]
        target = ET.SubElement(
            root,
            f"{{{_CUSTOM_PROPERTY_NS}}}property",
            {
                "fmtid": "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}",
                "pid": str(max(used, default=1) + 1),
                "name": _SOURCE_SHA256_PROPERTY,
            },
        )
    else:
        for child in tuple(target):
            target.remove(child)
    text = ET.SubElement(target, f"{{{_VALUE_TYPE_NS}}}lpwstr")
    text.text = fingerprint
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _package_relationships_xml(existing: bytes) -> bytes:
    ET.register_namespace("", _PACKAGE_REL_NS)
    root = ET.fromstring(existing)
    if any(item.attrib.get("Type") == _CUSTOM_REL_TYPE for item in root):
        return existing
    used = {item.attrib.get("Id") for item in root}
    index = 1
    while f"rId{index}" in used:
        index += 1
    ET.SubElement(
        root,
        f"{{{_PACKAGE_REL_NS}}}Relationship",
        {
            "Id": f"rId{index}",
            "Type": _CUSTOM_REL_TYPE,
            "Target": "docProps/custom.xml",
        },
    )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _content_types_xml(existing: bytes) -> bytes:
    ET.register_namespace("", _CONTENT_TYPE_NS)
    root = ET.fromstring(existing)
    if any(item.attrib.get("PartName") == "/docProps/custom.xml" for item in root):
        return existing
    ET.SubElement(
        root,
        f"{{{_CONTENT_TYPE_NS}}}Override",
        {
            "PartName": "/docProps/custom.xml",
            "ContentType": _CUSTOM_CONTENT_TYPE,
        },
    )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)
