"""Template-driven openpyxl writer for IR/DWV record workbooks."""

from __future__ import annotations

from copy import copy, deepcopy
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from math import ceil
from pathlib import Path
import re
from typing import Any, Iterable
from zipfile import BadZipFile
from unicodedata import east_asian_width

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.formula.translate import Translator
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import DEFAULT_COLUMN_WIDTH, pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet
from backend.application.matrix_editor_ir_dwv_record_projection import (
    IrDwvEquipment, IrDwvRecordStep, IrDwvRecordGroup, IrDwvRecordProjection,
)

from backend.infrastructure.office.ir_dwv_record_workbook_layout import (
    IrDwvBlockLayout,
    TEMPLATE_PAIR_ROW_COUNT,
    TEMPLATE_SAMPLE_COUNT,
    block_layout,
    dynamic_merge_ranges,
    plan_block_placements,
)


class IrDwvRecordWorkbookGateway:
    """Build a blank record workbook without mutating its registered template."""

    def __init__(self, template_path: Path | None) -> None:
        self._template_path = Path(template_path) if template_path is not None else None

    def template_fingerprint(self) -> str:
        """Read and validate one consistent template snapshot without changing it."""
        workbook, payload = self._load_template()
        try:
            _validate_template_fingerprint(workbook.worksheets[0])
        finally:
            workbook.close()
        return f"{self._template_path.resolve()}:{sha256(payload).hexdigest()}"

    def _load_template(self):
        self._validate_template_available()
        payload = self._template_path.read_bytes()
        try:
            return load_workbook(BytesIO(payload), data_only=False), payload
        except (BadZipFile, InvalidFileException, SyntaxError, KeyError) as exc:
            raise ValueError(f"IR/DWV template is not a readable XLSX workbook: {exc}") from exc

    def validate_projection(self, projection: IrDwvRecordProjection) -> None:
        """Check layout capacity without opening or changing any output file."""
        for group in projection.groups:
            if not group.steps:
                continue
            sample_count = _positive_integer(group.sample_count)
            pairs = tuple(item.strip() for item in group.steps[0].measurement_pairs if item.strip())
            try:
                plan_block_placements(sample_count=sample_count, pair_count=len(pairs), step_count=len(group.steps))
            except ValueError as exc:
                raise ValueError(f"{group.label}: {exc} Sample count: {sample_count}.") from exc

    def write(
        self,
        *,
        output_path: Path,
        projection: IrDwvRecordProjection,
    ) -> Path:
        target = Path(output_path)
        if target.suffix.lower() != ".xlsx":
            raise ValueError("IR/DWV record output must be .xlsx.")
        if not target.parent.is_dir():
            raise FileNotFoundError(f"Output directory does not exist: {target.parent}")
        self.validate_projection(projection)
        workbook, _payload = self._load_template()
        try:
            prototype = workbook.worksheets[0]
            _validate_template_fingerprint(prototype)
            logos = template_block_logos(prototype)
            generated: list[Worksheet] = []
            generated_names: list[tuple[Worksheet, str]] = []
            used_names: set[str] = set()
            for group in projection.groups:
                sample_count = _positive_integer(group.sample_count)
                if not group.steps:
                    continue
                pair_sets = {
                    tuple(item.strip() for item in step.measurement_pairs if item.strip())
                    for step in group.steps
                }
                if len(pair_sets) != 1:
                    raise ValueError(
                        f"{group.label} IR/DWV steps must share one measurement-pair list."
                    )
                pairs = next(iter(pair_sets))
                for step in group.steps:
                    if (
                        step.expected_fee_quantity is not None
                        and step.expected_fee_quantity != sample_count * len(pairs)
                    ):
                        raise ValueError(
                            f"Fee quantity mismatch for {group.label} {step.label}: "
                            f"expected {sample_count * len(pairs)}, got "
                            f"{step.expected_fee_quantity}."
                        )
                placements = plan_block_placements(
                    sample_count=sample_count,
                    pair_count=len(pairs),
                    step_count=len(group.steps),
                )
                group_sheets: dict[int, Worksheet] = {}
                for placement, step in zip(placements, group.steps, strict=True):
                    sheet = group_sheets.get(placement.sheet_index)
                    if sheet is None:
                        suffix = (
                            ""
                            if placement.sheet_index == 0
                            else f" ({placement.sheet_index + 1})"
                        )
                        final_name = _unique_sheet_name(
                            f"{group.label}{suffix}", used_names
                        )
                        sheet = workbook.create_sheet(
                            f"__ir_dwv_generated_{len(generated) + 1}"
                        )
                        _copy_sheet_settings(prototype, sheet)
                        group_sheets[placement.sheet_index] = sheet
                        generated.append(sheet)
                        generated_names.append((sheet, final_name))
                    layout = block_layout(
                        origin_column=placement.origin_column,
                        sample_count=sample_count,
                        pair_count=len(pairs),
                    )
                    for offset, column in enumerate(
                        range(sheet.max_column + 1, layout.origin_column)
                    ):
                        _copy_column_dimension(prototype, sheet, 13 + offset, column)
                    previous_remarks_height = sheet.row_dimensions[11].height or 0
                    _copy_resized_template_block(prototype, sheet, layout)
                    copy_block_logos(logos, sheet, layout)
                    _fill_block(
                        sheet=sheet,
                        layout=layout,
                        template=prototype,
                        projection=projection,
                        group=group,
                        step=step,
                        pairs=pairs,
                    )
                    # Side-by-side forms share row 11. A later short/blank
                    # requirement must not reset the height needed earlier.
                    if previous_remarks_height:
                        sheet.row_dimensions[11].height = max(
                            previous_remarks_height, sheet.row_dimensions[11].height or 0,
                        )
            if not generated:
                raise ValueError("IR/DWV record projection contains no test steps.")
            for sheet in tuple(workbook.worksheets):
                if sheet not in generated:
                    workbook.remove(sheet)
            for sheet, final_name in generated_names:
                sheet.title = final_name
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.save(target)
        finally:
            workbook.close()
        return target

    def _validate_template_available(self) -> None:
        if self._template_path is None:
            raise ValueError("Configure the IR/DWV record template in File Locations before generating a form.")
        if not self._template_path.is_file():
            raise FileNotFoundError(
                f"IR/DWV record template does not exist: {self._template_path}"
            )
        lock = self._template_path.with_name(f"~${self._template_path.name}")
        if lock.exists():
            raise ValueError("blocked: 请关闭 Excel 中的 IR&DWV Template.xlsx")


def _positive_integer(value: Any) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValueError("IR/DWV sample count must be a positive integer.")
    return value


def _validate_template_fingerprint(sheet: Worksheet) -> None:
    anchors = {
        "B1": "AP Product Test Laboratory",
        "E1": "Equipment Used",
        "B13": "Test/Test Condition",
        "H13": "Test/Test Condition",
        "B20": "UNITS",
        "L39": "FDQF-E-033",
        "L40": "G",
    }
    changed = [
        coordinate
        for coordinate, expected in anchors.items()
        if sheet[coordinate].value != expected
    ]
    required_merges = {
        "B1:D2",
        "E1:J1",
        "B10:H10",
        "I10:L10",
        "B13:G13",
        "H13:L13",
        "H37:L37",
    }
    actual_merges = {str(item) for item in sheet.merged_cells.ranges}
    ir_statistics_topologies = (
        {"C34:G34", "C35:G35", "C36:G37"},
        {"C34:G34", "C35:G35", "C36:G36", "C37:G37"},
    )
    if (changed or not required_merges.issubset(actual_merges)
            or not any(topology.issubset(actual_merges) for topology in ir_statistics_topologies)):
        detail = ", ".join(changed) or "merged-cell topology"
        raise ValueError(f"IR/DWV template fingerprint mismatch: {detail}.")
    if not _first_block_logo_images(sheet):
        raise ValueError("IR/DWV template must contain a supported first-block LOGO anchored to the form header.")


def copy_block(
    source: Worksheet,
    target: Worksheet,
    *,
    source_min_row: int,
    source_max_row: int,
    source_min_column: int,
    source_max_column: int,
    target_min_row: int,
    target_min_column: int,
) -> None:
    """Copy a rectangular block with styles, formulas, merges and dimensions."""
    row_offset = target_min_row - source_min_row
    column_offset = target_min_column - source_min_column
    for row in range(source_min_row, source_max_row + 1):
        _copy_row_dimension(source, target, row, row + row_offset)
        for column in range(source_min_column, source_max_column + 1):
            _copy_cell(
                source.cell(row, column),
                target.cell(row + row_offset, column + column_offset),
            )
    for column in range(source_min_column, source_max_column + 1):
        _copy_column_dimension(source, target, column, column + column_offset)
    for merged in source.merged_cells.ranges:
        if (
            source_min_row <= merged.min_row <= merged.max_row <= source_max_row
            and source_min_column
            <= merged.min_col
            <= merged.max_col
            <= source_max_column
        ):
            target.merge_cells(
                start_row=merged.min_row + row_offset,
                start_column=merged.min_col + column_offset,
                end_row=merged.max_row + row_offset,
                end_column=merged.max_col + column_offset,
            )


def _copy_cell(source: Any, target: Any) -> None:
    value = source.value
    if isinstance(value, str) and value.startswith("="):
        value = Translator(value, origin=source.coordinate).translate_formula(
            target.coordinate
        )
    target.value = value
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def _copy_row_dimension(
    source: Worksheet,
    target: Worksheet,
    source_row: int,
    target_row: int,
) -> None:
    if source_row not in source.row_dimensions:
        return
    source_dimension = source.row_dimensions[source_row]
    target_dimension = target.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed


def _copy_column_dimension(
    source: Worksheet,
    target: Worksheet,
    source_column: int,
    target_column: int,
) -> None:
    source_dimension = _effective_column_dimension(source, source_column)
    if source_dimension is None:
        return
    target_dimension = target.column_dimensions[get_column_letter(target_column)]
    target_dimension.width = source_dimension.width
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed


def _effective_column_dimension(source: Worksheet, source_column: int):
    # Excel can store one dimension for a whole column range. Indexing a
    # missing letter creates openpyxl's width-13 default, not the Excel width.
    source_dimension = None
    selected_span = None
    for letter, dimension in source.column_dimensions.items():
        first = dimension.min or column_index_from_string(letter)
        last = dimension.max or first
        span = last - first
        if first <= source_column <= last and (
            selected_span is None or span <= selected_span
        ):
            source_dimension = dimension
            selected_span = span
    return source_dimension


def _copy_sheet_settings(source: Worksheet, target: Worksheet) -> None:
    target.sheet_format = copy(source.sheet_format)
    _copy_column_dimension(source, target, 1, 1)
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.freeze_panes = source.freeze_panes
    target.oddHeader = copy(source.oddHeader)
    target.oddFooter = copy(source.oddFooter)
    target.evenHeader = copy(source.evenHeader)
    target.evenFooter = copy(source.evenFooter)
    target.firstHeader = copy(source.firstHeader)
    target.firstFooter = copy(source.firstFooter)


def _copy_resized_template_block(
    source: Worksheet,
    target: Worksheet,
    layout: IrDwvBlockLayout,
) -> None:
    if (
        layout.sample_slot_count == TEMPLATE_SAMPLE_COUNT
        and layout.pair_count <= TEMPLATE_PAIR_ROW_COUNT
    ):
        copy_block(
            source,
            target,
            source_min_row=1,
            source_max_row=40,
            source_min_column=2,
            source_max_column=12,
            target_min_row=1,
            target_min_column=layout.origin_column,
        )
        return

    source_columns = [2]
    source_columns.extend(
        3 + min(index, TEMPLATE_SAMPLE_COUNT - 1)
        for index in range(layout.sample_slot_count)
    )
    source_columns.extend(
        8 + min(index, TEMPLATE_SAMPLE_COUNT - 1)
        for index in range(layout.sample_slot_count)
    )
    source_columns.extend(range(2 + len(source_columns), 2 + layout.width))
    source_rows = _source_rows(layout)
    for target_row, source_row in source_rows:
        _copy_row_dimension(source, target, source_row, target_row)
        for target_column, source_column in zip(
            range(layout.origin_column, layout.last_column + 1),
            source_columns,
            strict=True,
        ):
            _copy_cell(
                source.cell(source_row, source_column),
                target.cell(target_row, target_column),
            )
    for target_column, source_column in zip(
        range(layout.origin_column, layout.last_column + 1),
        source_columns,
        strict=True,
    ):
        _copy_column_dimension(source, target, source_column, target_column)
    # Retain the laboratory's static identity at the left; everything else in
    # the header is rebuilt. Repeating sample columns can duplicate historical
    # calibration/environment values even when final anchors are overwritten.
    for row in range(1, 13):
        first_column = layout.origin_column + 3 if row <= 7 else layout.origin_column
        for column in range(first_column, layout.last_column + 1):
            target.cell(row, column).value = None
    _copy_dynamic_field_styles(source, target, layout)
    separate_ir_statistics = "C36:G36" in {str(merged) for merged in source.merged_cells.ranges}
    for merged in dynamic_merge_ranges(layout, ir_statistics_separate_rows=separate_ir_statistics):
        target.merge_cells(merged)


def template_block_logos(source: Worksheet) -> tuple:
    """Snapshot first-block images once: openpyxl consumes their image streams."""
    return tuple((image._data(), deepcopy(image.anchor), image.width, image.height,
                  _logo_horizontal_span(source, image.anchor))
                 for image in _first_block_logo_images(source))


def _first_block_logo_images(source: Worksheet) -> tuple:
    return tuple(image for image in source._images
                 if hasattr(image.anchor, "_from")
                 and 1 <= image.anchor._from.col <= 11
                 and image.anchor._from.row < 7)


def copy_block_logos(logos, target: Worksheet, layout: IrDwvBlockLayout) -> None:
    """Give each form its own streams and preserve the logo's right alignment."""
    offset = layout.origin_column - 2 + layout.width - 11
    for payload, prototype_anchor, width, height, horizontal_span in logos:
        image = Image(BytesIO(payload))
        image.width, image.height = width, height
        image.anchor = deepcopy(prototype_anchor)
        image.anchor._from.col += offset
        if hasattr(image.anchor, "to"):
            image.anchor.to.col += offset
            # Column-index translation can shrink a TwoCellAnchor when its
            # original wide column becomes one of the narrow repeated slots.
            # Keep the translated right edge; find the left edge by its span.
            if _logo_horizontal_span(target, image.anchor) != horizontal_span:
                column = image.anchor.to.col
                column_offset = image.anchor.to.colOff - horizontal_span
                while column_offset < 0 and column > 0:
                    column -= 1
                    column_offset += _column_width_emu(target, column)
                if column_offset < 0:
                    raise ValueError("IR/DWV template LOGO cannot fit before its right anchor.")
                image.anchor._from.col = column
                image.anchor._from.colOff = column_offset
        target.add_image(image)


def _logo_horizontal_span(sheet: Worksheet, anchor):
    if not hasattr(anchor, "to"):
        return None
    return (sum(_column_width_emu(sheet, column) for column in range(anchor._from.col, anchor.to.col))
            + anchor.to.colOff - anchor._from.colOff)


def _column_width_emu(sheet: Worksheet, column: int) -> int:
    dimension = _effective_column_dimension(sheet, column + 1)
    if dimension is not None and dimension.hidden:
        return 0
    width = dimension.width if dimension is not None else sheet.sheet_format.defaultColWidth
    if width is None:
        width = DEFAULT_COLUMN_WIDTH
    # Standard 96-dpi / seven-pixel OpenXML metric for marker rebasing. The
    # current template's correction is verified at both seven and eight pixels;
    # this is not a font rasterizer or a claim about arbitrary template fonts.
    # https://learn.microsoft.com/en-us/dotnet/api/documentformat.openxml.spreadsheet.column
    pixels = int(((256 * width + int(128 / 7)) / 256) * 7)
    return pixels_to_EMU(pixels)


def _copy_dynamic_field_styles(source: Worksheet, target: Worksheet, layout: IrDwvBlockLayout) -> None:
    # Measurement-column repetition must not select merged-cell placeholders
    # for independently positioned header fields. Copy their semantic anchors
    # before merging so openpyxl can propagate the original outer borders.
    origin, last = layout.origin_column, layout.last_column
    instrument = origin + 3
    target.cell(1, instrument).value = source["E1"].value
    target.cell(1, instrument)._style = copy(source["E1"]._style)
    gage = origin + 6
    for row in range(2, 8):
        for source_column, target_column in zip((5, 8, 9, 10), (instrument, gage, gage + 1, gage + 2), strict=True):
            target.cell(row, target_column)._style = copy(source.cell(row, source_column)._style)
    target.cell(3, instrument).value = source["E3"].value
    target.cell(3, gage).value = source["H3"].value
    second_header = min(last, layout.ir_end_column + 2)
    fields = [(8, 2, 8, origin), (8, 11, 8, last - 1), (9, 2, 9, origin), (9, 11, 9, last - 1),
              (10, 2, 10, origin), (10, 9, 10, second_header), (11, 2, 11, origin), (11, 9, 11, second_header)]
    fields.extend((12, source_column, 12, target_column)
                  for source_column, target_column in zip((2, 5, 8, 11), _personnel_columns(layout), strict=True))
    fields.extend((source_row, source_column, target_row, target_column)
                  for source_row, target_row in ((39, layout.footer_form_row), (40, layout.footer_revision_row))
                  for source_column, target_column in ((2, origin), (11, last - 1), (12, last)))
    for source_row, source_column, target_row, target_column in fields:
        target.cell(target_row, target_column)._style = copy(source.cell(source_row, source_column)._style)


def _source_rows(layout: IrDwvBlockLayout) -> tuple[tuple[int, int], ...]:
    rows: list[tuple[int, int]] = [(row, row) for row in range(1, 14)]
    rows.extend(
        (
            row,
            14 + min(row - layout.conditions_start_row, TEMPLATE_SAMPLE_COUNT - 1),
        )
        for row in range(
            layout.conditions_start_row,
            layout.conditions_end_row + 1,
        )
    )
    rows.extend(
        (target, source)
        for target, source in (
            (layout.sequence_row, 19),
            (layout.units_row, 20),
            (layout.sample_id_row, 21),
        )
    )
    rows.extend(
        (
            row,
            22 + min(row - layout.data_start_row, TEMPLATE_PAIR_ROW_COUNT - 1),
        )
        for row in range(layout.data_start_row, layout.data_end_row + 1)
    )
    rows.extend(
        (layout.stats_min_row + offset, 34 + offset)
        for offset in range(7)
    )
    return tuple(rows)


def _fill_block(
    *,
    sheet: Worksheet,
    layout: IrDwvBlockLayout,
    template: Worksheet,
    projection: IrDwvRecordProjection,
    group: IrDwvRecordGroup,
    step: IrDwvRecordStep,
    pairs: tuple[str, ...],
) -> None:
    origin = layout.origin_column
    last = layout.last_column
    for row in range(13, layout.stats_stdev_row + 1):
        for column in range(layout.dwv_end_column + 1, last + 1):
            cell = sheet.cell(row, column)
            if not isinstance(cell, MergedCell):
                cell.value = None
    _fill_equipment(sheet, layout, projection.equipment)
    sheet.cell(8, origin, f"Start Date:{_text(projection.start_date)}")
    sheet.cell(9, origin, f"Finish Date: {_text(projection.finish_date)}")
    sheet.cell(8, max(origin, last - 1), f"Amb Temp: {projection.ambient_temperature}")
    sheet.cell(9, max(origin, last - 1), f"Rel. Hum.:{projection.relative_humidity}")

    first_header_end = min(
        last,
        layout.ir_end_column + 1,
    )
    second_header_start = min(last, first_header_end + 1)
    sheet.cell(10, origin, f"Item/Process: {group.label}  IR&DWV-{step.label}")
    sheet.cell(10, second_header_start, f"Request  No.:  {projection.request_number}")
    _fill_remarks(sheet, origin, first_header_end, step)
    sheet.cell(11, second_header_start, f"Product Name: {projection.product_name}")

    personnel_columns = _personnel_columns(layout)
    for column, text in zip(
        personnel_columns,
        (
            f"Tested By: {projection.tested_by}",
            f"Checked By:{projection.checked_by}",
            f"Approved By:{projection.approved_by}",
            f"Requestor:{projection.requestor}",
        ),
        strict=True,
    ):
        sheet.cell(12, column, text)

    sheet.cell(13, origin).value = "Test/Test Condition" if step.ir_enabled else None
    sheet.cell(13, layout.dwv_start_column).value = "Test/Test Condition" if step.dwv_enabled else None
    for index in range(layout.sample_slot_count):
        row = layout.conditions_start_row + index
        sheet.cell(row, origin).value = (
            f"{index + 1}. IR testing/ {index + 1}# {step.ir_condition}"
            if step.ir_enabled and index < layout.sample_count else None
        )
        sheet.cell(row, layout.dwv_start_column).value = (
            (
                f"{layout.sample_count + index + 1}. DWV testing/ "
                f"{index + 1}# {step.dwv_condition}"
            ) if step.dwv_enabled and index < layout.sample_count else None
        )

    sheet.cell(layout.sequence_row, origin).value = None
    ir_unit = sheet.cell(layout.units_row, layout.ir_start_column).value
    dwv_unit = sheet.cell(layout.units_row, layout.dwv_start_column).value
    for start, enabled, unit, number_offset in (
        (layout.ir_start_column, step.ir_enabled, ir_unit, 0),
        (layout.dwv_start_column, step.dwv_enabled, dwv_unit, layout.sample_count),
    ):
        for index in range(layout.sample_slot_count):
            active = enabled and index < layout.sample_count
            sheet.cell(layout.sequence_row, start + index).value = number_offset + index + 1 if active else None
            sheet.cell(layout.units_row, start + index).value = unit if active else None
            sheet.cell(layout.sample_id_row, start + index).value = f"{index + 1}#" if active else None
    sheet.cell(layout.units_row, origin, "UNITS")
    sheet.cell(layout.sample_id_row, origin, "SAMPLE ID      \n +&-")

    for row in range(layout.data_start_row, layout.data_end_row + 1):
        pair_index = row - layout.data_start_row
        sheet.cell(row, origin).value = pairs[pair_index] if pair_index < len(pairs) else None
        for column in range(layout.ir_start_column, layout.dwv_end_column + 1):
            sheet.cell(row, column).value = None

    for row, label in (
        (layout.stats_min_row, "Min"),
        (layout.stats_max_row, "Max"),
        (layout.stats_average_row, "Average"),
        (layout.stats_stdev_row, "Stdev"),
    ):
        sheet.cell(row, origin).value = label if step.ir_enabled or step.dwv_enabled else None
    for source_column, source_range, start, enabled in (
        (3, "C22:G33", layout.ir_start_column, step.ir_enabled),
        (8, "H22:L33", layout.dwv_start_column, step.dwv_enabled),
    ):
        actual_range = (
            f"{get_column_letter(start)}{layout.data_start_row}:"
            f"{get_column_letter(start + layout.sample_count - 1)}{layout.data_end_row}"
        )
        for offset in range(4):
            cell = sheet.cell(layout.stats_min_row + offset, start)
            if not isinstance(cell, MergedCell):
                cell.value = (_template_statistic(template.cell(34 + offset, source_column).value,
                                                 source_range, actual_range) if enabled else None)
    for row in (layout.footer_form_row, layout.footer_revision_row):
        for column in range(origin, last + 1):
            cell = sheet.cell(row, column)
            if not isinstance(cell, MergedCell):
                cell.value = None
    sheet.cell(layout.footer_form_row, origin, "此表为DGLAB-LNP-11之附件")
    sheet.cell(layout.footer_form_row, last - 1, "Form No:")
    sheet.cell(layout.footer_form_row, last, "FDQF-E-033")
    sheet.cell(layout.footer_revision_row, last - 1, "Rev:")
    sheet.cell(layout.footer_revision_row, last, "G")


def _template_statistic(value, source_range: str, actual_range: str):
    """Keep template formula intent; never copy historical numeric results."""
    if value == "/":
        return value
    if not isinstance(value, str) or not value.startswith("="):
        return None
    tokens = Tokenizer(value).items
    for token in tokens:
        if token.type == "OPERAND" and token.subtype == "RANGE" and token.value.replace("$", "") == source_range:
            token.value = actual_range
    return "=" + "".join(token.value for token in tokens)


def _fill_remarks(sheet: Worksheet, first_column: int, last_column: int, step: IrDwvRecordStep) -> None:
    lines = ["Remarks:"]
    for label, enabled, source in (("IR", step.ir_enabled, step.ir_source_step),
                                   ("DWV", step.dwv_enabled, step.dwv_source_step)):
        if enabled and source is not None and source.requirement and source.requirement.strip():
            lines.append(f"{label} Requirement: {source.requirement}")
    cell = sheet.cell(11, first_column)
    if len(lines) == 1:
        cell.value = lines[0]
        return
    text = "\n".join(lines)
    # openpyxl truncates strings above Excel's cell limit; fail explicitly.
    if len(text) > 32767:
        raise ValueError("Matrix requirements exceed the Remarks cell's 32767-character limit.")
    font_size = cell.font.sz or 11
    available_pixels = sum(_column_width_emu(sheet, column - 1) for column in range(first_column, last_column + 1)) / 9525
    # Budget one font em per Latin character and two for wide characters,
    # allowing more space than ordinary glyphs without changing template fonts.
    characters_per_line = max(1, int((available_pixels - 8) / (font_size * 4 / 3)))
    wrapped_lines = sum(max(1, ceil(sum(2 if east_asian_width(char) in ("W", "F") else 1 for char in line) / characters_per_line))
                        for line in text.split("\n"))
    height = wrapped_lines * font_size * 1.3 + 4
    if height > 409:
        raise ValueError("Matrix requirements cannot fit visibly within the Remarks row's 409-point height limit.")
    cell.value = text
    alignment = copy(cell.alignment)
    alignment.wrap_text = True
    cell.alignment = alignment
    sheet.row_dimensions[11].height = max(sheet.row_dimensions[11].height or 0, height)


def _fill_equipment(
    sheet: Worksheet,
    layout: IrDwvBlockLayout,
    equipment: Iterable[IrDwvEquipment],
) -> None:
    origin = layout.origin_column
    instrument_column = min(layout.last_column, origin + 3)
    gage_column = min(layout.last_column, origin + 6)
    last_cal_column = min(layout.last_column, gage_column + 1)
    due_column = min(layout.last_column, last_cal_column + 1)
    defaults = (sheet.cell(3, instrument_column).value, sheet.cell(3, gage_column).value)
    for column, label in zip((instrument_column, gage_column, last_cal_column, due_column),
                             ("Instrument", "Gage ID", "Last Cal.", "Cal Due."), strict=True):
        sheet.cell(2, column).value = label
    for row in range(3, 8):
        for column in {
            instrument_column,
            gage_column,
            last_cal_column,
            due_column,
        }:
            sheet.cell(row, column).value = None
    sheet.cell(3, instrument_column).value, sheet.cell(3, gage_column).value = defaults
    for row, item in zip(range(3, 8), equipment):
        sheet.cell(row, instrument_column, item.instrument)
        sheet.cell(row, gage_column, item.gage_id)
        sheet.cell(row, last_cal_column, item.last_calibration)
        sheet.cell(row, due_column, item.calibration_due)


def _personnel_columns(layout: IrDwvBlockLayout) -> tuple[int, int, int, int]:
    base, remainder = divmod(layout.width, 4)
    sizes = [base + (1 if index < remainder else 0) for index in range(4)]
    columns: list[int] = []
    cursor = layout.origin_column
    for size in sizes:
        columns.append(min(layout.last_column, cursor))
        cursor += size
    return tuple(columns)  # type: ignore[return-value]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y/%m/%d")
    return str(value)


def _unique_sheet_name(value: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]+", "_", value.strip())[:31] or "Group"
    candidate = base
    suffix = 2
    while candidate in used:
        tail = f" ({suffix})"
        candidate = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate
