"""Excel gateway for fee-evaluation workbook generation."""

from __future__ import annotations

from collections.abc import Callable, Iterator
from contextlib import contextmanager
from decimal import Decimal
from pathlib import Path
from typing import Any
import os
import shutil
import tempfile

from backend.application.confirmed_matrix_fee_draft_service import FeeEvaluationDraft
from backend.application.confirmed_matrix_fee_template_basic_fill_service import (
    MatrixBasicFillWorkbook,
)
from backend.application.fee_evaluation_edited_export_values import (
    FeeEvaluationEditedExportValues,
)
from backend.application.test_record_fee_dataset_preview_service import (
    TestRecordFeeDatasetPreview,
)
from backend.infrastructure.office.fee_evaluation_identity_header_writer import (
    write_basic_information_identity,
)
from backend.infrastructure.office.fee_evaluation_anchor_snapshot import (
    FeeEvaluationAnchorSnapshot,
)
from backend.infrastructure.office.fee_evaluation_matrix_basic_fill_writer import (
    write_matrix_basic_fill,
)
from backend.infrastructure.office.fee_evaluation_openpyxl_adapter import (
    OpenpyxlFeeSheetAdapter,
)
from backend.infrastructure.office.models import FeeEvaluationWorkbookWriteResult
from backend.infrastructure.office.office_lifecycle import OfficeAutomationUnavailable
from backend.shared.operation_diagnostics import stage, emit


_EXCEL_FILE_FORMATS = {
    ".xlsx": 51,
    ".xls": 56,
}
_EXCEL_CALCULATION_MANUAL = -4135


class FeeEvaluationWorkbookGateway:
    """Generate native XLSX Fee Forms, retaining COM only for legacy XLS files."""

    def __init__(self, excel_app_factory: Callable[[], Any] | None = None) -> None:
        self._excel_app_factory = excel_app_factory

    def generate(
        self,
        *,
        template_path: Path,
        output_path: Path,
        preview: TestRecordFeeDatasetPreview,
    ) -> FeeEvaluationWorkbookWriteResult:
        """Write fee-evaluation content through the selected workbook format path."""
        template = Path(template_path)
        target = Path(output_path)

        def write(workbook: Any) -> None:
            summary = preview.fee_dataset.summary if preview.fee_dataset is not None else None
            sheet = _first_sheet(workbook)
            sheet.Cells(1, 1).Value = "ConnLab Generated Fee Evaluation"
            sheet.Cells(2, 1).Value = f"Project ID: {preview.project_id}"
            sheet.Cells(3, 1).Value = f"Draft ID: {preview.draft_id}"
            if summary is not None:
                sheet.Cells(4, 1).Value = f"Group count: {summary.group_count}"
                sheet.Cells(5, 1).Value = f"Step count: {summary.step_count}"
                sheet.Cells(6, 1).Value = (
                    f"Explicit duration days: {summary.explicit_duration_days}"
                )

        self._generate_workbook(template, target, write)

        return FeeEvaluationWorkbookWriteResult(
            output_path=target,
            status="generated",
            warnings=("Pricing values are not calculated by ConnLab.",),
        )

    def generate_from_draft(
        self,
        *,
        template_path: Path,
        output_path: Path,
        draft: FeeEvaluationDraft,
        prepared_by: str | None,
        approved_by: str | None,
    ) -> FeeEvaluationWorkbookWriteResult:
        """Write structured fee draft rows to the workbook Testing Prices sheet."""
        template = Path(template_path)
        target = Path(output_path)

        def write(workbook: Any) -> None:
            sheet = _testing_prices_sheet(workbook)
            _write_structured_fee_draft(
                sheet=sheet,
                draft=draft,
                prepared_by=prepared_by,
                approved_by=approved_by,
            )

        self._generate_workbook(template, target, write)

        return FeeEvaluationWorkbookWriteResult(
            output_path=target,
            status="generated",
            warnings=(),
        )

    def generate_matrix_basic_fill(
        self,
        *,
        template_path: Path,
        output_path: Path,
        basic_fill: MatrixBasicFillWorkbook,
        review_required: bool,
        prepared_by: str | None,
        approved_by: str | None,
        edited_values: FeeEvaluationEditedExportValues | None = None,
        basic_information_values: dict[str, str] | None = None,
    ) -> FeeEvaluationWorkbookWriteResult:
        """Write Matrix basic-fill A/C rows to the Testing Prices sheet."""
        template = Path(template_path)
        target = Path(output_path)

        def write(workbook: Any) -> tuple[str, ...]:
            sheet = _testing_prices_sheet(workbook)
            anchors = FeeEvaluationAnchorSnapshot.from_sheet(sheet)
            write_basic_information_identity(
                sheet, basic_information_values, anchors=anchors
            )
            return write_matrix_basic_fill(
                sheet=sheet,
                basic_fill=basic_fill,
                edited_values=edited_values,
                anchors=anchors,
            )

        gateway_warnings = self._generate_workbook(template, target, write)

        warnings = ["Matrix basic fill only."]
        warnings.extend(gateway_warnings)
        if review_required:
            warnings.append("Pricing still requires review.")
        return FeeEvaluationWorkbookWriteResult(
            output_path=target,
            status="generated",
            warnings=tuple(warnings),
        )

    def _generate_workbook(
        self, template: Path, target: Path, write: Callable[[Any], Any],
    ) -> Any:
        if template.suffix.lower() not in _EXCEL_FILE_FORMATS:
            raise ValueError(f"Unsupported fee template type: {template}")
        if target.suffix.lower() not in _EXCEL_FILE_FORMATS:
            raise ValueError(f"Unsupported fee output type: {target}")
        if not template.is_file():
            raise FileNotFoundError(f"Template does not exist: {template}")
        if not target.parent.is_dir():
            raise FileNotFoundError(f"Output directory does not exist: {target.parent}")
        if template.resolve() == target.resolve():
            raise ValueError("Fee output path must not replace the approved template.")

        if template.suffix.lower() == ".xlsx" and target.suffix.lower() == ".xlsx":
            return self._generate_native_xlsx(template, target, write)

        # Excel SaveAs has a stricter path limit than Python filesystem operations.
        with _short_excel_output(target.suffix.lower()) as staged:
            with stage("excel_start"):
                excel, pythoncom_module = self._open_excel_application()
                try:
                    emit("office_environment", version=str(excel.Version), operating_system=str(excel.OperatingSystem))
                except Exception:
                    emit("office_environment_unavailable")
            workbook = None
            excel_state = None
            with _cleanup_on_exit((
                ("Close fee workbook", lambda: workbook.Close(SaveChanges=False)
                 if workbook is not None else None),
                ("Restore Excel settings", lambda: _restore_excel_batch(excel, excel_state)),
                ("Quit Excel", lambda: excel.Quit()),
                ("Release COM", lambda: _uninitialize_com(pythoncom_module)),
            )):
                excel.Visible = False
                excel.DisplayAlerts = False
                excel_state = _begin_excel_batch(excel)
                with stage("excel_open_template", template=template):
                    workbook = excel.Workbooks.Open(str(template))
                with stage("excel_fill_workbook"):
                    result = write(workbook)
                with stage("excel_save_workbook", output=staged):
                    _save_as(workbook, staged)
            with stage("excel_publish_staged_workbook", staging=staged, output=target):
                _publish_workbook(staged, target)
            return result

    def _generate_native_xlsx(
        self, template: Path, target: Path, write: Callable[[Any], Any],
    ) -> Any:
        from openpyxl import load_workbook

        with _native_xlsx_output() as staged:
            with stage("xlsx_open_template", template=template):
                workbook = load_workbook(template, data_only=False, keep_links=True)
            try:
                with stage("xlsx_fill_workbook"):
                    result = write(workbook)
                workbook.calculation.calcMode = "auto"
                workbook.calculation.fullCalcOnLoad = True
                workbook.calculation.forceFullCalc = True
                with stage("xlsx_save_workbook", output=staged):
                    workbook.save(staged)
            finally:
                workbook.close()
            with stage("xlsx_publish_staged_workbook", staging=staged, output=target):
                _publish_workbook(staged, target)
            return result

    def _open_excel_application(self) -> tuple[Any, Any | None]:
        if self._excel_app_factory is not None:
            return self._excel_app_factory(), None
        try:
            import pythoncom  # type: ignore[import-not-found]
            import win32com.client  # type: ignore[import-not-found]
        except ImportError as exc:  # pragma: no cover
            raise OfficeAutomationUnavailable(
                "Excel COM automation is required for fee template generation."
            ) from exc
        pythoncom.CoInitialize()
        try:
            excel = win32com.client.DispatchEx("Excel.Application")
        except Exception:
            pythoncom.CoUninitialize()
            raise
        return excel, pythoncom


def _uninitialize_com(pythoncom_module: Any | None) -> None:
    if pythoncom_module is not None:
        pythoncom_module.CoUninitialize()


def _publish_workbook(staged: Path, target: Path) -> None:
    if not staged.is_file() or staged.stat().st_size == 0:
        raise RuntimeError(
            "Generated fee workbook is missing or empty; output was not replaced."
        )
    # The sibling is private to this invocation; the old target survives failed copies.
    descriptor, name = tempfile.mkstemp(
        prefix=".clfee-", suffix=target.suffix, dir=target.parent,
    )
    sibling = Path(name)
    with _cleanup_on_exit((
        ("Remove fee publication temporary file", lambda: sibling.unlink(missing_ok=True)),
    )):
        with os.fdopen(descriptor, "wb") as destination, staged.open("rb") as source:
            shutil.copyfileobj(source, destination)
        os.replace(sibling, target)


@contextmanager
def _short_excel_output(suffix: str) -> Iterator[Path]:
    root = Path(tempfile.gettempdir())
    # Do not silently send SaveAs to a redirected network or overlong TEMP path.
    if root.anchor.startswith("\\\\") or len(str(root / "clfee-xxxxxxxx" / "fee.xlsx")) > 200:
        raise ValueError(
            "Fee generation requires a short local TEMP directory; configure Windows TEMP locally."
        )
    directory = Path(tempfile.mkdtemp(prefix="clfee-", dir=root))
    with _cleanup_on_exit((("Remove fee staging directory", lambda: shutil.rmtree(directory)),)):
        yield directory / ("fee" + suffix)


@contextmanager
def _cleanup_on_exit(actions: tuple[tuple[str, Callable[[], Any]], ...]) -> Iterator[None]:
    """Attempt every owned cleanup, preserving the first failure and its context."""
    primary = None
    try:
        yield
    except BaseException as exc:
        primary = exc
        raise
    finally:
        error = primary
        for label, action in actions:
            try:
                with stage(label.lower().replace(" ", "_")):
                    action()
            except BaseException as exc:
                if error is None:
                    error = exc
                else:
                    error.add_note(f"{label} also failed: {exc}")
        if primary is None and error is not None:
            raise error


def _testing_prices_sheet(workbook: Any) -> Any:
    if hasattr(workbook, "__getitem__") and "Testing Prices" in workbook.sheetnames:
        return OpenpyxlFeeSheetAdapter(workbook["Testing Prices"])
    try:
        return workbook.Worksheets.Item("Testing Prices")
    except Exception as exc:
        raise ValueError("Workbook sheet 'Testing Prices' was not found.") from exc


def _first_sheet(workbook: Any) -> Any:
    if hasattr(workbook, "worksheets"):
        return OpenpyxlFeeSheetAdapter(workbook.worksheets[0])
    return workbook.Worksheets.Item(1)


def _begin_excel_batch(excel: Any) -> dict[str, Any]:
    state: dict[str, Any] = {}
    for name, value in (
        ("ScreenUpdating", False),
        ("EnableEvents", False),
        ("Calculation", _EXCEL_CALCULATION_MANUAL),
    ):
        try:
            state[name] = getattr(excel, name)
            setattr(excel, name, value)
        except Exception:
            state.pop(name, None)
    return state


def _restore_excel_batch(excel: Any, state: dict[str, Any] | None) -> None:
    if not state:
        return
    for name, value in state.items():
        try:
            setattr(excel, name, value)
        except Exception:
            continue


def _write_structured_fee_draft(
    *,
    sheet: Any,
    draft: FeeEvaluationDraft,
    prepared_by: str | None,
    approved_by: str | None,
) -> None:
    sheet.Cells(1, 1).Value = "ConnLab Generated Fee Evaluation"
    sheet.Cells(2, 1).Value = f"Project ID: {draft.header.project_id}"
    sheet.Cells(3, 1).Value = (
        f"Confirmed Matrix: {draft.header.confirmed_matrix_id} / "
        f"rev {draft.header.confirmed_revision}"
    )
    sheet.Cells(4, 1).Value = (
        f"Fee rule version: {draft.header.pricing_rule_version_id}"
    )
    sheet.Cells(5, 1).Value = (
        f"Pricing effective from: {draft.header.pricing_effective_from or ''}"
    )
    sheet.Cells(6, 1).Value = f"Prepared by: {prepared_by or ''}"
    sheet.Cells(7, 1).Value = f"Approved by: {approved_by or ''}"

    headers = (
        "Group",
        "Spend time",
        "Description",
        "Unit price",
        "Units",
        "Base fee",
        "Discount",
        "Testing fee",
        "Line ID",
        "Confirmed group ID",
        "Confirmed row ID",
        "Source row ID",
        "Matched rule ID",
        "Matched rule version ID",
        "Step tokens",
    )
    for index, header in enumerate(headers, start=1):
        sheet.Cells(9, index).Value = header

    row_index = 10
    for group in draft.groups:
        for line in group.line_items:
            sheet.Cells(row_index, 1).Value = group.group_label
            sheet.Cells(row_index, 2).Value = line.spend_time
            sheet.Cells(row_index, 3).Value = line.test_item
            sheet.Cells(row_index, 4).Value = _decimal_text(line.unit_price)
            sheet.Cells(row_index, 5).Value = _decimal_text(line.units)
            sheet.Cells(row_index, 6).Value = _decimal_text(line.base_fee)
            sheet.Cells(row_index, 7).Value = _decimal_text(line.discount_percent)
            sheet.Cells(row_index, 8).Value = _decimal_text(line.testing_fee)
            sheet.Cells(row_index, 9).Value = line.line_id
            sheet.Cells(row_index, 10).Value = line.confirmed_group_id
            sheet.Cells(row_index, 11).Value = line.confirmed_row_id
            sheet.Cells(row_index, 12).Value = line.source_row_id or ""
            sheet.Cells(row_index, 13).Value = line.matched_rule_id or ""
            sheet.Cells(row_index, 14).Value = line.matched_rule_version_id or ""
            sheet.Cells(row_index, 15).Value = ", ".join(line.step_tokens)
            row_index += 1

    sheet.Cells(row_index, 7).Value = "Total"
    sheet.Cells(row_index, 8).Value = _decimal_text(draft.total_fee)


def _save_as(workbook: Any, target: Path) -> None:
    file_format = _EXCEL_FILE_FORMATS.get(target.suffix.lower())
    if file_format is None:
        raise ValueError(f"Unsupported fee output type: {target}")
    workbook.SaveAs(str(target), FileFormat=file_format)


@contextmanager
def _native_xlsx_output() -> Iterator[Path]:
    directory = Path(tempfile.mkdtemp(prefix="clfee-xlsx-"))
    with _cleanup_on_exit((("Remove fee XLSX staging directory", lambda: shutil.rmtree(directory)),)):
        yield directory / "fee.xlsx"


def _decimal_text(value: Decimal | None) -> str:
    if value is None:
        return ""
    return format(value, "f")
