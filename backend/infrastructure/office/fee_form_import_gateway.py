"""Read ConnLab Fee Form data without starting Office or evaluating workbook code."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
import xlrd

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
UNIT_TYPES = {f"per {unit}" for unit in ("sample", "reading", "contact", "cycle", "time", "hour", "day", "photo", "report")}


def read_fee_form(data: bytes, filename: str) -> dict:
    """Return editable values only; reject unknown layouts and ambiguous worksheets."""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xls", ".xlsx"}:
        raise ValueError("Choose a ConnLab Fee Form .xls or .xlsx file; macro files are not supported.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Fee Form file exceeds the 10 MB limit.")
    try:
        sheets = _read_sheets(data, suffix)
        candidates = []
        for cells, notes in sheets:
            if len(cells) < 5:
                continue
            headers = ["".join(_text(value).casefold().split()) for value in cells[3]]
            allowed = ({"group"}, {"man-hour", "spendtime", "time(unit:hour)"}, {"description"},
                       {"unitprice", "unitprice(a)"}, {"", "unittype"}, {"units", "units(b)"},
                       {"basefee", "basefee(c)"}, {"discount", "pricepercentoff(d)"},
                       {"testingfee", "testingfee(rmb)=a*b*(1-d)+c"})
            if all(value in choices for value, choices in zip(headers, allowed)):
                candidates.append((cells, notes))
        if len(candidates) != 1:
            raise ValueError("Expected exactly one ConnLab Fee Form sheet with the exported columns in row 4.")
        cells, notes = candidates[0]
        result = []
        group = ""
        seen_groups = set()
        for index, row in enumerate(cells[4:], start=5):
            description = _text(row[2])
            first = _text(row[0])
            if first.casefold() in {"total", "条件确认", "condition confirmation"}:
                break
            if not description:
                continue
            if description.casefold() == "sample preparation":
                group = first.removeprefix("Group ").strip()
                if not group or group.casefold() in seen_groups:
                    raise ValueError(f"Fee Form row {index}: missing or duplicate Group heading.")
                seen_groups.add(group.casefold())
                kind = "sample_preparation"
            elif description.casefold() == "report preparation":
                kind = "manual_trailing"
            else:
                kind = "matrix_step"
                if not group:
                    raise ValueError(f"Fee Form row {index}: no Sample preparation Group heading.")
            values = {key: _number(row[column], index, key, percent=key == "discount") for key, column in
                      (("spendTime", 1), ("unitPrice", 3), ("units", 5), ("baseFee", 6), ("discount", 7))}
            unit = _text(row[4]).casefold()
            if unit and unit not in UNIT_TYPES:
                raise ValueError(f"Fee Form row {index}: unsupported Unit Type '{unit}'.")
            values.update(unitType=unit, notes=notes.get(index, ""))
            result.append({"group": "" if kind == "manual_trailing" else group,
                           "description": description, "rowKind": kind, "values": values})
        if not result or not seen_groups:
            raise ValueError("Fee Form contains no grouped fee rows.")
        return {"rows": result}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unable to read Fee Form. Use an unencrypted ConnLab .xls or .xlsx export.") from exc


def _read_sheets(data, suffix):
    if suffix == ".xlsx":
        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                    raise ValueError("Fee Form expanded workbook is too large.")
                if any("vbaproject" in item.filename.casefold() for item in archive.infolist()):
                    raise ValueError("Macro-enabled workbooks are not supported.")
        except BadZipFile as exc:
            raise ValueError("Invalid .xlsx Fee Form file.") from exc
        book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
        try:
            if len(book.worksheets) > 20 or any(sheet.max_row > MAX_ROWS for sheet in book.worksheets):
                raise ValueError("Fee Form has too many sheets or rows.")
            return [(list(sheet.iter_rows(max_col=9, values_only=True)),
                     {row: sheet.cell(row, 9).comment.text for row in range(5, sheet.max_row + 1)
                      if sheet.cell(row, 9).comment}) for sheet in book.worksheets]
        finally:
            book.close()
    book = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
    try:
        if book.nsheets > 20:
            raise ValueError("Fee Form has too many sheets.")
        result = []
        for sheet in book.sheets():
            if sheet.nrows > MAX_ROWS:
                raise ValueError("Fee Form has too many rows.")
            rows = [(sheet.row_values(row, 0, 9) + [None] * 9)[:9] for row in range(sheet.nrows)]
            result.append((rows, {row + 1: note.text for (row, col), note in sheet.cell_note_map.items() if col == 8}))
        return result
    finally:
        book.release_resources()


def _text(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return format(Decimal(str(value)).normalize(), "f")
    return "" if value is None else str(value).strip()


def _number(value, row, field, *, percent=False):
    text = _text(value)
    if not text:
        return ""
    try:
        if len(text) > 128:
            raise InvalidOperation
        number = Decimal(text.rstrip("%"))
        if not number.is_finite() or abs(number.as_tuple().exponent) > 100:
            raise InvalidOperation
        if percent and not text.endswith("%"):
            number *= 100
        if not number.is_finite() or number < 0 or number > Decimal("1e15") or (percent and number > 100):
            raise InvalidOperation
        return format(number.normalize(), "f")
    except InvalidOperation as exc:
        raise ValueError(f"Fee Form row {row}: {field} must be a non-negative number (discount 0–100%). Replace input formulas with values before importing.") from exc
