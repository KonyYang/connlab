"""Report-ready temperature-rise tables and editable native Excel chart labels."""

import re
from math import ceil

from openpyxl.chart import Reference, Series
from openpyxl.chart.data_source import NumData, NumVal, StrRef, StrData, StrVal
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.text import Text, RichText
from openpyxl.chart.trendline import Trendline, TrendlineLabel
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName

from backend.domain.temperature_rise import calculate_current
from backend.domain.temperature_data import UNPOWERED_CURRENT_THRESHOLD, measurement_name
from backend.application.temperature_report_metadata import TemperatureReportMetadata

MAX_COLOR = 'FFA500'
AVG_COLOR = '00008B'


def build_rise_report(sheet, *, table, selection, analysis, maximum, average, target, caches, chart,
                      report_metadata=TemperatureReportMetadata(), auxiliary_current_columns=()):
    stages = [point for point in analysis.points if point.source_row is not None]
    chart_column = max(len(analysis.points) + 3, 11)
    info_end = _test_info(sheet, report_metadata)
    summary_header = info_end + 2
    current_row = summary_header + len(auxiliary_current_columns)
    samples = len(analysis.points[0].sample_maxima)
    calculator_header = current_row + len(selection.temperature_columns) + samples + 4
    print_end = max(calculator_header + 1, current_row + 36)
    # Write supporting data at its final location; inserting rows later would leave
    # formula/chart references stale. Keep the printable report ahead of the detail.
    raw_header = max(print_end, calculator_header + 8) + 3
    raw_rows, raw_columns = _key_readings(sheet, table, selection, stages, raw_header)
    rise_header = raw_header + len(stages) + 5
    rise_rows = _stage_rises(sheet, selection, stages, raw_rows, raw_columns, rise_header, caches)
    maximum_row, average_row = _summary(sheet, table, selection, analysis, rise_rows, current_row, caches)
    _auxiliary_currents(sheet, table, selection, analysis, auxiliary_current_columns,
                       raw_rows, raw_columns, summary_header, caches)
    _calculator(sheet, maximum, analysis.maximum_fit, target, calculator_header, caches)
    _average_inputs(sheet, average, analysis.average_fit, calculator_header + 3, calculator_header + 8, caches)
    label_rows = _fit_equations(sheet, analysis, current_row, maximum_row, average_row,
                               calculator_header + 6, caches)
    # Keep live fit/AVG dependencies inspectable without duplicating the visible MAX calculator.
    sheet.row_dimensions.group(calculator_header + 3, calculator_header + 8, hidden=True)
    _rise_chart(sheet, analysis, chart, current_row, maximum_row, average_row, label_rows, caches)
    _name(sheet, 'TemperatureRiseSummary', summary_header, average_row, len(analysis.points) + 1)
    sheet.freeze_panes = None
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 1
    sheet.print_area = f'A1:{get_column_letter(chart_column + 12)}{print_end}'
    for column in range(1, max(table.width + 2, chart_column + 13)):
        sheet.column_dimensions[get_column_letter(column)].width = 10
    sheet.column_dimensions['A'].width = 22
    for row in range(summary_header, average_row + 1):
        label = str(sheet.cell(row, 1).value or '')
        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 15, ceil(len(label) / 20) * 15)


def _test_info(sheet, info):
    rows = [('LTR', info.ltr_number), ('Test Date', info.test_date), ('Tested By', info.tested_by)]
    environment = []
    if info.ambient_temperature.strip():
        environment.append(f'{info.ambient_temperature.strip()} °C')
    if info.relative_humidity is not None:
        environment.append(f'{info.relative_humidity:g} %RH')
    rows.append(('Amb Temp. and Rel. Hum.', ' / '.join(environment)))
    rows.append(('Test Equipment ID', info.equipment_id))
    if info.test_condition.strip():
        rows.append(('Test Condition', info.test_condition))
    rows.append(('Remark', info.remark))
    edge = Side(style='thin', color='D8D8B0')
    for row, (label, value) in enumerate(rows, 1):
        for start, end, text in ((1, 2, label), (3, 9, value)):
            for column in range(start, end + 1):
                cell = sheet.cell(row, column)
                cell.fill = PatternFill('solid', fgColor='FFFFCC')
                cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
            _literal(sheet, row, start, text)
            cell = sheet.cell(row, start)
            cell.font = Font(name='Arial', size=10, bold=start == 1, color='172033')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.row_dimensions[row].height = max(22, ceil(len(value) / 60) * 15 + 6)
    end = len(rows)
    _name(sheet, 'TemperatureTestInfo', 1, end, 9)
    return end


def _literal(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = 's'  # Scanner labels/values never become executable Excel formulas.


def _formula(sheet, row, column, formula, value, caches):
    cell = sheet.cell(row, column, formula)
    caches[cell.coordinate] = value


def _name(sheet, name, start, end, width):
    reference = f'{quote_sheetname(sheet.title)}!$A${start}:${get_column_letter(width)}${end}'
    sheet.parent.defined_names.add(DefinedName(name, attr_text=reference))


def _table_style(sheet, start, end, width):
    edge = Side(style='thin', color='A6A6A6')
    for row in sheet.iter_rows(min_row=start, max_row=end, max_col=width):
        for cell in row:
            cell.font = Font(name='Arial', size=10, color='000000')
            cell.alignment = Alignment(horizontal='center',
                                       vertical='center', wrap_text=cell.row == start or cell.column == 1)
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            cell.number_format = '0.0' if cell.column != 1 else '0'
        sheet.row_dimensions[row[0].row].height = 15
    for cell in sheet[start][:width]:
        cell.fill = PatternFill('solid', fgColor='D9D9D9')
        cell.font = Font(name='Arial', size=10, bold=True, color='000000')
    sheet.row_dimensions[start].height = 36


def _reading_column_order(table, selection):
    headers = table.rows[selection.header_row - 1]
    measured = {*selection.temperature_columns, selection.ambient_column, selection.current_column}
    identifiers = []
    for pattern in (r'\bscan\b|扫描|序号', r'\b(?:time|timestamp|date)\b|时间|日期'):
        matches = [column for column, label in enumerate(headers, 1)
                   if column not in measured and re.search(pattern, str(label).casefold())]
        identifiers.append(matches[0] if len(matches) == 1 else None)
    if identifiers[0] == identifiers[1]:
        identifiers = [None, None]
    # Missing/ambiguous metadata stays blank; all original source columns remain available.
    return (*identifiers, *(column for column in range(1, table.width + 1) if column not in identifiers))


def _key_readings(sheet, table, selection, stages, header):
    order = _reading_column_order(table, selection)
    source_columns = {source: column for column, source in enumerate(order, 1) if source is not None}
    headers = table.rows[selection.header_row - 1]
    for column, source in enumerate(order, 1):
        label = ('扫描', '时间')[column - 1] if column <= 2 else (headers[source - 1] if source <= len(headers) else None)
        _literal(sheet, header, column, label)
    locations = {}
    for row, point in enumerate(stages, header + 1):
        locations[point.source_row] = row
        values = table.rows[point.source_row - 1]
        for source, column in source_columns.items():
            _literal(sheet, row, column, values[source - 1] if source <= len(values) else None)
    _table_style(sheet, header, header + len(stages), len(order))
    measured_columns = {*selection.temperature_columns, selection.ambient_column, selection.current_column}
    for row in range(header + 1, header + len(stages) + 1):
        for column, source in enumerate(order, 1):
            sheet.cell(row, column).number_format = '0.000' if source in measured_columns else 'General'
    _name(sheet, 'KeyStageReadings', header, header + len(stages), len(order))
    return locations, source_columns


def _stage_rises(sheet, selection, stages, raw_rows, raw_columns, header, caches):
    channels = len(selection.temperature_columns)
    ambient_column = get_column_letter(raw_columns[selection.ambient_column])
    for column, label in enumerate(('扫描', '时间'), 1):
        _literal(sheet, header, column, label)
    for slot in range(channels):
        _literal(sheet, header, slot + 3,
                 f'{slot // selection.thermocouples_per_sample + 1}#- T{slot % selection.thermocouples_per_sample + 1} (°C)')
    _literal(sheet, header, channels + 3, 'Ambient (°C)')
    _literal(sheet, header, channels + 4, 'Current (A)')
    locations = {}
    for row, point in enumerate(stages, header + 1):
        locations[point.source_row] = row
        source = raw_rows[point.source_row]
        for column in (1, 2):
            identifier = sheet.cell(source, column)
            if isinstance(identifier.value, bool):
                _literal(sheet, row, column, identifier.value)
            elif identifier.value is not None:
                _formula(sheet, row, column, f'={identifier.coordinate}', identifier.value, caches)
        for slot, column in enumerate(selection.temperature_columns):
            _formula(sheet, row, slot + 3,
                     f'={get_column_letter(raw_columns[column])}{source}-${ambient_column}{source}', point.rises[slot], caches)
        # Lookup ambient through the selected role, never through physical last-column assumptions.
        ambient = sheet[f'{ambient_column}{source}'].value
        _formula(sheet, row, channels + 3, f'={ambient_column}{source}', ambient, caches)
        current = f'{get_column_letter(raw_columns[selection.current_column])}{source}*{selection.current_multiplier!r}'
        _formula(sheet, row, channels + 4,
                 f'=IF(ABS({current})<{UNPOWERED_CURRENT_THRESHOLD},0,{current})',
                 point.current, caches)
    _table_style(sheet, header, header + len(stages), channels + 4)
    for row in range(header + 1, header + len(stages) + 1):
        for column in (1, 2):
            sheet.cell(row, column).number_format = 'General'
    _name(sheet, 'StageTemperatureRise', header, header + len(stages), channels + 4)
    return locations


def _summary(sheet, table, selection, analysis, rise_rows, header, caches):
    channels = len(selection.temperature_columns)
    samples, count = len(analysis.points[0].sample_maxima), selection.thermocouples_per_sample
    _literal(sheet, header, 1, _current_name(table, selection, selection.current_column))
    for slot, column in enumerate(selection.temperature_columns):
        row = header + 1 + slot + slot // count
        name = _report_measurement_name(measurement_name(table, selection, column)) or f'T{slot % count + 1}'
        _literal(sheet, row, 1, name)
    max_rows = [header + (sample + 1) * (count + 1) for sample in range(samples)]
    for sample, row in enumerate(max_rows, 1):
        _literal(sheet, row, 1, f'{sample}# Single Max')
    maximum_row, average_row = header + channels + samples + 1, header + channels + samples + 2
    _literal(sheet, maximum_row, 1, 'Max T-Rise')
    _literal(sheet, average_row, 1, 'Avg of max T-Rise on each sample')
    for column, point in enumerate(analysis.points, 2):
        letter = get_column_letter(column)
        if point.source_row is None:
            _literal(sheet, header, column, 0)
        else:
            source = rise_rows[point.source_row]
            _formula(sheet, header, column, f'={get_column_letter(channels + 4)}{source}', point.current, caches)
        for slot, value in enumerate(point.rises):
            row = header + 1 + slot + slot // count
            if point.source_row is None:
                _literal(sheet, row, column, 0)
            else:
                _formula(sheet, row, column, f'={get_column_letter(slot + 3)}{source}', value, caches)
        for sample, row in enumerate(max_rows):
            _formula(sheet, row, column, f'=MAX({letter}{row-count}:{letter}{row-1})', point.sample_maxima[sample], caches)
        references = ','.join(f'{letter}{row}' for row in max_rows)
        _formula(sheet, maximum_row, column, f'=MAX({references})', point.maximum, caches)
        _formula(sheet, average_row, column, f'=AVERAGE({references})', point.average, caches)
    _table_style(sheet, header, average_row, len(analysis.points) + 1)
    for cell in sheet[header][1:len(analysis.points) + 1]:
        cell.number_format = '0'  # Display whole amperes without rounding chart/fit inputs.
    sheet.row_dimensions[header].height = 22
    for row in (*max_rows, maximum_row, average_row):
        for cell in sheet[row][:len(analysis.points) + 1]:
            cell.fill = PatternFill('solid', fgColor='D2B48C' if row in max_rows else 'ADD8E6')
            if cell.column == 1:
                cell.font = Font(name='Arial', size=10, bold=True, color='000000')
    sheet.row_dimensions[average_row].height = 32
    return maximum_row, average_row


def _current_name(table, selection, column):
    name = _report_measurement_name(measurement_name(table, selection, column))
    if not name:
        name = 'Current' if column == selection.current_column else f'Current {column}'
    return f'{name}(A)'


def _auxiliary_currents(sheet, table, selection, analysis, columns, raw_rows, raw_columns, header, caches):
    if not columns:
        return
    width = len(analysis.points) + 1
    for row, source_column in enumerate(columns, header):
        _literal(sheet, row, 1, _current_name(table, selection, source_column))
        for column, point in enumerate(analysis.points, 2):
            if point.source_row is None:
                continue  # An inserted origin has no auxiliary reading to report.
            source = sheet.cell(raw_rows[point.source_row], raw_columns[source_column])
            # CSV values stay literal in the raw table; coerce numeric text only in the report.
            _formula(sheet, row, column, f'={source.coordinate}*1', float(source.value), caches)
    _table_style(sheet, header, header + len(columns) - 1, width)
    for row in range(header, header + len(columns)):
        sheet.row_dimensions[row].height = 22
        for cell in sheet[row][:width]:
            cell.fill = PatternFill('solid', fgColor='D9D9D9')
            cell.font = Font(name='Arial', size=10, bold=True, color='000000')
            cell.number_format = '0'


def _report_measurement_name(name):
    """Remove numeric instrument/sample prefixes without expanding measurement names."""
    clean = name.strip()
    if clean and clean[0].isdigit():
        first_letter = re.search(r'[^\W\d_]', clean)
        clean = clean[first_letter.start():] if first_letter else ''
    return clean


def _calculator(sheet, maximum, fit, target, header, caches):
    for column, label in enumerate(('y (°C)', 'a', 'b', 'c', 'x (A)'), 1):
        _literal(sheet, header, column, label)
    row = header + 1
    for column, value in enumerate((target, maximum.a, maximum.b, maximum.c), 1):
        _literal(sheet, row, column, value)
        if column > 1 and maximum == fit.coefficients:
            # Default coefficients follow the fit; explicit UI overrides remain operator-owned values.
            _formula(sheet, row, column, f'=ROUND({get_column_letter(column)}{header+7},6)', value, caches)
    formula = f'=IF(A{row}="","",(-C{row}+SQRT(C{row}^2-4*B{row}*(D{row}-A{row})))/(2*B{row}))'
    _formula(sheet, row, 5, formula, calculate_current(maximum, target) if target is not None else '', caches)
    _table_style(sheet, header, row, 5)
    sheet.row_dimensions[header].height = 18
    for column in range(1, 5):
        cell = sheet.cell(row, column)
        cell.font = Font(name='Arial', size=10, color='000000' if cell.data_type == 'f' else '0000FF')
        cell.fill = PatternFill('solid', fgColor='E2F0D9')
        cell.number_format = '0.000000' if column != 1 else '0.0'
    sheet.cell(row, 5).number_format = '0.0'
    sheet.cell(row, 5).font = Font(name='Arial', size=10, bold=True, color=MAX_COLOR)
    _name(sheet, 'CurrentCalculator', header, row, 5)


def _equation_text(name, fit):
    coefficients = fit.fitted_coefficients
    text = f'{name}: y = {coefficients.a:.6f}x²'
    for value, suffix in ((coefficients.b, 'x'), (coefficients.c, '')):
        if suffix or value:
            text += f' {"+" if value >= 0 else "-"} {abs(value):.6f}{suffix}'
    return f'{text}\nR² = {fit.r_squared:.6f}'


def _average_inputs(sheet, coefficients, fit, header, fit_row, caches):
    for column, label in enumerate(('Avg of Max', 'a', 'b', 'c'), 1):
        _literal(sheet, header, column, label)
    _literal(sheet, header + 1, 1, 'Effective Coefficients')
    for column, value in enumerate((coefficients.a, coefficients.b, coefficients.c), 2):
        _literal(sheet, header + 1, column, value)
        if coefficients == fit.coefficients:
            _formula(sheet, header + 1, column, f'=ROUND({get_column_letter(column)}{fit_row},6)', value, caches)
    _table_style(sheet, header, header + 1, 4)
    sheet.row_dimensions[header].height = 18
    for cell in sheet[header + 1][1:4]:
        cell.number_format = '0.000000'
    _name(sheet, 'AverageCoefficients', header, header + 1, 4)


def _fit_equations(sheet, analysis, current_row, maximum_row, average_row, header, caches):
    """LINEST/linked labels avoid stale hardcoded equations after Excel cell edits."""
    end = get_column_letter(len(analysis.points) + 1)
    x_range = f'$B${current_row}:${end}${current_row}'
    for column, label in enumerate(('Chart Fit', 'a', 'b', 'c', 'R²', 'Equation'), 1):
        _literal(sheet, header, column, label)
    rows = []
    for row, (source, fit) in enumerate(((maximum_row, analysis.maximum_fit),
                                        (average_row, analysis.average_fit)), header + 1):
        rows.append(row)
        name = sheet.cell(source, 1).value
        _formula(sheet, row, 1, f'=A{source}', name, caches)
        y_range = f'$B${source}:${end}${source}'
        intercept = 'FALSE' if analysis.zero_intercept else 'TRUE'
        regression = f'LINEST({y_range},{x_range}^{{1;2}},{intercept})'
        for column, value in enumerate((fit.fitted_coefficients.a, fit.fitted_coefficients.b, fit.fitted_coefficients.c), 2):
            formula = '=0' if column == 4 and analysis.zero_intercept else f'=INDEX({regression},1,{column-1})'
            _formula(sheet, row, column, formula, value, caches)
            sheet.cell(row, column).number_format = '0.000000'
        predicted = f'{x_range}^2*B{row}+{x_range}*C{row}+D{row}'
        _formula(sheet, row, 5,
                 f'=IFERROR(CORREL({y_range},{predicted})^2,IF(SUMXMY2({y_range},{predicted})<1E-20,1,0))',
                 fit.r_squared, caches)
        formula = (f'=A{row}&": y = "&TEXT(B{row},"0.000000")&"x²"&IF(C{row}<0," - "," + ")'
                   f'&TEXT(ABS(C{row}),"0.000000")&"x"&IF(D{row}=0,"",IF(D{row}<0," - "," + ")'
                   f'&TEXT(ABS(D{row}),"0.000000"))&CHAR(10)&"R² = "&TEXT(E{row},"0.000000")')
        _formula(sheet, row, 6, formula, _equation_text(name, fit), caches)
    _table_style(sheet, header, header + 2, 6)
    for row in rows:
        for column in range(2, 6):
            sheet.cell(row, column).number_format = '0.000000'
        sheet.merge_cells(start_row=row, start_column=6, end_row=row, end_column=14)
        sheet.cell(row, 6).alignment = Alignment(wrap_text=True, vertical='center')
        sheet.cell(row, 6).font = Font(name='Arial', size=10, bold=True, color=MAX_COLOR if row == rows[0] else AVG_COLOR)
        sheet.row_dimensions[row].height = 32
    return rows


def _rise_chart(sheet, analysis, chart, current_row, maximum_row, average_row, label_rows, caches):
    end = len(analysis.points) + 1
    for index, (row, name, color, marker) in enumerate(((maximum_row, 'Max', MAX_COLOR, 'diamond'),
                                                       (average_row, 'Avg of Max', AVG_COLOR, 'square'))):
        x = Reference(sheet, min_col=2, max_col=end, min_row=current_row)
        y = Reference(sheet, min_col=2, max_col=end, min_row=row)
        series = Series(y, x)
        series.tx = SeriesLabel(strRef=StrRef(
            f=f'{quote_sheetname(sheet.title)}!$A${row}',
            strCache=StrData(ptCount=1, pt=[StrVal(idx=0, v=sheet.cell(row, 1).value)])))
        for source, reference in ((current_row, series.xVal.numRef), (row, series.yVal.numRef)):
            values = [caches.get(sheet.cell(source, col).coordinate, sheet.cell(source, col).value) for col in range(2, end + 1)]
            reference.numCache = NumData(formatCode='0.0', ptCount=len(values), pt=[NumVal(idx=i, v=v) for i, v in enumerate(values)])
        series.graphicalProperties.line.solidFill = None
        series.graphicalProperties.line.noFill = True
        series.marker.symbol, series.marker.size = marker, 6
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color
        properties = CharacterProperties(b=True, sz=1100, solidFill=color)
        label_coordinate = f'F{label_rows[index]}'
        label = TrendlineLabel(tx=Text(strRef=StrRef(
            f=f'{quote_sheetname(sheet.title)}!${label_coordinate[0]}${label_rows[index]}',
            strCache=StrData(ptCount=1, pt=[StrVal(idx=0, v=caches[label_coordinate])]))),
            txPr=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=properties), endParaRPr=properties)]),
            layout=Layout(manualLayout=ManualLayout(x=.15,
                y=.14 if index == 0 else .27, xMode='edge', yMode='edge')))
        series.trendline = Trendline(trendlineType='poly', order=2, name=f'Fit ({name})',
            intercept=0 if analysis.zero_intercept else None, dispEq=True, dispRSqr=True, trendlineLbl=label)
        series.trendline.spPr = GraphicalProperties()
        series.trendline.spPr.line.solidFill = color
        chart.series.append(series)
    chart.width, chart.height = 23.5, 14
    chart.roundedCorners = True
    chart.layout = Layout(manualLayout=ManualLayout(layoutTarget='inner', xMode='edge', yMode='edge',
        x=.1, y=.09, w=.87, h=.76))
    chart.graphical_properties = GraphicalProperties(solidFill='FFFFE0')
    chart.plot_area.spPr = GraphicalProperties(solidFill='FFFFFF')
    sheet.add_chart(chart, f'{get_column_letter(max(end + 2, 11))}{current_row}')
