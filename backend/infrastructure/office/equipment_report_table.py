"""Locate and rebuild Section 7 without accepting revisions elsewhere in a report."""

from copy import deepcopy
import re
import unicodedata

from docx.oxml.ns import nsmap, qn
from docx.table import Table, _Cell


HEADERS = ("Item", "Manufacturer", "ID Number", "Last Cal.", "Cal. Due")
_WORD_NAMESPACE = "{" + nsmap["w"] + "}"
_DELETED = {qn("w:del"), qn("w:moveFrom")}
_INSERTED = {qn("w:ins"), qn("w:moveTo")}
_CONTAINERS = _INSERTED | {qn("w:sdt"), qn("w:sdtContent"), qn("w:customXml")}


def _canonical(text: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip()


def _current_text(element) -> str:
    if element.tag in _DELETED:
        return ""
    if element.tag == qn("w:t"):
        return element.text or ""
    if element.tag in {qn("w:tab"), qn("w:br"), qn("w:cr")}:
        return " "
    return "".join(_current_text(child) for child in element)


def _current_blocks(element):
    for child in element:
        if child.tag in {qn("w:p"), qn("w:tbl")}:
            yield child
        elif child.tag in _CONTAINERS:
            yield from _current_blocks(child)


def _equipment_heading(text: str) -> bool:
    return bool(re.fullmatch(r"7\s*[.、:)]?\s*EQUIPMENTS?\s*[.:]?", _canonical(text), re.I))


def _chapter_heading(text: str) -> bool:
    return bool(re.match(r"\d+\s*(?:[.、:)]\s*|\s+)[A-Za-z]", _canonical(text)))


def _resolve_table_revisions(table) -> None:
    # Interpret current values only within this owned table. Never change settings.xml
    # or run a document-wide AcceptAllRevisions operation.
    for row in list(table.iter(qn("w:tr"))):
        if row.find("./" + qn("w:trPr") + "/" + qn("w:del")) is not None:
            row.getparent().remove(row)
    for cell in list(table.iter(qn("w:tc"))):
        if cell.find("./" + qn("w:tcPr") + "/" + qn("w:cellDel")) is not None:
            cell.getparent().remove(cell)
    for element in list(table.iter())[::-1]:
        parent = element.getparent()
        if parent is None:
            continue
        if (element.tag in _DELETED
                or element.tag in {qn("w:delText"), qn("w:delInstrText"), qn("w:cellIns")}
                or (element.tag.startswith(_WORD_NAMESPACE)
                    and element.tag.endswith("Change"))):
            parent.remove(element)
        elif element.tag in _INSERTED:
            index = parent.index(element)
            for child in list(element):
                parent.insert(index, child)
                index += 1
            parent.remove(element)


def equipment_row_cells(row) -> tuple[_Cell, ...]:
    """Return visible fields once, instead of repeating horizontally merged cells."""
    return tuple(_Cell(cell, row._parent) for cell in row._tr.tc_lst)


def _usable_row(row, grid_width: int) -> bool:
    cells = equipment_row_cells(row)
    return (len(cells) == len(HEADERS)
            and sum(cell._tc.grid_span for cell in cells) == grid_width
            and all(cell._tc.vMerge is None for cell in cells))


def prepare_equipment_table(document) -> Table:
    """Find one current equipment table in Section 7 and resolve its local revisions."""
    blocks = list(_current_blocks(document.element.body))
    headings = [i for i, block in enumerate(blocks)
                if block.tag == qn("w:p") and _equipment_heading(_current_text(block))]
    if len(headings) != 1:
        raise ValueError(
            "Equipment List requires one identifiable Section 7 EQUIPMENTS heading. "
            "Select a report containing that section and its Equipment List table. "
            "Customer reports without an Equipment List cannot be updated. "
            "Required columns: Item, Manufacturer, ID Number, Last Cal., Cal. Due."
        )
    candidates = []
    for block in blocks[headings[0] + 1:]:
        if block.tag == qn("w:p"):
            text = _canonical(_current_text(block))
            if _chapter_heading(text) or text.upper().startswith(("APPENDIX ", "REVISION RECORD")):
                break
        if block.tag != qn("w:tbl"):
            continue
        clean = deepcopy(block)
        _resolve_table_revisions(clean)
        table = Table(clean, document._body)
        if not table.rows:
            continue
        header = tuple(_canonical(cell.text).casefold() for cell in equipment_row_cells(table.rows[0]))
        if header == tuple(value.casefold() for value in HEADERS):
            candidates.append((block, table))
    if len(candidates) != 1:
        raise ValueError(
            "Section 7 Equipment List must contain exactly one compatible table with "
            "Item, Manufacturer, ID Number, Last Cal., and Cal. Due columns."
        )
    original, table = candidates[0]
    if table._tbl.tblGrid is None or not _usable_row(table.rows[0], len(table._tbl.tblGrid.gridCol_lst)):
        raise ValueError("Section 7 Equipment List header has unsupported merged cells.")
    original.getparent().replace(original, table._tbl)
    return table


def _write_field(cell: _Cell, value: str) -> None:
    paragraph = deepcopy(cell.paragraphs[0]._p) if cell.paragraphs else None
    run_properties = next((deepcopy(run._r.rPr) for p in cell.paragraphs
                           for run in p.runs if run._r.rPr is not None), None)
    for child in list(cell._tc):
        if child.tag != qn("w:tcPr"):
            cell._tc.remove(child)
    target = cell.add_paragraph()
    if paragraph is not None and paragraph.pPr is not None:
        target._p.insert(0, deepcopy(paragraph.pPr))
    run = target.add_run(value)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def replace_equipment_rows(table: Table, values: tuple[tuple[str, ...], ...]) -> None:
    """Remove all old records, retaining only reusable row/cell formatting."""
    grid_width = len(table._tbl.tblGrid.gridCol_lst)
    templates = [deepcopy(row._tr) for row in table.rows[1:] if _usable_row(row, grid_width)]
    fallback = templates[0] if templates else deepcopy(table.rows[0]._tr)
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for index, fields in enumerate(values):
        template = templates[index] if index < len(templates) else fallback
        row = deepcopy(template)
        for marker in row.xpath("./w:trPr/w:tblHeader"):
            marker.getparent().remove(marker)
        table._tbl.append(row)
        for cell, value in zip(equipment_row_cells(table.rows[-1]), fields, strict=True):
            _write_field(cell, value)
