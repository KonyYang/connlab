"""Minimal worksheet adapter for native XLSX Fee Form generation."""

from __future__ import annotations

from copy import copy
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import Border, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries


_NUMERIC_COLUMNS = frozenset({2, 4, 6, 7, 8})


class OpenpyxlFeeSheetAdapter:
    """Expose the small worksheet seam used by the existing Fee Form writers."""

    def __init__(self, worksheet: Any) -> None:
        self.worksheet = worksheet

    def read_anchor_region(
        self, max_row: int, max_column: int,
    ) -> tuple[tuple[tuple[Any, ...], ...], tuple[tuple[Any, ...], ...]]:
        values: list[tuple[Any, ...]] = []
        formulas: list[tuple[Any, ...]] = []
        for row in range(1, max_row + 1):
            row_values: list[Any] = []
            row_formulas: list[Any] = []
            for column in range(1, max_column + 1):
                value = self.worksheet.cell(row, column).value
                row_values.append(value)
                row_formulas.append(value)
            values.append(tuple(row_values))
            formulas.append(tuple(row_formulas))
        return tuple(values), tuple(formulas)

    def Cells(self, row: int, column: int) -> "_CellAdapter":  # noqa: N802
        return _CellAdapter(self, row, column)

    def insert_rows(self, row: int, count: int) -> None:
        if count <= 0:
            return
        worksheet = self.worksheet
        merged_ranges = tuple(worksheet.merged_cells.ranges)
        for merged in merged_ranges:
            worksheet.unmerge_cells(str(merged))
        formulas = {
            (cell.row, cell.column): cell.value
            for sheet_row in worksheet.iter_rows(min_row=row)
            for cell in sheet_row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
        source_styles = tuple(
            (
                copy(worksheet.cell(row - 1, column)._style),
                copy(worksheet.cell(row - 1, column).number_format),
            )
            for column in range(1, worksheet.max_column + 1)
        )
        source_height = worksheet.row_dimensions[row - 1].height
        moved_dimensions = {
            index: copy(dimension)
            for index, dimension in worksheet.row_dimensions.items()
            if index >= row
        }
        for index in moved_dimensions:
            del worksheet.row_dimensions[index]

        worksheet.insert_rows(row, count)

        for old_row, dimension in moved_dimensions.items():
            new_row = old_row + count
            dimension.index = new_row
            worksheet.row_dimensions[new_row] = dimension
        for inserted_row in range(row, row + count):
            worksheet.row_dimensions[inserted_row].height = source_height
            for column, (style, number_format) in enumerate(source_styles, start=1):
                cell = worksheet.cell(inserted_row, column)
                cell._style = copy(style)
                cell.number_format = number_format

        for merged in merged_ranges:
            min_col, min_row, max_col, max_row = merged.bounds
            if min_row >= row:
                min_row += count
                max_row += count
            elif max_row >= row:
                max_row += count
            worksheet.merge_cells(
                start_row=min_row,
                start_column=min_col,
                end_row=max_row,
                end_column=max_col,
            )

        for (old_row, column), formula in formulas.items():
            new_row = old_row + count
            origin = f"{get_column_letter(column)}{old_row}"
            destination = f"{get_column_letter(column)}{new_row}"
            try:
                translated = Translator(formula, origin=origin).translate_formula(destination)
            except Exception:
                translated = formula
            worksheet.cell(new_row, column).value = translated

        self._shift_print_area(row, count)
        self._shift_image_anchors(row, count)

    def set_row_values(
        self, row: int, start_column: int, values: tuple[Any | None, ...],
    ) -> None:
        for offset, value in enumerate(values):
            self._set_value(row, start_column + offset, value)

    def set_block_values(
        self,
        start_row: int,
        start_column: int,
        rows: tuple[tuple[Any | None, ...], ...],
    ) -> None:
        for row_offset, values in enumerate(rows):
            self.set_row_values(start_row + row_offset, start_column, values)

    def set_formula_block(
        self, start_row: int, column: int, formulas: tuple[str, ...],
    ) -> None:
        for offset, formula in enumerate(formulas):
            self.worksheet.cell(start_row + offset, column).value = formula

    def set_a_column_fill(self, start_row: int, end_row: int, color: int) -> None:
        for row in range(start_row, end_row + 1):
            self.set_cell_fill(row, 1, color)

    def cell_fill_color(self, row: int, column: int) -> int:
        color = self.worksheet.cell(row, column).fill.fgColor
        rgb = color.rgb if color.type == "rgb" else None
        if isinstance(rgb, str) and len(rgb) >= 6:
            return int(rgb[-6:], 16)
        return 0xFFFFFF

    def set_cell_fill(self, row: int, column: int, color: int) -> None:
        self.worksheet.cell(row, column).fill = PatternFill(
            fill_type="solid", fgColor=f"FF{color:06X}",
        )

    def clear_cell_fill(self, row: int, column: int) -> None:
        self.worksheet.cell(row, column).fill = PatternFill(fill_type=None)

    def clear_cell_fill_range(self, start_row: int, end_row: int, column: int) -> None:
        for row in range(start_row, end_row + 1):
            self.clear_cell_fill(row, column)

    def set_a_column_bold(self, start_row: int, end_row: int, bold: bool) -> None:
        for row in range(start_row, end_row + 1):
            cell = self.worksheet.cell(row, 1)
            font = copy(cell.font)
            font.bold = bold
            cell.font = font

    def apply_a_column_group_borders(self, start_row: int, end_row: int) -> None:
        for row in range(start_row, end_row + 1):
            cell = self.worksheet.cell(row, 1)
            border = cell.border
            left = copy(border.left)
            left.style = left.style or "thin"
            top = copy(border.top)
            if row == start_row:
                top.style = top.style or "thin"
            cell.border = Border(
                left=left,
                right=copy(border.right),
                top=top,
                bottom=copy(border.bottom),
                diagonal=copy(border.diagonal),
                diagonal_direction=border.diagonal_direction,
                diagonalUp=border.diagonalUp,
                diagonalDown=border.diagonalDown,
                outline=border.outline,
                vertical=copy(border.vertical),
                horizontal=copy(border.horizontal),
            )

    def set_cell_comment(self, row: int, column: int, text: str) -> None:
        self.worksheet.cell(row, column).comment = Comment(text, "ConnLab")

    def _set_value(self, row: int, column: int, value: Any) -> None:
        self.worksheet.cell(row, column).value = _coerce_numeric(value, column)

    def _shift_print_area(self, row: int, count: int) -> None:
        print_area = self.worksheet.print_area
        if print_area is None:
            return
        ranges: list[str] = []
        raw_ranges = (
            tuple(print_area.ranges)
            if hasattr(print_area, "ranges")
            else tuple(part.strip() for part in str(print_area).split(","))
        )
        for cell_range in raw_ranges:
            address = str(cell_range).split("!", 1)[-1].replace("$", "")
            min_col, min_row, max_col, max_row = range_boundaries(address)
            if min_row >= row:
                min_row += count
                max_row += count
            elif max_row >= row:
                max_row += count
            ranges.append(
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{max_row}"
            )
        self.worksheet.print_area = ",".join(ranges)

    def _shift_image_anchors(self, row: int, count: int) -> None:
        zero_based_row = row - 1
        for image in self.worksheet._images:
            anchor = image.anchor
            if isinstance(anchor, str):
                continue
            start = getattr(anchor, "_from", None)
            end = getattr(anchor, "to", None)
            if start is not None and start.row >= zero_based_row:
                start.row += count
            if end is not None and end.row >= zero_based_row:
                end.row += count


class _CellAdapter:
    def __init__(self, sheet: OpenpyxlFeeSheetAdapter, row: int, column: int) -> None:
        self._sheet = sheet
        self._row = row
        self._column = column

    @property
    def Value(self) -> Any:  # noqa: N802
        return self._sheet.worksheet.cell(self._row, self._column).value

    @Value.setter
    def Value(self, value: Any) -> None:  # noqa: N802
        self._sheet._set_value(self._row, self._column, value)

    @property
    def Formula(self) -> Any:  # noqa: N802
        return self.Value

    @Formula.setter
    def Formula(self, value: Any) -> None:  # noqa: N802
        self._sheet.worksheet.cell(self._row, self._column).value = value


def _coerce_numeric(value: Any, column: int) -> Any:
    if column not in _NUMERIC_COLUMNS or not isinstance(value, str):
        return value
    normalized = value.strip()
    if not normalized:
        return None
    try:
        decimal = Decimal(normalized)
    except InvalidOperation:
        return value
    if decimal == decimal.to_integral_value():
        return int(decimal)
    return float(decimal)
