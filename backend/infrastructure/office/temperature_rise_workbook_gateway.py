"""Read-only tabular imports and new temperature-rise XLSX outputs; never runs macros."""
from __future__ import annotations

import csv
import re
from itertools import islice
from zipfile import ZipFile
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class TemperatureRiseWorkbookGateway:
    def read(self, source: Path) -> list[dict]:
        suffix = source.suffix.lower()
        try:
            if source.stat().st_size > 32 * 1024 * 1024:
                raise ValueError("Select a measurement file smaller than 32 MB.")
            if suffix == ".csv":
                try:
                    content = source.read_text(encoding="utf-8-sig")
                except UnicodeDecodeError:
                    content = source.read_text(encoding="gb18030")
                try:
                    dialect = csv.Sniffer().sniff(content[:16384], delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                sheets = [("CSV", _bounded_rows(csv.reader(content.splitlines(), dialect)))]
            elif suffix == ".xls":
                import xlrd
                book = xlrd.open_workbook(str(source), on_demand=True)
                try:
                    sheets = []
                    for sheet in book.sheets():
                        if sheet.nrows > 100000 or sheet.ncols > 256:
                            raise ValueError("Measurement tables support up to 100000 rows and 256 columns.")
                        rows = []
                        for row in range(sheet.nrows):
                            rows.append([xlrd.xldate_as_datetime(cell.value, book.datemode)
                                         if cell.ctype == xlrd.XL_CELL_DATE else cell.value for cell in sheet.row(row)])
                        sheets.append((sheet.name, rows))
                finally:
                    book.release_resources()
            elif suffix in {".xlsx", ".xlsm"}:
                with ZipFile(source) as archive:
                    if sum(entry.file_size for entry in archive.infolist()) > 256 * 1024 * 1024:
                        raise ValueError("Workbook expanded content exceeds 256 MB; export a smaller measurement table.")
                book = load_workbook(source, read_only=True, data_only=True, keep_links=False)
                try:
                    sheets = []
                    for sheet in book:
                        sheet.reset_dimensions()  # Do not allocate using potentially inflated cached dimensions.
                        sheets.append((sheet.title, _bounded_rows(sheet.iter_rows(values_only=True))))
                finally:
                    book.close()
            else:
                raise ValueError("Select a .csv, .xls, .xlsx or .xlsm measurement file.")
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Cannot read the measurement file. Check its format and remove workbook encryption.") from exc
        blocks = []
        for sheet_name, rows in sheets:
            if len(rows) > 100000 or any(len(row) > 256 for row in rows):
                raise ValueError("Measurement files support up to 100000 rows and 256 columns per sheet.")
            normalized = [[_cell(value) for value in row] for row in rows]
            for index, row in enumerate(normalized):
                headers = [str(value or "").strip() for value in row]
                ambient = next((column for column, value in enumerate(headers) if re.search(r"ambient|环境|環境", value, re.I)), None)
                current = next((column for column, value in enumerate(headers) if re.search(r"current|电流|電流", value, re.I)), None)
                if ambient is None or current is None or ambient == current or len(headers) < 3:
                    continue
                has_trace = any(re.search(r"scan|time|扫描|掃描|时间|時間", header, re.I) for header in headers)
                temperature_columns = [column for column, header in enumerate(headers)
                                       if column not in {ambient, current} and
                                       re.search(r"\(\s*[°℃Cc]\s*\)|temp|温度|溫度|(?:^|[_\s<])(?:HS|H|C|RS)(?:>|\s|$)", header, re.I)]
                if not has_trace and not temperature_columns:
                    continue
                data = []
                for offset, values in enumerate(normalized[index + 1:], index + 2):
                    padded = (values + [None] * len(headers))[:len(headers)]
                    if all(value in (None, "") for value in padded):
                        if data:
                            break
                        continue
                    try:
                        float(padded[ambient]); float(padded[current])
                    except (ValueError, TypeError):
                        if data and any(re.search(r"ambient|current|环境|电流", str(value), re.I) for value in padded):
                            break  # A repeated header is a real block boundary.
                        trace = [column for column, header in enumerate(headers) if re.search(r"scan|time|扫描|时间", header, re.I)]
                        if ((data and any(padded[column] not in (None, "") for column in trace)) or
                                any(_numeric(padded[column]) for column in temperature_columns)):
                            raise ValueError(f"{sheet_name}, row {offset}: Ambient or Current is missing or non-numeric. Correct this measurement row.")
                        if data:
                            break
                        continue
                    data.append({"row": offset, "values": padded})
                if not data:
                    continue
                channels = []
                for column, header in enumerate(headers):
                    if column in {ambient, current} or re.search(r"scan|time|扫描|掃描|时间|時間|alarm|警报", header, re.I):
                        continue
                    if not header or not any(_numeric(item["values"][column]) for item in data[:5]):
                        continue
                    label = re.search(r"<([^>]+)>", header)
                    label = label.group(1) if label else re.sub(r"\s*\([^)]*\)\s*$", "", header).strip()
                    pieces = re.split(r"[_\s]+", label, maxsplit=1)
                    channels.append({"column": column, "sample": pieces[0], "point": pieces[1] if len(pieces) == 2 else label})
                if not channels:
                    continue
                metadata = []
                current_id = re.match(r"\d+", headers[current])
                if current_id:
                    for values in normalized[:index]:
                        if values and _numeric(values[0]) and float(values[0]) == float(current_id.group(0)):
                            metadata = values
                            break
                blocks.append({"id": f"{len(blocks)}", "sheet": sheet_name, "header_row": index + 1,
                               "headers": headers, "rows": data, "source_rows": normalized,
                               "suggested_mapping": {"ambient": ambient, "current": current, "channels": channels},
                               "current_metadata": metadata})
        if not blocks:
            raise ValueError("No measurement table found. Include Ambient and Current channel labels above numeric data.")
        return sorted(blocks, key=lambda block: (block["sheet"].casefold() != "initial data", -len(block["rows"])))

    def write(self, source: Path, output: Path, block: dict, result: dict) -> Path:
        if source.resolve() == output.resolve() or output.exists():
            raise ValueError("Export requires a new output file; the source cannot be overwritten.")
        book = Workbook()
        setup = book.active
        setup.title = "Setup"
        chart_sheet = book.create_sheet("T-riseChart")
        derating = book.create_sheet("Derating")
        raw = book.create_sheet("Initial Data")
        options = result["settings"]
        setup_rows = [
            ["Temperature rise and derating", "Independent calculation tool"],
            ["Source", source.name], ["Data block", f'{block["sheet"]}, header row {block["header_row"]}'],
            ["Current input", options["current_mode"]], ["Applied current gain", options["applied_gain"]],
            ["Zero intercept constraint", options["zero_intercept"]],
            ["R²", "Centered total sum of squares; " + ("measured rows + display origin" if options["include_origin"] else "measured rows only")],
            ["Display origin", "Non-measured (0 A, 0 C), display/metric only" if options["include_origin"] else "Not included"],
            ["Candidate suggestion", "Contiguous current range <= 1%; not a thermal stability determination"],
            ["Confirmed source rows", ", ".join(str(point["row"]) for point in result["points"])],
            ["Target rise (C)", options["target_rise"]], ["Max-curve current (A)", result["target_current"]],
            ["Maximum temperature (C)", options["max_temperature"]], ["Derating factor", options["derating_factor"]],
            ["Interpretation", "No compliance or pass/fail assessment; extrapolation requires review."],
            ["Sample", "Point", "Source channel"]]
        for channel in options["mapping"]["channels"]:
            setup_rows.append([channel["sample"], channel["point"], block["headers"][channel["column"]]])
        coefficient_start = len(setup_rows) + 2
        setup_rows.extend([[], ["Curve", "a (full precision)", "b (full precision)", "c (full precision)", "Centered R²"],
                           ["Max"] + [result["max_curve"][key] for key in ("a", "b", "c", "r_squared")],
                           ["Avg of Max"] + [result["avg_curve"][key] for key in ("a", "b", "c", "r_squared")]])
        _rows(setup, setup_rows)
        _rows(raw, [block["headers"]] + [row["values"] for row in block["rows"]])
        _rows(chart_sheet, [block["headers"]] + [point["raw"] for point in result["points"]])
        rise_start = max(12, len(result["points"]) + 5)
        rise_headers = list(block["headers"])
        rise_headers[options["mapping"]["current"]] = "Applied Current (A)"
        rise_rows = []
        for point in result["points"]:
            values = list(point["raw"])
            values[options["mapping"]["current"]] = point["current"]
            for channel, rise in zip(options["mapping"]["channels"], point["rises"]):
                values[channel["column"]] = rise[2]
                rise_headers[channel["column"]] = f'{channel["sample"]}_{channel["point"]} ΔT (C)'
            rise_rows.append(values)
        _rows(chart_sheet, [rise_headers] + rise_rows, rise_start)
        summary_start = max(24, rise_start + len(result["points"]) + 6)
        samples = list(result["points"][0]["single_max"])
        display_points = ([{"current": 0, "single_max": {sample: 0 for sample in samples}, "rises": [], "maximum": 0, "average_of_max": 0}] if options["include_origin"] else []) + result["points"]
        summary = [["Applied Current (A)"] + [point["current"] for point in display_points]]
        if options["include_origin"]:
            summary.append(["Origin status", "Non-measured"])
        for sample in samples:
            for channel in [channel for channel in options["mapping"]["channels"] if channel["sample"] == sample]:
                summary.append([f'{sample}_{channel["point"]}'] +
                               [next((rise[2] for rise in point["rises"] if rise[0] == sample and rise[1] == channel["point"]), 0) for point in display_points])
            summary.append([f"{sample} Single Max"] + [point["single_max"][sample] for point in display_points])
        summary.extend([["Max"] + [point["maximum"] for point in display_points],
                        ["Avg of Max"] + [point["average_of_max"] for point in display_points]])
        _rows(chart_sheet, summary, summary_start)
        helper_start = summary_start + len(summary) + 4
        curve_data = [["Applied Current (A)", "Max", "Avg of Max", "Max fit", "Avg of Max fit"]]
        max_current = max(point["current"] for point in result["points"])
        for index in range(51):
            current = max_current * index / 50
            curve_data.append([current, None, None, _evaluate(result["max_curve"], current), _evaluate(result["avg_curve"], current)])
        for point in result["points"]:
            curve_data.append([point["current"], point["maximum"], point["average_of_max"], None, None])
        _rows(chart_sheet, curve_data, helper_start)
        chart_column = max(12, len(display_points) + 4)
        chart = _chart(chart_sheet, helper_start, len(curve_data), "Temperature Rise", "Applied Current (A)", "T-rise (C)", 4)
        chart_sheet.add_chart(chart, f"{get_column_letter(chart_column)}{summary_start}")
        for offset, (name, curve) in enumerate([( "Max", result["max_curve"]), ("Avg of Max", result["avg_curve"])]):
            chart_sheet.cell(summary_start + 21 + offset * 2, chart_column, f'{name}: ΔT = {curve["a"]:.6f} I² {curve["b"]:+.6f} I {curve["c"]:+.6f}')
            chart_sheet.cell(summary_start + 22 + offset * 2, chart_column, f'Centered R² = {curve["r_squared"]:.6f}')
        _rows(derating, [["Ambient (C)", "Allowable Rise (C)", "Basic Current (A)", "Derated Current (A)"]] +
              [[row["ambient"], row["allowable_rise"], row["basic_current"], row["derated_current"]] for row in result["derating"]])
        derating.add_chart(_chart(derating, 1, len(result["derating"]) + 1, "Derating", "Ambient Temperature (C)", "Applied Current (A)", 3, skip_first=True), "F2")
        for sheet in book:
            _style(sheet)
        for row in setup.iter_rows(min_row=coefficient_start + 1, max_row=coefficient_start + 2, min_col=2, max_col=5):
            for cell in row:
                cell.number_format = "0.000000000000000"
        for start in (1, rise_start):
            for cell in chart_sheet[start]:
                cell.fill = PatternFill("solid", fgColor="D9D9D9")
                cell.font = Font(bold=True)
        for index, row in enumerate(summary, summary_start):
            label = str(row[0])
            color = "C6E0F5" if label in {"Max", "Avg of Max"} else "D5AD88" if "Single Max" in label else "D9D9D9"
            for cells in chart_sheet.iter_rows(min_row=index, max_row=index, max_col=len(display_points) + 1):
                for cell in cells:
                    cell.fill = PatternFill("solid", fgColor=color)
        chart_sheet.print_area = f"A1:{get_column_letter(max(len(block['headers']), chart_column + 11))}{max(helper_start - 2, summary_start + 28)}"
        derating.print_area = f"A1:Q{max(28, len(result['derating']) + 2)}"
        setup.column_dimensions["B"].width = 85
        chart_sheet.column_dimensions["A"].width = 23
        output.parent.mkdir(parents=True, exist_ok=True)
        try:
            with output.open("xb") as stream:
                book.save(stream)
        finally:
            book.close()
        return output


def _cell(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return value


def _bounded_rows(iterator):
    rows = []
    cells = 0
    for row in islice(iterator, 100001):
        if len(row) > 256 or len(rows) == 100000:
            raise ValueError("Measurement tables support up to 100000 rows and 256 columns.")
        cells += len(row)
        if cells > 2000000:
            raise ValueError("Measurement sheet exceeds 2 million cells; export a smaller table.")
        rows.append(list(row))
    return rows


def _numeric(value):
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def _rows(sheet, rows, start=1):
    for index, row in enumerate(rows, start):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(index, column, value)
            if isinstance(value, str):
                cell.data_type = "s"  # CSV labels and timestamps never become executable formulas.


def _evaluate(curve, current):
    return (curve["a"] * current + curve["b"]) * current + curve["c"]


def _chart(sheet, start, count, title, x_title, y_title, series_count, skip_first=False):
    chart = ScatterChart()
    chart.title, chart.x_axis.title, chart.y_axis.title = title, x_title, y_title
    chart.width, chart.height = 23, 13
    for column in range(3 if skip_first else 2, series_count + 2):
        series = Series(Reference(sheet, min_col=column, min_row=start + 1, max_row=start + count - 1),
                        Reference(sheet, min_col=1, min_row=start + 1, max_row=start + count - 1), title=str(sheet.cell(start, column).value))
        series.graphicalProperties.line.solidFill = "ED7D31" if column % 2 == 0 else "203864"
        if not skip_first and column < 4:
            series.marker.symbol = "circle"
            series.marker.size = 5
            series.marker.graphicalProperties.solidFill = "ED7D31" if column % 2 == 0 else "203864"
            series.marker.graphicalProperties.line.solidFill = "ED7D31" if column % 2 == 0 else "203864"
            series.graphicalProperties.line.noFill = True
        chart.series.append(series)
    return chart


def _style(sheet):
    thin = Side(style="thin", color="000000")
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if isinstance(cell.value, float):
                    cell.number_format = "0.0"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
        cell.font = Font(bold=True)
    for column in range(1, min(sheet.max_column, 40) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.freeze_panes = "B2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_area = sheet.dimensions
