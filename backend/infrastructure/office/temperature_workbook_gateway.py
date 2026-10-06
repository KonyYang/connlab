"""Read scanner values and produce independent, macro-free native Excel chart workbooks."""

from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO, StringIO
from itertools import islice
import csv
from math import isfinite
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, BadZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.data_source import NumData, NumVal, NumFmt
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter
import xlrd

from backend.application.temperature_data_preparation import PreparedData
from backend.domain.temperature_data import DataSelection, WorkbookTable, column_letter
from backend.domain.temperature_rise import Coefficients, TemperatureRiseAnalysis, DeratingAnalysis
from backend.infrastructure.office.temperature_rise_report_sheet import build_rise_report

MAX_ROWS = 20000
MAX_COLUMNS = 256
MAX_CELLS = 1000000
MAX_FILE_BYTES = 25 * 1024 * 1024


class TemperatureWorkbookGateway:
    def read_upload(self, content: bytes, file_name: str, *, sheet_name: str | None = None) -> WorkbookTable:
        if len(content) > MAX_FILE_BYTES:
            raise ValueError('Select a data file smaller than 25 MB.')
        with TemporaryDirectory(prefix='connlab-temperature-import-') as directory:
            source = Path(directory) / Path(file_name).name
            source.write_bytes(content)
            return self.read(source, sheet_name=sheet_name)

    def render(self, file_name: str, *, table: WorkbookTable, selection: DataSelection,
               prepared: PreparedData, analysis: TemperatureRiseAnalysis,
               maximum_coefficients: Coefficients, average_coefficients: Coefficients,
               target_rise: float | None, derating: DeratingAnalysis | None) -> bytes:
        with TemporaryDirectory(prefix='connlab-temperature-export-') as directory:
            output = Path(directory) / Path(file_name).name
            self.write(output, table=table, selection=selection, prepared=prepared, analysis=analysis,
                maximum_coefficients=maximum_coefficients, average_coefficients=average_coefficients,
                target_rise=target_rise, derating=derating)
            return output.read_bytes()

    def read(self, source_path: Path, *, sheet_name: str | None = None) -> WorkbookTable:
        source = Path(source_path)
        if source.suffix.lower() not in ('.xlsx', '.xlsm', '.xls', '.csv'):
            raise ValueError('Select an Excel .xlsx, .xlsm, .xls or CSV .csv file.')
        if source.stat().st_size > MAX_FILE_BYTES:
            raise ValueError('Select a data file smaller than 25 MB.')
        try:
            if source.suffix.lower() == '.csv':
                return _read_csv(source, sheet_name)
            if source.suffix.lower() == '.xls':
                return self._read_xls(source, sheet_name)
            with ZipFile(source) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('The expanded workbook is too large. Export only the scanner data sheet.')
            book = load_workbook(source, read_only=True, data_only=True, keep_links=False)
            try:
                name = _choose_sheet(book.sheetnames, sheet_name)
                sheet = book[name]
                _check_dimensions(sheet.max_row or 0, sheet.max_column or 0)
                rows = tuple(tuple(_cell(value) for value in row) for row in sheet.iter_rows(values_only=True))
                return _table(source.name, tuple(book.sheetnames), name, rows)
            finally:
                book.close()
        except (BadZipFile, KeyError, xlrd.XLRDError, OSError) as exc:
            raise ValueError('The workbook could not be read. Check that it is an unencrypted Excel file.') from exc

    def _read_xls(self, source: Path, requested: str | None) -> WorkbookTable:
        book = xlrd.open_workbook(source, on_demand=True)
        try:
            name = _choose_sheet(book.sheet_names(), requested)
            sheet = book.sheet_by_name(name)
            _check_dimensions(sheet.nrows, sheet.ncols)
            def value(row: int, col: int):
                cell = sheet.cell(row, col)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).isoformat(sep=' ')
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    return None
                if cell.ctype == xlrd.XL_CELL_ERROR:
                    return xlrd.error_text_from_code.get(cell.value, '#ERROR')
                return _cell(cell.value)
            rows = tuple(tuple(value(r, c) for c in range(sheet.ncols)) for r in range(sheet.nrows))
            return _table(source.name, tuple(book.sheet_names()), name, rows)
        finally:
            book.release_resources()

    def write(self, output_path: Path, *, table: WorkbookTable, selection: DataSelection,
              prepared: PreparedData, analysis: TemperatureRiseAnalysis,
              maximum_coefficients: Coefficients, average_coefficients: Coefficients,
              target_rise: float | None, derating: DeratingAnalysis | None) -> None:
        output = Path(output_path)
        if output.exists() or output.suffix.lower() != '.xlsx':
            raise ValueError('Choose a new .xlsx output file; existing files are preserved.')
        book = Workbook()
        initial = book.active
        initial.title = 'Initial Data'
        rise = book.create_sheet('T-riseChart')
        derated = book.create_sheet('Derating')
        caches: dict[str, dict[str, float | str]] = {name: {} for name in book.sheetnames}
        _initial_sheet(initial, table, selection, prepared)
        build_rise_report(rise, table=table, selection=selection, analysis=analysis,
            maximum=maximum_coefficients, average=average_coefficients, target=target_rise,
            caches=caches[rise.title], chart=_chart('Temperature Rise vs Applied Current', 'Applied Current (A)', 'T-rise (°C)'))
        _derating_sheet(derated, derating, caches[derated.title])
        for sheet in (initial, derated):
            _style_sheet(sheet)
        book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
        memory = BytesIO()
        book.save(memory)
        book.close()
        # Cache formula values as well as chart points so viewers can display the file immediately.
        content = _formula_caches(memory.getvalue(), caches)
        with output.open('xb') as stream:
            stream.write(content)


def _choose_sheet(names, requested):
    if requested and requested not in names:
        raise ValueError('The selected sheet is no longer available. Choose a listed sheet.')
    return requested or next((name for name in names if name.replace(' ', '').lower() == 'initialdata'), names[0])


def _read_csv(source: Path, requested: str | None) -> WorkbookTable:
    name = _choose_sheet(('Initial Data',), requested)
    text = _decode_csv(source.read_bytes())
    delimiter = _csv_delimiter(text)
    rows = []
    width = 0
    try:
        for row in csv.reader(StringIO(text, newline=''), delimiter=delimiter, strict=True):
            width = max(width, len(row))
            _check_dimensions(len(rows) + 1, width)
            # Keep decimal readings as literal source text; preparation parses them without rounding.
            rows.append(tuple(value if value != '' else None for value in row))
    except csv.Error as exc:
        raise ValueError('The CSV could not be read. Check its separators and quoted fields.') from exc
    return _table(source.name, ('Initial Data',), name, rows)


def _decode_csv(content: bytes) -> str:
    if content.startswith((b'\xff\xfe', b'\xfe\xff')):
        encodings = ('utf-16',)
    else:
        encodings = ('utf-8-sig', 'gb18030')
    for encoding in encodings:
        try:
            text = content.decode(encoding)
        except UnicodeError:
            continue
        if '\x00' in text:
            break
        return text
    raise ValueError('The CSV encoding could not be read. Export it as UTF-8, UTF-16 or GB18030 text.')


def _csv_delimiter(text: str) -> str:
    # Scanner metadata is not a uniform table, so whole-file dialect sniffing is unreliable.
    # Compare parsed records, not raw separator counts inside quoted labels.
    scores = []
    for delimiter in ('\t', ',', ';'):
        try:
            widths = [len(row) for row in islice(csv.reader(StringIO(text, newline=''), delimiter=delimiter), 100)]
        except csv.Error:
            continue
        scores.append(((sum(width > 1 for width in widths), sum(widths)), delimiter))
    if not scores or max(scores)[0][0] == 0:
        raise ValueError('The CSV has no readable columns. Use comma, tab or semicolon separators.')
    return max(scores)[1]


def _check_dimensions(rows: int, columns: int) -> None:
    if rows > MAX_ROWS or columns > MAX_COLUMNS or rows * columns > MAX_CELLS:
        raise ValueError('Limit the sheet to 20,000 rows, 256 columns and 1,000,000 cells before import.')


def _cell(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not isfinite(value):
        return None
    return value if value is None or isinstance(value, (str, float, int, bool)) else str(value)


def _table(file_name, names, name, rows):
    # Strip trailing empty cells/rows, keeping original indices within the used region.
    trimmed = [list(row) for row in rows]
    while trimmed and all(value is None for value in trimmed[-1]):
        trimmed.pop()
    for row in trimmed:
        while row and row[-1] is None:
            row.pop()
    if not trimmed:
        raise ValueError('The selected sheet is empty. Choose a sheet containing scanner readings.')
    return WorkbookTable(file_name, names, name, tuple(tuple(row) for row in trimmed))


def _text(sheet, coordinate, value):
    # Uploaded labels remain literal text even if they begin with =, +, - or @.
    cell = sheet[coordinate]
    cell.value = str(value)
    cell.data_type = 's'


def _initial_sheet(sheet, table, selection, prepared):
    _text(sheet, 'A1', 'Confirmed Initial Data')
    _text(sheet, 'A2', f'Source: {table.file_name} / {table.sheet_name}')
    _text(sheet, 'A3', f'Source rows {selection.start_row}–{selection.end_row}; excluded: '
          + (', '.join(map(str, selection.excluded_rows)) or 'none')
          + f'; current scale: {selection.current_multiplier:g} A per source unit')
    _text(sheet, 'A4', 'Warnings acknowledged: ' + ('; '.join(issue.message for issue in prepared.issues) or 'none'))
    headers = ['Original Row']
    original_header = table.rows[selection.header_row - 1]
    for slot, column in enumerate(selection.temperature_columns):
        label = original_header[column - 1] if column <= len(original_header) else ''
        headers.append(f'Sample {slot // selection.thermocouples_per_sample + 1} / TC {slot % selection.thermocouples_per_sample + 1} [{column_letter(column)}: {label}] (°C)')
    headers += [f'Ambient [{column_letter(selection.ambient_column)}] (°C)',
                f'Current [{column_letter(selection.current_column)}] (A)']
    headers += [f'Retained Current [{column_letter(column)}: {original_header[column - 1]}]'
                for column in prepared.retained_current_columns]
    for col, header in enumerate(headers, 1):
        _text(sheet, f'{get_column_letter(col)}5', header)
    for row in prepared.measurements:
        sheet.append([row.source_row, *row.temperatures, row.ambient, row.current])
        source = table.rows[row.source_row - 1]
        first_retained = len(selection.temperature_columns) + 4
        for offset, column in enumerate(prepared.retained_current_columns):
            value = source[column - 1] if column <= len(source) else None
            coordinate = f'{get_column_letter(first_retained + offset)}{sheet.max_row}'
            if isinstance(value, str):
                _text(sheet, coordinate, value)
            else:
                sheet[coordinate] = value
    sheet.auto_filter.ref = f'A5:{get_column_letter(len(headers))}{sheet.max_row}'
    sheet.freeze_panes = 'B6'
    sheet.row_dimensions[5].height = 48


def _derating_sheet(sheet, result, caches):
    _text(sheet, 'A1', 'Current vs Ambient Temperature')
    if result is None:
        _text(sheet, 'A3', 'Generate Derating in ConnLab and download again to include this chart.')
        return
    _text(sheet, 'A2', f'Max working temperature: {result.max_temperature:g} °C; step: {result.step:g} °C; current factor: 80%')
    for coordinate, value in {'F5': 'AVG a', 'G5': 'AVG b', 'H5': 'AVG c', 'I5': 'Max Temp (°C)',
                              'F9': 'Ambient Point (°C)', 'G9': 'Basic (A)', 'H9': '80% Derating (A)'}.items():
        _text(sheet, coordinate, value)
    for coordinate, value in {'F6': result.coefficients.a, 'G6': result.coefficients.b, 'H6': result.coefficients.c,
                              'I6': result.max_temperature, 'F10': result.annotation.ambient}.items():
        sheet[coordinate] = value
    for coordinate in ('F6', 'G6', 'H6'):
        sheet[coordinate].number_format = '0.000000'
    for col, value in enumerate(('Ambient (°C)', 'Basic (A)', '80% Derating (A)'), 1):
        _text(sheet, f'{get_column_letter(col)}5', value)
    for row, point in enumerate(result.points, 6):
        sheet.cell(row, 1, point.ambient)
        sheet.cell(row, 2, f'=(-$G$6+SQRT($G$6^2+4*$F$6*($I$6-A{row})))/(2*$F$6)')
        sheet.cell(row, 3, f'=B{row}*0.8')
        caches[f'B{row}'], caches[f'C{row}'] = point.basic, point.derated
    sheet['G10'] = '=(-$G$6+SQRT($G$6^2+4*$F$6*($I$6-F10)))/(2*$F$6)'
    sheet['H10'] = '=G10*0.8'
    # Excel can inherit label formatting from these source cells despite chart numFmt.
    for coordinate in ('G10', 'H10'):
        sheet[coordinate].number_format = '0.0'
    caches['G10'], caches['H10'] = result.annotation.basic, result.annotation.derated
    chart = _chart('Current Carrying Capacity De-rating Curve', 'Ambient Temperature [°C]', 'Current [Amp]')
    _style_derating_chart(chart)
    end = 5 + len(result.points)
    for col, title, color in ((2, 'Basic', '8B0000'), (3, '80% Derating', 'FFA500')):
        series = _series(sheet, 1, col, 6, end, title, color, caches)
        series.smooth = True
        series.marker.symbol = 'none'
        chart.series.append(series)
    for col, color, title in ((7, '8B0000', 'Basic At Ambient Point'), (8, 'FFA500', 'Derated At Ambient Point')):
        series = _series(sheet, 6, col, 10, 10, title, color, caches)
        series.marker.symbol = 'circle'
        series.marker.size = 7
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.solidFill = None
        series.graphicalProperties.line.noFill = True
        series.dLbls = DataLabelList(showVal=True, showLegendKey=False, showCatName=False,
            showSerName=False, showPercent=False, numFmt='0.0', dLblPos='r',
            txPr=_chart_text_properties(color=color, bold=True))
        chart.series.append(series)
    # The guide ends at the selected Basic point and follows edits to the ambient input.
    for coordinate, value in {'F13': '=$F$10', 'F14': '=$F$10', 'G13': 0, 'G14': '=$G$10'}.items():
        sheet[coordinate] = value
    caches.update(F13=result.annotation.ambient, F14=result.annotation.ambient, G14=result.annotation.basic)
    guide = _series(sheet, 6, 7, 13, 14, 'Ambient Point', '0000FF', caches)
    guide.graphicalProperties.line.prstDash = 'dash'
    guide.marker.symbol = 'none'
    chart.series.append(guide)
    chart.x_axis.scaling.max = result.max_temperature
    chart.legend.position = 'b'
    chart.legend.legendEntry = [LegendEntry(idx=index, delete=True) for index in (2, 3, 4)]
    sheet.add_chart(chart, 'E17')
    sheet.freeze_panes = None


def _chart_text_properties(*, color='000000', bold=False, size=1100):
    properties = CharacterProperties(b=bold, sz=size, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=properties), endParaRPr=properties)])


def _style_derating_chart(chart):
    chart.graphical_properties = GraphicalProperties(solidFill='FFFFE0')
    chart.plot_area.spPr = GraphicalProperties(solidFill='FFFFFF')
    chart.title.txPr = _chart_text_properties(bold=True, size=1400)
    for paragraph in chart.title.tx.rich.p:
        paragraph.pPr = ParagraphProperties(defRPr=CharacterProperties(b=True, sz=1400))
    for axis in (chart.x_axis, chart.y_axis):
        axis.title.txPr = _chart_text_properties(bold=True)
        for paragraph in axis.title.tx.rich.p:
            paragraph.pPr = ParagraphProperties(defRPr=CharacterProperties(b=True, sz=1100))
        axis.txPr = _chart_text_properties()
        axis.majorGridlines.spPr.line.solidFill = '808080'
        axis.majorGridlines.spPr.line.prstDash = 'dash'
    chart.x_axis.majorUnit = 5
    chart.x_axis.numFmt = NumFmt(formatCode='0', sourceLinked=False)
    chart.layout = Layout(manualLayout=ManualLayout(layoutTarget='inner', xMode='edge', yMode='edge',
        x=.1, y=.09, w=.87, h=.76))
    chart.legend.txPr = _chart_text_properties()
    chart.legend.layout = Layout(manualLayout=ManualLayout(x=.07, y=.94, w=.45, h=.04,
        xMode='edge', yMode='edge'))


def _chart(title, x_title, y_title):
    chart = ScatterChart()
    chart.title = title
    chart.title.overlay = False
    chart.x_axis.title, chart.y_axis.title = x_title, y_title
    chart.x_axis.scaling.min = chart.y_axis.scaling.min = 0
    chart.x_axis.axPos, chart.y_axis.axPos = 'b', 'l'
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.tickLblPos = 'nextTo'
        axis.numFmt = NumFmt(formatCode='0.0', sourceLinked=False)
        axis.majorTickMark = 'out'
        axis.crosses = 'autoZero'
        grid = GraphicalProperties()
        grid.line.solidFill = 'D8E0EA'
        grid.line.width = 6350
        axis.majorGridlines = ChartLines(spPr=grid)
    chart.layout = Layout(manualLayout=ManualLayout(layoutTarget='inner', xMode='edge', yMode='edge',
        x=.14, y=.12, w=.80, h=.66))
    chart.legend.overlay = False
    chart.legend.layout = Layout(manualLayout=ManualLayout(x=.2, y=.94, w=.65, h=.04, xMode='edge', yMode='edge'))
    chart.width, chart.height = 24, 14
    chart.style = 2
    return chart


def _series(sheet, x_col, y_col, start, end, title, color, caches=None):
    x = Reference(sheet, min_col=x_col, min_row=start, max_row=end)
    y = Reference(sheet, min_col=y_col, min_row=start, max_row=end)
    series = Series(y, x, title=title)
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = 22000
    for column, reference in ((x_col, series.xVal.numRef), (y_col, series.yVal.numRef)):
        values = [(caches or {}).get(sheet.cell(r, column).coordinate, sheet.cell(r, column).value) for r in range(start, end + 1)]
        reference.numCache = NumData(formatCode='0.000000', ptCount=len(values),
                                    pt=[NumVal(idx=i, v=value) for i, value in enumerate(values)])
    return series


def _style_sheet(sheet):
    for row in sheet:
        for cell in row:
            if cell.value is None:
                continue
            cell.font = Font(name='Aptos', size=11, color='172033')
            cell.alignment = Alignment(vertical='center')
            if isinstance(cell.value, (int, float)) or cell.data_type == 'f':
                if cell.number_format == 'General':
                    cell.number_format = '0.000'
    sheet['A1'].font = Font(name='Aptos', size=18, bold=True, color='164AA3')
    sheet.row_dimensions[1].height = 30
    for cell in sheet[5]:
        if cell.value is not None:
            cell.fill = PatternFill('solid', fgColor='E8EEF6')
            cell.font = Font(name='Aptos', size=11, bold=True, color='172033')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    for index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 20
    sheet.sheet_view.showGridLines = False


def _formula_caches(content, caches):
    namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('', namespace)
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            for index, values in enumerate(caches.values(), 1):
                if entry.filename == f'xl/worksheets/sheet{index}.xml' and values:
                    root = ET.fromstring(data)
                    for cell in root.iter(f'{{{namespace}}}c'):
                        value = values.get(cell.get('r'))
                        if value is not None:
                            node = cell.find(f'{{{namespace}}}v')
                            if node is None:
                                node = ET.SubElement(cell, f'{{{namespace}}}v')
                            if isinstance(value, str):
                                cell.set('t', 'str')
                                node.text = value
                            else:
                                node.text = repr(value)
                    data = ET.tostring(root, encoding='utf-8')
            target.writestr(entry, data)
    return output.getvalue()
