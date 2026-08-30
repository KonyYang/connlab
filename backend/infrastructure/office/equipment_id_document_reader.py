"""Read equipment references from a project EquipmentID.docx without mutation."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document

from backend.application.equipment_report_update_service import equipment_reference_key


@dataclass(frozen=True, slots=True)
class EquipmentIdDocumentReadResult:
    source_path: Path
    references: tuple[str, ...]


class EquipmentIdDocumentReader:
    """Extract ordered, case-insensitively unique equipment references."""

    def read(self, source_path: Path) -> EquipmentIdDocumentReadResult:
        path = Path(source_path)
        if path.suffix.casefold() != ".docx" or not path.is_file():
            raise FileNotFoundError(f"EquipmentID.docx does not exist: {path}")
        document = Document(path)
        candidates = [paragraph.text for paragraph in document.paragraphs]
        candidates.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        references: list[str] = []
        seen: set[str] = set()
        for candidate in candidates:
            value = _clean_reference(candidate)
            if not value:
                continue
            key = equipment_reference_key(value)
            if key in seen:
                continue
            seen.add(key)
            references.append(value)
        if not references:
            raise ValueError("EquipmentID.docx does not contain equipment references.")
        return EquipmentIdDocumentReadResult(path, tuple(references))
def _clean_reference(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\x07", " ").strip())
